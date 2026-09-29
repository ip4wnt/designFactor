"""Адаптер к движку `pptx_template_parser_generator` (парсер шаблона → Template
JSON 3.0 → генерация слайдов в РЕАЛЬНЫХ макетах шаблона).

Почему отдельный модуль. Прежний сборщик (app/pipeline/assembly.py) очищал
шаблон до пустого layout и рисовал синтетические текстовые боксы — из-за этого
на выходе терялись фон, логотипы, декоративные фигуры и фирменные макеты
шаблона, а «три варианта» отличались только геометрией зон. Движок коллеги
делает ровно то, что требует ТЗ (п.1 «парсинг шаблона для воспроизводимости»,
п.3 «генерация слайдов в контексте шаблона»): он раскладывает шаблон на
варианты-макеты (`variants`) с семантическими ролями слотов, контрактами
вместимости и каталогом выбора для модели, а затем создаёт слайды из
оригинальных `slideLayout`-ов, заполняя только редактируемые слоты.

Этот модуль:
  * подключает пакет из соседней папки репозитория (`../pptx_template_parser_generator`);
  * кэширует Template JSON по пути и mtime файла шаблона;
  * оборачивает наш OpenAI-совместимый клиент под интерфейс `model.json(...)`,
    который ожидает планировщик коллеги (`pptx_template_parser.planning`);
  * извлекает из Template JSON расширенную дизайн-систему (шкалу кеглей по
    ролям, гарнитуры, цвета текста, каталог макетов с геометрией слотов);
  * из одного плана строит ТРИ визуально разных варианта вёрстки
    (`VARIANT_STRATEGIES`) — ось различия задокументирована в описании
    каждой стратегии и в ARCHITECTURE.md;
  * рендерит план через `scripts.generate.create_presentation` и
    пост-обрабатывает файл: вставляет нативные графики/таблицы в слоты
    макета, убирает пустые плейсхолдеры, удаляет неиспользованные макеты
    (чтобы файл не тащил все медиа шаблона).
"""
from __future__ import annotations

import logging
import re
import sys
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu, Pt

from app.schemas.content_plan import ContentBlock, ContentBlockType, ContentPlan
from app.schemas.design_manifest import DesignManifest

logger = logging.getLogger(__name__)

# backend/app/pipeline/template_engine.py -> designfactor/
REPO_ROOT = Path(__file__).resolve().parents[3]
ENGINE_DIR = REPO_ROOT / "pptx_template_parser_generator"
PROMPTS_DIR = ENGINE_DIR / "prompts"


def _ensure_engine_on_path() -> None:
    if not ENGINE_DIR.is_dir():
        raise RuntimeError(
            f"Не найден движок шаблонов {ENGINE_DIR} — ожидается папка "
            "pptx_template_parser_generator рядом с backend/ (см. README)."
        )
    if str(ENGINE_DIR) not in sys.path:
        sys.path.insert(0, str(ENGINE_DIR))


_ensure_engine_on_path()

from pptx_template_parser import build_template, validate_generation_plan  # noqa: E402
from pptx_template_parser.model import ModelError, parse_json_response  # noqa: E402
from pptx_template_parser.planning import generate_plan  # noqa: E402
from scripts.generate import create_presentation  # noqa: E402

EMU_PER_CM = 360000

# --------------------------------------------------------------------------
# Template JSON: построение и кэш
# --------------------------------------------------------------------------

_TEMPLATE_CACHE: dict[tuple[str, float], dict[str, Any]] = {}


def build_template_json(pptx_path: str | Path) -> dict[str, Any]:
    """Template JSON 3.0 по файлу шаблона (кэш по пути + mtime)."""
    path = Path(pptx_path)
    key = (str(path.resolve()), path.stat().st_mtime)
    cached = _TEMPLATE_CACHE.get(key)
    if cached is not None:
        return cached
    template = build_template(str(path))
    _enrich_slot_typography(template, path)
    _fix_metric_roles(template)
    _collect_layout_decor(template, path)
    # Источник дизайна: макеты / слайды-образцы / смешанный / нет примеров.
    # Слайды-образцы становятся композициями (source=sample); при полном
    # отсутствии примеров подмешиваем собственную библиотеку компоновок.
    try:
        from app.pipeline.sample_slides import augment_template

        augment_template(template, path)
    except Exception:  # noqa: BLE001 — разбор образцов не должен ломать парсинг
        logger.exception("Не удалось разобрать слайды-образцы шаблона")
        template.setdefault("design_source", "layouts")
    try:
        from app.pipeline.library import build_library_variants

        # Шаблон без единого слота заголовка (все «макеты» — авторские слайды с
        # текстовыми блоками без плейсхолдеров): своих компоновок для контента
        # с заголовком у него нет — берём всю библиотеку на его шрифтах и цветах
        has_title_slot = any(
            sl.get("semantic_role") in {"slide_title", "presentation_title", "section_title"} and not sl.get("virtual")
            for v in template.get("variants") or []
            for sl in (v.get("slots") or {}).values()
        )
        if template.get("design_source") == "library" or not has_title_slot:
            if not has_title_slot and template.get("design_source") != "library":
                logger.info("В шаблоне нет слотов заголовка — подмешиваем всю библиотеку компоновок")
                template["design_source"] = "library"
                info = template.setdefault("design_source_info", {})
                if isinstance(info, dict):
                    info["design_source"] = "library"
                    info["reason"] = "нет слотов заголовка"
            only_kinds = None
        else:
            # Таблица, график и крупный тезис — если среди образцов их нет,
            # дорисуем библиотечные на базовом макете (иначе таблица попадает
            # в случайный макет с телом, например «Визитку»)
            sample_kinds = {
                (v.get("composition") or {}).get("kind")
                for v in template.get("variants") or []
                if v.get("source") == "sample"
            }
            only_kinds = {"table", "chart", "statement"} - sample_kinds
            # Разделитель: если ни образцы, ни макеты шаблона не дают композицию
            # раздела — рисуем библиотечный (иначе раздел собирался бы как «тезис»)
            has_section = any(
                (v.get("composition") or {}).get("kind") == "section" or ((v.get("inferred") or {}).get("layout") or {}).get("narrative_role") == "section"
                for v in template.get("variants") or []
            )
            if not has_section:
                only_kinds.add("section")
            # Метрики: образцы на 1–2 цифры не вместят 3–4 показателя (третья цифра
            # терялась или уходила мелким текстом в карточки) — добавляем библиотечные
            max_sample_metrics = max(
                (sum(1 for sl in (v.get("slots") or {}).values() if sl.get("semantic_role") == "metric_value" and not sl.get("virtual"))
                 for v in template.get("variants") or [] if v.get("source") == "sample" and (v.get("composition") or {}).get("kind") == "metrics"),
                default=0,
            )
            if max_sample_metrics < 3:
                only_kinds.add("metrics")
            # Обычные тезисы и текст: если ни образцы, ни макеты шаблона не дают
            # композиции с телом/карточками (VK WorkSpace: все макеты — «Титульный
            # слайд» и «Заголовок», образцы — обложки, галереи, проблемы), слайду
            # с 2–4 тезисами достаётся «только заголовок» и содержание пропадает
            all_kinds = {(v.get("composition") or {}).get("kind") for v in template.get("variants") or []}
            layout_has_body = any(
                v.get("source") not in {"sample", "library"}
                and any(sl.get("semantic_role") in {"body", "column_body"} for sl in (v.get("slots") or {}).values())
                for v in template.get("variants") or []
            )
            if not ({"text", "image_text"} & all_kinds) and not layout_has_body:
                only_kinds.add("text")
            if not ({"cards", "columns"} & all_kinds) and not layout_has_body:
                only_kinds.update({"cards", "columns"})
            # Таблица/график внутри образца другого вида (таблица-линейка на
            # обложке VK WorkSpace) не делают его композицией таблицы/графика:
            # драматургия отдаёт такие слайды только видам table/chart, поэтому
            # библиотечные добавляются, пока среди образцов нет этих видов
        if only_kinds is None or only_kinds:
            for variant in build_library_variants(template, only_kinds):
                template["variants"].append(variant)
                template.setdefault("variant_catalog", []).append(_library_catalog_entry(variant))
    except Exception:  # noqa: BLE001
        logger.exception("Не удалось построить композиции библиотеки")
    _TEMPLATE_CACHE[key] = template
    return template


def _picture_busyness(shape) -> dict[str, Any]:
    """Насколько фоновая картинка макета «занята» в зоне контента (под
    заголовком, над подвалом).

    Возвращает `busy` (0 — ровный фон; больше — контуры плашек, крупные числа,
    объекты, поверх которых контент библиотеки будет мешаться), `faint`/`strong`
    (доли пикселей с бледными/заметными отличиями от медианы; тонкие контуры
    учитываются через max/min-фильтр), `flat` (доля ровного фона), `bg_hex`
    (цвет ровного фона; отсутствует, если картинка в основном прозрачна).
    """
    try:
        from io import BytesIO

        from PIL import Image, ImageFilter

        source = Image.open(BytesIO(shape.image.blob))
        result: dict[str, Any] = {}
        if "A" in source.getbands():
            alpha = source.getchannel("A")
            alpha.thumbnail((128, 128))
            alpha_pixels = list(alpha.getdata())
            if alpha_pixels and sorted(alpha_pixels)[len(alpha_pixels) // 2] < 200:
                result["transparent"] = True
            source = Image.alpha_composite(Image.new("RGBA", source.size, (128, 128, 128, 255)), source.convert("RGBA"))
        gray = source.convert("L")
        while gray.width > 1100:
            gray = gray.reduce(2)
        target = (256, 144)
        box = (int(target[0] * 0.02), int(target[1] * 0.22), int(target[0] * 0.98), int(target[1] * 0.9))
        base = list(gray.resize(target).crop(box).getdata())
        if not base:
            return {"busy": 0.0}
        median = sorted(base)[len(base) // 2]
        mx = list(gray.filter(ImageFilter.MaxFilter(5)).resize(target, Image.NEAREST).crop(box).getdata())
        mn = list(gray.filter(ImageFilter.MinFilter(5)).resize(target, Image.NEAREST).crop(box).getdata())
        deviation = [max(abs(a - median), abs(b - median)) for a, b in zip(mx, mn)]
        faint = sum(1 for v in deviation if v > 12) / len(deviation)
        strong = sum(1 for v in deviation if v > 30) / len(deviation)
        flat = sum(1 for v in base if abs(v - median) <= 6) / len(base)
        result.update({"busy": round(0.5 * faint + strong, 3), "faint": round(faint, 3), "strong": round(strong, 3), "flat": round(flat, 3)})
        if not result.get("transparent"):
            rgb = source.convert("RGB").resize(target).crop(box)
            channels = list(zip(*rgb.getdata()))
            result["bg_hex"] = "".join(f"{sorted(ch)[len(ch) // 2]:02X}" for ch in channels)
        return result
    except Exception:  # noqa: BLE001
        return {"busy": 0.0}


def _collect_layout_decor(template: dict[str, Any], path: Path) -> None:
    """Статические фигуры макетов (не плейсхолдеры): декоративный текст, плашки,
    линии — в `variant["decor"]` (см, с признаком текста). Нужны библиотеке,
    чтобы не рисовать контент поверх декоративных «03 04» и плашек макета.
    """
    try:
        prs = Presentation(str(path))
    except Exception:  # noqa: BLE001
        return
    by_name: dict[str, list[dict[str, Any]]] = {}
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            decor: list[dict[str, Any]] = []
            for shape in layout.shapes:
                if shape.is_placeholder:
                    continue
                try:
                    x, y, w, h = (int(v) / 360000 for v in (shape.left, shape.top, shape.width, shape.height))
                except (TypeError, ValueError):
                    continue
                if w <= 0 or h <= 0:
                    continue
                has_text = bool(getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip())
                is_picture = "PICTURE" in str(shape.shape_type)
                item: dict[str, Any] = {"x": round(x, 2), "y": round(y, 2), "w": round(w, 2), "h": round(h, 2), "text": has_text, "picture": is_picture}
                if is_picture and w * h >= 0.75 * (int(prs.slide_width) / 360000) * (int(prs.slide_height) / 360000):
                    item.update(_picture_busyness(shape))
                decor.append(item)
            by_name.setdefault(layout.name.strip(), decor)
    for variant in template.get("variants") or []:
        if variant.get("source") in {"sample", "library"}:
            continue
        name = (variant.get("layout_name") or variant["id"]).strip()
        if name in by_name:
            variant["decor"] = by_name[name]


def _library_catalog_entry(variant: dict[str, Any]) -> dict[str, Any]:
    from app.pipeline.sample_slides import catalog_entry

    entry = catalog_entry(variant)
    entry["family"] = f"library.{variant['composition']['kind']}"
    return entry


_METRIC_ROLES = {"metric_value", "metric_label"}


def _fix_metric_roles(template: dict[str, Any]) -> None:
    """Парсер размечает слоты фактоидов по геометрии и иногда называет крупное
    число (88pt) «подписью». Переразмечаем по кеглю: внутри макета слоты
    фактоидов с кеглем ≥ 2× минимального — значения, остальные — подписи.
    Правим и `variants[].slots`, и контракты, и каталог (его читает планировщик),
    иначе модель кладёт слово «вузов» в 88-пунктовый слот числа."""
    catalog = {item["id"]: item for item in template.get("variant_catalog") or []}
    for variant in template.get("variants", []):
        slots = variant.get("slots") or {}
        metric = {
            name: (slot.get("resolved_font") or {}).get("size_pt")
            for name, slot in slots.items()
            if slot.get("semantic_role") in _METRIC_ROLES
        }
        sizes = [size for size in metric.values() if size]
        if len(sizes) < 2 or max(sizes) < 2 * min(sizes):
            continue
        threshold = 2 * min(sizes)
        new_roles = {name: ("metric_value" if size and size >= threshold else "metric_label") for name, size in metric.items()}
        for name, role in new_roles.items():
            slots[name]["semantic_role"] = role
        for contract_slot in (variant.get("generation_contract") or {}).get("text_slots") or []:
            if contract_slot.get("slot") in new_roles:
                contract_slot["role"] = new_roles[contract_slot["slot"]]
        for cat_slot in (catalog.get(variant["id"]) or {}).get("text_slots") or []:
            if cat_slot.get("slot") in new_roles:
                role = new_roles[cat_slot["slot"]]
                cat_slot["role"] = role
                cat_slot["allowed_content_kinds"] = ["number_or_short_value"] if role == "metric_value" else ["short_text"]
                if role == "metric_value":
                    cat_slot["recommended_characters"] = min(cat_slot.get("recommended_characters") or 8, 8)


# --------------------------------------------------------------------------
# Типографика: уточнение шрифтов слотов по XML макета/мастера/темы
# --------------------------------------------------------------------------

_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}

_TITLE_ROLES = {"presentation_title", "slide_title", "section_title", "closing_title", "speaker_name"}


def _enrich_slot_typography(template: dict[str, Any], pptx_path: Path) -> None:
    """Дописывает каждому слоту `resolved_font` {family,size_pt,bold,color,source}.

    Парсер коллеги отдаёт шрифт слота только когда он задан явно на фигуре
    макета; для унаследованных значений он ставит None и оценку из
    capacity (role_fallback). Для дизайн-системы этого мало: дополняем
    цепочкой наследования OOXML — lstStyle фигуры макета → плейсхолдер того же
    типа в мастере → шрифты темы (major для заголовков, minor для остального).
    """
    import xml.etree.ElementTree as ET

    try:
        archive = zipfile.ZipFile(str(pptx_path))
    except (OSError, zipfile.BadZipFile):
        return
    with archive:
        names = set(archive.namelist())
        theme_fonts = _theme_fonts(archive, names)
        palette = template.get("colors") or {}
        master_style_cache: dict[str, dict[str, dict]] = {}
        layout_root_cache: dict[str, ET.Element] = {}

        for variant in template.get("variants", []):
            layout_file = variant.get("layout_file") or variant.get("slide_file")
            if not layout_file or layout_file not in names:
                continue
            root = layout_root_cache.get(layout_file)
            if root is None:
                try:
                    root = ET.fromstring(archive.read(layout_file))
                except ET.ParseError:
                    continue
                layout_root_cache[layout_file] = root
            master_file = _master_for_layout(archive, names, layout_file)
            if master_file and master_file not in master_style_cache:
                master_style_cache[master_file] = _master_placeholder_styles(archive, master_file)
            master_styles = master_style_cache.get(master_file or "", {})

            for slot_name, slot in (variant.get("slots") or {}).items():
                shape_id = slot.get("shape_id")
                sp = _find_shape(root, shape_id)
                observed_font = slot.get("font") or {}
                family = observed_font.get("family") or None
                size = observed_font.get("size")
                bold = observed_font.get("weight") == "bold"
                color = None
                source = "layout_shape" if (family or size) else None

                if sp is not None:
                    xml_style = _shape_text_style(sp, palette)
                    family = family or xml_style.get("family")
                    if size is None:
                        size = xml_style.get("size")
                    if xml_style.get("bold") is not None and not bold:
                        bold = xml_style["bold"]
                    color = xml_style.get("color")
                    if source is None and (xml_style.get("family") or xml_style.get("size")):
                        source = "layout_style"

                ph_type = (slot.get("placeholder_type") or "").lower()
                master_style = master_styles.get("title" if ph_type in {"title", "ctrtitle", "center_title"} else "body", {})
                if family is None and master_style.get("family"):
                    family, source = master_style["family"], source or "master"
                if size is None and master_style.get("size"):
                    size, source = master_style["size"], source or "master"
                if color is None and master_style.get("color"):
                    color = master_style["color"]

                role = slot.get("semantic_role") or ""
                if family is None or family.startswith("+mj") or family.startswith("+mn"):
                    family = theme_fonts["major" if role in _TITLE_ROLES or ph_type.startswith("title") or ph_type == "ctrtitle" else "minor"] or family or "—"
                    source = source or "theme"
                if size is None:
                    size = (slot.get("capacity") or {}).get("font_size_pt")
                    source = "estimated"

                slot["resolved_font"] = {
                    "family": family,
                    "size_pt": float(size) if size else None,
                    "bold": bool(bold),
                    "color": color,
                    "source": source or "estimated",
                }
        template["theme_fonts"] = theme_fonts


def _theme_fonts(archive: zipfile.ZipFile, names: set[str]) -> dict[str, str | None]:
    import xml.etree.ElementTree as ET

    for name in sorted(names):
        if name.startswith("ppt/theme/") and name.endswith(".xml"):
            try:
                root = ET.fromstring(archive.read(name))
            except ET.ParseError:
                continue
            major = root.find(".//a:majorFont/a:latin", _NS)
            minor = root.find(".//a:minorFont/a:latin", _NS)
            return {
                "major": major.attrib.get("typeface") if major is not None else None,
                "minor": minor.attrib.get("typeface") if minor is not None else None,
            }
    return {"major": None, "minor": None}


def _master_for_layout(archive: zipfile.ZipFile, names: set[str], layout_file: str) -> str | None:
    rels = f"ppt/slideLayouts/_rels/{Path(layout_file).name}.rels"
    if rels not in names:
        return None
    text = archive.read(rels).decode("utf-8", errors="ignore")
    match = re.search(r'Target="\.\./slideMasters/(slideMaster\d+\.xml)"', text)
    return f"ppt/slideMasters/{match.group(1)}" if match else None


def _master_placeholder_styles(archive: zipfile.ZipFile, master_file: str) -> dict[str, dict]:
    """Шрифт/кегль/цвет плейсхолдеров title и body в мастере (lvl1)."""
    import xml.etree.ElementTree as ET

    try:
        root = ET.fromstring(archive.read(master_file))
    except (KeyError, ET.ParseError):
        return {}
    result: dict[str, dict] = {}
    for sp in root.iter(f"{{{_NS['p']}}}sp"):
        ph = sp.find("./p:nvSpPr/p:nvPr/p:ph", _NS)
        if ph is None:
            continue
        kind = (ph.attrib.get("type") or "body").lower()
        key = "title" if kind in {"title", "ctrtitle"} else "body" if kind == "body" else None
        if key is None or key in result:
            continue
        result[key] = _shape_text_style(sp, {})
    tx = root.find("./p:txStyles", _NS)
    if tx is not None:
        for tag, key in (("p:titleStyle", "title"), ("p:bodyStyle", "body")):
            node = tx.find(f"./{tag}/a:lvl1pPr/a:defRPr", _NS)
            if node is None:
                continue
            style = result.setdefault(key, {})
            if not style.get("size") and node.attrib.get("sz"):
                style["size"] = int(node.attrib["sz"]) / 100
            latin = node.find("./a:latin", _NS)
            if not style.get("family") and latin is not None and latin.attrib.get("typeface"):
                style["family"] = latin.attrib["typeface"]
    return result


def _find_shape(root, shape_id):
    if shape_id is None:
        return None
    for sp in root.iter(f"{{{_NS['p']}}}sp"):
        cnv = sp.find("./p:nvSpPr/p:cNvPr", _NS)
        if cnv is not None and cnv.attrib.get("id") == str(shape_id):
            return sp
    return None


def _shape_text_style(sp, palette: dict[str, str]) -> dict[str, Any]:
    """Первый явный стиль текста фигуры: lstStyle/lvl1pPr/defRPr → первый run → endParaRPr."""
    style: dict[str, Any] = {}
    candidates = []
    lst = sp.find("./p:txBody/a:lstStyle/a:lvl1pPr/a:defRPr", _NS)
    if lst is not None:
        candidates.append(lst)
    for tag in ("a:rPr", "a:endParaRPr"):
        node = sp.find(f"./p:txBody/a:p//{tag}", _NS)
        if node is not None:
            candidates.append(node)
    for node in candidates:
        if "size" not in style and node.attrib.get("sz"):
            style["size"] = int(node.attrib["sz"]) / 100
        if "bold" not in style and node.attrib.get("b") is not None:
            style["bold"] = node.attrib["b"] in {"1", "true"}
        latin = node.find("./a:latin", _NS)
        if "family" not in style and latin is not None and latin.attrib.get("typeface"):
            style["family"] = latin.attrib["typeface"]
        if "color" not in style:
            srgb = node.find("./a:solidFill/a:srgbClr", _NS)
            scheme = node.find("./a:solidFill/a:schemeClr", _NS)
            if srgb is not None and srgb.attrib.get("val"):
                style["color"] = "#" + srgb.attrib["val"].upper()
            elif scheme is not None and scheme.attrib.get("val"):
                style["color"] = palette.get(scheme.attrib["val"]) or f"scheme:{scheme.attrib['val']}"
    return style


# --------------------------------------------------------------------------
# Дизайн-система из Template JSON
# --------------------------------------------------------------------------

ROLE_LABELS: dict[str, str] = {
    "presentation_title": "Название презентации",
    "presentation_subtitle": "Подзаголовок обложки",
    "section_title": "Заголовок раздела",
    "slide_title": "Заголовок слайда",
    "closing_title": "Заголовок финального слайда",
    "subtitle": "Подзаголовок",
    "body": "Основной текст",
    "column_body": "Текст колонки",
    "metric_value": "Фактоид — значение",
    "metric_label": "Фактоид — подпись",
    "quote": "Цитата",
    "quote_author": "Автор цитаты",
    "speaker_name": "Имя спикера",
    "speaker_bio": "Описание спикера",
    "image_caption": "Подпись к изображению",
    "image_source": "Источник изображения",
    "contact_text": "Контакты / призыв",
    "footer": "Колонтитул",
    "generic_text": "Универсальный текст",
    "decorative": "Декоративный текст (не заполняется)",
}

ROLE_ORDER = list(ROLE_LABELS.keys())

NARRATIVE_LABELS: dict[str, str] = {
    "cover": "Обложка",
    "section": "Разделитель",
    "content": "Контент",
    "evidence": "Доказательство / данные",
    "closing": "Финал",
    "speaker": "Спикер",
}


def typography_scale(template: dict[str, Any]) -> list[dict[str, Any]]:
    """Полная типографическая шкала шаблона: по каждой роли — гарнитура,
    кегли (все встретившиеся), насыщенность, цвет и число макетов."""
    buckets: dict[tuple[str, str, bool], dict[str, Any]] = {}
    for variant in template.get("variants", []):
        for slot_name, slot in (variant.get("slots") or {}).items():
            font = slot.get("resolved_font") or {}
            role = slot.get("semantic_role") or "generic_text"
            if role == "decorative":
                continue
            family = font.get("family") or "—"
            key = (role, family, bool(font.get("bold")))
            bucket = buckets.setdefault(
                key,
                {
                    "role": role,
                    "label": ROLE_LABELS.get(role, role),
                    "family": family,
                    "bold": bool(font.get("bold")),
                    "sizes_pt": set(),
                    "estimated": False,
                    "colors": set(),
                    "layouts": set(),
                    "content_kind": slot.get("content_kind"),
                },
            )
            if font.get("size_pt"):
                bucket["sizes_pt"].add(round(float(font["size_pt"]), 1))
            if font.get("source") == "estimated":
                bucket["estimated"] = True
            if font.get("color") and str(font["color"]).startswith("#"):
                bucket["colors"].add(font["color"])
            bucket["layouts"].add(variant["id"])

    rows = []
    for bucket in buckets.values():
        rows.append(
            {
                **bucket,
                "sizes_pt": sorted(bucket["sizes_pt"], reverse=True),
                "colors": sorted(bucket["colors"]),
                "layouts": sorted(bucket["layouts"]),
                "layouts_count": len(bucket["layouts"]),
            }
        )
    rows.sort(key=lambda r: (ROLE_ORDER.index(r["role"]) if r["role"] in ROLE_ORDER else 99, -(r["sizes_pt"][0] if r["sizes_pt"] else 0)))
    return rows


def font_families(template: dict[str, Any]) -> list[str]:
    families: set[str] = set()
    for variant in template.get("variants", []):
        for slot in (variant.get("slots") or {}).values():
            fam = (slot.get("resolved_font") or {}).get("family")
            if fam and fam != "—":
                families.add(fam)
    theme = template.get("theme_fonts") or {}
    for key in ("major", "minor"):
        if theme.get(key):
            families.add(theme[key])
    return sorted(families)


def font_sizes(template: dict[str, Any]) -> set[float]:
    sizes: set[float] = set()
    for variant in template.get("variants", []):
        for slot in (variant.get("slots") or {}).values():
            if slot.get("semantic_role") == "decorative" or slot.get("decorative"):
                continue
            size = (slot.get("resolved_font") or {}).get("size_pt")
            if size:
                sizes.add(float(size))
    return sizes


def layout_cards(template: dict[str, Any]) -> list[dict[str, Any]]:
    """Каталог макетов для UI: имя, роль, композиция, фон, слоты с геометрией в %."""
    size = template.get("slide_size") or {}
    sw = float(size.get("width") or 33.867)
    sh = float(size.get("height") or 19.05)
    catalog = {item["id"]: item for item in template.get("variant_catalog") or []}
    cards = []
    for variant in template.get("variants", []):
        cat = catalog.get(variant["id"], {})
        elements = []
        for slot_name, slot in (variant.get("slots") or {}).items():
            if slot.get("virtual"):
                continue
            elements.append(_pct_box(slot, sw, sh, "text", slot.get("semantic_role") or "generic_text", slot_name))
        for slot_name, slot in (variant.get("image_slots") or {}).items():
            elements.append(_pct_box(slot, sw, sh, "image", slot.get("semantic_role") or "photo", slot_name))
        for image in variant.get("images") or []:
            elements.append(_pct_box(image, sw, sh, "decor_image", "decorative_image", image.get("name")))
        for art in ((variant.get("observed") or {}).get("fixed_artwork") or {}).get("shapes", []) if isinstance((variant.get("observed") or {}).get("fixed_artwork"), dict) else []:
            if isinstance(art, dict) and all(k in art for k in ("x", "y", "w", "h")):
                elements.append(_pct_box(art, sw, sh, "decor", "decorative", art.get("name")))
        background = (variant.get("background") or {}).get("color")
        cards.append(
            {
                "layout_id": variant["id"],
                "name": (variant.get("layout_name") or variant["id"]).strip(),
                "family": cat.get("family"),
                "narrative_role": cat.get("narrative_role") or (variant.get("semantic") or {}).get("layout_role"),
                "narrative_label": NARRATIVE_LABELS.get(cat.get("narrative_role"), cat.get("narrative_role") or ""),
                "layout_role": (variant.get("semantic") or {}).get("layout_role"),
                "composition": cat.get("composition"),
                "content_pattern": cat.get("content_pattern"),
                "best_for": cat.get("best_for") or [],
                "text_slots": len(variant["generation_contract"]["text_slots"]),
                "image_slots": len(variant["generation_contract"]["image_slots"]),
                "background_hex": background,
                "elements": elements,
                "source_mode": template.get("source_mode"),
                "source": variant.get("source") or "layout",
                "kind": (variant.get("composition") or {}).get("kind"),
                "units": (variant.get("composition") or {}).get("units"),
            }
        )
    return cards


def _pct_box(obj: dict[str, Any], sw: float, sh: float, kind: str, role: str, name: str | None) -> dict[str, Any]:
    font = obj.get("resolved_font") or {}
    return {
        "kind": kind,
        "role": role,
        "name": name,
        "left_pct": round(float(obj.get("x") or 0) / sw * 100, 2),
        "top_pct": round(float(obj.get("y") or 0) / sh * 100, 2),
        "width_pct": round(float(obj.get("w") or 0) / sw * 100, 2),
        "height_pct": round(float(obj.get("h") or 0) / sh * 100, 2),
        "font_size_pt": font.get("size_pt"),
        "color": font.get("color") if str(font.get("color") or "").startswith("#") else None,
        "align": obj.get("horizontal_align"),
    }


# --------------------------------------------------------------------------
# Модель: адаптер нашего OpenAI-совместимого клиента под planning.generate_plan
# --------------------------------------------------------------------------


class BackendModel:
    """Синхронный `model.json(system, user, temperature)` для планировщика коллеги.

    Планировщик — синхронный код; оркестратор запускает его в `asyncio.to_thread`,
    поэтому здесь используется синхронный клиент openai. Промпты планировщика
    живут отдельными файлами в pptx_template_parser_generator/prompts/*.txt.
    """

    def __init__(self, base_url: str, model: str, api_key: str | None, timeout: float = 150.0) -> None:
        from openai import OpenAI

        self.model = model
        self._client = OpenAI(base_url=base_url, api_key=api_key or "unused", timeout=timeout, max_retries=1)

    def json(self, system: str, user: str, temperature: float = 0.2) -> dict[str, Any]:
        messages = [{"role": "system", "content": system}, {"role": "user", "content": user}]
        attempts: list[dict[str, Any]] = [
            {"response_format": {"type": "json_object"}, "reasoning_effort": "none"},
            {"reasoning_effort": "none"},
            {},
        ]
        last_error: Exception | None = None
        for extra in attempts:
            try:
                response = self._client.chat.completions.create(
                    model=self.model,
                    messages=messages,
                    temperature=temperature,
                    max_completion_tokens=8192,
                    **extra,
                )
                content = response.choices[0].message.content
                if not content:
                    raise ModelError("Модель вернула пустой content")
                return parse_json_response(content)
            except Exception as exc:  # noqa: BLE001 — пробуем следующий набор параметров
                last_error = exc
                logger.warning("Вызов модели %s не удался (%s), пробую упрощённые параметры", self.model, exc)
        raise ModelError(f"Не удалось вызвать модель {self.model}: {last_error}")


def relaxed_template(template: dict[str, Any]) -> dict[str, Any]:
    """Копия Template JSON, где «безопасный» бюджет символов заменён на
    максимальный. Планировщик коллеги режет текст по recommended_characters —
    для заголовков это слишком консервативно (обложка: 17 символов при
    максимуме 25), и названия вроде «Стратегия VK Education 2026» обрезались.
    Максимальный бюджет остаётся жёстким ограничением валидатора."""
    copy = deepcopy(template)
    for variant in copy.get("variants", []):
        for slot in variant["generation_contract"]["text_slots"]:
            constraints = slot.get("constraints") or {}
            if constraints.get("maximum_characters"):
                constraints["recommended_characters"] = constraints["maximum_characters"]
            if constraints.get("maximum_lines"):
                constraints["recommended_lines"] = constraints["maximum_lines"]
    for item in copy.get("variant_catalog", []):
        for slot in item.get("text_slots", []):
            if slot.get("maximum_characters"):
                slot["recommended_characters"] = slot["maximum_characters"]
    return copy


def planner_brief(brief: str, content_plan: ContentPlan, case_prompt_addition: str = "") -> str:
    """Бриф для планировщика макетов, построенный из готового ContentPlan.

    Планировщик коллеги принимает свободный текст; отдаём ему исходный бриф и
    пронумерованную структуру слайдов (заголовок, тезисы, что за объект будет
    на слайде), требуя сохранить порядок и число слайдов. Так layout_plan и
    ContentPlan описывают одни и те же слайды, а графики/таблицы вставляются по
    индексу на нужный слайд."""
    lines = [
        "ЗАДАЧА ПРЕЗЕНТАЦИИ:",
        brief.strip(),
        "",
        "СТРУКТУРА ПРЕЗЕНТАЦИИ (утверждена, следуй ей строго: тот же порядок, то же число слайдов,",
        "слайд N структуры = слайд N презентации; заголовки сохраняй по смыслу, текст пиши по тезисам):",
    ]
    total = len(content_plan.slides)
    for index, slide in enumerate(content_plan.slides, start=1):
        purpose = slide.purpose.strip()
        kind = ""
        if index == 1:
            kind = " [ОБЛОЖКА — выбери титульный макет]"
        elif index == total:
            kind = " [ФИНАЛ — выбери финальный макет (контакты/QR/спасибо), если он есть в каталоге; заголовок — «Спасибо» или короткий призыв, в подписи — контакты]"
        elif purpose.lower().startswith("раздел") or not slide.content_blocks:
            kind = " [РАЗДЕЛИТЕЛЬ — выбери макет раздела/перебивки только с заголовком]"
        elif purpose.lower().startswith("цифры"):
            kind = " [КЛЮЧЕВЫЕ ЦИФРЫ — выбери макет с фактоидами/метриками, если он есть в каталоге]"
        lines.append(f"Слайд {index}. {slide.title.strip()} — {purpose}{kind}")
        for block in slide.content_blocks:
            if block.type == ContentBlockType.BULLETS and block.bullets:
                for item in block.bullets:
                    if item and item.strip():
                        lines.append(f"  • {item.strip()}")
            elif block.type == ContentBlockType.TEXT and block.text:
                lines.append(f"  {block.text.strip()}")
            elif block.type == ContentBlockType.CHART and block.chart is not None:
                lines.append(
                    f"  [на слайде будет график: {', '.join(block.chart.categories[:6])} — "
                    f"{', '.join(series.name for series in block.chart.series)}; оставь место, текст короче]"
                )
            elif block.type == ContentBlockType.TABLE and block.table is not None:
                lines.append(f"  [на слайде будет таблица: {', '.join(block.table.headers)}; оставь место, текст короче]")
    if case_prompt_addition:
        lines.extend(["", f"Дополнительно: {case_prompt_addition.strip()}"])
    return "\n".join(lines)


def plan_presentation(
    template: dict[str, Any],
    brief: str,
    purpose: str,
    slide_count: int,
    model: BackendModel,
) -> tuple[dict[str, Any], dict[str, Any]]:
    """outline + plan от планировщика коллеги (2–4 вызова модели, промпты из файлов)."""
    return generate_plan(model, relaxed_template(template), brief, purpose=purpose, slide_count=slide_count)


# --------------------------------------------------------------------------
# Три варианта вёрстки из одного плана
# --------------------------------------------------------------------------

VARIANT_STRATEGIES: dict[str, dict[str, Any]] = {
    "variant_a": {
        "label": "Классический",
        "description": (
            "Контентные слайды в одноколоночных макетах «заголовок + текст»; "
            "первые визуальные версии обложки, разделителей и финала. "
            "Графики и таблицы занимают нижнюю часть текстовой зоны."
        ),
        "prefer": [("composition", "title_body")],
        "visual_index": 0,
    },
    "variant_b": {
        "label": "Структурный",
        "description": (
            "Контентные слайды в двухколоночных макетах: тезисы разбиты на две "
            "колонки, графики и таблицы — в правой колонке; вторые визуальные "
            "версии обложки, разделителей и финала."
        ),
        "prefer": [("content_pattern", "parallel_content")],
        "visual_index": 1,
    },
    "variant_c": {
        "label": "Визуальный",
        "description": (
            "Контентные слайды в макетах с изображением (фото/паттерн из "
            "медиатеки шаблона) и фактоидами; третьи визуальные версии обложки, "
            "разделителей и финала."
        ),
        "prefer": [("composition", "text_with_media"), ("has_image", True), ("composition", "image_caption_pair")],
        "visual_index": 2,
    },
}

VALID_VARIANTS = tuple(VARIANT_STRATEGIES.keys())

_STRUCTURAL_ROLES = {"cover", "section", "closing", "speaker"}
_KEEP_PATTERNS = {"person_profile", "quote", "metrics", "section_transition", "introduction", "call_to_action"}


def derive_variant_plan(
    template: dict[str, Any],
    base_plan: dict[str, Any],
    variant_key: str,
    data_slides: set[int] | None = None,
) -> dict[str, Any]:
    """Строит план варианта из базового плана модели по стратегии `variant_key`.

    Ось различия: (1) семейство макета контентных слайдов — одноколоночные /
    двухколоночные / с изображением; (2) визуальная версия внутри семейства для
    обложки, разделителей и финала (visual_index). Текст переносится по
    семантическим ролям слотов и сокращается до maximum_characters нового
    контракта, заголовки сохраняются. `data_slides` — индексы слайдов, на
    которых будет нативный график/таблица: им нужен макет с широкой текстовой
    зоной (A, C) или симметричными колонками (B), иначе объект не поместится.
    """
    data_slides = data_slides or set()
    strategy = VARIANT_STRATEGIES[variant_key]
    catalog = {item["id"]: item for item in template.get("variant_catalog") or []}
    variants = {v["id"]: v for v in template.get("variants", [])}
    families: dict[str, list[str]] = {}
    for item in template.get("variant_catalog") or []:
        families.setdefault(item["family"], []).append(item["id"])

    new_slides = []
    for index, slide in enumerate(base_plan["slides"]):
        original_id = slide["variant"]
        cat = catalog.get(original_id, {})
        family_members = families.get(cat.get("family"), [original_id])
        role = cat.get("narrative_role")
        pattern = cat.get("content_pattern")

        if role in _STRUCTURAL_ROLES or pattern in _KEEP_PATTERNS or len(catalog) < 4:
            # Структурные слайды: та же семья, другая визуальная версия.
            target_id = family_members[strategy["visual_index"] % len(family_members)] if family_members else original_id
        else:
            effective = strategy
            if index in data_slides and variant_key == "variant_c":
                effective = {**strategy, "prefer": VARIANT_STRATEGIES["variant_a"]["prefer"]}
            target_id = _pick_content_layout(template, catalog, families, original_id, effective, index)

        target = variants.get(target_id) or variants[original_id]
        if target_id == original_id:
            new_slides.append(deepcopy(slide))
            continue
        content = _remap_content(variants[original_id], target, slide.get("content") or {}, index)
        new_slides.append({"variant": target_id, "content": content, "image_content": {}})
    return {"slides": new_slides}


def _symmetric_columns(variant: dict[str, Any]) -> bool:
    """Двухколоночный макет считается «настоящим», если все колонки набраны
    одним шрифтом/кеглем/цветом. Иначе это, как правило, «текст + объект»
    (вторая колонка — подпись на тёмной плашке или к скриншоту)."""
    columns = [
        slot for slot in (variant.get("slots") or {}).values()
        if slot.get("semantic_role") == "column_body"
    ]
    if len(columns) < 2:
        return True
    # Колонки должны стоять на одной высоте и быть одного размера, иначе
    # одна из них — подпись к скриншоту/объекту, а не текстовая колонка.
    ys = [float(c.get("y") or 0) for c in columns]
    hs = [float(c.get("h") or 0) for c in columns]
    ws = [float(c.get("w") or 0) for c in columns]
    if max(ys) - min(ys) > 0.6 or max(hs) - min(hs) > 0.6 or max(ws) - min(ws) > 0.6:
        return False
    keys = {
        ((c.get("resolved_font") or {}).get("family"), (c.get("resolved_font") or {}).get("size_pt"), (c.get("resolved_font") or {}).get("color"))
        for c in columns
    }
    return len(keys) == 1


def _pick_content_layout(template, catalog, families, original_id, strategy, index) -> str:
    variants = {v["id"]: v for v in template.get("variants", [])}
    candidates = []
    for item in catalog.values():
        if item.get("narrative_role") != "content" or item.get("review_required"):
            continue
        if item.get("content_pattern") in _KEEP_PATTERNS:
            continue
        # Не тянуть спикерские/цитатные макеты в обычный контент.
        if not any(slot.get("role") == "slide_title" for slot in item.get("text_slots", [])):
            continue
        if item.get("content_pattern") == "parallel_content" and not _symmetric_columns(variants.get(item["id"], {})):
            continue
        candidates.append(item)
    for key, value in strategy["prefer"]:
        if key == "has_image":
            matched = [c for c in candidates if bool(c.get("image_slots")) == value]
        else:
            matched = [c for c in candidates if c.get(key) == value]
        if matched:
            # Чередуем семейства/визуальные версии по номеру слайда, чтобы
            # внутри варианта не повторялся один и тот же макет подряд.
            matched.sort(key=lambda c: (c["family"], c["id"]))
            return matched[index % len(matched)]["id"]
    # Стратегия не представлена в шаблоне — оставляем выбор модели, но берём
    # другую визуальную версию внутри семейства, если она есть.
    members = families.get(catalog.get(original_id, {}).get("family"), [original_id])
    return members[strategy["visual_index"] % len(members)]


def _remap_content(source_variant, target_variant, content: dict[str, str], index: int) -> dict[str, str]:
    by_role: dict[str, list[str]] = {}
    for slot_name, text in content.items():
        slot = (source_variant.get("slots") or {}).get(slot_name) or {}
        role = slot.get("semantic_role") or "generic_text"
        if isinstance(text, str) and text.strip():
            by_role.setdefault(role, []).append(text.strip())

    title = _first(by_role, ("slide_title", "section_title", "presentation_title", "closing_title", "generic_text"))
    body_parts: list[str] = []
    for role in ("body", "column_body", "quote", "speaker_bio", "generic_text"):
        body_parts.extend(by_role.get(role, []))
    body_text = "\n".join(body_parts)

    target_slots = target_variant["generation_contract"]["text_slots"]
    column_slots = [s for s in target_slots if s["role"] == "column_body"]
    columns = _split_columns(body_text, len(column_slots)) if column_slots else []

    result: dict[str, str] = {}
    column_i = 0
    for slot in target_slots:
        role = slot["role"]
        constraints = slot.get("constraints") or {}
        limit = constraints.get("maximum_characters") or constraints.get("recommended_characters")
        max_lines = constraints.get("maximum_lines") or constraints.get("recommended_lines")
        value: str | None = None
        if role in {"slide_title", "section_title", "presentation_title", "closing_title"}:
            value = title
        elif role == "body":
            value = body_text
        elif role == "column_body":
            value = columns[column_i] if column_i < len(columns) else None
            column_i += 1
        elif role in {"subtitle", "presentation_subtitle"}:
            value = _first(by_role, ("subtitle", "presentation_subtitle")) or (body_parts[0] if body_parts else None)
        elif role in by_role:
            value = by_role[role][0]
        if role in _TITLE_ROLES and limit:
            # Заголовок не режем до обрубка («Спасибо за доверие и») — даём
            # запас, а на рендере уменьшаем кегль (_fit_text_to_shape).
            limit = int(limit * _TITLE_OVERFLOW_FACTOR)
        if value and value.strip():
            result[slot["slot"]] = _shorten(value.strip(), limit, max_lines, is_title=role in _TITLE_ROLES)
        elif slot.get("required") and role not in {"footer"}:
            result[slot["slot"]] = _shorten(title or body_text or "—", limit, max_lines, is_title=role in _TITLE_ROLES)
    return result


_TITLE_ROLES = {"slide_title", "section_title", "presentation_title", "closing_title"}
_TITLE_OVERFLOW_FACTOR = 3.2


def _first(by_role: dict[str, list[str]], roles: tuple[str, ...]) -> str | None:
    for role in roles:
        if by_role.get(role):
            return by_role[role][0]
    return None


def _split_columns(text: str, count: int) -> list[str]:
    if count <= 0:
        return []
    units = [u.strip() for u in re.split(r"\n+", text) if u.strip()]
    if len(units) < count:
        sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]
        if len(sentences) >= count:
            units = sentences
    if not units:
        return [text] + [""] * (count - 1)
    per = max(1, -(-len(units) // count))
    columns = ["\n".join(units[i * per:(i + 1) * per]) for i in range(count)]
    return columns


def _shorten(text: str, limit: int | None, max_lines: int | None = None, is_title: bool = False) -> str:
    """Как planning._shorten: режем по границе слова, не оставляя обрубков;
    дополнительно ограничиваем число строк бюджетом слота. Заголовок не должен
    заканчиваться предлогом/союзом («…интеграцию в») — хвостовые короткие слова
    отбрасываем."""
    if max_lines:
        lines = [line for line in text.splitlines() if line.strip()]
        if len(lines) > max_lines:
            text = "\n".join(lines[:max_lines])
    if not limit or len(text) <= limit:
        return text
    clipped = text[: limit + 1]
    boundary = clipped.rfind(" ")
    if boundary >= max(4, limit // 2):
        clipped = clipped[:boundary]
    else:
        clipped = clipped[:limit]
    clipped = clipped.rstrip(" ,;:.—-")
    if is_title:
        words = clipped.split(" ")
        while len(words) > 2 and (len(words[-1].strip("«»\"'")) <= 3 or words[-1] in {"—", "-"}):
            words.pop()
        clipped = " ".join(words).rstrip(" ,;:.—-")
    return clipped


_BODY_ROLES = ("body", "column_body")


def supplement_plan_text(template: dict[str, Any], plan: dict[str, Any], content_plan: ContentPlan | None) -> dict[str, Any]:
    """Дозаполняет пустые текстовые зоны слайдов тезисами из ContentPlan.

    Планировщик по макетам иногда ставит контентному слайду макет «только
    заголовок» — после переноса в семейство варианта у слайда появляется
    текстовая зона, но заполнять её нечем. Контент-агент к этому моменту уже
    сформировал по каждому слайду тезисы/абзацы — берём их (bullets → строки
    с маркером, text → абзац) и раскладываем по слотам body/column_body с
    учётом бюджета символов и строк контракта."""
    if content_plan is None:
        return plan
    variants = {v["id"]: v for v in template.get("variants", [])}
    cp_slides = list(content_plan.slides)
    for index, slide in enumerate(plan["slides"]):
        variant = variants.get(slide["variant"]) or {}
        contract_slots = (variant.get("generation_contract") or {}).get("text_slots") or []
        body_slots = [cs for cs in contract_slots if cs.get("role") in _BODY_ROLES]
        if not body_slots:
            continue
        content = slide.setdefault("content", {})
        if index >= len(cp_slides):
            continue
        parts: list[str] = []
        for block in cp_slides[index].content_blocks:
            if block.type == ContentBlockType.BULLETS and block.bullets:
                parts.extend(f"• {item.strip()}" for item in block.bullets if item and item.strip())
            elif block.type == ContentBlockType.TEXT and block.text:
                parts.append(block.text.strip())
        if not parts:
            continue
        existing = sum(len((content.get(cs["slot"]) or "").strip()) for cs in body_slots)
        if existing:
            # Планировщик написал текст сам. Заменяем его тезисами ContentPlan
            # только если он заметно беднее: короче половины тезисов и меньше
            # 40 % бюджета зоны — иначе слайд выглядит полупустым.
            budget = sum(_limit(cs) or 0 for cs in body_slots) or 0
            cp_len = sum(len(part) for part in parts)
            if not (existing < cp_len * 0.5 and (not budget or existing < budget * 0.4)):
                continue
            logger.info("Слайд %d: текст планировщика (%d зн.) заменён тезисами ContentPlan (%d зн.)", index + 1, existing, cp_len)
        column_slots = [cs for cs in body_slots if cs.get("role") == "column_body"]
        text = "\n".join(parts)
        if column_slots and len(column_slots) >= 2:
            columns = _split_columns(text, len(column_slots))
            for cs, column in zip(column_slots, columns):
                if column.strip():
                    content[cs["slot"]] = _shorten(column, _limit(cs), _lines(cs))
        else:
            cs = body_slots[0]
            content[cs["slot"]] = _shorten(text, _limit(cs), _lines(cs))
    return plan


def _limit(contract_slot: dict[str, Any]) -> int | None:
    constraints = contract_slot.get("constraints") or {}
    return constraints.get("maximum_characters") or constraints.get("recommended_characters")


def _lines(contract_slot: dict[str, Any]) -> int | None:
    constraints = contract_slot.get("constraints") or {}
    return constraints.get("maximum_lines") or constraints.get("recommended_lines")


def plan_signature(plan: dict[str, Any]) -> tuple[str, ...]:
    return tuple(slide["variant"] for slide in plan["slides"])


# --------------------------------------------------------------------------
# Изображения для обязательных image-слотов
# --------------------------------------------------------------------------


def extract_fallback_image(source: str | Path, output_dir: Path) -> Path | None:
    """Самое крупное растровое изображение из медиатеки шаблона — запасной
    вариант для обязательных фото-слотов, когда у макета нет образца."""
    with zipfile.ZipFile(str(source)) as archive:
        best = None
        for info in archive.infolist():
            if not info.filename.startswith("ppt/media/"):
                continue
            if Path(info.filename).suffix.lower() not in {".png", ".jpg", ".jpeg"}:
                continue
            if best is None or info.file_size > best.file_size:
                best = info
        if best is None:
            return None
        output_dir.mkdir(parents=True, exist_ok=True)
        target = output_dir / ("template_image" + Path(best.filename).suffix.lower())
        target.write_bytes(archive.read(best.filename))
        return target.resolve()


def extract_sample_images(template: dict[str, Any], source: str | Path, output_dir: Path) -> dict[tuple[str, str], Path]:
    """Картинки из образцовых слайдов шаблона по (id макета, слот).

    Если в загруженном шаблоне есть слайды-образцы на нужном макете, берём
    изображение из того же picture-плейсхолдера (например, QR-код или фото
    из фирменного образца) — это точнее, чем случайный файл из медиатеки.
    """
    result: dict[tuple[str, str], Path] = {}
    try:
        prs = Presentation(str(source))
    except Exception:  # noqa: BLE001
        return result
    by_partname: dict[str, list[dict[str, Any]]] = {}
    for variant in template.get("variants", []):
        layout_file = variant.get("layout_file")
        if layout_file:
            by_partname.setdefault("/" + layout_file.lstrip("/"), []).append(variant)
    output_dir.mkdir(parents=True, exist_ok=True)
    for slide in prs.slides:
        variants = by_partname.get(str(slide.slide_layout.part.partname))
        if not variants:
            continue
        pictures = {}
        for shape in slide.placeholders:
            image = getattr(shape, "image", None)
            if image is not None:
                pictures[shape.placeholder_format.idx] = image
        if not pictures:
            continue
        for variant in variants:
            for slot_name, slot in (variant.get("image_slots") or {}).items():
                key = (variant["id"], slot_name)
                image = pictures.get(slot.get("placeholder_idx"))
                if image is None or key in result:
                    continue
                ext = "." + (image.ext or "png").lstrip(".")
                target = output_dir / f"sample_{abs(hash(key)) % 10**8}{ext}"
                target.write_bytes(image.blob)
                result[key] = target.resolve()
    return result


_URL_RE = re.compile(r"https?://[^\s)\]>«»\"']+")


def make_qr_image(text: str, output_dir: Path) -> Path | None:
    """PNG с настоящим QR-кодом (segno, чистый Python) для слотов qr_code."""
    try:
        import segno
    except ImportError:  # pragma: no cover
        return None
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / f"qr_{abs(hash(text)) % 10**8}.png"
    if not target.exists():
        segno.make(text, error="m").save(str(target), scale=12, border=1, dark="#000000", light="#FFFFFF")
    return target.resolve()


def fill_image_slots(
    template: dict[str, Any],
    plan: dict[str, Any],
    source: str | Path,
    work_dir: Path,
    brief: str = "",
    fill_optional: bool = False,
) -> None:
    """Заполняет image-слоты по их семантической роли.

    qr_code → настоящий QR (ссылка из брифа/контактов, иначе образец из
    шаблона); прочие роли → картинка из образцового слайда того же макета,
    иначе самое крупное изображение медиатеки. Необязательные слоты
    заполняются только в «визуальном» варианте; пустые плейсхолдеры затем
    удаляются при пост-обработке.
    """
    variants = {v["id"]: v for v in template["variants"]}
    samples = extract_sample_images(template, source, work_dir)
    fallback: Path | None = None
    urls = _URL_RE.findall(brief or "")
    for slide in plan["slides"]:
        variant = variants[slide["variant"]]
        contract = variant["generation_contract"]
        images = slide.setdefault("image_content", {})
        required = set(contract.get("required_slots") or [])
        for item in contract["image_slots"]:
            slot_name = item["slot"]
            current = images.get(slot_name)
            if isinstance(current, str) and Path(current).is_file():
                continue
            images.pop(slot_name, None)
            if slot_name not in required and not fill_optional:
                continue
            role = (variant.get("image_slots", {}).get(slot_name) or {}).get("semantic_role") or item.get("role")
            chosen: Path | None = None
            if role == "qr_code":
                link = None
                for text in list((slide.get("content") or {}).values()) + urls:
                    found = _URL_RE.findall(str(text))
                    if found:
                        link = found[0]
                        break
                if link is None and urls:
                    link = urls[0]
                if link is None:
                    domain_match = re.search(r"\b[\w.-]+\.(?:ru|com|org|io|dev|company|net)\b", " ".join(str(v) for v in (slide.get("content") or {}).values()))
                    link = "https://" + domain_match.group(0) if domain_match else None
                chosen = make_qr_image(link, work_dir) if link else samples.get((variant["id"], slot_name))
            else:
                chosen = samples.get((variant["id"], slot_name))
                if chosen is None:
                    if fallback is None:
                        fallback = extract_fallback_image(source, work_dir)
                    chosen = fallback
            if chosen is not None:
                images[slot_name] = str(chosen)


# --------------------------------------------------------------------------
# Рендер + пост-обработка (нативные графики/таблицы, чистка)
# --------------------------------------------------------------------------

_CHART_TYPE_MAP = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}


def render_variant(
    template: dict[str, Any],
    plan: dict[str, Any],
    source_path: str | Path,
    output_path: str | Path,
    manifest: DesignManifest | None,
    content_plan: ContentPlan | None,
    variant_key: str,
) -> None:
    """План → .pptx в макетах шаблона + графики/таблицы из ContentPlan."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    hard_errors = []
    for error in validate_generation_plan(template, plan, check_image_files=True):
        code = error.get("code") or ""
        if code.startswith("text_exceeds_recommended"):
            logger.info("Мягкое превышение бюджета слота: %s", error)
            continue
        if code == "text_too_long" and _title_overflow_allowed(template, plan, error):
            logger.info("Заголовок длиннее бюджета слота — кегль будет уменьшен: %s", error)
            continue
        if code in {"text_too_long", "too_many_lines"} and _section_body_allowed(template, plan, error):
            logger.info("Тело разделителя длиннее однострочного слота — рамка вырастет вниз: %s", error)
            continue
        if code == "missing_required_slot" and error.get("field") == "image_content":
            continue
        if code == "duplicate_authored_slide_variant":
            # Композиции библиотеки и образцов рисуются нашим рендером и могут
            # повторяться; запрет движка касается только авторских слайдов
            src = next((v.get("source") for v in template.get("variants") or [] if v.get("id") == error.get("variant")), None)
            if src in {"library", "sample"}:
                continue
        hard_errors.append(error)
    if hard_errors:
        raise ValueError(
            "План варианта не прошёл валидацию движка: "
            + "; ".join(f"слайд {e.get('slide_index')}/{e.get('slot')}: {e.get('message')}" for e in hard_errors[:5])
        )

    plan = _dedupe_metrics(template, plan)
    plan = _glue_short_words_in_titles(template, plan)
    variants = {v["id"]: v for v in template["variants"]}
    if any((variants[s["variant"]].get("source") or "layout") != "layout" for s in plan["slides"]):
        _render_mixed(template, plan, source_path, output_path, manifest, content_plan, variant_key)
        return
    tmp_path = output_path.with_suffix(".raw.pptx")
    tmp_path.unlink(missing_ok=True)
    create_presentation(template, plan["slides"], tmp_path, None, source_file=str(source_path))

    prs = Presentation(str(tmp_path))
    objects_by_index = _visual_objects(content_plan, plan) if content_plan is not None else {}
    for index, (slide, slide_plan) in enumerate(zip(prs.slides, plan["slides"])):
        variant = variants[slide_plan["variant"]]
        blocks = objects_by_index.get(index, [])
        if blocks and manifest is not None:
            _inject_objects(slide, variant, slide_plan, blocks, manifest, variant_key)
        _fit_slide_text(slide, variant, slide_plan)
        _remove_empty_placeholders(slide)
    _drop_unused_layouts(prs)
    prs.save(str(output_path))
    tmp_path.unlink(missing_ok=True)


def _render_mixed(
    template: dict[str, Any],
    plan: dict[str, Any],
    source_path: str | Path,
    output_path: Path,
    manifest: DesignManifest | None,
    content_plan: ContentPlan | None,
    variant_key: str,
) -> None:
    """Сборка, где слайды берутся из разных источников: слайды-образцы
    клонируются (slide_clone), композиции библиотеки рисуются (library),
    макеты шаблона заполняются как раньше (scripts.generate)."""
    from app.pipeline.library import render_library_slide
    from app.pipeline.slide_clone import render_sample_slide
    from scripts.generate import create_slide_from_layout

    prs = Presentation(str(source_path))
    original_ids = list(prs.slides._sldIdLst)
    variants = {v["id"]: v for v in template["variants"]}
    layouts = {str(layout.part.partname).lstrip("/"): layout for master in prs.slide_masters for layout in master.slide_layouts}
    objects_by_index = _visual_objects(content_plan, plan) if content_plan is not None else {}
    for index, slide_plan in enumerate(plan["slides"]):
        variant = variants[slide_plan["variant"]]
        blocks = objects_by_index.get(index, [])
        source = variant.get("source") or "layout"
        if source == "sample":
            render_sample_slide(prs, template, variant, slide_plan, blocks, manifest, _CHART_TYPE_MAP, _fit_font_size)
        elif source == "library":
            render_library_slide(prs, template, variant, slide_plan, blocks, manifest, layouts, _CHART_TYPE_MAP, _shorten)
        else:
            create_slide_from_layout(
                prs, template, variant["id"], slide_plan.get("content") or {}, slide_plan.get("image_content") or {},
                variant=variant, layouts=layouts,
            )
            slide = prs.slides[len(prs.slides) - 1]
            if blocks and manifest is not None:
                _inject_objects(slide, variant, slide_plan, blocks, manifest, variant_key)
            _fit_slide_text(slide, variant, slide_plan)
            _remove_empty_placeholders(slide)
    # Исходные слайды шаблона (образцы) убираем — остаются только собранные
    for slide_id in original_ids:
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    _drop_unused_layouts(prs)
    prs.save(str(output_path))


def data_slide_indices(content_plan: ContentPlan, plan: dict[str, Any]) -> set[int]:
    """Индексы слайдов, на которые попадёт нативный график/таблица."""
    return set(_visual_objects(content_plan, plan).keys())


def _visual_objects(content_plan: ContentPlan, plan: dict[str, Any]) -> dict[int, list[ContentBlock]]:
    """Графики/таблицы из ContentPlan по индексу слайда (первый/последний слайд —
    обложка и финал — объектов не получают)."""
    result: dict[int, list[ContentBlock]] = {}
    total = len(plan["slides"])
    pending: list[list[ContentBlock]] = []
    for spec in content_plan.slides:
        blocks = [b for b in spec.content_blocks if b.type in {ContentBlockType.CHART, ContentBlockType.TABLE}]
        pending.append(blocks)
    for index, slide in enumerate(plan["slides"]):
        if index == 0 or index == total - 1:
            continue
        # Варианты без разделителей короче плана содержания: слайд варианта → слайд плана по source_index
        src = int(slide.get("source_index", index))
        if src < len(pending) and pending[src]:
            result[index] = pending[src][:1]
    return result


def _extend_bbox_down(slide, bbox) -> tuple:
    """Область под график/таблицу продлевается вниз до верхнего края ближайшей
    фигуры, которая останется на слайде (заполненные плейсхолдеры, обычные
    фигуры не на весь слайд), либо до нижнего поля слайда."""
    left, top, width, height = bbox
    try:
        slide_h = int(slide.part.package.presentation_part.presentation.slide_height)
        slide_w = int(slide.part.package.presentation_part.presentation.slide_width)
    except Exception:  # noqa: BLE001
        return bbox
    limit = slide_h - Emu(360000)
    right = left + width
    for shape in slide.shapes:
        if None in (shape.left, shape.top, shape.width, shape.height):
            continue
        if shape.is_placeholder and not (shape.has_text_frame and shape.text_frame.text.strip()) and getattr(shape, "image", None) is None:
            continue  # пустой плейсхолдер будет удалён
        if shape.width >= slide_w * 0.9 and shape.height >= slide_h * 0.9:
            continue  # фон на весь слайд
        if shape.top <= top or shape.left >= right or shape.left + shape.width <= left:
            continue
        limit = min(limit, int(shape.top) - Emu(120000))
    if limit - top > height:
        return (left, top, width, limit - top)
    return bbox


def _slot_shape(slide, variant, slot_name):
    slot = (variant.get("slots") or {}).get(slot_name) or {}
    idx = slot.get("placeholder_idx")
    if idx is not None:
        for shape in slide.placeholders:
            if shape.placeholder_format.idx == idx:
                return shape
    for shape in slide.shapes:
        if shape.name == slot.get("shape_name") or shape.name == slot_name:
            return shape
    return None


def _inject_objects(slide, variant, slide_plan, blocks: list[ContentBlock], manifest: DesignManifest, variant_key: str) -> None:
    from app.pipeline.styling import style_chart, style_table

    contract = variant["generation_contract"]["text_slots"]
    body_slots = [s["slot"] for s in contract if s["role"] in {"body", "column_body"}]
    if not body_slots:
        return
    block = blocks[0]
    slots = variant.get("slots") or {}
    if len(body_slots) >= 2:
        # Двухколоночный макет: объект занимает правую колонку, текст остаётся в левой.
        target_slot = max(body_slots, key=lambda name: float((slots.get(name) or {}).get("x") or 0))
        shape = _slot_shape(slide, variant, target_slot)
        if shape is None:
            return
        bbox = (shape.left, shape.top, shape.width, shape.height)
        shape._element.getparent().remove(shape._element)
        # Текст правой колонки дописываем в левую, чтобы тезисы не потерялись.
        moved = slide_plan.get("content", {}).get(target_slot)
        left_slot = min(body_slots, key=lambda name: float((slots.get(name) or {}).get("x") or 0))
        left_shape = _slot_shape(slide, variant, left_slot)
        if moved and left_shape is not None and left_shape.has_text_frame:
            _append_paragraphs(left_shape.text_frame, moved)
    else:
        shape = _slot_shape(slide, variant, body_slots[0])
        if shape is None:
            return
        text = slide_plan.get("content", {}).get(body_slots[0]) or ""
        if text.strip():
            left0, top0, width0, height0 = shape.left, shape.top, shape.width, shape.height
            font_pt = ((slots.get(body_slots[0]) or {}).get("resolved_font") or {}).get("size_pt") or 16
            text_h = min(int(height0 * 0.5), _estimate_text_height(text, width0, font_pt))
            gap = Emu(150000)
            bbox = (left0, top0 + text_h + gap, width0, height0 - text_h - gap)
            # Плейсхолдер наследует геометрию от макета: задаём все четыре
            # значения явно, иначе python-pptx создаст xfrm с нулевой шириной.
            shape.left, shape.top, shape.width, shape.height = left0, top0, width0, text_h
        else:
            bbox = (shape.left, shape.top, shape.width, shape.height)
            shape._element.getparent().remove(shape._element)
    if bbox[3] < Emu(1500000):
        # Под объектом часто свободно (пустые плейсхолдеры картинок будут
        # удалены): растягиваем область вниз до ближайшей занятой фигуры.
        bbox = _extend_bbox_down(slide, bbox)
    if bbox[3] < Emu(1500000):
        # Объекту нужно хотя бы ~4 см по высоте — иначе оставляем текст как есть.
        return

    left, top, width, height = bbox
    if block.type == ContentBlockType.CHART and block.chart is not None:
        chart_data = CategoryChartData()
        chart_data.categories = block.chart.categories
        for series in block.chart.series:
            chart_data.add_series(series.name, series.values)
        frame = slide.shapes.add_chart(_CHART_TYPE_MAP[block.chart.chart_type.value], left, top, width, height, chart_data)
        style_chart(frame.chart, manifest, slide=slide)
    elif block.type == ContentBlockType.TABLE and block.table is not None:
        headers = block.table.headers[:5]
        rows = [r[:5] for r in block.table.rows[:7]]
        frame = slide.shapes.add_table(len(rows) + 1, max(len(headers), 1), left, top, width, min(height, Emu(int(720000 * (len(rows) + 1)))))
        table = frame.table
        for c, header in enumerate(headers):
            table.cell(0, c).text = header
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                if c < len(headers):
                    table.cell(r, c).text = str(value)
        style_table(table, manifest)


def _append_paragraphs(text_frame, text: str) -> None:
    """Добавляет абзацы в конец текстового фрейма, копируя стиль последнего абзаца."""
    for line in [l for l in text.splitlines() if l.strip()]:
        if not text_frame.text.strip():
            text_frame.text = line
            continue
        last = text_frame.paragraphs[-1]
        paragraph = text_frame.add_paragraph()
        paragraph.text = line
        if last.runs and paragraph.runs:
            src, dst = last.runs[0].font, paragraph.runs[0].font
            dst.size, dst.bold, dst.name = src.size, src.bold, src.name


def _estimate_text_height(text: str, width_emu: int, font_pt: float) -> int:
    """Грубая оценка высоты текста: средняя ширина кириллического знака ≈ 0.55 кегля,
    межстрочный интервал 1.25 + отступ абзаца."""
    char_w_emu = font_pt * 0.55 * 12700
    chars_per_line = max(8, int(width_emu / char_w_emu))
    lines = 0
    for paragraph in text.splitlines() or [text]:
        lines += max(1, -(-len(paragraph) // chars_per_line))
    line_h = font_pt * 1.25 * 12700
    paragraphs = max(1, len(text.splitlines()))
    return int(lines * line_h + paragraphs * font_pt * 0.5 * 12700 + 120000)


_SHORT_WORD_RE = re.compile(r"(?<![\w-])([А-Яа-яЁёA-Za-z]{1,2}|для|или|при|над|под|без|про|как|что|это)\s+(?=\S)", re.UNICODE)


def _glue_short_words(text: str) -> str:
    """Союзы и предлоги не должны висеть в конце строки: приклеиваем их к
    следующему слову неразрывным пробелом (U+00A0)."""
    return _SHORT_WORD_RE.sub(lambda m: m.group(1) + "\u00a0", text)


_METRIC_IN_BULLET_RE = re.compile(r"^\s*[•\-–—]?\s*(\d[\d\s\u00a0]*(?:[.,]\d+)?\s*%?)\s+(.+?)\s*$")


def _dedupe_metrics(template: dict[str, Any], plan: dict[str, Any]) -> dict[str, Any]:
    """Убирает повторяющиеся фактоиды на слайдах с метриками.

    Планировщик иногда ставит в оба слота metric_value одно и то же число
    (и дублирует его в заголовке). Если во втором фактоиде то же значение —
    берём другую цифру из буллетов body-слота («• 320 преподавателей …» →
    значение «320», подпись «преподавателей …») и убираем этот буллет из
    текста; если подходящего буллета нет — очищаем дубль, чтобы на слайде не
    стояли две одинаковые цифры."""
    variants = {v["id"]: v for v in template.get("variants", [])}
    for slide in plan.get("slides", []):
        contract = (variants.get(slide.get("variant")) or {}).get("generation_contract") or {}
        slots = contract.get("text_slots") or []
        values = [cs["slot"] for cs in slots if cs.get("role") == "metric_value"]
        labels = [cs["slot"] for cs in slots if cs.get("role") == "metric_label"]
        if len(values) < 2:
            continue
        content = slide.get("content") or {}
        seen: set[str] = set()
        for pos, value_slot in enumerate(values):
            value = str(content.get(value_slot) or "").strip()
            key = re.sub(r"\s+", "", value)
            if not value or key not in seen:
                seen.add(key)
                continue
            label_slot = labels[pos] if pos < len(labels) else None
            replacement = None
            for body_slot in [cs["slot"] for cs in slots if cs.get("role") in _BODY_ROLES]:
                lines = str(content.get(body_slot) or "").split("\n")
                for i, line in enumerate(lines):
                    m = _METRIC_IN_BULLET_RE.match(line)
                    if m and re.sub(r"\s+", "", m.group(1)) not in seen:
                        replacement = (m.group(1).strip(), m.group(2).strip())
                        del lines[i]
                        content[body_slot] = "\n".join(lines)
                        break
                if replacement:
                    break
            if replacement:
                content[value_slot] = replacement[0]
                if label_slot:
                    content[label_slot] = replacement[1]
                seen.add(re.sub(r"\s+", "", replacement[0]))
                logger.info("Дубль метрики заменён на «%s %s»", replacement[0], replacement[1])
            else:
                content[value_slot] = ""
                if label_slot:
                    content[label_slot] = ""
                logger.info("Дубль метрики «%s» очищен", value)
    return plan


def _glue_short_words_in_titles(template: dict[str, Any], plan: dict[str, Any]) -> dict[str, Any]:
    plan = deepcopy(plan)
    variants = {v["id"]: v for v in template.get("variants", [])}
    for slide in plan["slides"]:
        variant = variants.get(slide.get("variant")) or {}
        contract = {t["slot"]: t for t in (variant.get("generation_contract") or {}).get("text_slots") or []}
        for slot_name, text in list((slide.get("content") or {}).items()):
            if isinstance(text, str) and (contract.get(slot_name) or {}).get("role") in _TITLE_ROLES:
                slide["content"][slot_name] = _glue_short_words(text)
    return plan


def _section_body_allowed(template: dict[str, Any], plan: dict[str, Any], error: dict[str, Any]) -> bool:
    """Лид и пункты раздела в теле макета-разделителя: слот там обычно
    однострочный, но под ним пусто — рамка вырастет вниз (см. _fit_slide_text)."""
    try:
        slide_plan = plan["slides"][error["slide_index"]]
        if slide_plan.get("composition_kind") not in {"section", "statement", "title_only"}:
            return False
        variant = next(v for v in template["variants"] if v["id"] == slide_plan["variant"])
        slot = (variant.get("slots") or {})[error["slot"]]
    except (KeyError, IndexError, StopIteration, TypeError):
        return False
    if slot.get("semantic_role") not in {"body", "subtitle", "subheading"}:
        return False
    text = (slide_plan.get("content") or {}).get(error["slot"]) or ""
    sh = float((template.get("slide_size") or {}).get("height") or 0)
    room = sh - float(slot.get("y") or 0) - 1.2 if sh else 0
    return len(text.splitlines()) <= 6 and len(text) <= 420 and room >= 3.0


def _title_overflow_allowed(template: dict[str, Any], plan: dict[str, Any], error: dict[str, Any]) -> bool:
    """`text_too_long` для заголовка допустим в пределах _TITLE_OVERFLOW_FACTOR:
    текст останется целым, а кегль уменьшится на рендере."""
    try:
        slide_plan = plan["slides"][error["slide_index"]]
        variant = next(v for v in template["variants"] if v["id"] == slide_plan["variant"])
        contract = next(t for t in variant["generation_contract"]["text_slots"] if t["slot"] == error["slot"])
    except (KeyError, IndexError, StopIteration, TypeError):
        return False
    limit = (contract.get("constraints") or {}).get("maximum_characters")
    text = (slide_plan.get("content") or {}).get(error["slot"]) or ""
    if contract.get("role") in {"contact_text", "speaker_name", "url"} and " " not in text.split(" · ")[0]:
        return bool(limit) and len(text) <= int(limit * 3) + 1  # e-mail/ссылка целиком, кегль уменьшится
    if contract.get("role") == "unit_title":
        return bool(limit) and len(text) <= int(limit * 2.0) + 1
    if contract.get("role") in {"metric_value", "unit_number"} and "\n" not in text:
        return len(text) <= 24  # значение метрики целиком («150 млн»), кегль уменьшится до ширины бокса
    if contract.get("role") not in _TITLE_ROLES:
        return False
    return bool(limit) and len(text) <= int(limit * _TITLE_OVERFLOW_FACTOR) + 1


_MIN_FIT_SCALE = 0.55


def _fit_slide_text(slide, variant: dict[str, Any], slide_plan: dict[str, Any]) -> None:
    """Shrink-to-fit для текстовых слотов: если самое длинное слово не помещается
    в строку (LibreOffice/PowerPoint ломают его посередине — «Распределени-е»)
    или текст выше зоны слота, уменьшаем кегль всех ранов, но не ниже 55 %
    от исходного. Шрифт и цвет шаблона не трогаем."""
    slots = variant.get("slots") or {}
    for slot_name, text in (slide_plan.get("content") or {}).items():
        if not isinstance(text, str) or not text.strip():
            continue
        shape = _slot_shape(slide, variant, slot_name)
        if shape is None or not shape.has_text_frame:
            continue
        size_pt = _shape_font_pt(shape) or ((slots.get(slot_name) or {}).get("resolved_font") or {}).get("size_pt")
        if not size_pt or not shape.width or not shape.height:
            continue
        _clip_to_decor(slide, shape, slot_name)
        role = (slots.get(slot_name) or {}).get("semantic_role")
        is_title = role in _TITLE_ROLES
        if is_title:
            _widen_title(slide, shape, text, float(size_pt))
        if role == "unit_title" and slide_plan.get("metric_units") and "\n" not in text and float(size_pt) < 24:
            # Значение метрики в заголовке карточки: 13 pt «20%» теряется — укрупняем,
            # насколько позволяют ширина бокса и свободное место под ним до следующей фигуры
            room = _available_height(slide, shape)
            inner_w = max(1, int(shape.width) - 2 * 91440)
            target = float(size_pt)
            for candidate in (32.0, 28.0, 24.0, 20.0, 18.0):
                if candidate <= float(size_pt):
                    break
                if len(text) * candidate * 0.6 * 12700 <= inner_w and candidate * 1.25 * 12700 + 2 * 45720 <= room:
                    target = candidate
                    break
            if target > float(size_pt):
                l, t, w = int(shape.left), int(shape.top), int(shape.width)
                shape.left, shape.top, shape.width, shape.height = Emu(l), Emu(t), Emu(w), Emu(int(target * 1.25 * 12700 + 2 * 45720))
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(target)
                        run.font.bold = True
                size_pt = target
        if role in {"body", "subtitle", "subheading"} and slide_plan.get("composition_kind") in {"section", "statement", "title_only"} and "\n" in text:
            # Лид + пункты раздела в однострочном слоте: рамка растёт вниз до свободного места
            room = _available_height(slide, shape)
            if room > int(shape.height):
                # плейсхолдер наследует геометрию от макета — задаём все четыре значения
                l, t, w = int(shape.left), int(shape.top), int(shape.width)
                shape.left, shape.top, shape.width, shape.height = Emu(l), Emu(t), Emu(w), Emu(room)
                # подпись разделителя в шаблоне мелкая (8–10 pt) — лид с пунктами укрупняем
                if float(size_pt) < 12 and _title_height(text, max(1, w - 2 * 91440), min(14.0, float(size_pt) * 1.6)) <= room:
                    size_pt = min(14.0, float(size_pt) * 1.6)
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(size_pt)
        height = _available_height(slide, shape) if is_title else int(shape.height)
        if role in {"metric_value", "unit_number"} and "\n" not in text:
            # Значение метрики — одной строкой («150 млн», а не «150 / млн»): меряем как одно слово
            fitted = _fit_font_size(text.replace(" ", "\u00a0"), int(shape.width), height, float(size_pt), check_height=False)
        else:
            fitted = _fit_font_size(text, int(shape.width), height, float(size_pt), check_height=is_title)
        if fitted is None or fitted >= size_pt - 0.5:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(fitted)
        logger.info("Кегль слота %s уменьшен %.0f → %.0f pt (shrink-to-fit)", slot_name, size_pt, fitted)


def _widen_title(slide, shape, text: str, size_pt: float) -> None:
    """Заголовок в узком плейсхолдере, справа от которого на слайде и в макете
    ничего нет, растягиваем до правого поля — вместо переноса на 3 строки или
    уменьшения кегля."""
    prs_w = int(slide.part.package.presentation_part.presentation.slide_width)
    left, top, width, height = int(shape.left), int(shape.top), int(shape.width), int(shape.height)
    if width >= prs_w * 0.6:
        return
    if _title_height(text, max(1, width - 2 * 91440), size_pt) <= height:
        return  # и так помещается
    right_margin = max(left, int(prs_w * 0.04))
    new_right = prs_w - right_margin
    if new_right - left <= width * 1.15:
        return
    for other in list(slide.shapes) + list(slide.slide_layout.shapes):
        if other is shape or other.left is None or other.width is None:
            continue
        if other.shape_type == MSO_SHAPE_TYPE.PICTURE and int(other.width) * int(other.height) > 0.3 * 12192000 * 6858000:
            continue
        o_left, o_top, o_bottom = int(other.left), int(other.top), int(other.top) + int(other.height)
        if o_left < left + width - 91440:
            continue  # не справа
        if o_bottom <= top or o_top >= top + height:
            continue  # не в полосе заголовка
        new_right = min(new_right, o_left - 182880)
    if new_right - left > width * 1.15:
        shape.width = Emu(new_right - left)
        logger.info("Заголовок расширен %.1f → %.1f см (справа пусто)", width / 360000, (new_right - left) / 360000)


_DECOR_GAP_EMU = 228600  # 0.25"
_MIN_CLIP_SCALE = 0.45


def _clip_to_decor(slide, shape, slot_name: str) -> None:
    """Сужает текстовую фигуру, если справа от неё на макете лежит декоративная
    картинка/фигура (паттерн на обложке VK и т.п.), перекрывающая её по
    вертикали. В шаблоне текстовый бокс часто уходит под паттерн — короткий
    образец текста туда не доходил, а сгенерированный заголовок доходит и
    прячется за картинкой. После сужения текст переносится до паттерна."""
    left, top, width, height = int(shape.left), int(shape.top), int(shape.width), int(shape.height)
    right = left + width
    limit = right
    candidates = list(slide.slide_layout.shapes) + [sh for sh in slide.shapes if sh is not shape]
    for other in candidates:
        if other.is_placeholder or other.has_text_frame and other.text_frame.text.strip():
            continue
        if other.shape_type not in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.GROUP):
            continue
        if other.width is None or other.left is None:
            continue
        o_left, o_top, o_h = int(other.left), int(other.top), int(other.height or 0)
        if o_top >= top + height or o_top + o_h <= top:
            continue  # не пересекается по вертикали
        if o_left <= left + width * 0.3 or o_left >= right:
            continue  # фон/подложка слева или объект правее бокса
        limit = min(limit, o_left - _DECOR_GAP_EMU)
    if limit >= right:
        return
    new_width = limit - left
    if new_width < width * _MIN_CLIP_SCALE:
        return
    # У плейсхолдера без собственного <a:xfrm> геометрия наследуется от макета;
    # присвоение одной только ширины создало бы xfrm с нулями в остальных
    # полях — фиксируем все четыре значения явно.
    shape.left, shape.top, shape.height = Emu(left), Emu(top), Emu(height)
    shape.width = Emu(new_width)
    logger.info("Слот %s сужен %.2f → %.2f\" — справа декоративный элемент макета", slot_name, width / 914400, new_width / 914400)


def _shape_font_pt(shape) -> float | None:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                return run.font.size.pt
    return None


def _available_height(slide, shape) -> int:
    """Высота, доступная заголовку: от верха его бокса до ближайшей фигуры ниже
    (тело слайда, объект), минус зазор. Боксы заголовков в шаблонах обычно
    рассчитаны на одну строку, но под ними есть свободное поле — двухстрочный
    заголовок в него помещается без уменьшения кегля."""
    top, height = int(shape.top), int(shape.height)
    left, right = int(shape.left), int(shape.left) + int(shape.width)
    bottom_limit = None
    for other in list(slide.shapes) + list(slide.slide_layout.shapes):
        if other is shape or other.top is None or other.width is None:
            continue
        if other.shape_type == MSO_SHAPE_TYPE.PICTURE and int(other.width) * int(other.height) > 0.3 * 12192000 * 6858000:
            continue  # фоновые картинки/паттерны не ограничивают
        o_top, o_left, o_right = int(other.top), int(other.left), int(other.left) + int(other.width)
        if o_top < top + height * 0.5 or o_right <= left or o_left >= right:
            continue
        bottom_limit = o_top if bottom_limit is None else min(bottom_limit, o_top)
    if bottom_limit is None:
        return height
    return max(height, bottom_limit - top - 91440)


def _title_height(text: str, width_emu: int, font_pt: float) -> int:
    """Высота заголовка: строки × 1.2 кегля, без абзацных отступов."""
    char_w_emu = font_pt * 0.55 * 12700
    chars_per_line = max(8, int(width_emu / char_w_emu))
    lines = 0
    for paragraph in text.splitlines() or [text]:
        # жадный перенос по словам (как в PowerPoint): неразрывный пробел
        # не является точкой переноса, хвосты строк не заполняются
        current = 0
        lines += 1
        for word in re.split(r"[ \t]+", paragraph.strip()):
            if not word:
                continue
            need = len(word) if current == 0 else current + 1 + len(word)
            if need <= chars_per_line:
                current = need
            else:
                lines += 1
                current = len(word)
    return int(lines * font_pt * 1.2 * 12700)


def _fit_font_size(text: str, width_emu: int, height_emu: int, size_pt: float, check_height: bool = True) -> float | None:
    """Наибольший кегль ≤ size_pt (шаг 1 pt), при котором самое длинное слово
    помещается в строку, а оценка высоты не превышает высоту фигуры."""
    inner_w = max(1, width_emu - 2 * 91440)
    inner_h = max(1, height_emu - 2 * 45720)
    longest = max((len(word) for word in re.split(r"\s+", text) if word), default=0)
    floor = size_pt * _MIN_FIT_SCALE
    size = size_pt
    while size >= floor:
        char_w = size * 0.55 * 12700
        fits_h = (not check_height) or _title_height(text, inner_w, size) <= inner_h
        if longest * char_w <= inner_w and fits_h:
            return size
        size -= 1.0
    return floor


def _remove_empty_placeholders(slide) -> None:
    """Пустые текстовые/картиночные плейсхолдеры не должны оставаться на
    слайде: в PowerPoint они показывают подсказку «Вставьте текст», а в PDF —
    пустую область (аудит «пустой слайд/незаполненный слот»)."""
    for shape in list(slide.placeholders):
        if shape.has_text_frame and shape.text_frame.text.strip():
            continue
        if shape.shape_type is not None and "PICTURE" in str(shape.shape_type) and not shape.is_placeholder:
            continue
        # Плейсхолдер с вставленной картинкой становится PlaceholderPicture без text_frame — не трогаем.
        if not shape.has_text_frame and getattr(shape, "image", None) is not None:
            continue
        if shape.has_text_frame or getattr(shape, "placeholder_format", None) is not None:
            shape._element.getparent().remove(shape._element)


def _drop_unused_layouts(prs) -> None:
    used = {slide.slide_layout.part.partname for slide in prs.slides}
    for master in prs.slide_masters:
        for layout in list(master.slide_layouts):
            if layout.part.partname in used:
                continue
            try:
                master.slide_layouts.remove(layout)
            except Exception as exc:  # noqa: BLE001 — макет остаётся, файл просто чуть больше
                logger.debug("Не удалось удалить макет %s: %s", layout.name, exc)
