"""Собственная библиотека композиций — для шаблонов без примеров дизайна
(`design_source == "library"`: макеты только с заголовком и нет слайдов-образцов).

Берём из шаблона шрифты темы, палитру и самый чистый макет «только заголовок»,
а компоновку (карточки, метрики, таймлайн, колонки, тезис, оглавление, раздел,
финал, таблица, график) рисуем сами средствами python-pptx. Композиции
описываются как варианты `source: "library"` в том же формате, что и
слайды-образцы (см. sample_slides.py), поэтому драматургия и рендер работают с
ними единообразно.
"""
from __future__ import annotations

import logging
import math
from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Emu, Pt

from app.pipeline.sample_slides import _slot_constraints

logger = logging.getLogger(__name__)

_MARGIN = 1.3  # см
_TITLE_ROLES = {"slide_title", "presentation_title", "section_title", "closing_title"}


# ---------------------------------------------------------------------------
# Описание композиций
# ---------------------------------------------------------------------------

_LIBRARY: list[dict[str, Any]] = [
    {"key": "cards", "label": "Карточки", "kind": "cards", "units": 4, "min_units": 2, "max_units": 8, "unit_roles": ["unit_number", "unit_title", "unit_text"], "extra": ["note"], "role": "content"},
    {"key": "columns", "label": "Две колонки", "kind": "columns", "units": 2, "min_units": 2, "max_units": 3, "unit_roles": ["unit_title", "unit_text"], "extra": [], "role": "content"},
    {"key": "metrics", "label": "Ключевые цифры", "kind": "metrics", "units": 3, "min_units": 1, "max_units": 4, "unit_roles": ["metric_value", "metric_label"], "extra": ["body"], "role": "content"},
    {"key": "timeline", "label": "Таймлайн", "kind": "timeline", "units": 4, "min_units": 3, "max_units": 6, "unit_roles": ["unit_label", "unit_title", "unit_text"], "extra": [], "role": "content"},
    {"key": "problems", "label": "Проблемы и решения", "kind": "problems", "units": 3, "min_units": 2, "max_units": 6, "unit_roles": ["unit_number", "unit_title", "unit_text"], "extra": [], "role": "content"},
    {"key": "text", "label": "Текст", "kind": "text", "units": 0, "min_units": 0, "max_units": 0, "unit_roles": [], "extra": ["body", "note"], "role": "content"},
    {"key": "statement", "label": "Тезис", "kind": "statement", "units": 0, "min_units": 0, "max_units": 0, "unit_roles": [], "extra": ["body"], "role": "content"},
    {"key": "agenda", "label": "Оглавление", "kind": "agenda", "units": 5, "min_units": 2, "max_units": 8, "unit_roles": ["unit_number", "agenda_item"], "extra": [], "role": "content"},
    {"key": "section", "label": "Раздел", "kind": "section", "units": 0, "min_units": 0, "max_units": 0, "unit_roles": [], "extra": ["label", "subtitle", "body"], "role": "section"},
    {"key": "closing", "label": "Финал", "kind": "closing", "units": 0, "min_units": 0, "max_units": 0, "unit_roles": [], "extra": ["subtitle", "contact"], "role": "closing"},
    {"key": "table", "label": "Таблица", "kind": "table", "units": 0, "min_units": 0, "max_units": 0, "unit_roles": [], "extra": ["note"], "role": "content"},
    {"key": "chart", "label": "График", "kind": "chart", "units": 0, "min_units": 0, "max_units": 0, "unit_roles": [], "extra": ["body"], "role": "content"},
]

_SUFFIX = {"unit_number": "num", "unit_title": "title", "unit_text": "text", "unit_label": "label", "metric_value": "value", "metric_label": "caption", "agenda_item": "item"}


def _decor_images(variant: dict[str, Any], sw: float, sh: float) -> list[dict[str, float]]:
    """Декоративные картинки макета, кроме фоновых (почти на весь слайд)."""
    result = []
    for image in variant.get("images") or []:
        try:
            x, y, w, h = float(image.get("x") or 0), float(image.get("y") or 0), float(image.get("w") or 0), float(image.get("h") or 0)
        except (TypeError, ValueError):
            continue
        if w * h >= 0.8 * sw * sh:
            continue
        result.append({"x": x, "y": y, "w": w, "h": h})
    # Статические фигуры макета (декоративный текст «03 04», плашки, линии)
    for shape in variant.get("decor") or []:
        w, h = float(shape.get("w") or 0), float(shape.get("h") or 0)
        if w * h >= 0.8 * sw * sh or (not shape.get("text") and h < 0.15 and not shape.get("picture")):
            continue  # фон или тонкая линия
        result.append({"x": float(shape.get("x") or 0), "y": float(shape.get("y") or 0), "w": w, "h": h, "text": bool(shape.get("text"))})
    return result


def _intersection(a: tuple[float, float, float, float], b: dict[str, float]) -> float:
    ax, ay, aw, ah = a
    w = min(ax + aw, b["x"] + b["w"]) - max(ax, b["x"])
    h = min(ay + ah, b["y"] + b["h"]) - max(ay, b["y"])
    return max(w, 0) * max(h, 0)


def pick_base_layout(template: dict[str, Any]) -> dict[str, Any] | None:
    """Макет с заголовком, у которого декор меньше всего залезает в зону контента."""
    sw = float(template["slide_size"]["width"])
    sh = float(template["slide_size"]["height"])
    best = None
    best_score = None
    for variant in template.get("variants") or []:
        if variant.get("source") in {"sample", "library"}:
            continue
        slots = variant.get("slots") or {}
        title = next((s for s in slots.values() if s.get("semantic_role") in {"slide_title", "presentation_title", "section_title"}), None)
        if title is None:
            continue
        area = _content_area(variant, sw, sh)
        overlap = sum(_intersection(area, img) * (2.0 if img.get("text") else 1.0) for img in _decor_images(variant, sw, sh))
        others = [s for s in slots.values() if s is not title]
        busy = sum(float(d.get("busy") or 0) for d in variant.get("decor") or [])
        score = busy * 60 + (4 if area[3] < 0.45 * sh else 0) + overlap / (sw * sh) * 100 + len(others) * 1.5 + (0 if float(title.get("y") or 0) < 4 else 3) + (0 if title.get("semantic_role") == "slide_title" else 2) + (0 if area[2] >= 0.55 * sw else 4) + (1 - area[2] / (sw - 2 * _MARGIN)) * 6
        if best_score is None or score < best_score:
            best, best_score = variant, score
    if best is None:
        # Ни в одном макете нет слота заголовка (авторские слайды с текстовыми
        # блоками): берём самый «чистый» макет и рисуем заголовок сами —
        # библиотека подставит геометрию по умолчанию (см. build_library_variants)
        candidates = [v for v in template.get("variants") or [] if v.get("source") not in {"sample", "library"}]
        if candidates:
            best = min(
                candidates,
                key=lambda v: (
                    sum(float(d.get("busy") or 0) for d in v.get("decor") or []) * 60
                    + len(v.get("slots") or {}) * 1.5
                    + len(v.get("image_slots") or []) * 2
                ),
            )
    return best


def _content_area(base: dict[str, Any], sw: float, sh: float) -> tuple[float, float, float, float]:
    """Зона контента: под заголовком, до нижнего декора, в стороне от боковых колонок декора."""
    slots = base.get("slots") or {}
    title = next((s for s in slots.values() if s.get("semantic_role") in {"slide_title", "presentation_title", "section_title"}), None)
    top = float(title["y"]) + float(title["h"]) + 0.5 if title else 3.5
    top = min(max(top, 3.0), sh * 0.4)
    left = float(title["x"]) if title else _MARGIN
    left = min(max(left, _MARGIN), sw * 0.2)
    right = sw - _MARGIN
    bottom = sh - 1.4
    for img in _decor_images(base, sw, sh):
        tall = img["h"] >= 0.6 * sh
        if tall and img["x"] <= 0.5 and img["w"] < 0.5 * sw:      # левая колонка декора
            left = max(left, img["x"] + img["w"] + 0.6)
        elif tall and img["x"] + img["w"] >= sw - 0.5 and img["w"] < 0.5 * sw:  # правая колонка
            right = min(right, img["x"] - 0.6)
        elif img["y"] >= sh * 0.7 and img["w"] >= 0.5 * sw:      # нижняя полоса
            bottom = min(bottom, img["y"] - 0.4)
        elif img["y"] >= sh * 0.82:                                # логотипы в подвале
            bottom = min(bottom, img["y"] - 0.4)
    if right - left < 0.45 * sw:
        left, right = _MARGIN, sw - _MARGIN
    return left, top, right - left, max(bottom - top, 4.0)


def build_library_variants(template: dict[str, Any], only_kinds: set[str] | None = None) -> list[dict[str, Any]]:
    """Композиции библиотеки на базовом макете шаблона.

    `only_kinds` — построить только указанные виды (так в шаблон с образцами
    добавляются таблица/график/тезис, которых среди образцов не нашлось).
    """
    base = pick_base_layout(template)
    if base is None:
        return []
    sw = float(template["slide_size"]["width"])
    sh = float(template["slide_size"]["height"])
    theme = template.get("theme_fonts") or {}
    major = theme.get("major") or "Arial"
    minor = theme.get("minor") or "Arial"
    title_slot_name, title_slot = next(((n, s) for n, s in (base.get("slots") or {}).items() if s.get("semantic_role") in {"slide_title", "presentation_title", "section_title"}), ("title", {}))
    left, top, width, height = _content_area(base, sw, sh)
    result = []
    for spec in _LIBRARY:
        if only_kinds is not None and spec["kind"] not in only_kinds:
            continue
        slots: dict[str, Any] = {}
        contract: list[dict[str, Any]] = []

        def add(name, role, x, y, w, h, size, family, required=False, unit=None, virtual=False, virtual_from=None):
            constraints = _slot_constraints(role, w, h, size)
            slots[name] = {
                "shape_id": None, "shape_name": name, "placeholder_idx": None, "paragraph": None, "line": None,
                "unit": unit, "x": round(x, 3), "y": round(y, 3), "w": round(w, 3), "h": round(h, 3),
                "font": {"family": family, "size": size}, "resolved_font": {"family": family, "size_pt": size, "bold": role in _TITLE_ROLES or role == "unit_title", "color": None, "source": "library"},
                "semantic_role": role, "required": required, "template_text": "", "editable": True,
                **({"virtual": True, "virtual_from": virtual_from} if virtual else {}),
            }
            contract.append({"slot": name, "generator_field": "content", "accepted_value": "string", "role": role, "required": required,
                             "semantic_scope": "specific", "allowed_content_kinds": [role], "constraints": constraints, **({"virtual": True} if virtual else {})})

        # Заголовок — плейсхолдер макета
        t_size = float((title_slot.get("resolved_font") or {}).get("size_pt") or 32)
        add("title", "section_title" if spec["kind"] == "section" else ("closing_title" if spec["kind"] == "closing" else "slide_title"),
            float(title_slot.get("x", left)), float(title_slot.get("y", 1.2)), float(title_slot.get("w", width)), float(title_slot.get("h", 2.2)), t_size, major, required=True)
        slots["title"]["placeholder_idx"] = title_slot.get("placeholder_idx")
        slots["title"]["shape_id"] = title_slot.get("shape_id")

        unit_slots: list[list[str]] = []
        units = spec["units"]
        if units:
            geometry = unit_geometry(spec["kind"], units, left, top, width, height)
            for k in range(spec["max_units"]):
                names = []
                g = geometry[k] if k < len(geometry) else geometry[-1]
                for role in spec["unit_roles"]:
                    name = f"u{k + 1}_{_SUFFIX[role]}"
                    box, size, family = _unit_slot_box(spec["kind"], role, g, major, minor)
                    add(name, role, *box, size, family, unit=k, virtual=k >= units, virtual_from=f"u1_{_SUFFIX[role]}" if k >= units else None)
                    names.append(name)
                unit_slots.append(names)
        for role in spec["extra"]:
            if role == "body":
                if spec["kind"] == "text":
                    add("body", "body", left, top, width * 0.62, height, 18, minor)
                elif spec["kind"] == "statement":
                    add("body", "body", left, top + 0.4, width, height * 0.7, 28, major)
                elif spec["kind"] == "metrics":
                    add("body", "body", left, top + height * 0.62, width, height * 0.38, 16, minor)
                elif spec["kind"] == "chart":
                    add("body", "body", left + width * 0.66, top, width * 0.34, height, 15, minor)
                elif spec["kind"] == "section":
                    add("body", "body", left + width * 0.66, top + height * 0.5, width * 0.34, height * 0.5, 15, minor)
            elif role == "note":
                add("note", "note", left, top + height - 1.6, width, 1.6, 14, minor)
            elif role == "subtitle":
                if spec["kind"] == "section":
                    add("subtitle", "subtitle", left, top + height * 0.5, width * 0.62, 2.6, 20, minor)
                else:
                    add("subtitle", "subtitle", left, top, width * 0.7, 2.4, 20, minor)
            elif role == "label":
                add("label", "section_label", left, top, width * 0.5, 0.9, 14, minor)
            elif role == "contact":
                add("contact", "contact_text", left, top + 3.0, width * 0.7, 2.0, 16, minor)

        variant_id = f"Библиотека · {spec['label']}" + (f" · {units} блока" if units and units < 5 else f" · {units} блоков" if units else "")
        composition = {
            "kind": spec["kind"], "units": units, "min_units": spec["min_units"], "max_units": spec["max_units"],
            "grid": {}, "unit_members": [], "unit_slots": unit_slots, "title_slot": "title",
            "has_body": "body" in spec["extra"], "has_note": "note" in spec["extra"], "has_image": False,
            "table_shape_id": "native" if spec["kind"] == "table" else None, "chart_pictures": ["native"] if spec["kind"] == "chart" else [],
            "removable_when_empty": ["note", "subtitle", "contact", "label", "body"], "library_key": spec["key"],
        }
        narrative = {"content": "content", "section": "section", "closing": "closing"}[spec["role"]]
        result.append({
            "id": variant_id,
            "source": "library",
            "layout_name": base.get("layout_name"),
            "type": "library",
            "layout_file": base.get("layout_file"),
            "master_index": base.get("master_index", 0),
            "base_variant": base["id"],
            "slots": slots,
            "image_slots": {},
            "editable_image_slots": [],
            "background": base.get("background"),
            "images": [],
            "constraints": {},
            "semantic": {"layout_role": narrative, "composition": spec["kind"]},
            "inferred": {
                "layout": {"narrative_role": narrative, "content_pattern": {"cards": "repeated_blocks", "metrics": "metrics", "timeline": "horizontal_sequence", "columns": "parallel_content"}.get(spec["kind"], "title_body"), "confidence": 0.7},
                "composition": {"primary": {"cards": "repeated_blocks", "columns": "parallel_content", "metrics": "metrics"}.get(spec["kind"], "title_body")},
            },
            "composition": composition,
            "generation_contract": {
                "text_slots": contract,
                "image_slots": [],
                "required_slots": ["title"],
                "content_budget": {"maximum_characters": sum((c["constraints"].get("maximum_characters") or 0) for c in contract if not c.get("virtual"))},
                "selection": {"best_for": [spec["label"]]},
            },
        })
    return result


def unit_geometry(kind: str, n: int, left: float, top: float, width: float, height: float) -> list[dict[str, float]]:
    """Прямоугольники блоков (см) для n блоков."""
    gap = 0.6
    n = max(n, 1)
    if kind in {"cards", "problems", "columns"}:
        cols = n if n <= 4 else math.ceil(n / 2)
        rows = math.ceil(n / cols)
        w = (width - gap * (cols - 1)) / cols
        h = (height - gap * (rows - 1)) / rows
        return [{"x": left + (i % cols) * (w + gap), "y": top + (i // cols) * (h + gap), "w": w, "h": h} for i in range(n)]
    if kind == "metrics":
        w = (width - gap * (n - 1)) / n
        h = height * 0.58
        return [{"x": left + i * (w + gap), "y": top, "w": w, "h": h} for i in range(n)]
    if kind == "timeline":
        w = (width - gap * (n - 1)) / n
        return [{"x": left + i * (w + gap), "y": top, "w": w, "h": height} for i in range(n)]
    if kind == "agenda":
        h = min(1.7, (height - 0.3 * (n - 1)) / n)
        return [{"x": left, "y": top + i * (h + 0.3), "w": width * 0.75, "h": h} for i in range(n)]
    return [{"x": left, "y": top, "w": width, "h": height}]


def _unit_slot_box(kind: str, role: str, g: dict[str, float], major: str, minor: str) -> tuple[tuple[float, float, float, float], float, str]:
    x, y, w, h = g["x"], g["y"], g["w"], g["h"]
    pad = 0.5
    if kind in {"cards", "problems"}:
        if role == "unit_number":
            return (x + pad, y + pad, 1.4, 1.0), 14, major
        if role == "unit_title":
            return (x + pad, y + pad + 1.2, w - 2 * pad, 1.6), 18, major
        return (x + pad, y + pad + 2.9, w - 2 * pad, max(h - 2.9 - 2 * pad, 1.0)), 14, minor
    if kind == "columns":
        if role == "unit_title":
            return (x, y, w, 1.6), 20, major
        return (x, y + 1.9, w, h - 1.9), 16, minor
    if kind == "metrics":
        if role == "metric_value":
            return (x, y, w, h * 0.6), 54, major
        return (x, y + h * 0.6, w, h * 0.4), 15, minor
    if kind == "timeline":
        if role == "unit_label":
            return (x, y, w - 0.4, 1.0), 14, major
        if role == "unit_title":
            return (x, y + 2.0, w - 0.4, 1.6), 16, major
        return (x, y + 3.7, w - 0.4, max(h - 3.7, 1.0)), 13, minor
    if kind == "agenda":
        if role == "unit_number":
            return (x, y, 1.8, h), 22, major
        return (x + 2.0, y, w - 2.0, h), 18, minor
    return (x, y, w, h), 16, minor


# ---------------------------------------------------------------------------
# Рендер
# ---------------------------------------------------------------------------


def _rgb(hex_value: str | None, fallback: str) -> RGBColor:
    value = (hex_value or fallback).lstrip("#")
    if len(value) != 6:
        value = fallback
    return RGBColor.from_string(value.upper())


def _luminance(color: RGBColor) -> float:
    r, g, b = color[0] / 255, color[1] / 255, color[2] / 255
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _blend(a: RGBColor, b: RGBColor, t: float) -> RGBColor:
    return RGBColor(*(int(round(a[i] * (1 - t) + b[i] * t)) for i in range(3)))


def palette_for(template: dict[str, Any], variant: dict[str, Any], manifest) -> dict[str, RGBColor]:
    pal = (manifest.palette if manifest is not None else {}) or {}
    dk1 = _rgb(pal.get("dk1"), "1B1B1F")
    lt1 = _rgb(pal.get("lt1"), "FFFFFF")
    accent = _rgb(pal.get("accent1"), "2F80ED")
    # Тёмный ли фон макета — по цвету заголовка базового макета
    base_id = variant.get("base_variant")
    base = next((v for v in template.get("variants") or [] if v["id"] == base_id), None)
    title_color = None
    if base:
        title = next((s for s in (base.get("slots") or {}).values() if s.get("semantic_role") in _TITLE_ROLES), None)
        title_color = ((title or {}).get("resolved_font") or {}).get("color")
    dark = False
    if title_color:
        dark = _luminance(_rgb(title_color, "000000")) > 0.6
    bg = dk1 if dark else lt1
    text = lt1 if dark else dk1
    if _luminance(accent) < 0.25 and dark:
        accent = _blend(accent, lt1, 0.45)
    if _luminance(accent) > 0.8 and not dark:
        accent = _blend(accent, dk1, 0.4)
    return {
        "bg": bg, "text": text, "accent": accent,
        "surface": _blend(bg, text, 0.07), "surface_line": _blend(bg, text, 0.16),
        "muted": _blend(text, bg, 0.35),
    }


def _textbox(slide, name, x, y, w, h, text, size, family, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.1)
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    lines = [line for line in str(text).split("\n") if line.strip()] or [""]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if i > 0:
            p.space_before = Pt(size * 0.35)
        run = p.add_run()
        run.text = line.strip()
        run.font.size = Pt(size)
        run.font.name = family
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def _rect(slide, name, x, y, w, h, fill, line=None, rounded=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    shape.name = name
    if rounded:
        shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    shape.text_frame.text = ""
    return shape


def _decor_right_of_title(base: dict[str, Any], title: dict[str, Any], sw: float, sh: float) -> bool:
    """Есть ли декор макета справа от рамки заголовка в его полосе."""
    tx, ty, tw, th = (float(title.get(k) or 0) for k in ("x", "y", "w", "h"))
    for d in _decor_images(base, sw, sh):
        if d["x"] >= tx + tw - 0.2 and d["y"] < ty + th and d["y"] + d["h"] > ty and d["w"] * d["h"] < 0.8 * sw * sh:
            return True
    return False


def _fit_size(text: str, w_cm: float, h_cm: float, size: float, min_size: float = 10.0) -> float:
    """Грубая оценка: уменьшаем кегль, пока текст не поместится в рамку."""
    while size > min_size:
        chars_per_line = max(6, int((w_cm * 28.35) / (size * 0.52)))
        lines = 0
        for paragraph in text.split("\n"):
            lines += max(1, math.ceil(len(paragraph) / chars_per_line))
        if lines * size * 1.25 / 28.35 <= h_cm:
            return size
        size -= 1
    return min_size


def render_library_slide(prs, template, variant, slide_plan, blocks, manifest, layouts, chart_type_map, shorten) -> Any:
    layout = layouts[variant["layout_file"]]
    slide = prs.slides.add_slide(layout)
    content: dict[str, str] = {k: v for k, v in (slide_plan.get("content") or {}).items() if isinstance(v, str) and v.strip()}
    slots = variant["slots"]
    comp = variant["composition"]
    kind = comp["kind"]
    colors = palette_for(template, variant, manifest)
    theme = template.get("theme_fonts") or {}
    major, minor = theme.get("major") or "Arial", theme.get("minor") or "Arial"
    sw, sh = prs.slide_width / 360000, prs.slide_height / 360000

    base = next((v for v in template["variants"] if v["id"] == variant.get("base_variant")), variant)
    # Заголовок — в плейсхолдер макета; прочие плейсхолдеры убираем
    title_text = content.get("title", "")
    title_done = False
    for shape in list(slide.placeholders):
        role_idx = slots["title"].get("placeholder_idx")
        if not title_done and shape.placeholder_format.idx == role_idx and shape.has_text_frame:
            shape.text_frame.text = title_text
            shape.name = "title"
            title_done = True
            # Длинный заголовок не должен наезжать на контент: кегль под рамку,
            # а рамку расширяем до правого поля, если справа пусто
            t = slots["title"]
            base_size = float((t.get("resolved_font") or {}).get("size_pt") or 32)
            box_w = max(float(t.get("w") or 0), min(sw - float(t.get("x") or 0) - _MARGIN, sw * 0.75)) if not _decor_right_of_title(base, t, sw, sh) else float(t.get("w") or 0)
            if box_w > float(t.get("w") or 0) + 0.1:
                shape.width = Cm(box_w)
            fitted = _fit_size(title_text, box_w - 0.5, max(float(t.get("h") or 0) - 0.2, 1.0), base_size, min_size=max(16.0, base_size * 0.55))
            if fitted < base_size - 0.5:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(fitted)
            continue
        shape._element.getparent().remove(shape._element)
    if not title_done and title_text:
        t = slots["title"]
        _textbox(slide, "title", t["x"], t["y"], t["w"], t["h"], title_text, t["resolved_font"]["size_pt"], major, colors["text"], bold=True)

    left, top, width, height = _content_area(base, sw, sh)
    # Фоновая картинка макета с бледными плашками/водяными знаками: закрываем
    # зону контента плашкой цвета фона, чтобы контент не мешался с декором
    busy_bg = next((d for d in base.get("decor") or [] if d.get("picture") and float(d.get("faint") or 0) > 0.02 and float(d.get("flat") or 0) > 0.6), None)
    cover_hex = (busy_bg or {}).get("bg_hex") or ((base.get("background") or {}).get("color") if busy_bg else None)
    if busy_bg is not None and cover_hex:
        cover = _rect(slide, "bg_cover", 0.3, max(top - 0.5, 0.2), sw - 0.6, min(height + 0.5 + 0.15, sh - top), _rgb(cover_hex, "FFFFFF"), rounded=False)
        tree = slide.shapes._spTree
        tree.remove(cover._element)
        tree.insert(2, cover._element)  # под остальные фигуры слайда

    unit_slots: list[list[str]] = comp.get("unit_slots") or []
    n = 0
    for k, names in enumerate(unit_slots):
        if any(content.get(nm) for nm in names):
            n = k + 1
    if kind in {"cards", "problems", "columns", "metrics", "timeline", "agenda"} and n:
        geometry = unit_geometry(kind, n, left, top, width, height)
        _draw_units(slide, kind, n, geometry, unit_slots, slots, content, colors, major, minor)
    if kind == "section":
        _render_section(slide, slide_plan, content, slots, colors, major, minor, left, top, width, height, sw, sh, base)
        return slide
    if kind == "statement" and content.get("body"):
        b = slots["body"]
        _rect(slide, "accent_bar", left, top + 0.3, 0.25, height * 0.7 - 0.2, colors["accent"], rounded=False)
        size = _fit_size(content["body"], b["w"] - 1.2, b["h"], b["resolved_font"]["size_pt"], 16)
        _textbox(slide, "body", b["x"] + 0.9, b["y"], b["w"] - 1.2, b["h"], content["body"], size, major, colors["text"], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)
    elif kind == "text" and content.get("body"):
        b = slots["body"]
        size = _fit_size(content["body"], b["w"], b["h"] - 1.8, b["resolved_font"]["size_pt"], 12)
        _textbox(slide, "body", b["x"], b["y"], b["w"], b["h"] - 1.8, content["body"], size, minor, colors["text"], line_spacing=1.2)
        _rect(slide, "accent_bar", left + width * 0.7, top, 0.25, height - 1.8, colors["accent"], rounded=False)
    elif kind == "metrics" and content.get("body"):
        b = slots["body"]
        _textbox(slide, "body", b["x"], b["y"], b["w"], b["h"] - 0.2, content["body"], 15, minor, colors["muted"])
    elif kind == "chart" and content.get("body"):
        b = slots["body"]
        _textbox(slide, "body", b["x"], b["y"], b["w"], b["h"], content["body"], 14, minor, colors["text"])
    for name in ("subtitle", "contact"):
        if content.get(name) and name in slots:
            s = slots[name]
            _textbox(slide, name, s["x"], s["y"], s["w"], s["h"], content[name], s["resolved_font"]["size_pt"], minor, colors["muted"] if name == "contact" else colors["text"])
    if content.get("note") and "note" in slots:
        s = slots["note"]
        _rect(slide, "note_plate", s["x"], s["y"], s["w"], s["h"], colors["surface"], colors["surface_line"])
        _textbox(slide, "note", s["x"] + 0.3, s["y"], s["w"] - 0.6, s["h"], content["note"], 13, minor, colors["muted"], anchor=MSO_ANCHOR.MIDDLE)

    # Таблица / график
    table_blocks = [b for b in blocks if getattr(b, "table", None) is not None]
    chart_blocks = [b for b in blocks if getattr(b, "chart", None) is not None]
    if kind == "table" and table_blocks:
        from app.pipeline.styling import style_table
        tb = table_blocks[0].table
        rows, cols = 1 + min(len(tb.rows), 8), min(len(tb.headers), 6)
        area_h = height - (1.9 if content.get("note") else 0)
        frame = slide.shapes.add_table(rows, cols, Cm(left), Cm(top), Cm(width), Cm(min(area_h, 1.1 * rows)))
        frame.name = "table"
        table = frame.table
        for c in range(cols):
            table.cell(0, c).text = str(tb.headers[c])
        for r in range(1, rows):
            for c in range(cols):
                row = tb.rows[r - 1]
                table.cell(r, c).text = str(row[c]) if c < len(row) else ""
        if manifest is not None:
            try:
                style_table(table, manifest)
            except Exception:  # noqa: BLE001
                logger.exception("Не удалось стилизовать таблицу")
    if kind == "chart" and chart_blocks:
        from app.pipeline.styling import style_chart
        block = chart_blocks[0].chart
        data = CategoryChartData()
        data.categories = block.categories
        for series in block.series:
            data.add_series(series.name, series.values)
        chart_w = width * (0.63 if content.get("body") else 1.0)
        frame = slide.shapes.add_chart(chart_type_map[block.chart_type.value], Cm(left), Cm(top), Cm(chart_w), Cm(height), data)
        frame.name = "chart"
        if manifest is not None:
            try:
                style_chart(frame.chart, manifest, slide=slide)
            except Exception:  # noqa: BLE001
                logger.exception("Не удалось стилизовать график")
    return slide


def _render_section(slide, slide_plan, content, slots, colors, major, minor, left, top, width, height, sw, sh, base) -> None:
    """Слайд-разделитель библиотеки. Отличается от контентных слайдов: название
    части крупно в средней полосе, над ним — номер части, ниже — лид приглушённым
    и (если частей внутри две и больше) список слайдов раздела. Два построения
    чередуются по порядковому номеру раздела: с акцентной полосой и с крупной
    цифрой слева."""
    ordinal = int(slide_plan.get("section_ordinal") or 1)
    total = int(slide_plan.get("sections_total") or 0)
    numbered = ordinal % 2 == 0
    title_text = content.get("title", "")
    kicker = content.get("label") or (f"Раздел {ordinal:02d}" if total >= 2 else "")
    lead = content.get("subtitle", "")
    inside = [t for t in (content.get("body") or "").split("\n") if t.strip()]

    # Заголовок-плейсхолдер переносим в среднюю полосу и увеличиваем
    title_shape = next((sh_ for sh_ in slide.shapes if sh_.name == "title"), None)
    band_top = top + height * 0.16
    col_x = left
    col_w = width * 0.68 if not _decor_right_of_title(base, slots["title"], sw, sh) else min(width * 0.68, float(slots["title"].get("w") or width))
    if numbered:
        num_w = 4.6
        _textbox(slide, "section_number", left, band_top - 0.6, num_w, 3.6, f"{ordinal:02d}", 96, major, colors["accent"], bold=True, anchor=MSO_ANCHOR.TOP, line_spacing=0.9)
        col_x = left + num_w + 0.4
        col_w = col_w - num_w - 0.4
    else:
        _rect(slide, "accent_bar", left, band_top - 0.9, 2.4, 0.35, colors["accent"], rounded=False)
    y = band_top
    if kicker and not numbered:
        _textbox(slide, "label", col_x, y, col_w, 0.9, kicker.upper(), 13, minor, colors["accent"], bold=True)
        y += 1.0
    t_size = _fit_size(title_text, col_w - 0.3, 3.6, 48, min_size=30)
    title_h = 3.6 if t_size >= 40 else 4.2
    if title_shape is not None:
        title_shape.left, title_shape.top, title_shape.width, title_shape.height = Cm(col_x), Cm(y), Cm(col_w), Cm(title_h)
        tf = title_shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for paragraph in tf.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT
            for run in paragraph.runs:
                run.font.size = Pt(t_size)
                run.font.bold = True
    else:
        _textbox(slide, "title", col_x, y, col_w, title_h, title_text, t_size, major, colors["text"], bold=True)
    y += title_h + 0.2
    if lead:
        l_size = _fit_size(lead, col_w - 0.3, 2.4, 20, min_size=14)
        _textbox(slide, "subtitle", col_x, y, col_w, 2.4, lead, l_size, minor, colors["muted"], line_spacing=1.25)
        y += 2.4
    if inside:
        # список слайдов раздела — справа, колонкой, с тонкой линией-отступом
        list_x = left + width * 0.72
        list_w = width * 0.28
        list_y = band_top + 0.2
        _rect(slide, "list_rule", list_x - 0.5, list_y, 0.06, min(1.2 * len(inside[:4]) + 0.7, height - (list_y - top)), colors["surface_line"], rounded=False)
        _textbox(slide, "list_caption", list_x, list_y, list_w, 0.6, "В ЭТОМ РАЗДЕЛЕ", 10, minor, colors["muted"], bold=True)
        for k, item in enumerate(inside[:4]):
            _textbox(slide, f"list_item_{k + 1}", list_x, list_y + 0.7 + 1.2 * k, list_w, 1.15, item, 14, minor, colors["text"], line_spacing=1.15)


def _draw_units(slide, kind, n, geometry, unit_slots, slots, content, colors, major, minor) -> None:
    for k in range(n):
        g = geometry[k]
        names = unit_slots[k] if k < len(unit_slots) else []
        by_role = {slots[nm]["semantic_role"]: nm for nm in names if nm in slots}
        get = lambda role: content.get(by_role.get(role, ""), "")  # noqa: E731
        if kind in {"cards", "problems"}:
            _rect(slide, f"card_{k + 1}", g["x"], g["y"], g["w"], g["h"], colors["surface"], colors["surface_line"])
            pad = 0.5
            number = get("unit_number") or f"{k + 1:02d}"
            _textbox(slide, by_role.get("unit_number", f"u{k + 1}_num"), g["x"] + pad, g["y"] + pad, 2.0, 1.0, number, 14, major, colors["accent"], bold=True)
            if kind == "problems":
                _rect(slide, f"mark_{k + 1}", g["x"], g["y"] + 0.9, 0.18, g["h"] - 1.8, colors["accent"], rounded=False)
            title_h = 1.6 if g["h"] > 5 else 1.2
            if get("unit_title"):
                size = _fit_size(get("unit_title"), g["w"] - 2 * pad, title_h, 18 if g["w"] > 7 else 15, 12)
                _textbox(slide, by_role.get("unit_title", f"u{k + 1}_title"), g["x"] + pad, g["y"] + pad + 1.1, g["w"] - 2 * pad, title_h, get("unit_title"), size, major, colors["text"], bold=True)
            if get("unit_text"):
                y = g["y"] + pad + 1.1 + title_h + 0.2
                h = g["y"] + g["h"] - pad - y
                if h > 0.8:
                    size = _fit_size(get("unit_text"), g["w"] - 2 * pad, h, 14 if g["w"] > 7 else 12, 10)
                    _textbox(slide, by_role.get("unit_text", f"u{k + 1}_text"), g["x"] + pad, y, g["w"] - 2 * pad, h, get("unit_text"), size, minor, colors["muted"], line_spacing=1.15)
        elif kind == "columns":
            _rect(slide, f"col_bar_{k + 1}", g["x"], g["y"], g["w"] * 0.35, 0.2, colors["accent"], rounded=False)
            if get("unit_title"):
                _textbox(slide, by_role.get("unit_title", f"u{k + 1}_title"), g["x"], g["y"] + 0.5, g["w"], 1.6, get("unit_title"), 20, major, colors["text"], bold=True)
            if get("unit_text"):
                size = _fit_size(get("unit_text"), g["w"], g["h"] - 2.4, 16, 11)
                _textbox(slide, by_role.get("unit_text", f"u{k + 1}_text"), g["x"], g["y"] + 2.3, g["w"], g["h"] - 2.4, get("unit_text"), size, minor, colors["text"], line_spacing=1.2)
        elif kind == "metrics":
            value = get("metric_value")
            if value:
                size = _fit_size(value, g["w"], g["h"] * 0.6, 54 if n <= 3 else 44, 24)
                if "\n" not in value:
                    # значение — одной строкой («120 млн», не «120 / млн»): кегль по ширине блока
                    one_line = (g["w"] - 0.3) * 28.35 / max(1, len(value)) / 0.62
                    size = max(24.0, min(size, math.floor(one_line)))
                _textbox(slide, by_role.get("metric_value", f"u{k + 1}_value"), g["x"], g["y"], g["w"], g["h"] * 0.6, value.replace(" ", "\u00a0"), size, major, colors["accent"], bold=True, anchor=MSO_ANCHOR.BOTTOM)
            _rect(slide, f"metric_bar_{k + 1}", g["x"], g["y"] + g["h"] * 0.6 + 0.15, 1.6, 0.14, colors["accent"], rounded=False)
            if get("metric_label"):
                _textbox(slide, by_role.get("metric_label", f"u{k + 1}_caption"), g["x"], g["y"] + g["h"] * 0.6 + 0.5, g["w"], g["h"] * 0.4 - 0.5, get("metric_label"), 15, minor, colors["muted"])
        elif kind == "timeline":
            line_y = g["y"] + 1.5
            if k == 0:
                total_w = geometry[-1]["x"] + geometry[-1]["w"] - g["x"]
                _rect(slide, "timeline_line", g["x"], line_y - 0.03, total_w, 0.06, colors["surface_line"], rounded=False)
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(g["x"]), Cm(line_y - 0.35), Cm(0.7), Cm(0.7))
            dot.name = f"timeline_dot_{k + 1}"
            dot.fill.solid()
            dot.fill.fore_color.rgb = colors["accent"]
            dot.line.fill.background()
            label = get("unit_label") or get("unit_title")
            if label:
                _textbox(slide, by_role.get("unit_label", f"u{k + 1}_label"), g["x"], g["y"], g["w"] - 0.4, 1.0, label, 14, major, colors["accent"], bold=True)
            if get("unit_title") and get("unit_label"):
                _textbox(slide, by_role.get("unit_title", f"u{k + 1}_title"), g["x"], g["y"] + 2.0, g["w"] - 0.4, 1.6, get("unit_title"), 16, major, colors["text"], bold=True)
            if get("unit_text"):
                y = g["y"] + (3.7 if get("unit_title") and get("unit_label") else 2.0)
                size = _fit_size(get("unit_text"), g["w"] - 0.4, g["h"] - y + g["y"], 13, 10)
                _textbox(slide, by_role.get("unit_text", f"u{k + 1}_text"), g["x"], y, g["w"] - 0.4, g["h"] - (y - g["y"]), get("unit_text"), size, minor, colors["muted"], line_spacing=1.15)
        elif kind == "agenda":
            number = get("unit_number") or f"{k + 1:02d}"
            _textbox(slide, by_role.get("unit_number", f"u{k + 1}_num"), g["x"], g["y"], 1.8, g["h"], number, 22, major, colors["accent"], bold=True, anchor=MSO_ANCHOR.MIDDLE)
            if get("agenda_item"):
                _textbox(slide, by_role.get("agenda_item", f"u{k + 1}_item"), g["x"] + 2.0, g["y"], g["w"] - 2.0, g["h"], get("agenda_item"), 18, minor, colors["text"], anchor=MSO_ANCHOR.MIDDLE)
            _rect(slide, f"agenda_line_{k + 1}", g["x"], g["y"] + g["h"] - 0.02, g["w"], 0.04, colors["surface_line"], rounded=False)


_ = (Emu,)
