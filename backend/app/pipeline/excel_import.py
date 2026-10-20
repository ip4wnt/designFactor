"""Import Excel using data properties, preserving missing observations as gaps."""
from __future__ import annotations

import datetime
import math
import re
from pathlib import Path
from typing import Any, Literal

import openpyxl

from app.schemas.content_plan import ChartData, ChartSeries, ChartType, TableData

MAX_ROWS = 500
MAX_COLS = 50


class ExcelImportError(Exception):
    """Лист Excel не удалось разобрать или он пуст."""


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        return str(int(value)) if value == int(value) else str(round(value, 6))
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return str(value)


def _cell_to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    try:
        number = float(str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", "."))
        return number if math.isfinite(number) else None
    except ValueError:
        return None


def read_sheet_rows(path: str | Path, sheet_name: str | None = None) -> tuple[str, list[list[Any]]]:
    """Читает сырые значения ячеек одного листа (по умолчанию — первый лист)."""
    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]

        rows: list[list[Any]] = []
        for row in worksheet.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS):
            rows.append([cell.value for cell in row])

        while rows and all(v is None or v == "" for v in rows[-1]):
            rows.pop()

        max_col_used = 0
        for row in rows:
            for i in range(len(row) - 1, -1, -1):
                if row[i] not in (None, ""):
                    max_col_used = max(max_col_used, i + 1)
                    break
        rows = [row[:max_col_used] for row in rows]

        if not rows or max_col_used == 0:
            raise ExcelImportError(f"Лист «{worksheet.title}» пуст")

        return worksheet.title, rows
    finally:
        workbook.close()


def list_sheet_names(path: str | Path) -> list[str]:
    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def sheet_to_table_data(sheet_name: str, rows: list[list[Any]]) -> TableData:
    header_row, *data_rows = rows
    headers = [_cell_to_str(v) or f"col_{i + 1}" for i, v in enumerate(header_row)]
    table_rows = [[_cell_to_str(v) for v in row] + [""] * (len(headers) - len(row)) for row in data_rows]
    return TableData(headers=headers, rows=table_rows)


def sheet_to_chart_data(sheet_name: str, rows: list[list[Any]], chart_type: ChartType = ChartType.BAR) -> ChartData | None:
    """Возвращает None для нечисловых данных или повторяющихся имён серий —
    вызывающая сторона откатывается на TableData либо сообщает ошибку."""
    header_row, *data_rows = rows
    series_names = [_cell_to_str(v) or f"series_{i + 1}" for i, v in enumerate(header_row[1:], start=1)]
    if len(series_names) != len(set(series_names)):
        return None  # не объединяем столбцы с одинаковыми заголовками

    categories: list[str] = []
    series_values: dict[str, list[float | None]] = {name: [] for name in series_names}

    for row in data_rows:
        category = _cell_to_str(row[0]) if row else ""
        categories.append(category)
        for col_idx, series_name in enumerate(series_names, start=1):
            raw = row[col_idx] if col_idx < len(row) else None
            value = _cell_to_float(raw)
            if value is None and raw not in (None, ""):
                return None  # Invalid values must not silently become gaps.
            series_values[series_name].append(value)

    if not categories or not series_names or any(
        not any(value is not None for value in values) for values in series_values.values()
    ):
        return None

    return ChartData(
        chart_type=chart_type,
        category_axis_title=_cell_to_str(header_row[0]),
        value_axis_title=series_names[0] if len(series_names) == 1 else "",
        categories=categories,
        series=[
            ChartSeries(name=name, values=values)
            for name, values in series_values.items()
        ],
    )


def _date_key(value: Any) -> datetime.datetime | None:
    if isinstance(value, datetime.datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, datetime.date):
        return datetime.datetime.combine(value, datetime.time())
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m", "%d.%m.%Y", "%Y/%m/%d"):
        try:
            return datetime.datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def select_chart_type(rows: list[list[Any]], goal: str = "") -> ChartType:
    """Conservative automatic choice; explicit comparison overrides time series."""
    labels = [row[0] for row in rows[1:] if row]
    goal = goal.lower()
    if re.search(r"сравн|compare|comparison", goal):
        return ChartType.BAR
    if len(labels) >= 2 and all(_date_key(label) is not None for label in labels):
        return ChartType.LINE
    if re.search(r"дол[яьи]|структур|share|composition", goal) and len(rows[0]) == 2:
        values = [_cell_to_float(row[1]) if len(row) > 1 else None for row in rows[1:]]
        if 2 <= len(values) <= 6 and all(v is not None and v >= 0 for v in values) and sum(values) > 0:
            return ChartType.PIE
    return ChartType.BAR


def import_excel_block(
    path: str | Path,
    sheet_name: str | None = None,
    kind: Literal["table", "chart"] | None = None,
    chart_type: ChartType | None = None,
    goal: str = "",
) -> tuple[Literal["table", "chart"], TableData | ChartData]:
    """Choose by column values and user goal, never by file or worksheet name.

    Large reference tables remain tables; chart conversion does not truncate
    categories or aggregate records. Sheet reading retains MAX_ROWS/MAX_COLS limits. Explicit chart/table requests take priority.
    """
    resolved_name, rows = read_sheet_rows(path, sheet_name)
    if kind == "table":
        return "table", sheet_to_table_data(resolved_name, rows)
    selected = chart_type or select_chart_type(rows, goal)
    if selected == ChartType.LINE and all(_date_key(row[0]) is not None for row in rows[1:] if row):
        rows = [rows[0]] + sorted(rows[1:], key=lambda row: _date_key(row[0]))
    chart = sheet_to_chart_data(resolved_name, rows, selected)
    if chart is not None and (kind == "chart" or (len(rows) <= 61 and len(rows[0]) <= 7)):
        return "chart", chart
    if kind == "chart":
        raise ExcelImportError(
            f"Лист «{resolved_name}»: нужны числовые ряды с уникальными заголовками; "
            "пустые ячейки допустимы, текст в числовых столбцах — нет."
        )
    return "table", sheet_to_table_data(resolved_name, rows)
