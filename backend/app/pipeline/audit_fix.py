"""Точечные исправления замечаний аудита прямо в .pptx выбранного варианта.

Аудит (app/pipeline/audit.py) отдаёт каждому замечанию машиночитаемый `code`
и `shape_id` фигуры на слайде. Здесь по коду выбирается детерминированное
(или, для текста, LLM-) исправление, которое правит ТОЛЬКО эту фигуру, не
пересобирая колоду: пользователь кликает красную рамку на превью → «Исправить»
→ фигура правится → аудит варианта запускается заново → рамка зеленеет.

Что исправляется:

| code                  | действие                                                       |
| --------------------- | -------------------------------------------------------------- |
| bullet_too_long       | LLM сокращает абзац до ≤ 14 слов (фолбэк: обрез по знаку препинания) |
| too_many_bullets      | LLM сжимает список до 6 тезисов (фолбэк: оставить первые 6)   |
| text_overflow         | пропорциональное уменьшение кегля всех ранов фигуры (не ниже 60 %) |
| font_size_off_scale   | кегль ранов → ближайший из шкалы шаблона                       |
| font_off_template     | гарнитура ранов → основная гарнитура шаблона                   |
| color_off_palette     | цвет текста → ближайший цвет палитры шаблона                   |
| low_contrast          | цвет текста → цвет палитры с максимальным контрастом к фону    |
| out_of_bounds         | фигура сдвигается/ужимается внутрь слайда                      |
| edge_margin           | фигура сдвигается внутрь на величину поля                      |
| image_distorted       | высота картинки приводится к исходным пропорциям               |
| chart_no_annotations  | включается легенда диаграммы                                   |

Остальные коды (таблица/диаграмма слишком большие, растровый слайд, пустой
слайд, плотность слайда, VLM-вопросы) не имеют безопасного точечного
исправления — фронтенд показывает их как «требует ручной правки».
"""
from __future__ import annotations

import logging
import math
import re
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.util import Emu, Pt

from app.skills import skill_text
from pydantic import BaseModel

from app.config import get_settings
from app.llm_client import LLMClient
from app.schemas.design_manifest import DesignManifest
from app.schemas.job import AuditIssue

logger = logging.getLogger(__name__)

FIXABLE_CODES: frozenset[str] = frozenset(
    {
        "bullet_too_long",
        "too_many_bullets",
        "text_overflow",
        "font_size_off_scale",
        "font_off_template",
        "color_off_palette",
        "low_contrast",
        "out_of_bounds",
        "edge_margin",
        "image_distorted",
        "chart_no_annotations",
    }
)

_MAX_WORDS = 14
_MAX_BULLETS = 6
_EDGE_MARGIN_RATIO = 0.03
_MIN_SHRINK = 0.6


class _Rewrite(BaseModel):
    texts: list[str]


class _Condense(BaseModel):
    bullets: list[str]


async def apply_fixes(
    pptx_path: str | Path,
    issues: Iterable[AuditIssue],
    manifest: DesignManifest | None,
    allowed_fonts: Iterable[str] | None = None,
    allowed_sizes: Iterable[float] | None = None,
) -> list[str]:
    """Применяет исправления к .pptx на месте. Возвращает issue_id, для которых
    исправление реально выполнено."""
    pptx_path = Path(pptx_path)
    prs = Presentation(str(pptx_path))
    slides = {f"slide-{index + 1}": slide for index, slide in enumerate(prs.slides)}
    fonts = [f for f in (allowed_fonts or []) if f] or _manifest_fonts(manifest)
    sizes = sorted({float(s) for s in (allowed_sizes or []) if s}) or _manifest_sizes(manifest)
    palette = _manifest_palette(manifest)

    applied: list[str] = []
    text_jobs: list[tuple[AuditIssue, object]] = []
    for issue in issues:
        if issue.code not in FIXABLE_CODES or issue.shape_id is None:
            continue
        slide = slides.get(issue.slide_id)
        shape = _find_shape(slide, issue.shape_id) if slide is not None else None
        if shape is None:
            continue
        try:
            if issue.code in {"bullet_too_long", "too_many_bullets"}:
                text_jobs.append((issue, shape))
                continue
            ok = _apply_deterministic(issue.code, shape, prs, fonts, sizes, palette)
        except Exception:  # noqa: BLE001 — одно неудачное исправление не должно ронять остальные
            logger.exception("Исправление %s (%s) не применилось", issue.code, issue.issue_id)
            ok = False
        if ok:
            applied.append(issue.issue_id)

    if text_jobs:
        applied.extend(await _apply_text_fixes(text_jobs))

    if applied:
        prs.save(str(pptx_path))
    return applied


# --------------------------------------------------------------------------
# Детерминированные исправления
# --------------------------------------------------------------------------


def _apply_deterministic(code: str, shape, prs, fonts: list[str], sizes: list[float], palette: list[RGBColor]) -> bool:
    if code == "text_overflow":
        return _shrink_text(shape)
    if code == "font_size_off_scale":
        return _snap_sizes(shape, sizes)
    if code == "font_off_template":
        return _set_font(shape, fonts)
    if code == "color_off_palette":
        return _snap_colors(shape, palette)
    if code == "low_contrast":
        return _fix_contrast(shape, palette)
    if code == "out_of_bounds":
        return _move_into_bounds(shape, prs.slide_width, prs.slide_height, margin_ratio=0.0)
    if code == "edge_margin":
        return _move_into_bounds(shape, prs.slide_width, prs.slide_height, margin_ratio=_EDGE_MARGIN_RATIO)
    if code == "image_distorted":
        return _fix_aspect(shape)
    if code == "chart_no_annotations":
        return _enable_legend(shape)
    return False


def _runs(shape):
    if not getattr(shape, "has_text_frame", False):
        return []
    return [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]


def _run_size_pt(run, paragraph_default: float = 18.0) -> float:
    return run.font.size.pt if run.font.size is not None else paragraph_default


def _shrink_text(shape) -> bool:
    """Оценка та же, что в audit._check_text_overflow: нужное число строк vs.
    доступное; кегль уменьшаем на sqrt(available/needed) — и символов в строке,
    и строк по высоте становится больше."""
    tf = shape.text_frame
    width_cm = Emu(shape.width).cm
    height_cm = Emu(shape.height).cm
    if width_cm <= 0 or height_cm <= 0:
        return False
    sizes = [_run_size_pt(run) for run in _runs(shape)] or [18.0]
    avg = sum(sizes) / len(sizes)
    needed = 0
    for paragraph in tf.paragraphs:
        text = paragraph.text
        if not text.strip():
            needed += 1
            continue
        chars_per_line = max(4.3 * (10.0 / avg) * width_cm, 1.0)
        needed += max(1, -(-len(text) // int(chars_per_line)))
    line_h_cm = (avg / 72) * 2.54 * 1.2
    available = height_cm / line_h_cm if line_h_cm > 0 else 0
    if needed <= available:
        return False
    factor = max(_MIN_SHRINK, math.sqrt(available / needed) * 0.95)
    for run in _runs(shape):
        run.font.size = Pt(max(8.0, round(_run_size_pt(run) * factor)))
    tf.word_wrap = True
    return True


def _snap_sizes(shape, sizes: list[float]) -> bool:
    if not sizes:
        return False
    changed = False
    for run in _runs(shape):
        if run.font.size is None:
            continue
        current = run.font.size.pt
        if current in sizes:
            continue
        lower = [s for s in sizes if s <= current]
        nearest = min(sizes, key=lambda s: abs(s - current))
        # Больший кегль допустим, если он не более чем на 8 % крупнее — иначе
        # текст, который только что подогнали, снова переполнит рамку.
        target = nearest if nearest <= current * 1.08 else (max(lower) if lower else nearest)
        run.font.size = Pt(target)
        changed = True
    return changed


def _set_font(shape, fonts: list[str]) -> bool:
    if not fonts:
        return False
    changed = False
    for run in _runs(shape):
        if run.font.name and run.font.name not in fonts:
            run.font.name = fonts[0]
            changed = True
    return changed


def _snap_colors(shape, palette: list[RGBColor]) -> bool:
    if not palette:
        return False
    changed = False
    for run in _runs(shape):
        rgb = _run_rgb(run)
        if rgb is None or rgb in palette:
            continue
        run.font.color.rgb = min(palette, key=lambda c: _color_distance(c, rgb))
        changed = True
    return changed


def _fix_contrast(shape, palette: list[RGBColor]) -> bool:
    background = _shape_fill_rgb(shape) or RGBColor(0xFF, 0xFF, 0xFF)
    candidates = list(palette) + [RGBColor(0x1A, 0x1A, 0x1A), RGBColor(0xFF, 0xFF, 0xFF)]
    best = max(candidates, key=lambda c: _contrast(c, background))
    changed = False
    for run in _runs(shape):
        rgb = _run_rgb(run)
        if rgb is not None and _contrast(rgb, background) < 4.5:
            run.font.color.rgb = best
            changed = True
    return changed


def _move_into_bounds(shape, slide_w: int, slide_h: int, margin_ratio: float) -> bool:
    mx, my = int(slide_w * margin_ratio), int(slide_h * margin_ratio)
    left, top, width, height = int(shape.left), int(shape.top), int(shape.width), int(shape.height)
    new_w = min(width, slide_w - 2 * mx)
    new_h = min(height, slide_h - 2 * my)
    new_left = min(max(left, mx), slide_w - mx - new_w)
    new_top = min(max(top, my), slide_h - my - new_h)
    if (new_left, new_top, new_w, new_h) == (left, top, width, height):
        return False
    shape.left, shape.top, shape.width, shape.height = Emu(new_left), Emu(new_top), Emu(new_w), Emu(new_h)
    return True


def _fix_aspect(shape) -> bool:
    try:
        px_w, px_h = shape.image.size
    except (AttributeError, ValueError):
        return False
    if not px_w or not px_h or not shape.width:
        return False
    target_h = int(shape.width * px_h / px_w)
    if abs(target_h - int(shape.height)) < 12700:
        return False
    shape.height = Emu(target_h)
    return True


def _enable_legend(shape) -> bool:
    if not getattr(shape, "has_chart", False):
        return False
    chart = shape.chart
    if chart.has_legend:
        return False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    return True


# --------------------------------------------------------------------------
# Текстовые исправления через LLM (с детерминированным фолбэком)
# --------------------------------------------------------------------------

# Промпты лежат в skills/audit_text_fix.yaml (не в коде).
_SKILL = "audit_text_fix"


def _rewrite_system() -> str:
    return skill_text(_SKILL, "rewrite_system_prompt")


def _condense_system() -> str:
    return skill_text(_SKILL, "condense_system_prompt")


async def _apply_text_fixes(jobs: list[tuple[AuditIssue, object]]) -> list[str]:
    applied: list[str] = []
    settings = get_settings()
    client: LLMClient | None = None
    try:
        client = LLMClient(
            base_url=settings.LLM_BASE_URL,
            model_name=settings.LLM_MODEL_NAME,
            api_key=settings.LLM_API_KEY,
            max_retries=0,
            total_timeout=settings.LLM_AUDIT_FIX_TIMEOUT_S,
        )
    except Exception:  # noqa: BLE001
        client = None

    try:
        for issue, shape in jobs:
            try:
                if issue.code == "bullet_too_long":
                    ok = await _fix_long_bullets(shape, client)
                else:
                    ok = await _fix_too_many_bullets(shape, client)
            except Exception:  # noqa: BLE001
                logger.exception("Текстовое исправление %s не применилось", issue.issue_id)
                ok = False
            if ok:
                applied.append(issue.issue_id)
    finally:
        if client is not None:
            await client.aclose()
    return applied


def _word_count(text: str) -> int:
    return len([w for w in re.split(r"\s+", text.strip()) if w])


def _set_paragraph_text(paragraph, text: str) -> None:
    """Меняет текст абзаца, сохраняя форматирование первого рана."""
    runs = paragraph.runs
    if not runs:
        paragraph.text = text
        return
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


async def _fix_long_bullets(shape, client: LLMClient | None) -> bool:
    paragraphs = [p for p in shape.text_frame.paragraphs if _word_count(p.text) > _MAX_WORDS + 1]
    if not paragraphs:
        return False
    originals = [p.text.strip() for p in paragraphs]
    rewritten: list[str] | None = None
    if client is not None:
        try:
            result = await client.chat_structured(
                system_prompt=_rewrite_system(),
                user_input="\n".join(f"{i + 1}. {t}" for i, t in enumerate(originals)),
                output_model=_Rewrite,
                max_completion_tokens=1024,
            )
            if len(result.texts) == len(originals) and all(t.strip() for t in result.texts):
                rewritten = [t.strip() for t in result.texts]
        except Exception as exc:  # noqa: BLE001
            logger.warning("LLM-сокращение тезисов недоступно, фолбэк: %s", exc)
    if rewritten is None:
        rewritten = [_trim_words(t) for t in originals]
    for paragraph, text in zip(paragraphs, rewritten):
        if _word_count(text) > _MAX_WORDS + 1:
            text = _trim_words(text)
        _set_paragraph_text(paragraph, text)
    return True


async def _fix_too_many_bullets(shape, client: LLMClient | None) -> bool:
    paragraphs = [p for p in shape.text_frame.paragraphs if p.text.strip()]
    if len(paragraphs) <= _MAX_BULLETS:
        return False
    originals = [p.text.strip() for p in paragraphs]
    bullets: list[str] | None = None
    if client is not None:
        try:
            result = await client.chat_structured(
                system_prompt=_condense_system(),
                user_input="\n".join(f"- {t}" for t in originals),
                output_model=_Condense,
                max_completion_tokens=1024,
            )
            if 1 <= len(result.bullets) <= _MAX_BULLETS:
                bullets = [b.strip() for b in result.bullets if b.strip()]
        except Exception as exc:  # noqa: BLE001
            logger.warning("LLM-сжатие списка недоступно, фолбэк: %s", exc)
    if not bullets:
        bullets = originals[:_MAX_BULLETS]
    for paragraph, text in zip(paragraphs, bullets):
        _set_paragraph_text(paragraph, text)
    for paragraph in paragraphs[len(bullets):]:
        paragraph._p.getparent().remove(paragraph._p)
    return True


def _trim_words(text: str) -> str:
    """Фолбэк без LLM: обрезаем по ближайшему знаку препинания до лимита слов,
    иначе — по границе слова."""
    words = [w for w in re.split(r"\s+", text.strip()) if w]
    if len(words) <= _MAX_WORDS:
        return text.strip()
    head = " ".join(words[:_MAX_WORDS])
    for separator in (". ", "; ", " — ", ": ", ", "):
        cut = head.rfind(separator)
        if cut >= len(head) * 0.5:
            return head[:cut].rstrip(" ,;:—-") + ("." if separator == ". " else "")
    return head.rstrip(" ,;:—-")


# --------------------------------------------------------------------------
# Вспомогательные
# --------------------------------------------------------------------------


def _find_shape(slide, shape_id: int):
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    return None


def _manifest_fonts(manifest: DesignManifest | None) -> list[str]:
    if manifest is None:
        return []
    fonts: list[str] = []
    for style in (manifest.typography.body, manifest.typography.title):
        if style.font and style.font not in fonts:
            fonts.append(style.font)
    return fonts


def _manifest_sizes(manifest: DesignManifest | None) -> list[float]:
    if manifest is None:
        return []
    return sorted({float(style.size) for style in (manifest.typography.title, manifest.typography.body) if style.size})


def _manifest_palette(manifest: DesignManifest | None) -> list[RGBColor]:
    if manifest is None:
        return []
    colors: list[RGBColor] = []
    for value in manifest.palette.values():
        try:
            rgb = RGBColor.from_string(value.lstrip("#").upper())
        except (ValueError, TypeError):
            continue
        if rgb not in colors:
            colors.append(rgb)
    return colors


def _run_rgb(run) -> RGBColor | None:
    try:
        if run.font.color is not None and run.font.color.type is not None:
            return run.font.color.rgb
    except (AttributeError, TypeError):
        return None
    return None


def _shape_fill_rgb(shape) -> RGBColor | None:
    try:
        if shape.fill.type == 1:
            return shape.fill.fore_color.rgb
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def _color_distance(a: RGBColor, b: RGBColor) -> float:
    return sum((int(x) - int(y)) ** 2 for x, y in zip(a, b))


def _luminance(rgb: RGBColor) -> float:
    def channel(c: int) -> float:
        c = c / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4

    r, g, b = (channel(int(x)) for x in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _contrast(a: RGBColor, b: RGBColor) -> float:
    la, lb = _luminance(a), _luminance(b)
    lighter, darker = max(la, lb), min(la, lb)
    return (lighter + 0.05) / (darker + 0.05)
