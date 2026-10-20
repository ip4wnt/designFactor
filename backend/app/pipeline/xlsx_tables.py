"""XLSX worksheets -> server-owned tables shared with the DOCX analysis pipeline."""
from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from app.pipeline.docx_tables import MAX_COLUMNS, MAX_DATA_ROWS, MAX_TABLES
from app.pipeline.materials import MaterialError
from app.schemas.content_plan import TableData
from app.schemas.source_tables import DocumentMaterial, SourceTable

MAX_ROWS = 100_000
MAX_ANALYSIS_COLUMNS = 100
MAX_EXPANDED_BYTES = 100 * 1024 * 1024


def _value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def _sheet_table(sheet, formulas, material_id, filename, ordinal):
    origin = f'Лист «{sheet.title}»'
    table = SourceTable(table_id=f"{material_id}.table_{ordinal}", source_name=filename,
                        ordinal=ordinal, sheet_name=sheet.title, summary=origin, available=False)
    # Some producers omit dimensions or leave stale/formatting-only dimensions.
    sheet.reset_dimensions()
    formulas.reset_dimensions()
    rows, formats, formula_gaps, errors = [], {}, [], []
    first_row = last_row = None
    first_col, last_col = MAX_ANALYSIS_COLUMNS, -1
    for row_number, (values, expressions) in enumerate(zip(sheet.iter_rows(), formulas.iter_rows()), 1):
        if row_number > MAX_ROWS + 1:
            table.limitation = "Превышен лимит анализа: 100000 строк данных или 100 столбцов"
            return table
        populated = [i for i, (cell, expression) in enumerate(zip(values, expressions))
                     if cell.value is not None or expression.value is not None]
        if not populated:
            if first_row is not None:
                rows.append([])
            continue
        if max(populated) >= MAX_ANALYSIS_COLUMNS:
            table.limitation = "Превышен лимит анализа: 100000 строк данных или 100 столбцов"
            return table
        if first_row is None:
            first_row = row_number
        last_row = row_number
        first_col, last_col = min(first_col, min(populated)), max(last_col, max(populated))
        rows.append([_value(cell.value) for cell in values])
        for i in populated:
            cell, expression = values[i], expressions[i]
            if expression.data_type == "f" and cell.value is None:
                if len(formula_gaps) < 3:
                    formula_gaps.append(expression.coordinate)
            if cell.data_type == "e":
                if len(errors) < 3:
                    errors.append(cell.coordinate)
            if row_number > first_row and cell.value is not None and cell.number_format != "General":
                formats.setdefault(i, set()).add(cell.number_format)
    if first_row is None:
        return None
    rows = [(row + [""] * (last_col + 1 - len(row)))[first_col:last_col + 1]
            for row in rows[:last_row - first_row + 1]]
    table.cell_range = f"{get_column_letter(first_col + 1)}{first_row}:{get_column_letter(last_col + 1)}{last_row}"
    table.summary = origin + ", диапазон " + table.cell_range + ". Колонки: " + "; ".join(rows[0])
    if formula_gaps:
        table.limitation = ("Нет сохраненных результатов формул в " + ", ".join(formula_gaps) +
                            "; пересчитайте книгу в Excel или LibreOffice и загрузите заново")
    elif errors:
        table.limitation = "В ячейках ошибки Excel: " + ", ".join(errors) + "; исправьте их и загрузите книгу заново"
    elif len(rows) < 2 or any(not h.strip() for h in rows[0]) or len(set(rows[0])) != len(rows[0]):
        table.limitation = "Первая непустая строка листа должна содержать непустые уникальные заголовки и хотя бы одну строку данных"
    else:
        table.data = TableData(headers=rows[0], rows=rows[1:])
        table.number_formats = {rows[0][i - first_col]: sorted(values)
                                for i, values in formats.items() if first_col <= i <= last_col}
        if len(rows) - 1 > MAX_DATA_ROWS or len(rows[0]) > MAX_COLUMNS:
            table.limitation = f"Для слайда допустимо до {MAX_DATA_ROWS} строк данных и {MAX_COLUMNS} столбцов; таблица не обрезана"
        elif any(len(c) > 160 for row in rows for c in row) or sum(len(c) for row in rows for c in row) > 2000:
            table.limitation = "Слишком длинный текст в ячейках для одного слайда; таблица не сокращена"
        else:
            table.available = True
    return table


def extract_xlsx(filename: str, content: bytes) -> DocumentMaterial:
    """Read stored values, never execute formulas or replace missing caches with zero."""
    values = expressions = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise MaterialError("Распакованное содержимое XLSX больше 100 МБ")
            merged_sheets = set()
            # Read-only openpyxl does not expose merged ranges; detect them explicitly.
            for name in archive.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    with archive.open(name) as xml:
                        for _, element in etree.iterparse(xml, events=("end",), resolve_entities=False, no_network=True):
                            if etree.QName(element).localname == "mergeCell":
                                merged_sheets.add(name)
                            element.clear()
        values = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        expressions = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        material_id = "material_" + uuid4().hex
        tables, text = [], []
        for sheet, formulas in zip(values.worksheets, expressions.worksheets):
            table = _sheet_table(sheet, formulas, material_id, filename, len(tables) + 1)
            if table is None:
                continue
            if len(tables) >= MAX_TABLES:
                raise MaterialError(f"В XLSX больше {MAX_TABLES} непустых листов; разделите книгу")
            if sheet._worksheet_path.lstrip("/") in merged_sheets:
                table.available, table.data = False, None
                table.limitation = "Объединенные ячейки пока не поддерживаются; нужен лист с одной прямоугольной таблицей"
            tables.append(table)
            text.extend([f'Лист «{sheet.title}»', f"[[{table.table_id}]]"])
        if not tables:
            raise MaterialError("В XLSX не найдено непустых листов с данными")
        return DocumentMaterial(material_id=material_id, filename=filename, text="\n\n".join(text), tables=tables)
    except MaterialError:
        raise
    except (BadZipFile, InvalidFileException, KeyError, ValueError, TypeError, OSError, etree.XMLSyntaxError) as exc:
        raise MaterialError("Файл .xlsx поврежден или не является книгой Excel") from exc
    finally:
        if values is not None:
            values.close()
        if expressions is not None:
            expressions.close()
