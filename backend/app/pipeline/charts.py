"""Shared native chart rendering for every slide assembly path."""
from __future__ import annotations

import logging
import math
import re

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_AXIS_CROSSES, XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.util import Inches, Pt

from app.pipeline.styling import slide_is_dark, style_chart
from app.schemas.content_plan import ChartData, ChartType
from app.schemas.design_manifest import DesignManifest


logger = logging.getLogger(__name__)


def _label_every_category(category_axis, category_count: int, series_count: int) -> None:
    from pptx.oxml.ns import qn
    from lxml import etree
    element = category_axis._element
    for tag in ("c:tickLblSkip", "c:tickMarkSkip"):
        old = element.find(qn(tag))
        if old is not None:
            element.remove(old)
    if category_count * max(1, series_count) <= 48:
        skip = etree.SubElement(element, qn("c:tickLblSkip"))
        skip.set("val", "1")
        # Схема OOXML требует порядок: noMultiLvlLbl идет последним.
        tail = element.find(qn("c:noMultiLvlLbl"))
        if tail is not None:
            element.remove(tail)
            element.append(tail)


def use_horizontal_bars(data: ChartData) -> bool:
    if data.chart_type != ChartType.BAR:
        return False
    if data.orientation != "auto":
        return data.orientation == "horizontal"
    return any(len(label) > 18 for label in data.categories)


def add_chart(slide, data: ChartData, bbox, manifest: DesignManifest | None, chart_type_map=None):
    """Keep source text inside the allocated box; retain all labels and values."""
    left, top, width, height = (int(value) for value in bbox)
    horizontal = use_horizontal_bars(data)
    types = chart_type_map or {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else types[data.chart_type.value]
    font = manifest.typography.body.font if manifest else "Arial"
    base_size = manifest.typography.body.size if manifest else 18
    if manifest and manifest.chart_style:
        font = manifest.chart_style.font or font
        base_size = manifest.chart_style.font_size or base_size
    axis_size = max(10, min(18, base_size - 4))
    source_size = max(10, min(12, base_size - 7))
    try:
        dark = slide_is_dark(slide)
    except Exception:
        logger.exception("Не удалось определить фон графика, используется светлый фон")
        dark = False
    text_color = RGBColor.from_string("FFFFFF" if dark else "18233B")
    footer_height = 0
    if data.source:
        # Account for wrapping before allocating chart height, never cover axes.
        chars_per_line = max(15, int(width / Inches(1) * 9))
        line_count = sum(max(1, math.ceil(len(line) / chars_per_line)) for line in ("Источник: " + data.source).splitlines())
        footer_height = min(int(height * .3), int(Inches(.10 + .19 * line_count)))
    chart_data = CategoryChartData()
    chart_data.categories = data.categories
    for series in data.series:
        chart_data.add_series(series.name, series.values)
    frame = slide.shapes.add_chart(kind, left, top, width, height - footer_height, chart_data)
    chart = frame.chart
    if manifest:
        try:
            style_chart(chart, manifest, dark=dark, slide=slide)
            text_color = chart.font.color.rgb
        except Exception:
            logger.exception("Не удалось применить стиль графика, сохранён нативный график")
    if data.chart_type != ChartType.PIE:
        value_title = data.value_axis_title or (data.series[0].name if len(data.series) == 1 else "")
        if data.units and not re.search(r"(?:^|[\s,(])" + re.escape(data.units) + r"(?:$|[\s,)])", value_title, re.IGNORECASE):
            value_title = f"{value_title}, {data.units}" if value_title else data.units
        for axis, title in ((chart.category_axis, data.category_axis_title), (chart.value_axis, value_title)):
            if not title:
                continue
            axis.has_title = True
            axis.axis_title.text_frame.text = title
            for paragraph in axis.axis_title.text_frame.paragraphs:
                paragraph.font.name = font
                paragraph.font.size = Pt(axis_size)
                paragraph.font.color.rgb = text_color
        # Все категории подписаны: без tickLblSkip=1 PowerPoint и LibreOffice
        # прореживают подписи (каждый второй регион без названия).
        _label_every_category(chart.category_axis, len(data.categories), len(data.series))
        if horizontal:
            chart.category_axis.reverse_order = True
            # ValueAxis.crosses controls the OTHER axis. Keep the category
            # axis at zero; move the value axis below reversed categories.
            chart.value_axis.crosses = XL_AXIS_CROSSES.AUTOMATIC
            from pptx.oxml.ns import qn
            crossing = chart.value_axis._element.find(qn("c:crosses"))
            if crossing is not None:
                crossing.set("val", "max")
            if chart.plots[0].has_data_labels:
                chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    if footer_height:
        footer = slide.shapes.add_textbox(left, top + height - footer_height, width, footer_height)
        footer.name = "chart_source"
        footer.text_frame.word_wrap = True
        footer.text_frame.margin_left = 0
        footer.text_frame.margin_right = 0
        footer.text_frame.margin_top = 0
        footer.text_frame.margin_bottom = 0
        paragraph = footer.text_frame.paragraphs[0]
        paragraph.text = "Источник: " + data.source
        paragraph.font.name = font
        paragraph.font.size = Pt(source_size)
        paragraph.font.color.rgb = text_color
    return frame
