"""Композиции из слайдов-образцов шаблона.

Многие корпоративные шаблоны (VK WorkSpace, VK Tech и т.п.) устроены так:
макеты в мастере «пустые» — фон, декор и один заголовок, — а настоящий дизайн
живёт на слайдах-образцах: карточки, таймлайны, фактоиды, таблицы, «Спасибо»
с QR. Парсер коллеги читает макеты, поэтому для таких шаблонов каталог
получается из одинаковых «заголовок + ничего».

Этот модуль читает слайды-образцы и превращает каждый пригодный слайд в
вариант каталога (`source: "sample"`), совместимый с остальным пайплайном:
`slots` с геометрией и ролями, `generation_contract`, `variant_catalog`.
Дополнительно каждый вариант несёт `composition` — сводку для драматургии:
вид композиции (карточки/таймлайн/метрики/…), число повторяемых блоков и их
слоты, чтобы рендер (`slide_clone.py`) мог убирать лишние блоки или
достраивать недостающие.

Определение источника дизайна (`design_source`):
- `layouts` — макеты содержат полноценные плейсхолдеры тела: работает старый путь;
- `slides` — макеты почти пустые, дизайн на слайдах-образцах;
- `mixed` — есть и то и другое;
- `library` — ни макетов с телом, ни образцов: компоновки берём из своей
  библиотеки (`library.py`), а шрифты и цвета — из шаблона.
"""
from __future__ import annotations

import logging
import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

logger = logging.getLogger(__name__)

CM = 360000  # EMU в сантиметре
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

# ---------------------------------------------------------------------------
# Распознавание текста-заглушки
# ---------------------------------------------------------------------------

_PLACEHOLDER_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("presentation_title", re.compile(r"^(название|наименование|заголовок)\s+презентации$|^presentation title$", re.I)),
    ("closing_title", re.compile(r"^спасибо(\s+за\s+внимание)?[!.]?$|^thank(s| you)", re.I)),
    ("qa_title", re.compile(r"^q\s*&\s*a$|^вопросы(\s+и\s+ответы)?[?]?$", re.I)),
    ("cta_title", re.compile(r"^call\s*to\s*action$|^призыв", re.I)),
    ("agenda_title", re.compile(r"^(оглавление|содержание|agenda|план(\s+выступления)?)$", re.I)),
    ("timeline_title", re.compile(r"^(таймлайн|timeline|дорожная карта|roadmap|этапы)$", re.I)),
    ("chart_title", re.compile(r"^(пример\s+(оформления\s+)?)?(график(и|а|ов)?|диаграмм[аы]|chart(s)?)$", re.I)),
    ("problems_title", re.compile(r"^проблем", re.I)),
    ("case_tag", re.compile(r"^(кейс|case|тег|tag)$", re.I)),
    ("speaker", re.compile(r"^имя\s+спикера|^имя\s+фамилия|^фамилия\s+имя|^спикер|^speaker|^фио|должность|^name\s+surname", re.I)),
    ("qr", re.compile(r"^qr(\s*-?\s*code|-код)?$|^(вставить|вставьте|добавить|добавьте|место\s+для|поместите)\s+qr", re.I)),
    ("illustration", re.compile(r"^(вставить|вставьте|добавить|добавьте|место\s+для|поместите|загрузите|разместите)\s+(фото|изображени|картинк|иллюстраци|логотип|скриншот)|^(фото|логотип)\s+(сюда|здесь)$|^(your\s+)?(photo|image|logo)\s+here$", re.I)),
    ("url", re.compile(r"^(ссылка|link|url|сайт)$", re.I)),
    ("date", re.compile(r"^(дата|date|год|квартал|q\d)$", re.I)),
    ("number", re.compile(r"^0?\d{1,2}$")),
    ("metric_value", re.compile(r"^[xхX]{1,4}\s*%?$|^[xх]+\s*%|^\d+\s*%$|^№?\s*[xх]{2,}", re.I)),
    ("metric_label", re.compile(r"^данные\s+показателя|^показател|^подпись\s+(метрики|показателя)|^(объяснение|описание|пояснение)\s+(этого\s+)?показателя|что\s+это\s+за\s+показатель", re.I)),
    ("heading", re.compile(r"^(заголовок|подзаголовок|название|название раздела|heading|title|headline)(\s+\d+)?$|^(название|заголовок)\s+(блока|карточки|этапа|пункта|раздела|слайда)$", re.I)),
    ("body", re.compile(r"^(основной\s+)?текст(\s+описания)?(\s+\d+)?$|^описание$|^text(\s+\d+)?$|^lorem ipsum|^body$"
                        r"|^здесь\s+(можно|нужно|будет|вы\s+можете)|^сюда\s+(можно|нужно)|^описание\b|^пояснение\b|^объяснение\b|^ваш\s+текст|^текст\s+(слайда|блока|карточки|тезиса)"
                        r"|^(первый|второй|третий|четв[её]ртый|пятый|шестой|седьмой|восьмой)\s+(этап|шаг|пункт|блок|тезис)|^место\s+для\s+текста|^добавьте\s+", re.I)),
    ("note", re.compile(r"^(заметка|примечание|note|сноска)$", re.I)),
    ("illustration", re.compile(r"^(иллюстрация|изображение|фото|image|picture|illustration)$", re.I)),
    ("group_label", re.compile(r"^[А-ЯA-Zа-яa-z][^:]{1,30}:$")),
]


def classify_placeholder_text(text: str) -> str | None:
    """Роль текста-заглушки или None, если это «живой» текст шаблона."""
    value = " ".join(text.replace("\x0b", " ").split()).strip()
    if not value:
        return None
    for role, pattern in _PLACEHOLDER_PATTERNS:
        if pattern.search(value):
            return role
    return None


def _is_placeholder_like(text: str) -> bool:
    return classify_placeholder_text(text) is not None


# ---------------------------------------------------------------------------
# Сбор фигур слайда с абсолютной геометрией
# ---------------------------------------------------------------------------


@dataclass
class Leaf:
    """Текстовая (или графическая) фигура с абсолютной геометрией в см."""

    shape_id: int
    kind: str  # text | plate | picture | line | table | chart | group | other
    x: float
    y: float
    w: float
    h: float
    paragraphs: list[str] = field(default_factory=list)  # строки: абзацы и сегменты после <a:br/>
    line_refs: list[tuple[int, int]] = field(default_factory=list)  # (индекс абзаца, индекс сегмента)
    line_fonts: list[float | None] = field(default_factory=list)  # кегль первого рана каждой строки
    font_pt: float | None = None
    bold: bool = False
    is_placeholder: bool = False
    ph_type: str | None = None
    top_id: int | None = None  # id верхнеуровневой фигуры (группы), в которой лежит
    has_fill: bool = False
    name: str = ""
    rotation: float = 0.0  # поворот фигуры в градусах (наклонённые ярлыки — часть декора)

    @property
    def text(self) -> str:
        return "\n".join(self.paragraphs)

    @property
    def cx(self) -> float:
        return self.x + self.w / 2

    @property
    def cy(self) -> float:
        return self.y + self.h / 2

    @property
    def area(self) -> float:
        return max(self.w, 0) * max(self.h, 0)


@dataclass
class TopShape:
    shape_id: int
    kind: str
    x: float
    y: float
    w: float
    h: float
    leaves: list[Leaf]  # все текстовые/графические листья внутри (для группы) или сам

    @property
    def cx(self) -> float:
        return self.x + self.w / 2

    @property
    def cy(self) -> float:
        return self.y + self.h / 2

    @property
    def area(self) -> float:
        return max(self.w, 0) * max(self.h, 0)

    def signature(self) -> tuple:
        texts = tuple(sorted(
            (classify_placeholder_text(leaf.text) or _norm(leaf.text))
            for leaf in self.leaves if leaf.kind == "text" and leaf.text.strip()
        ))
        kinds = tuple(sorted(leaf.kind for leaf in self.leaves))
        return (self.kind, round(self.w, 1), round(self.h, 1), texts, kinds)


def _norm(text: str) -> str:
    return " ".join(text.replace("\x0b", " ").lower().split())


def _shape_kind(shape) -> str:
    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if st == MSO_SHAPE_TYPE.PICTURE or (shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE):
        return "picture"
    if st == MSO_SHAPE_TYPE.LINE:
        return "line"
    prst = shape._element.find(f".//{_A}prstGeom")
    if prst is not None and "arrow" in (prst.get("prst") or "").lower() and not (shape.has_text_frame and shape.text_frame.text.strip()):
        return "arrow"
    if shape.has_text_frame and shape.text_frame.text.strip():
        return "text"
    if st == MSO_SHAPE_TYPE.TABLE or getattr(shape, "has_table", False):
        return "table"
    if getattr(shape, "has_chart", False):
        return "chart"
    if st in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.PLACEHOLDER):
        return "plate"
    return "other"


def _has_fill(shape) -> bool:
    try:
        return shape.fill.type is not None and str(shape.fill.type).split(".")[-1].split(" ")[0] != "BACKGROUND"
    except Exception:  # noqa: BLE001
        return False


def _font_of(shape) -> tuple[float | None, bool]:
    if not shape.has_text_frame:
        return None, False
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                return float(run.font.size.pt), bool(run.font.bold)
    # defRPr / endParaRPr
    for tag in ("defRPr", "endParaRPr", "rPr"):
        el = shape._element.find(f".//{_A}{tag}")
        if el is not None and el.get("sz"):
            return int(el.get("sz")) / 100.0, el.get("b") == "1"
    return None, False


def paragraph_segments(paragraph_element) -> list[str]:
    """Текст абзаца, разбитый по <a:br/> на сегменты (строки)."""
    segments: list[list[str]] = [[]]
    for child in paragraph_element:
        tag = child.tag
        if tag == f"{_A}br":
            segments.append([])
        elif tag == f"{_A}r":
            t = child.find(f"{_A}t")
            segments[-1].append(t.text or "" if t is not None else "")
        elif tag == f"{_A}fld":
            t = child.find(f"{_A}t")
            segments[-1].append(t.text or "" if t is not None else "")
    return ["".join(parts).strip() for parts in segments]


def _segment_fonts(paragraph_element) -> list[float | None]:
    """Кегль первого рана каждого сегмента абзаца (None — наследуется)."""
    sizes: list[float | None] = [None]
    for child in paragraph_element:
        if child.tag == f"{_A}br":
            sizes.append(None)
        elif child.tag in {f"{_A}r", f"{_A}fld"} and sizes[-1] is None:
            rpr = child.find(f"{_A}rPr")
            if rpr is not None and rpr.get("sz"):
                sizes[-1] = int(rpr.get("sz")) / 100.0
    return sizes


def _paragraphs_of(shape) -> tuple[list[str], list[tuple[int, int]], list[float | None]]:
    if not shape.has_text_frame:
        return [], [], []
    lines: list[str] = []
    refs: list[tuple[int, int]] = []
    fonts: list[float | None] = []
    for p_index, paragraph in enumerate(shape.text_frame.paragraphs):
        sizes = _segment_fonts(paragraph._p)
        for s_index, segment in enumerate(paragraph_segments(paragraph._p)):
            lines.append(segment)
            refs.append((p_index, s_index))
            fonts.append(sizes[s_index] if s_index < len(sizes) else None)
    return lines, refs, fonts


def _walk(shapes, ox: float, oy: float, sx: float, sy: float, top_id: int | None, out: list[Leaf]) -> None:
    for shape in shapes:
        x = ox + shape.left * sx
        y = oy + shape.top * sy
        w = shape.width * sx
        h = shape.height * sy
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = shape._element.grpSpPr.find(f"{_A}xfrm")
            ch_off = xfrm.find(f"{_A}chOff") if xfrm is not None else None
            ch_ext = xfrm.find(f"{_A}chExt") if xfrm is not None else None
            if ch_off is not None and ch_ext is not None:
                cx_, cy_ = int(ch_off.get("x")), int(ch_off.get("y"))
                cw, ch = int(ch_ext.get("cx")), int(ch_ext.get("cy"))
                nsx = (w / cw) if cw else sx
                nsy = (h / ch) if ch else sy
                _walk(shape.shapes, x - cx_ * nsx, y - cy_ * nsy, nsx, nsy, top_id if top_id is not None else shape.shape_id, out)
            continue
        kind = _shape_kind(shape)
        font_pt, bold = _font_of(shape)
        ph_type = None
        if shape.is_placeholder:
            try:
                ph_type = shape.placeholder_format.type.name if hasattr(shape.placeholder_format.type, "name") else str(shape.placeholder_format.type)
            except Exception:  # noqa: BLE001
                ph_type = None
        lines, refs, line_fonts = _paragraphs_of(shape) if kind == "text" else ([], [], [])
        out.append(
            Leaf(
                shape_id=shape.shape_id,
                kind=kind,
                x=x / CM,
                y=y / CM,
                w=w / CM,
                h=h / CM,
                paragraphs=lines,
                line_refs=refs,
                line_fonts=line_fonts,
                font_pt=font_pt,
                bold=bold,
                is_placeholder=bool(shape.is_placeholder),
                ph_type=ph_type,
                top_id=top_id if top_id is not None else shape.shape_id,
                has_fill=_has_fill(shape) if kind in {"text", "plate"} else False,
                name=shape.name or "",
                rotation=float(getattr(shape, "rotation", 0.0) or 0.0),
            )
        )


def group_transform(group_element) -> tuple[float, float, float, float] | None:
    """(ox, oy, sx, sy): abs = (ox, oy) + child × (sx, sy) для группы; None, если xfrm неполный."""
    xfrm = group_element.find(f"{_P}grpSpPr/{_A}xfrm")
    if xfrm is None:
        return None
    off, ext = xfrm.find(f"{_A}off"), xfrm.find(f"{_A}ext")
    ch_off, ch_ext = xfrm.find(f"{_A}chOff"), xfrm.find(f"{_A}chExt")
    if None in (off, ext, ch_off, ch_ext):
        return None
    try:
        x, y, cx, cy = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))
        chx, chy, ccx, ccy = int(ch_off.get("x")), int(ch_off.get("y")), int(ch_ext.get("cx")), int(ch_ext.get("cy"))
    except (TypeError, ValueError):
        return None
    if not ccx or not ccy:
        return None
    sx, sy = cx / ccx, cy / ccy
    return x - chx * sx, y - chy * sy, sx, sy


def _explodable(shape, slide_w_emu: int) -> bool:
    """Большая группа, объединяющая все блоки слайда: разбираем на детей."""
    if shape.shape_type != MSO_SHAPE_TYPE.GROUP:
        return False
    children = list(shape.shapes)
    return len(children) >= 4 and shape.width >= 0.4 * slide_w_emu and group_transform(shape._element) is not None


def collect_slide(slide, exploded: list[int] | None = None) -> list[TopShape]:
    prs = slide.part.package.presentation_part.presentation
    slide_w_emu = int(prs.slide_width)
    tops: list[TopShape] = []

    def add(shape, tr: tuple[float, float, float, float] = (0.0, 0.0, 1.0, 1.0)) -> None:
        ox, oy, sx, sy = tr
        leaves: list[Leaf] = []
        _walk([shape], ox, oy, sx, sy, None, leaves)
        kind = "group" if shape.shape_type == MSO_SHAPE_TYPE.GROUP else (leaves[0].kind if leaves else "other")
        tops.append(
            TopShape(
                shape_id=shape.shape_id,
                kind=kind,
                x=(ox + shape.left * sx) / CM,
                y=(oy + shape.top * sy) / CM,
                w=shape.width * sx / CM,
                h=shape.height * sy / CM,
                leaves=leaves,
            )
        )

    for shape in slide.shapes:
        if _explodable(shape, slide_w_emu):
            if exploded is not None:
                exploded.append(shape.shape_id)
            tr = group_transform(shape._element) or (0.0, 0.0, 1.0, 1.0)
            for child in shape.shapes:
                add(child, tr)
            continue
        add(shape)
    return tops


# ---------------------------------------------------------------------------
# Определение источника дизайна
# ---------------------------------------------------------------------------

_BODY_PH = {"BODY", "OBJECT", "SUBTITLE", "TABLE", "CHART", "PICTURE", "MEDIA"}


def detect_design_source(prs) -> dict[str, Any]:
    """`layouts` / `slides` / `mixed` / `library` + сводка признаков."""
    layouts = [layout for master in prs.slide_masters for layout in master.slide_layouts]
    with_body = 0
    for layout in layouts:
        kinds = set()
        for shape in layout.shapes:
            if shape.is_placeholder:
                try:
                    kinds.add(shape.placeholder_format.type.name)
                except Exception:  # noqa: BLE001
                    pass
        if kinds & _BODY_PH:
            with_body += 1
    layouts_rich = with_body >= max(2, len(layouts) // 4)

    sample_slides = 0
    for slide in prs.slides:
        tops = collect_slide(slide)
        free_placeholder_texts = [
            leaf for top in tops for leaf in top.leaves
            if leaf.kind == "text" and not leaf.is_placeholder and _is_placeholder_like(leaf.text)
        ]
        if len(free_placeholder_texts) >= 1 or any(top.kind == "table" for top in tops):
            sample_slides += 1
    slides_rich = sample_slides >= 2

    if layouts_rich and slides_rich:
        source = "mixed"
    elif layouts_rich:
        source = "layouts"
    elif slides_rich:
        source = "slides"
    else:
        source = "library"
    return {
        "design_source": source,
        "layouts_total": len(layouts),
        "layouts_with_body": with_body,
        "sample_slides": sample_slides,
        "slides_total": len(prs.slides),
    }


# ---------------------------------------------------------------------------
# Разбор одного слайда-образца в композицию
# ---------------------------------------------------------------------------

_UNIT_TEXT_ROLES = {
    "heading": "unit_title",
    "body": "unit_text",
    "number": "unit_number",
    "date": "unit_label",
    "metric_value": "metric_value",
    "metric_label": "metric_label",
    "note": "unit_text",
}

_SINGLE_TEXT_ROLES = {
    "body": "body",
    "note": "note",
    "heading": "subtitle",
    "group_label": "subheading",
    "speaker": "contact_text",
    "url": "contact_text",
    "date": "date",
    "case_tag": "tag",
    "metric_value": "metric_value",
    "metric_label": "metric_label",
    "number": "unit_number",
}

_KIND_LABELS = {
    "cover": "Обложка",
    "agenda": "Оглавление",
    "section": "Разделитель",
    "cards": "Карточки",
    "team": "Команда",
    "diagram": "Схема",
    "gallery": "Галерея",
    "timeline": "Таймлайн",
    "metrics": "Метрики",
    "table": "Таблица",
    "chart": "График",
    "statement": "Тезис",
    "text": "Текст",
    "image_text": "Текст и изображение",
    "closing": "Финал",
    "qa": "Вопросы",
    "cta": "Призыв к действию",
    "problems": "Проблемы и вызовы",
    "title_only": "Только заголовок",
    "speaker": "Спикер",
    "columns": "Колонки",
    "list": "Список",
    "quote": "Цитата",
    "comparison": "Сравнение",
}

_NARRATIVE_BY_KIND = {
    "cover": "cover", "agenda": "agenda", "section": "section", "closing": "closing",
    "qa": "closing", "cta": "closing",
}


def _cluster(values: list[float], tol: float) -> list[list[int]]:
    """Индексы значений, сгруппированные по близости (для рядов/колонок)."""
    order = sorted(range(len(values)), key=lambda i: values[i])
    groups: list[list[int]] = []
    for i in order:
        if groups and abs(values[i] - values[groups[-1][-1]]) <= tol:
            groups[-1].append(i)
        else:
            groups.append([i])
    return groups


def _title_leaf(tops: list[TopShape], slide_w: float, slide_h: float) -> Leaf | None:
    for top in tops:
        for leaf in top.leaves:
            if leaf.is_placeholder and leaf.ph_type in {"TITLE", "CENTER_TITLE"}:
                return leaf
    # Свободный заголовок: самый крупный кегль в верхней трети слайда
    candidates = [
        leaf for top in tops for leaf in top.leaves
        if leaf.kind == "text" and leaf.text.strip() and leaf.y < slide_h * 0.35 and leaf.w > slide_w * 0.25
    ]
    candidates.sort(key=lambda l: -(l.font_pt or 0))
    return candidates[0] if candidates else None


@dataclass
class Unit:
    index: int
    members: list[TopShape]
    row: int
    col: int

    @property
    def bbox(self) -> tuple[float, float, float, float]:
        x0 = min(m.x for m in self.members)
        y0 = min(m.y for m in self.members)
        x1 = max(m.x + m.w for m in self.members)
        y1 = max(m.y + m.h for m in self.members)
        return x0, y0, x1 - x0, y1 - y0


def _merge_similar_series(series: dict[tuple, list[TopShape]]) -> dict[tuple, list[TopShape]]:
    """Серии, отличающиеся только размером в пределах ~12 % (рука дизайнера), — одна серия."""
    merged: dict[tuple, list[TopShape]] = {}
    for key, items in sorted(series.items(), key=lambda kv: -len(kv[1])):
        kind, w, h, texts, kinds = key
        target = None
        for mkey in merged:
            mkind, mw, mh, mtexts, mkinds = mkey
            same_size = abs(mw - w) <= max(0.4, 0.12 * mw) and abs(mh - h) <= max(0.4, 0.12 * mh)
            # Одинаковые плашки с разным собственным текстом (6 карточек с
            # примерами) — одна серия, а не пары «заголовок + текст»
            same_shape_text = mtexts == texts or (len(mtexts) == len(texts) and same_size)
            if mkind == kind and same_shape_text and mkinds == kinds and same_size:
                target = mkey
                break
        if target is None:
            merged[key] = list(items)
        else:
            merged[target].extend(items)
    return merged


def _detect_units(tops: list[TopShape], title: Leaf | None, slide_w: float, slide_h: float) -> tuple[list[Unit], dict[str, Any], bool]:
    """Повторяемые блоки: анкеры одной серии + выровненные с ними серии.

    Возвращает (units, grid, has_conflicting_series)."""
    excluded = {title.top_id} if title else set()
    series: dict[tuple, list[TopShape]] = defaultdict(list)
    for top in tops:
        if top.shape_id in excluded or top.kind in {"line", "other", "table", "chart"}:
            continue
        # Растяжки на всю ширину/высоту — не блоки
        if top.w > slide_w * 0.8 or top.h > slide_h * 0.8:
            continue
        series[top.signature()].append(top)
    series = _merge_similar_series(series)
    repeated = {k: v for k, v in series.items() if len(v) >= 2}
    if not repeated:
        return [], {}, False

    def _score(items: list[TopShape]) -> tuple:
        has_text = any(leaf.kind == "text" and leaf.text.strip() for t in items for leaf in t.leaves)
        return (len(items), has_text, items[0].area)

    anchors_key = max(repeated, key=lambda k: _score(repeated[k]))
    anchors = sorted(repeated[anchors_key], key=lambda t: (round(t.y, 1), t.x))
    if anchors[0].area < 1.0 and not any(l.kind == "text" for t in anchors for l in t.leaves):
        return [], {}, False

    tol = max(0.6, anchors[0].h * 0.4)
    rows = _cluster([t.cy for t in anchors], tol)
    cols = _cluster([t.cx for t in anchors], max(0.6, anchors[0].w * 0.4))
    diagonal = False
    if len(anchors) >= 3 and len(rows) == len(anchors) and len(cols) == len(anchors):
        # «Лестница»/диагональ: каждый блок в своём ряду и своей колонке —
        # читается слева направо, как один ряд
        diagonal = True
        rows = [list(range(len(anchors)))]
        cols = [[i] for i in sorted(range(len(anchors)), key=lambda i: anchors[i].cx)]
    row_of = {i: r for r, group in enumerate(rows) for i in group}
    col_of = {i: c for c, group in enumerate(cols) for i in group}
    units: list[Unit] = []
    for i, anchor in enumerate(anchors):
        units.append(Unit(index=0, members=[anchor], row=row_of[i], col=col_of[i]))
    units.sort(key=lambda u: (u.row, u.col))
    for k, unit in enumerate(units):
        unit.index = k

    n = len(anchors)
    conflicting = False
    single_row = len(rows) == 1
    single_col = len(cols) == 1

    def _match_unit(top: TopShape) -> Unit | None:
        best, best_d = None, None
        for unit in units:
            a = unit.members[0]
            if single_row:
                ok = abs(top.cx - a.cx) <= max(a.w, top.w) * 0.55
                d = abs(top.cx - a.cx)
            elif single_col:
                ok = abs(top.cy - a.cy) <= max(a.h, top.h) * 0.55
                d = abs(top.cy - a.cy)
            else:
                ok = abs(top.cx - a.cx) <= max(a.w, top.w) * 0.55 and abs(top.cy - a.cy) <= max(a.h, top.h) * 1.6
                d = abs(top.cx - a.cx) + abs(top.cy - a.cy)
            if ok and (best_d is None or d < best_d):
                best, best_d = unit, d
        return best

    # Серии того же количества и той же сетки — части блока (пары по индексу ячейки)
    for key, items in repeated.items():
        if key == anchors_key:
            continue
        if len(items) == n:
            s_rows = _cluster([t.cy for t in items], max(0.6, items[0].h * 0.4))
            s_cols = _cluster([t.cx for t in items], max(0.6, items[0].w * 0.4))
            if len(s_rows) == len(rows) and len(s_cols) == len(cols):
                s_row_of = {i: r for r, group in enumerate(s_rows) for i in group}
                s_col_of = {i: c for c, group in enumerate(s_cols) for i in group}
                ordered = sorted(range(n), key=lambda i: (s_row_of[i], s_col_of[i]))
                by_cell = {(u.row, u.col): u for u in units}
                pairs = [(items[i], by_cell.get((s_row_of[i], s_col_of[i]))) for i in ordered]
                if all(u is not None for _, u in pairs):
                    for t, u in pairs:
                        u.members.append(t)
                    continue
            matched = [_match_unit(t) for t in items]
            if all(m is not None for m in matched) and len({m.index for m in matched}) == n:
                for t, m in zip(items, matched):
                    m.members.append(t)
                continue
        # Серия другого размера с текстом-заглушкой — конфликтующая (второй ряд блоков)
        if any(l.kind == "text" and _is_placeholder_like(l.text) for t in items for l in t.leaves):
            # Мелкие декоративные повторы (точки таймлайна) — не конфликт
            if items[0].area >= 1.0:
                conflicting = True

    # Одиночные фигуры внутри bbox блока (иконки, плашки) — его члены
    assigned = {m.shape_id for u in units for m in u.members}
    for top in tops:
        if top.shape_id in assigned or top.shape_id in excluded or top.kind in {"line", "table", "chart"}:
            continue
        if top.w > slide_w * 0.8:
            continue
        for unit in units:
            x0, y0, w, h = unit.bbox
            if x0 - 0.3 <= top.cx <= x0 + w + 0.3 and y0 - 0.3 <= top.cy <= y0 + h + 0.3:
                # Не забирать большие «фоновые» плашки, накрывающие несколько блоков
                if top.area <= (w * h) * 1.6:
                    unit.members.append(top)
                    assigned.add(top.shape_id)
                break

    grid = {"rows": len(rows), "cols": len(cols)}
    if diagonal:
        grid["diagonal"] = True
    if len(units) >= 2:
        xs = sorted({round(u.members[0].x, 1) for u in units})
        ys = sorted({round(u.members[0].y, 1) for u in units})
        grid["pitch_x"] = round(min(b - a for a, b in zip(xs, xs[1:])), 2) if len(xs) > 1 else None
        grid["pitch_y"] = round(min(b - a for a, b in zip(ys, ys[1:])), 2) if len(ys) > 1 else None
    return units, grid, conflicting


def _estimate_chars(w: float, h: float, font_pt: float, lines_cap: int | None = None) -> tuple[int, int]:
    pt_cm = 0.0352778
    line_h = max(0.1, font_pt * pt_cm * 1.2)
    char_w = max(0.05, font_pt * pt_cm * 0.52)
    lines = max(1, math.floor(max(h - 0.2, 0.1) / line_h))
    if lines_cap:
        lines = min(lines, lines_cap)
    cpl = max(1, math.floor(max(w - 0.3, 0.1) / char_w))
    return cpl * lines, lines


def _slot_constraints(role: str, w: float, h: float, font_pt: float | None, filled: bool = False) -> dict[str, Any]:
    fallback = {
        "slide_title": 30, "presentation_title": 40, "closing_title": 40, "section_title": 36,
        "unit_title": 20, "unit_text": 16, "unit_number": 16, "unit_label": 16,
        "metric_value": 60, "metric_label": 16, "body": 16, "note": 14, "subheading": 18,
        "subtitle": 18, "contact_text": 14, "tag": 16, "date": 16, "agenda_item": 18,
    }.get(role, 16)
    size = float(font_pt or fallback)
    maximum, lines = _estimate_chars(w, h, size)
    # Текстовые зоны могут расти вниз (autofit): не режем текст по высоте
    # однострочной заглушки — даём минимум строк по роли
    floor_lines = {"unit_text": 3, "body": 4, "note": 2, "unit_title": 1, "agenda_item": 2, "metric_label": 2, "contact_text": 1, "subtitle": 2}.get(role)
    if filled and floor_lines:
        floor_lines = min(floor_lines, 1)  # текст внутри залитой плашки: расти вниз некуда
    if floor_lines and lines < floor_lines:
        per_line = max(1, maximum // max(lines, 1))
        lines = floor_lines
        maximum = per_line * lines
    if role in {"unit_number"}:
        maximum, lines = 3, 1
    if role == "metric_value":
        maximum, lines = min(maximum, 8), 1
    if role in {"unit_label", "date", "tag"}:
        maximum, lines = min(maximum, 24), 1
    if role in {"slide_title", "presentation_title", "closing_title", "section_title"}:
        lines = min(lines, 3)
        maximum = min(maximum, 90)
    if role == "unit_title":
        lines = min(lines, 2)
        maximum = min(maximum, 60)
    return {
        "recommended_characters": max(1, math.floor(maximum * 0.8)),
        "maximum_characters": max(1, maximum),
        "recommended_lines": lines,
        "maximum_lines": lines,
        "font_size_pt": size,
    }


def _leaf_font(leaf: Leaf, theme_fonts: dict[str, str | None], is_title: bool) -> dict[str, Any]:
    family = theme_fonts.get("major" if is_title else "minor")
    return {"family": family, "size": leaf.font_pt, "weight": "bold" if leaf.bold else "regular"}


_LAYOUT_ROLE_RE: list[tuple[str, re.Pattern[str]]] = [
    ("cover", re.compile(r"титул|обложк|cover|title\s*slide", re.I)),
    ("closing", re.compile(r"спасибо|thank|финал|заключ|closing|the\s*end", re.I)),
    ("speaker", re.compile(r"визитк|спикер|speaker|докладчик|о\s+себе|контакт", re.I)),
    ("section", re.compile(r"раздел|section|divider|глава", re.I)),
]


def _layout_role(name: str) -> str | None:
    """Роль слайда по имени макета шаблона («2_Титульный слайд» → cover)."""
    for role, pattern in _LAYOUT_ROLE_RE:
        if pattern.search(name or ""):
            return role
    return None


def _plate_under(tops: list["TopShape"], caption: Leaf, unit_ids: set[int]) -> Leaf | None:
    """Пустая плашка (фигура с заливкой без текста), на которой лежит подпись-заглушка
    «Вставить фото»: центр подписи внутри плашки либо подпись накрывает ≥ 50 % плашки."""
    cx, cy = caption.cx, caption.cy
    best: Leaf | None = None
    for top in tops:
        if top.shape_id in unit_ids:
            continue
        for leaf in top.leaves:
            if leaf.shape_id == caption.shape_id or leaf.kind not in {"plate", "text"} or leaf.text.strip():
                continue
            if leaf.w < 1.0 or leaf.h < 1.0 or leaf.w * leaf.h < 1.5 * caption.w * caption.h * 0.5:
                continue
            inside = leaf.x - 0.2 <= cx <= leaf.x + leaf.w + 0.2 and leaf.y - 0.2 <= cy <= leaf.y + leaf.h + 0.2
            ox = max(0.0, min(leaf.x + leaf.w, caption.x + caption.w) - max(leaf.x, caption.x))
            oy = max(0.0, min(leaf.y + leaf.h, caption.y + caption.h) - max(leaf.y, caption.y))
            covers = ox * oy >= 0.5 * leaf.w * leaf.h
            if (inside or covers) and (best is None or leaf.area < best.area):
                best = leaf
    return best


def analyze_sample_slide(slide, slide_index: int, slide_w: float, slide_h: float, theme_fonts: dict[str, str | None]) -> dict[str, Any] | None:
    """Вариант каталога по слайду-образцу или None, если слайд не годится."""
    exploded: list[int] = []
    tops = collect_slide(slide, exploded)
    if not tops:
        return None
    title = _title_leaf(tops, slide_w, slide_h)
    title_role_hint = classify_placeholder_text(title.text) if title else None
    has_table = any(t.kind == "table" for t in tops)
    has_native_chart = any(t.kind == "chart" for t in tops)

    units, grid, conflicting = _detect_units(tops, title, slide_w, slide_h)
    layout_role = _layout_role(getattr(getattr(slide, "slide_layout", None), "name", "") or "")
    if layout_role in {"cover", "closing", "speaker"} and units and all(
        leaf.is_placeholder for u in units for m in u.members for leaf in m.leaves if leaf.kind == "text" and leaf.text.strip()
    ):
        # Титульный/финальный слайд или визитка: «Имя Фамилия / Должность» — не карточки,
        # а одиночные подписи; иначе обложка превращается в «Карточки · 3 блока»
        units, grid, conflicting = [], {}, False
    if units and any(abs(leaf.rotation % 360.0) > 1.0 and abs(leaf.rotation % 360.0 - 360.0) > 1.0 for u in units for m in u.members for leaf in m.leaves):
        # Наклонённые ярлыки лежат на декоре (нарисованные бирки 01–04): их нельзя
        # ни добавить, ни убрать, ни выровнять по сетке — число блоков фиксировано
        grid = dict(grid or {})
        grid["fixed"] = True
    unit_ids = {m.shape_id for u in units for m in u.members}

    slots: dict[str, Any] = {}
    contract: list[dict[str, Any]] = []
    composition: dict[str, Any] = {"source": "sample", "units": len(units), "grid": grid, "unit_members": [], "unit_slots": [], "title_slot": None, "ungroup": exploded}

    def _add_slot(name: str, leaf: Leaf, role: str, paragraph: int | None, unit: int | None, required: bool = False, is_title: bool = False, lines_share: float = 1.0) -> None:
        h = leaf.h * lines_share
        font_pt = leaf.font_pt
        if paragraph is not None and paragraph < len(leaf.line_fonts):
            fonts = [f or leaf.font_pt or 16.0 for f in leaf.line_fonts]
            font_pt = fonts[paragraph]
            if len(fonts) > 1 and sum(fonts) > 0:
                # Доля высоты фигуры пропорциональна кеглю строки
                h = leaf.h * fonts[paragraph] / sum(fonts)
        constraints = _slot_constraints(role, leaf.w, h, font_pt, filled=leaf.has_fill)
        line_ref = leaf.line_refs[paragraph] if paragraph is not None and paragraph < len(leaf.line_refs) else None
        slots[name] = {
            "shape_id": leaf.shape_id,
            "shape_name": name,
            "paragraph": line_ref[0] if line_ref else None,
            "line": line_ref[1] if line_ref else None,
            "unit": unit,
            "is_placeholder": leaf.is_placeholder,
            "placeholder_type": leaf.ph_type,
            "placeholder_idx": None,
            "x": round(leaf.x, 3), "y": round(leaf.y, 3), "w": round(leaf.w, 3), "h": round(h, 3),
            "font": _leaf_font(leaf, theme_fonts, is_title),
            "resolved_font": {"family": theme_fonts.get("major" if is_title else "minor"), "size_pt": constraints["font_size_pt"], "bold": leaf.bold, "color": None, "source": "sample_slide"},
            "semantic_role": role,
            "role_confidence": 0.9,
            "editable": True,
            "decorative": False,
            "required": required,
            "template_text": leaf.paragraphs[paragraph] if paragraph is not None and paragraph < len(leaf.paragraphs) else leaf.text,
            "horizontal_align": None,
            "capacity": {"font_size_pt": constraints["font_size_pt"], "maximum_characters": constraints["maximum_characters"], "maximum_lines": constraints["maximum_lines"]},
        }
        contract.append({
            "slot": name,
            "role": role,
            "required": required,
            "semantic_scope": "specific",
            "allowed_content_kinds": ["number_or_short_value"] if role in {"metric_value", "unit_number"} else (["paragraph_or_bullets"] if role in {"body", "unit_text", "note"} else ["short_text"]),
            "generator_field": "content",
            "unit": unit,
            "constraints": {k: v for k, v in constraints.items() if k != "font_size_pt"},
        })

    # --- заголовок
    kind: str
    if title is not None:
        if title_role_hint == "presentation_title" or slide_index == 0:
            kind, title_role = "cover", "presentation_title"
        elif title_role_hint == "closing_title":
            kind, title_role = "closing", "closing_title"
        elif title_role_hint == "qa_title":
            kind, title_role = "qa", "closing_title"
        elif title_role_hint == "cta_title":
            kind, title_role = "cta", "closing_title"
        elif title_role_hint == "agenda_title":
            kind, title_role = "agenda", "slide_title"
        elif title_role_hint == "timeline_title":
            kind, title_role = "timeline", "slide_title"
        elif title_role_hint == "chart_title":
            kind, title_role = "chart", "slide_title"
        elif title_role_hint == "problems_title":
            kind, title_role = "problems", "slide_title"
        elif title_role_hint == "case_tag":
            kind, title_role = "cards", "slide_title"
        elif layout_role == "cover":
            kind, title_role = "cover", "presentation_title"
        elif layout_role == "closing":
            kind, title_role = "closing", "closing_title"
        elif layout_role == "speaker":
            kind, title_role = "speaker", "slide_title"
        else:
            kind, title_role = "text", "slide_title"
        _add_slot("title", title, title_role, None, None, required=True, is_title=True)
        composition["title_slot"] = "title"
    else:
        kind = "text"

    # --- блоки
    if conflicting:
        logger.info("Слайд-образец %d: несколько несогласованных серий блоков — пропускаем", slide_index + 1)
        return None

    unit_slot_names: list[list[str]] = []
    has_unit_text = has_unit_title = has_metric = has_unit_label = has_unit_number = False
    for unit in units:
        names: list[str] = []
        leaves = sorted(
            [l for m in unit.members for l in m.leaves if l.kind == "text" and l.text.strip()],
            key=lambda l: (round(l.y, 1), l.x),
        )
        # Внутри блока роли по тексту-заглушке; неизвестный текст повторяемого
        # блока — тоже заглушка: самый крупный короткий текст — unit_title, остальное — unit_text
        max_font = max((l.font_pt or 0) for l in leaves) if leaves else 0
        for leaf in leaves:
            para_roles: list[tuple[int | None, str]] = []
            nonempty = [p for p in leaf.paragraphs if p.strip()]
            recognised = any(classify_placeholder_text(p) for p in nonempty)
            if 2 <= len(nonempty) <= 3 and recognised:
                fonts = [f or leaf.font_pt or 0 for f in leaf.line_fonts] or [0] * len(leaf.paragraphs)
                for p_index, para in enumerate(leaf.paragraphs):
                    if not para.strip():
                        continue
                    hint = classify_placeholder_text(para)
                    if hint is None:
                        first_bigger = p_index == 0 and len(fonts) > 1 and fonts[0] > max(fonts[1:]) + 0.5
                        hint = "heading" if (first_bigger and len(para.split()) <= 6) else "body"
                    para_roles.append((p_index, _UNIT_TEXT_ROLES.get(hint, "unit_text")))
            else:
                hint = classify_placeholder_text(leaf.text)
                if hint is None:
                    is_biggest = (leaf.font_pt or 0) >= max_font - 0.5 and len(leaves) >= 2
                    hint = "heading" if (is_biggest and len(leaf.text) <= 60) else "body"
                para_roles.append((None, _UNIT_TEXT_ROLES.get(hint, "unit_text")))
            for p_index, role in para_roles:
                if role == "unit_title" and any(slots[n]["semantic_role"] == "unit_title" for n in names):
                    role = "unit_text"
                if kind in {"closing", "qa", "cta"} and role in {"unit_title", "unit_text"}:
                    role = "contact_text"
                if kind == "agenda" and role in {"unit_title", "unit_text"}:
                    role = "agenda_item"
                suffix = {"unit_title": "title", "unit_text": "text", "unit_number": "num", "unit_label": "label", "metric_value": "value", "metric_label": "caption", "contact_text": "contact", "agenda_item": "item"}[role]
                name = f"u{unit.index + 1}_{suffix}"
                if name in slots:
                    name = f"{name}{len(names) + 1}"
                share = 1.0
                if p_index is not None and len(leaf.paragraphs) >= 2:
                    share = 0.4 if role in {"unit_title", "metric_value"} else 0.6
                _add_slot(name, leaf, role, p_index, unit.index, lines_share=share)
                names.append(name)
                has_unit_text |= role == "unit_text"
                has_unit_title |= role == "unit_title"
                has_metric |= role == "metric_value"
                has_unit_label |= role == "unit_label"
                has_unit_number |= role == "unit_number"
        unit_slot_names.append(names)
        composition["unit_members"].append([m.shape_id for m in unit.members])
    composition["unit_slots"] = unit_slot_names
    if units and not any(unit_slot_names):
        # Повторяются только декоративные элементы — блоков нет
        units, composition["units"], composition["unit_members"], composition["unit_slots"] = [], 0, [], []
        unit_ids = set()

    # --- одиночные текстовые слоты
    single_roles: Counter[str] = Counter()
    image_slots: dict[str, Any] = {}
    hinted_text = [
        leaf for top in tops for leaf in top.leaves
        if leaf.kind == "text" and leaf.text.strip() and classify_placeholder_text(leaf.text) is not None
        and not (title and leaf.shape_id == title.shape_id)
    ]

    def _under_text(pic: Leaf) -> bool:
        # Иллюстрация, на которой лежит текст-заглушка (число в круге диаграммы), — декор, не фото
        return any(pic.x <= t.cx <= pic.x + pic.w and pic.y <= t.cy <= pic.y + pic.h for t in hinted_text)

    for top in tops:
        if top.shape_id in unit_ids or (title and top.shape_id == title.top_id):
            continue
        for leaf in top.leaves:
            if leaf.kind == "picture":
                if leaf.w >= 4.0 and leaf.h >= 3.0 and top.kind != "group" and not _under_text(leaf):
                    square = abs(leaf.w - leaf.h) < 0.6 and leaf.w <= 7.5
                    is_qr = square and kind in {"closing", "qa", "cta", "cover"}
                    name = ("qr" if is_qr else "image") if not image_slots else (f"qr_{len(image_slots) + 1}" if is_qr else f"image_{len(image_slots) + 1}")
                    image_slots[name] = {
                        "shape_id": leaf.shape_id, "shape_name": name, "placeholder_idx": None,
                        "x": round(leaf.x, 3), "y": round(leaf.y, 3), "w": round(leaf.w, 3), "h": round(leaf.h, 3),
                        "semantic_role": "qr_code" if is_qr else "photo", "editable": not is_qr, "required": False, "source_kind": "sample_picture",
                    }
                continue
            if leaf.kind != "text" or not leaf.text.strip():
                continue
            if leaf.is_placeholder and title and leaf.shape_id == title.shape_id:
                continue
            hint = classify_placeholder_text(leaf.text)
            if hint is None and leaf.is_placeholder and kind in {"cover", "closing", "speaker"} and (leaf.ph_type or "").upper() in {"BODY", "SUBTITLE", "OBJECT"}:
                hint = "body"  # плейсхолдер титула с «живым» текстом («Разработчик корпоративного ПО») — подзаголовок
            if hint is None:
                continue  # живой текст шаблона (подписи, логотипы) — не трогаем
            if hint in {"qr", "illustration"}:
                # Плашка-заглушка под QR/иллюстрацию: слот изображения; без картинки плашку убираем.
                # Подпись «Вставить фото» часто лежит поверх отдельной белой плашки — слотом
                # становится плашка (её геометрия), подпись запоминаем, чтобы убрать в любом случае
                is_qr = hint == "qr"
                plate = _plate_under(tops, leaf, unit_ids)
                geom = plate or leaf
                name = ("qr" if is_qr else "image") if not image_slots else f"{'qr' if is_qr else 'image'}_{len(image_slots) + 1}"
                image_slots[name] = {
                    "shape_id": geom.shape_id, "shape_name": name, "placeholder_idx": None,
                    "x": round(geom.x, 3), "y": round(geom.y, 3), "w": round(geom.w, 3), "h": round(geom.h, 3),
                    "semantic_role": "qr_code" if is_qr else "photo", "editable": not is_qr, "required": False, "source_kind": "placeholder_shape",
                    "caption_shape_id": leaf.shape_id if plate is not None else None,
                }
                continue
            if hint in {"metric_value", "metric_label"} and len(leaf.paragraphs) >= 2:
                # «ххх% / данные показателя» одной фигурой
                for p_index, para in enumerate(leaf.paragraphs):
                    h2 = classify_placeholder_text(para)
                    if h2 in {"metric_value", "metric_label"}:
                        single_roles[h2] += 1
                        name = f"{'value' if h2 == 'metric_value' else 'caption'}_{single_roles[h2]}"
                        _add_slot(name, leaf, h2, p_index, None, lines_share=0.5 if h2 == "metric_value" else 0.5)
                        has_metric |= h2 == "metric_value"
                continue
            role = _SINGLE_TEXT_ROLES.get(hint, "body")
            if kind == "cover" and role in {"body", "subtitle"}:
                role = "presentation_subtitle"
            if kind in {"closing", "qa", "cta"} and role in {"body", "subtitle", "contact_text"}:
                role = "contact_text"
            single_roles[role] += 1
            base = {"body": "text", "note": "note", "subtitle": "subtitle", "subheading": "label", "contact_text": "contact", "date": "date", "tag": "tag", "metric_value": "value", "metric_label": "caption", "unit_number": "num", "presentation_subtitle": "subtitle"}.get(role, role)
            name = base if single_roles[role] == 1 else f"{base}_{single_roles[role]}"
            if name in slots:
                name = f"{name}_{len(slots)}"
            _add_slot(name, leaf, role, None, None)
            has_metric |= role == "metric_value"

    # --- вид композиции
    if kind == "cover" and (has_table or has_native_chart) and slide_index != 0:
        # «Титульный» макет с таблицей/графиком на слайде — это слайд данных,
        # а не обложка (VK WorkSpace: образец на макете «Титульный слайд» с таблицей
        # иначе становился обложкой и выводил пустую таблицу шаблона)
        kind = "table" if has_table else "chart"
        if "title" in slots:
            slots["title"]["semantic_role"] = "slide_title"
    if kind in {"cover", "closing", "qa", "cta", "agenda", "speaker"}:
        pass
    elif has_table:
        kind = "table"
    elif kind == "chart" or has_native_chart:
        kind = "chart"
    elif has_metric:
        kind = "metrics"
    elif kind == "timeline" or has_unit_label:
        kind = "timeline"
    elif kind == "problems":
        kind = "problems"
    elif units and sum(1 for t in tops for l in t.leaves if l.kind in {"arrow", "line"} and (l.kind == "arrow" or max(l.w, l.h) >= 1.0)) >= 2:
        kind = "diagram"  # схема со стрелками/связями — под произвольный контент не подходит
    elif units and sum(1 for u in units if any(l.kind == "picture" for m in u.members for l in m.leaves)) >= max(1, math.ceil(len(units) * 0.5)):
        kind = "team" if any(classify_placeholder_text(l.text) == "speaker" for u in units for m in u.members for l in m.leaves if l.kind == "text") else "gallery"
    elif units:
        kind = "cards"
    elif image_slots and any(s["semantic_role"] == "body" for s in slots.values()):
        kind = "image_text"
    elif any(s["semantic_role"] in {"body", "note"} for s in slots.values()):
        kind = "text"
    else:
        kind = "statement"

    # Слайд без единого редактируемого слота (чистый декор) — не вариант
    if not slots:
        return None

    tables = [t for t in tops if t.kind == "table"]
    chart_pics = [s for s in image_slots.values()] if kind == "chart" else []
    composition.update({
        "kind": kind,
        "has_unit_title": has_unit_title,
        "has_unit_text": has_unit_text,
        "has_unit_number": has_unit_number,
        "has_unit_label": has_unit_label,
        "has_metric": has_metric,
        "has_body": any(s["semantic_role"] in {"body"} for s in slots.values()),
        "has_image": any(s["semantic_role"] == "photo" for s in image_slots.values()) and kind != "chart",
        "min_units": (len(units) if grid.get("diagonal") or grid.get("fixed") else 2) if units else 0,
        "max_units": (len(units) if grid.get("diagonal") or grid.get("fixed") else min(8, max(len(units) + 2, math.ceil(len(units) * 1.5)))) if units else 0,
        "fixed_units": bool(units and grid.get("fixed")),
        "table_shape_id": tables[0].shape_id if tables else None,
        "chart_pictures": [s["shape_id"] for s in chart_pics],
        "removable_when_empty": [name for name, s in slots.items() if s["semantic_role"] in {"tag", "note", "subheading", "date"} or (s.get("unit") is None and s["semantic_role"] in {"metric_label", "metric_value"})],
    })

    # Виртуальные слоты для достраиваемых блоков (u{N+1}.. u{max})
    if units:
        template_names = unit_slot_names[-1]
        for extra in range(len(units) + 1, composition["max_units"] + 1):
            names = []
            for src_name in template_names:
                src = slots[src_name]
                suffix = src_name.split("_", 1)[1]
                name = f"u{extra}_{suffix}"
                slots[name] = {**src, "shape_name": name, "unit": extra - 1, "virtual": True, "virtual_from": src_name}
                src_contract = next(c for c in contract if c["slot"] == src_name)
                contract.append({**src_contract, "slot": name, "unit": extra - 1, "virtual": True})
                names.append(name)
            unit_slot_names.append(names)
        composition["unit_slots"] = unit_slot_names

    label = _KIND_LABELS.get(kind, kind)
    detail = ""
    if units:
        detail = f" · {len(units)} блок{'а' if 2 <= len(units) <= 4 else 'ов'}"
    if image_slots and kind not in {"chart"}:
        detail += " · фото"
    variant_id = f"Образец {slide_index + 1} · {label}{detail}"
    narrative = _NARRATIVE_BY_KIND.get(kind, "content")
    pattern = {
        "cards": "repeated_blocks", "problems": "repeated_blocks", "timeline": "sequence", "metrics": "metrics",
        "table": "table", "chart": "chart", "statement": "statement", "text": "explanation", "image_text": "explanation",
        "agenda": "navigation", "cover": "introduction", "closing": "call_to_action", "qa": "call_to_action", "cta": "call_to_action",
        "section": "section_transition",
    }.get(kind, "explanation")
    best_for = {
        "cards": ["3–6 однотипных пунктов: преимущества, шаги, направления"],
        "problems": ["перечень проблем/вызовов или рисков"],
        "timeline": ["этапы, план, хронология"],
        "metrics": ["2–4 ключевые цифры с подписями"],
        "table": ["сравнение по параметрам, табличные данные"],
        "chart": ["график/диаграмма с короткой подписью"],
        "statement": ["один сильный тезис на слайд"],
        "text": ["связный текст или список тезисов"],
        "image_text": ["текст с фото/иллюстрацией"],
        "agenda": ["оглавление презентации"],
        "cover": ["титульный слайд"],
        "closing": ["финальный слайд с контактами"],
        "qa": ["слайд вопросов"],
        "cta": ["призыв к действию, ссылка, QR"],
    }.get(kind, [])

    required = [name for name, s in slots.items() if s.get("required")]
    variant = {
        "id": variant_id,
        "layout_name": f"{slide.slide_layout.name} → образец {slide_index + 1}",
        "type": "sample",
        "source": "sample",
        "layout_file": str(slide.slide_layout.part.partname).lstrip("/"),
        "slide_file": str(slide.part.partname).lstrip("/"),
        "master_index": 1,
        "slide_index": slide_index,
        "slots": slots,
        "image_slots": image_slots,
        "editable_image_slots": [n for n, s in image_slots.items() if s.get("editable")],
        "background": None,
        "images": [],
        "constraints": {
            "slide_bounds_cm": {"x": 0, "y": 0, "w": slide_w, "h": slide_h},
            "text_regions_cm": {n: {k: s[k] for k in ("x", "y", "w", "h")} for n, s in slots.items()},
            "image_regions_cm": {n: {k: s[k] for k in ("x", "y", "w", "h")} for n, s in image_slots.items()},
        },
        "composition": composition,
        "semantic": {
            "layout_role": kind,
            "confidence": 0.85,
            "evidence": ["sample_slide"],
            "reading_order": list(slots.keys()),
            "selection_hints": {"best_for": best_for, "required_slots": required, "optional_slots": [n for n in slots if n not in required], "avoid_when": []},
            "ambiguous_slots": [],
            "flexible_slots": [],
            "semantic_scope": "specific",
            "slot_groups": [],
            "content_budget": {"text_slots": len(contract)},
            "review": {"recommended": False, "reasons": []},
            "downstream_choice": "auto",
        },
        "observed": {},
        "inferred": {
            "layout": {"narrative_role": narrative, "content_pattern": pattern, "semantic_scope": "specific", "confidence": 0.85},
            "composition": {"primary": {"cards": "repeated_blocks", "problems": "repeated_blocks", "timeline": "horizontal_sequence", "metrics": "metrics", "agenda": "repeated_blocks"}.get(kind, "single_focus")},
            "review": {"recommended": False},
        },
        "generation_contract": {
            "text_slots": contract,
            "image_slots": [
                {"slot": n, "role": s["semantic_role"], "required": False, "semantic_scope": "specific", "allowed_content_kinds": [s["semantic_role"]], "generator_field": "image_content", "constraints": {}}
                for n, s in image_slots.items()
            ],
            "required_slots": required,
            "selection": {"best_for": best_for, "avoid_when": []},
            "decision_owner": "host",
            "content_budget": {"text_slots": len(contract)},
        },
    }
    return variant


def catalog_entry(variant: dict[str, Any]) -> dict[str, Any]:
    layout = variant["inferred"]["layout"]
    contract = variant["generation_contract"]
    comp = variant.get("composition") or {}
    text_slots = [
        {"slot": s["slot"], "scope": s.get("semantic_scope", "specific"), "role": s["role"],
         "allowed_content_kinds": s.get("allowed_content_kinds", []), "required": s.get("required", False),
         "recommended_characters": (s.get("constraints") or {}).get("recommended_characters"),
         "maximum_lines": (s.get("constraints") or {}).get("maximum_lines")}
        for s in contract["text_slots"] if not s.get("virtual")
    ]
    return {
        "id": variant["id"],
        "family": f"{layout['narrative_role']}.{layout['content_pattern']}.{comp.get('kind', 'sample')}.{variant.get('slide_index', 0)}",
        "visual_variant": variant["id"],
        "semantic_scope": "specific",
        "narrative_role": layout["narrative_role"],
        "content_pattern": layout["content_pattern"],
        "composition": variant["inferred"]["composition"]["primary"],
        "confidence": layout["confidence"],
        "best_for": contract["selection"]["best_for"],
        "avoid_when": [],
        "text_slots": text_slots,
        "image_slots": [{"slot": s["slot"], "scope": "specific", "role": s["role"], "allowed_content_kinds": ["photo"], "required": False} for s in contract["image_slots"]],
        "required_slots": contract["required_slots"],
        "content_budget": contract["content_budget"],
        "review_required": False,
        "decision_owner": "host",
        "kind": comp.get("kind"),
        "units": comp.get("units", 0),
        "fixed_units": bool(comp.get("fixed_units")),
    }


def augment_template(template: dict[str, Any], pptx_path: str | Path) -> dict[str, Any]:
    """Дополняет Template JSON вариантами из слайдов-образцов и `design_source`."""
    prs = Presentation(str(pptx_path))
    info = detect_design_source(prs)
    template["design_source"] = info["design_source"]
    template["design_source_info"] = info
    theme_fonts = template.get("theme_fonts") or {}
    slide_w = float(template["slide_size"]["width"])
    slide_h = float(template["slide_size"]["height"])

    if info["design_source"] not in {"slides", "mixed"}:
        return template

    added = 0
    seen_signatures: set[tuple] = set()
    for index, slide in enumerate(prs.slides):
        try:
            variant = analyze_sample_slide(slide, index, slide_w, slide_h, theme_fonts)
        except Exception:  # noqa: BLE001 — один странный слайд не должен ломать разбор
            logger.exception("Не удалось разобрать слайд-образец %d", index + 1)
            continue
        if variant is None:
            continue
        comp = variant["composition"]
        signature = (comp["kind"], comp["units"], tuple(sorted(s["semantic_role"] for s in variant["slots"].values() if not s.get("virtual"))), bool(variant["image_slots"]))
        # Одинаковые по структуре образцы оставляем (это визуальные версии), но не более 3 на структуру
        if sum(1 for s in seen_signatures if s[:-1] == signature) >= 3:
            continue
        seen_signatures.add(signature + (index,))
        template["variants"].append(variant)
        template.setdefault("variant_catalog", []).append(catalog_entry(variant))
        added += 1
    logger.info("Слайды-образцы: добавлено %d композиций (design_source=%s)", added, info["design_source"])
    return template
