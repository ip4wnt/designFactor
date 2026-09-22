"""Компоновщик слайдов: DesignManifest + ContentPlan -> .pptx.

Мэтчинг макета — детерминированный rule-based скоринг ролей макета против
типов контент-блоков слайда, а НЕ вызов LLM на каждый слайд: при бюджете
5 минут на колоду LLM-вызов на каждый из ~10 слайдов слишком дорог и
недетерминирован по времени ответа.

Диаграммы и таблицы — нативные объекты python-pptx (add_chart/add_table),
растровые изображения слайдов результатом не считаются.
"""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu

from app.pipeline.capacity import fit_chart_data, fit_table_data
from app.pipeline.styling import style_chart, style_table
from app.schemas.content_plan import ContentBlock, ContentBlockType, ContentPlan, SlideSpec
from app.schemas.design_manifest import DesignManifest, Layout, LayoutRoleType

_CHART_TYPE_MAP = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
}

# Область по умолчанию для чарта/таблицы, если в макете вообще нет
# подходящего плейсхолдера (страховка от NPE, не основной путь).
_FALLBACK_BBOX = (Emu(838200), Emu(1600200), Emu(7620000), Emu(4525963))

# Конфиги скоринга под разные варианты вёрстки. ТЗ требует заложить
# расширение под variant_a/b/c — реализован сейчас только variant_a,
# остальные объявлены как явный TODO, чтобы не притворяться, что готовы.
SCORING_CONFIGS: dict[str, dict[str, float]] = {
    "variant_a": {"role_match": 2.0, "missing_role_penalty": -3.0, "unused_role_penalty": -0.5},
}
_NOT_IMPLEMENTED_VARIANTS = {"variant_b", "variant_c"}


def assemble_presentation(
    template_path: str | Path,
    manifest: DesignManifest,
    content_plan: ContentPlan,
    output_path: str | Path,
    variant: str = "variant_a",
) -> Path:
    if variant in _NOT_IMPLEMENTED_VARIANTS:
        raise NotImplementedError(f"Вариант вёрстки '{variant}' пока не реализован")
    if variant not in SCORING_CONFIGS:
        raise ValueError(f"Неизвестный вариант вёрстки: {variant}")

    config = SCORING_CONFIGS[variant]
    prs = Presentation(str(template_path))
    pptx_layouts = list(prs.slide_masters[0].slide_layouts)

    # manifest.layouts построен Парсером обходом того же master.slide_layouts
    # в том же порядке -> индексы совпадают 1-в-1. Это единственная связь
    # между "нашей" моделью Layout и реальным объектом python-pptx.
    if len(manifest.layouts) != len(pptx_layouts):
        raise ValueError("DesignManifest не соответствует шаблону: число макетов расходится")
    layout_pairs = list(zip(manifest.layouts, pptx_layouts))

    for slide_spec in content_plan.slides:
        layout, pptx_layout = _select_layout(layout_pairs, slide_spec, config)
        slide = prs.slides.add_slide(pptx_layout)
        _fill_slide(slide, layout, slide_spec, manifest)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path


def _select_layout(
    layout_pairs: list[tuple[Layout, object]],
    slide_spec: SlideSpec,
    config: dict[str, float],
) -> tuple[Layout, object]:
    best_score = None
    best_pair = None
    for layout, pptx_layout in layout_pairs:
        score = _score_layout(layout, slide_spec, config)
        if best_score is None or score > best_score:
            best_score = score
            best_pair = (layout, pptx_layout)
    return best_pair


def _score_layout(layout: Layout, slide_spec: SlideSpec, config: dict[str, float]) -> float:
    role_counts = Counter(role.role for role in layout.roles)
    match = config["role_match"]
    missing = config["missing_role_penalty"]
    unused = config["unused_role_penalty"]

    score = match if role_counts.get(LayoutRoleType.TITLE, 0) > 0 else missing

    body_needed = sum(
        1 for b in slide_spec.content_blocks if b.type in (ContentBlockType.BULLETS, ContentBlockType.TEXT)
    )
    body_available = role_counts.get(LayoutRoleType.BODY, 0)
    score += match * min(body_needed, body_available)
    if body_needed > body_available:
        score += missing * (body_needed - body_available)

    chart_needed = sum(1 for b in slide_spec.content_blocks if b.type == ContentBlockType.CHART)
    if chart_needed > 0:
        score += match if layout.supports_chart else missing

    table_needed = sum(1 for b in slide_spec.content_blocks if b.type == ContentBlockType.TABLE)
    if table_needed > 0:
        score += match if layout.supports_table else missing

    used_roles = 1 + min(body_needed, body_available) + (1 if chart_needed else 0) + (1 if table_needed else 0)
    total_roles = len(layout.roles)
    score += unused * max(0, total_roles - used_roles)

    return score


def _fill_slide(slide, layout: Layout, slide_spec: SlideSpec, manifest: DesignManifest) -> None:
    title_idx = next((r.placeholder_idx for r in layout.roles if r.role == LayoutRoleType.TITLE), None)
    if title_idx is not None:
        _set_placeholder_text(slide, title_idx, [slide_spec.title])

    body_idxs = [r.placeholder_idx for r in layout.roles if r.role == LayoutRoleType.BODY]
    chart_idxs = [r.placeholder_idx for r in layout.roles if r.role == LayoutRoleType.CHART]
    table_idxs = [r.placeholder_idx for r in layout.roles if r.role == LayoutRoleType.TABLE]

    for block in slide_spec.content_blocks:
        if block.type in (ContentBlockType.BULLETS, ContentBlockType.TEXT):
            if not body_idxs:
                continue
            idx = body_idxs.pop(0)
            texts = block.bullets if block.type == ContentBlockType.BULLETS else [block.text or ""]
            _set_placeholder_text(slide, idx, texts or [""])
        elif block.type == ContentBlockType.CHART:
            idx = chart_idxs.pop(0) if chart_idxs else (body_idxs.pop(0) if body_idxs else None)
            _add_native_chart(slide, idx, block, manifest)
        elif block.type == ContentBlockType.TABLE:
            idx = table_idxs.pop(0) if table_idxs else (body_idxs.pop(0) if body_idxs else None)
            _add_native_table(slide, idx, block, manifest)


def _set_placeholder_text(slide, idx: int, paragraphs: list[str]) -> None:
    placeholder = slide.placeholders[idx]
    text_frame = placeholder.text_frame
    text_frame.clear()
    text_frame.paragraphs[0].text = paragraphs[0]
    for extra in paragraphs[1:]:
        p = text_frame.add_paragraph()
        p.text = extra


def _placeholder_bbox(slide, idx: int | None) -> tuple[int, int, int, int]:
    if idx is None:
        return _FALLBACK_BBOX
    placeholder = slide.placeholders[idx]
    return placeholder.left, placeholder.top, placeholder.width, placeholder.height


def _remove_placeholder(slide, idx: int | None) -> None:
    if idx is None:
        return
    placeholder = slide.placeholders[idx]
    placeholder._element.getparent().remove(placeholder._element)


def _add_native_chart(slide, idx: int | None, block: ContentBlock, manifest: DesignManifest) -> None:
    if block.chart is None:
        return
    left, top, width, height = _placeholder_bbox(slide, idx)
    _remove_placeholder(slide, idx)

    categories, series, _truncated = fit_chart_data(
        block.chart.categories, block.chart.series, width
    )

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for series_name, values in series.items():
        chart_data.add_series(series_name, values)

    xl_chart_type = _CHART_TYPE_MAP[block.chart.chart_type.value]
    graphic_frame = slide.shapes.add_chart(xl_chart_type, left, top, width, height, chart_data)
    style_chart(graphic_frame.chart, manifest)


def _add_native_table(slide, idx: int | None, block: ContentBlock, manifest: DesignManifest) -> None:
    if block.table is None:
        return
    left, top, width, height = _placeholder_bbox(slide, idx)
    _remove_placeholder(slide, idx)

    headers, data_rows, _truncated = fit_table_data(
        block.table.headers,
        block.table.rows,
        height,
        width,
        manifest.typography.body.size,
    )

    rows = len(data_rows) + 1
    cols = max(len(headers), 1)
    graphic_frame = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = graphic_frame.table

    for col_idx, header in enumerate(headers):
        table.cell(0, col_idx).text = header
    for row_idx, row in enumerate(data_rows, start=1):
        for col_idx, value in enumerate(row):
            table.cell(row_idx, col_idx).text = str(value)

    style_table(table, manifest)
