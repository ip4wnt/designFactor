"""Рендер композиций из слайдов-образцов: клонирование слайда, заполнение
слотов по shape_id, адаптация повторяемых блоков (убрать лишние /
достроить недостающие), таблицы и графики, чистка заглушек.

Работает поверх python-pptx на уровне XML: слайд-образец копируется целиком
(фон, декор, группы, иконки), поэтому дизайн шаблона сохраняется дословно.
"""
from __future__ import annotations

import copy
import logging
import math
import re
from pathlib import Path
from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.pipeline.sample_slides import group_transform, CM, classify_placeholder_text, paragraph_segments

logger = logging.getLogger(__name__)

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_MIN_FONT_PT = 9.0


# ---------------------------------------------------------------------------
# Клонирование слайда
# ---------------------------------------------------------------------------


def clone_slide(prs, source_slide):
    """Копия слайда-образца в конце презентации (та же разметка, фон, все фигуры, картинки)."""
    layout = source_slide.slide_layout
    new_slide = prs.slides.add_slide(layout)
    # Убираем плейсхолдеры, которые python-pptx создал по макету, — копируем всё из образца.
    for shape in list(new_slide.shapes):
        shape._element.getparent().remove(shape._element)

    src_tree = source_slide.shapes._spTree
    dst_tree = new_slide.shapes._spTree
    rid_map: dict[str, str] = {}
    for child in src_tree:
        if child.tag in {f"{_P}nvGrpSpPr", f"{_P}grpSpPr"}:
            continue
        copied = copy.deepcopy(child)
        _remap_rels(copied, source_slide.part, new_slide.part, rid_map)
        dst_tree.append(copied)

    # Фон слайда
    src_bg = source_slide._element.cSld.find(f"{_P}bg")
    if src_bg is not None:
        copied_bg = copy.deepcopy(src_bg)
        _remap_rels(copied_bg, source_slide.part, new_slide.part, rid_map)
        dst_cSld = new_slide._element.cSld
        existing = dst_cSld.find(f"{_P}bg")
        if existing is not None:
            dst_cSld.remove(existing)
        dst_cSld.insert(0, copied_bg)
    return new_slide


def _remap_rels(element, src_part, dst_part, rid_map: dict[str, str]) -> None:
    for node in element.iter():
        for attr in (f"{_R}embed", f"{_R}link", f"{_R}id", f"{_R}pict"):
            rid = node.get(attr)
            if not rid:
                continue
            if rid not in rid_map:
                rel = src_part.rels.get(rid)
                if rel is None:
                    continue
                if rel.is_external:
                    rid_map[rid] = dst_part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
                else:
                    rid_map[rid] = dst_part.relate_to(rel.target_part, rel.reltype)
            node.set(attr, rid_map[rid])


# ---------------------------------------------------------------------------
# Поиск и геометрия
# ---------------------------------------------------------------------------


def _ungroup(slide, group_id: int) -> None:
    """Разбирает группу: дети встают в spTree на её место с пересчётом геометрии
    (масштаб группы применяется к координатам и размерам; кегль текста PowerPoint
    при группировке не масштабирует, поэтому шрифты не трогаем)."""
    group = find_by_id(slide, group_id)
    if group is None or group.tag != f"{_P}grpSp":
        return
    tr = group_transform(group) or (0.0, 0.0, 1.0, 1.0)
    ox, oy, sx, sy = tr
    parent = group.getparent()
    index = parent.index(group)
    children = [c for c in group if c.tag in {f"{_P}sp", f"{_P}pic", f"{_P}grpSp", f"{_P}cxnSp", f"{_P}graphicFrame"}]
    for offset, child in enumerate(children):
        xfrm = _xfrm(child)
        if xfrm is not None:
            off, ext = xfrm.find(f"{_A}off"), xfrm.find(f"{_A}ext")
            if off is not None and ext is not None:
                try:
                    x, y, cx, cy = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))
                    off.set("x", str(round(ox + x * sx)))
                    off.set("y", str(round(oy + y * sy)))
                    ext.set("cx", str(round(cx * sx)))
                    ext.set("cy", str(round(cy * sy)))
                except (TypeError, ValueError):
                    pass
        parent.insert(index + offset, child)
    parent.remove(group)


def find_by_id(slide, shape_id: int):
    """Элемент фигуры (p:sp/p:pic/p:grpSp/p:cxnSp/p:graphicFrame) с данным id — в любой глубине."""
    for node in slide.shapes._spTree.iter():
        if node.tag == f"{_P}cNvPr" and node.get("id") == str(shape_id):
            return node.getparent().getparent()
    return None


def _top_level(element, slide):
    tree = slide.shapes._spTree
    node = element
    while node is not None and node.getparent() is not tree:
        node = node.getparent()
    return node


def _xfrm(element):
    tag = element.tag
    if tag == f"{_P}grpSp":
        return element.find(f"{_P}grpSpPr/{_A}xfrm")
    if tag == f"{_P}graphicFrame":
        return element.find(f"{_P}xfrm")
    sppr = element.find(f"{_P}spPr")
    return sppr.find(f"{_A}xfrm") if sppr is not None else None


def _bbox(element) -> tuple[int, int, int, int] | None:
    xfrm = _xfrm(element)
    if xfrm is None:
        return None
    off, ext = xfrm.find(f"{_A}off"), xfrm.find(f"{_A}ext")
    if off is None or ext is None:
        return None
    return int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))


def _move(element, dx: int, dy: int) -> None:
    xfrm = _xfrm(element)
    if xfrm is None:
        return
    off = xfrm.find(f"{_A}off")
    if off is not None:
        off.set("x", str(int(off.get("x")) + dx))
        off.set("y", str(int(off.get("y")) + dy))


def _scale_x(element, origin_x: int, scale: float) -> None:
    """Масштаб по X относительно origin_x (ширины и отступы), высоты не меняются."""
    xfrm = _xfrm(element)
    if xfrm is None:
        return
    off, ext = xfrm.find(f"{_A}off"), xfrm.find(f"{_A}ext")
    if off is None or ext is None:
        return
    x = int(off.get("x"))
    w = int(ext.get("cx"))
    h = int(ext.get("cy"))
    # Квадратные мелкие элементы (иконки, кружки) не сплющиваем — сдвигаем
    if abs(w - h) < 0.15 * max(w, h) and w < 2.5 * CM:
        off.set("x", str(origin_x + int((x - origin_x) * scale)))
        return
    off.set("x", str(origin_x + int((x - origin_x) * scale)))
    ext.set("cx", str(max(int(w * scale), 1)))
    if element.tag == f"{_P}grpSp":
        ch_ext = xfrm.find(f"{_A}chExt")
        # Дочерние координаты группы остаются, масштаб задаётся через ext/chExt
        _ = ch_ext


def _scale_y(element, origin_y: int, scale: float) -> None:
    xfrm = _xfrm(element)
    if xfrm is None:
        return
    off, ext = xfrm.find(f"{_A}off"), xfrm.find(f"{_A}ext")
    if off is None or ext is None:
        return
    y = int(off.get("y"))
    w = int(ext.get("cx"))
    h = int(ext.get("cy"))
    if abs(w - h) < 0.15 * max(w, h) and h < 2.5 * CM:
        off.set("y", str(origin_y + int((y - origin_y) * scale)))
        return
    off.set("y", str(origin_y + int((y - origin_y) * scale)))
    ext.set("cy", str(max(int(h * scale), 1)))


def _scale_fonts(element, factor: float) -> None:
    for rpr in element.iter(f"{_A}rPr", f"{_A}endParaRPr", f"{_A}defRPr"):
        sz = rpr.get("sz")
        if sz:
            new = max(int(int(sz) * factor), int(_MIN_FONT_PT * 100))
            rpr.set("sz", str(new))


def _all_ids(slide) -> set[int]:
    return {int(n.get("id")) for n in slide.shapes._spTree.iter(f"{_P}cNvPr") if n.get("id", "").isdigit()}


def _renumber(element, next_id: int, id_map: dict[int, int]) -> int:
    for node in element.iter(f"{_P}cNvPr"):
        old = node.get("id")
        node.set("id", str(next_id))
        if old and old.isdigit():
            id_map[int(old)] = next_id
        next_id += 1
    return next_id


# ---------------------------------------------------------------------------
# Текст
# ---------------------------------------------------------------------------


def _txBody(element):
    body = element.find(f"{_P}txBody")
    if body is None:
        body = element.find(f".//{_A}txBody")
    return body


def _segment_runs(paragraph) -> list[list]:
    """Раны абзаца, сгруппированные по сегментам между <a:br/>."""
    segments: list[list] = [[]]
    for child in paragraph:
        if child.tag == f"{_A}br":
            segments.append([])
        elif child.tag in {f"{_A}r", f"{_A}fld"}:
            segments[-1].append(child)
    return segments


def _make_run(template_run, text: str):
    run = copy.deepcopy(template_run) if template_run is not None else None
    if run is None:
        run = paragraph_run_stub()
    if run.tag == f"{_A}fld":
        # Поле (номер слайда и т.п.) превращаем в обычный ран с тем же стилем
        new = paragraph_run_stub()
        rpr = run.find(f"{_A}rPr")
        if rpr is not None:
            new.insert(0, copy.deepcopy(rpr))
        run = new
    t = run.find(f"{_A}t")
    if t is None:
        from lxml import etree
        t = etree.SubElement(run, f"{_A}t")
    t.text = text
    return run


def paragraph_run_stub():
    from lxml import etree
    run = etree.Element(f"{_A}r")
    etree.SubElement(run, f"{_A}t")
    return run


def set_segment_text(element, paragraph_index: int, segment_index: int, text: str) -> None:
    """Заменяет текст сегмента (строки) абзаца, сохраняя стиль первого рана."""
    body = _txBody(element)
    if body is None:
        return
    paragraphs = body.findall(f"{_A}p")
    if paragraph_index >= len(paragraphs):
        return
    paragraph = paragraphs[paragraph_index]
    segments = _segment_runs(paragraph)
    if segment_index >= len(segments):
        return
    runs = segments[segment_index]
    template_run = runs[0] if runs else _first_run_anywhere(body)
    lines = [line for line in text.split("\n")] or [""]
    # Точка вставки — позиция первого рана сегмента (или после соответствующего <a:br/>)
    if runs:
        anchor = runs[0]
        insert_at = list(paragraph).index(anchor)
        for run in runs:
            paragraph.remove(run)
    else:
        brs = [c for c in paragraph if c.tag == f"{_A}br"]
        insert_at = list(paragraph).index(brs[segment_index - 1]) + 1 if segment_index > 0 and brs else 0
    from lxml import etree
    for i, line in enumerate(lines):
        if i > 0:
            br = etree.Element(f"{_A}br")
            rpr = template_run.find(f"{_A}rPr") if template_run is not None else None
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            paragraph.insert(insert_at, br)
            insert_at += 1
        paragraph.insert(insert_at, _make_run(template_run, line))
        insert_at += 1


def _first_run_anywhere(body):
    for tag in (f"{_A}r", f"{_A}fld"):
        found = body.find(f".//{tag}")
        if found is not None:
            return found
    return None


def set_shape_text(element, text: str) -> None:
    """Заменяет весь текст фигуры: строки → абзацы в стиле первого абзаца.

    Строки, начинающиеся с маркера «• », получают маркер-буллет, если у абзаца
    шаблона его нет; иначе маркер убираем (шаблон сам его рисует)."""
    body = _txBody(element)
    if body is None:
        return
    paragraphs = body.findall(f"{_A}p")
    if not paragraphs:
        return
    template_p = paragraphs[0]
    template_run = _first_run_anywhere(body)
    for extra in paragraphs[1:]:
        body.remove(extra)
    lines = [line for line in text.split("\n") if line.strip()] or [""]
    has_bullet_style = template_p.find(f"{_A}pPr/{_A}buChar") is not None or template_p.find(f"{_A}pPr/{_A}buAutoNum") is not None
    from lxml import etree
    for i, line in enumerate(lines):
        p = template_p if i == 0 else copy.deepcopy(template_p)
        for child in list(p):
            if child.tag in {f"{_A}r", f"{_A}br", f"{_A}fld"}:
                p.remove(child)
        end = p.find(f"{_A}endParaRPr")
        clean = line.strip()
        if clean.startswith(("• ", "- ", "– ", "— ")):
            if has_bullet_style:
                clean = clean[2:].strip()
        run = _make_run(template_run, clean)
        if end is not None:
            end.addprevious(run)
        else:
            p.append(run)
        if i > 0:
            body.append(p)
    _ = etree


def clear_shape_text(element) -> None:
    body = _txBody(element)
    if body is None:
        return
    for t in body.iter(f"{_A}t"):
        t.text = ""


def _shape_text(element) -> str:
    body = _txBody(element)
    if body is None:
        return ""
    return "\n".join("".join(t.text or "" for t in p.iter(f"{_A}t")) for p in body.findall(f"{_A}p"))


# ---------------------------------------------------------------------------
# Повторяемые блоки
# ---------------------------------------------------------------------------


def _units_needed(variant: dict[str, Any], content: dict[str, Any]) -> int:
    comp = variant.get("composition") or {}
    unit_slots = comp.get("unit_slots") or []
    needed = 0
    for k, names in enumerate(unit_slots):
        if any((content.get(n) or "").strip() for n in names):
            needed = k + 1
    return needed


def _bottom_limit(slide, slide_h_emu: int) -> int:
    """Нижняя граница зоны контента: над декором/логотипами макета и мастера."""
    limit = slide_h_emu - int(1.1 * CM)
    for holder in (slide.slide_layout, slide.slide_layout.slide_master):
        for shape in holder.shapes:
            if shape.is_placeholder:
                continue
            if shape.top > slide_h_emu * 0.8 and shape.width < slide_h_emu * 1.2:
                limit = min(limit, shape.top - int(0.25 * CM))
    return limit


def adapt_units(slide, variant: dict[str, Any], content: dict[str, Any], slide_w_emu: int, slide_h_emu: int) -> dict[str, int]:
    """Убирает лишние блоки или достраивает недостающие. Возвращает
    отображение имени виртуального слота → shape_id созданной фигуры."""
    comp = variant.get("composition") or {}
    members: list[list[int]] = comp.get("unit_members") or []
    N = len(members)
    if N == 0:
        return {}
    n = _units_needed(variant, content)
    n = max(n, 1) if n else 0
    if n == 0:
        # Ни один блок не заполнен. На обложке/финале/разделителе пустые рамки
        # карточек — мусор: убираем блоки целиком; на остальных видах заглушки
        # вычистит cleanup (текст блоков уходит в тело слайда)
        if (comp.get("kind") or "") in {"cover", "closing", "section", "statement", "qa", "cta", "chart", "table"}:
            for ids in members:
                for element in (find_by_id(slide, i) for i in ids):
                    if element is None:
                        continue
                    top = _top_level(element, slide)
                    parent = top.getparent()
                    if parent is not None:
                        parent.remove(top)
        return {}
    grid = comp.get("grid") or {}
    if n == N or grid.get("fixed"):
        # Блоков ровно столько, сколько в образце (или ярлыки лежат на декоре) —
        # геометрию образца не трогаем: наклонённые бирки 01–04 остаются на своих местах
        return {}
    R, C = int(grid.get("rows") or 1), int(grid.get("cols") or N)

    unit_elements: list[list] = []
    for ids in members:
        els = [e for e in (find_by_id(slide, i) for i in ids) if e is not None]
        unit_elements.append([_top_level(e, slide) for e in els])
    # Уникальные верхнеуровневые элементы каждого блока
    unit_elements = [list({id(e): e for e in els}.values()) for els in unit_elements]

    boxes = [_union_bbox(els) for els in unit_elements]
    if any(b is None for b in boxes):
        return {}
    uw = max(b[2] for b in boxes)
    uh = max(b[3] for b in boxes)
    xs = sorted({b[0] for b in boxes})
    ys = sorted({b[1] for b in boxes})
    px = min((b - a for a, b in zip(xs, xs[1:])), default=uw + int(0.5 * CM))
    py = min((b - a for a, b in zip(ys, ys[1:])), default=uh + int(0.5 * CM))
    # Разброс координат внутри одного ряда/колонки (плашки на 5.50 и 5.52 см)
    # — не шаг сетки: шаг не меньше размера блока с зазором
    if C > 1 and len(xs) > 1 and px >= 0.5 * uw:
        px = max(px, int(0.9 * uw))
    else:
        px = uw + int(0.5 * CM)
    if R > 1 and len(ys) > 1 and py >= 0.5 * uh:
        py = max(py, int(0.9 * uh))
    else:
        py = uh + int(0.5 * CM)
    x0, y0 = boxes[0][0], boxes[0][1]
    span_w = C * px - (px - uw)
    # Образец с выходом карточек за правый край (намеренный «bleed») — при
    # перестройке ряд должен уместиться в слайд с полем, равным левому
    slide_w_emu = int(slide.part.package.presentation_part.presentation.slide_width)
    span_w = min(span_w, slide_w_emu - x0 - max(x0, int(0.6 * CM)))
    bottom = _bottom_limit(slide, slide_h_emu)
    unit_ids = {id(e) for els in unit_elements for e in els}
    units_bottom = max(b[1] + b[3] for b in boxes)
    units_left, units_right = min(b[0] for b in boxes), max(b[0] + b[2] for b in boxes)
    for other in slide.shapes._spTree:
        if id(other) in unit_ids or other.tag not in {f"{_P}sp", f"{_P}pic", f"{_P}grpSp", f"{_P}graphicFrame", f"{_P}cxnSp"}:
            continue
        ob = _bbox(other)
        if ob is None or ob[2] < int(1.0 * CM) or ob[3] < int(0.4 * CM):
            continue
        if ob[0] + ob[2] <= units_left or ob[0] >= units_right:
            continue
        if ob[1] >= units_bottom - int(0.2 * CM):
            bottom = min(bottom, ob[1] - int(0.25 * CM))

    def _place(cell_positions: list[tuple[int, int]], els_by_unit: list[list]) -> None:
        for (tx, ty), els in zip(cell_positions, els_by_unit):
            box = _union_bbox(els)
            if box is None:
                continue
            dx, dy = tx - box[0], ty - box[1]
            if dx or dy:
                for e in els:
                    _move(e, dx, dy)

    virtual_map: dict[str, int] = {}
    if n < N:
        for els in unit_elements[n:]:
            for e in els:
                parent = e.getparent()
                if parent is not None:
                    parent.remove(e)
        kept = unit_elements[:n]
        if grid.get("diagonal"):
            return virtual_map  # лестницу не перестраиваем — только убираем лишние ступени
        if R == 1:
            sparse = uw < 0.55 * px
            if sparse and n > 1:
                step = (span_w - uw) // (n - 1)
                _place([(x0 + i * step, y0) for i in range(n)], kept)
            else:
                # плотные карточки в один ряд: растягиваем на всю ширину ряда
                gap = max(px - uw, int(0.3 * CM))
                uw_new = (span_w - gap * (n - 1)) / n
                scale = uw_new / uw
                if 1.0 < scale <= 1.7:
                    for els in kept:
                        box = _union_bbox(els)
                        for e in els:
                            _scale_x(e, box[0], scale)
                    _place([(x0 + i * int(uw_new + gap), y0) for i in range(n)], kept)
                elif scale > 1.7:
                    # слишком мало блоков для ряда — центрируем
                    total = n * uw + (n - 1) * (px - uw)
                    start = x0 + (span_w - total) // 2
                    _place([(start + i * px, y0) for i in range(n)], kept)
        elif C == 1:
            pass  # колонка: просто короче
        else:
            best_cols = C
            best_empty = None
            for cols in range(1, C + 1):
                rows = math.ceil(n / cols)
                if rows > R:
                    continue
                empty = cols * rows - n
                if best_empty is None or empty < best_empty or (empty == best_empty and cols > best_cols):
                    best_cols, best_empty = cols, empty
            positions = [(x0 + (i % best_cols) * px, y0 + (i // best_cols) * py) for i in range(n)]
            _place(positions, kept)
        return virtual_map

    # n > N: достраиваем блоки клонированием последнего. Кандидаты раскладки:
    # добавить ряды / добавить колонки (сжать по X) / сжать по Y / ряд+колонки.
    template_els = unit_elements[-1]
    extra = n - N
    gap_x = max(px - uw, int(0.3 * CM))
    avail_h = bottom - y0
    kind = comp.get("kind") or ""
    candidates: list[tuple[float, int, int, float, float]] = []

    def _consider(cols_new: int, rows_new: int, bias: float = 0.0) -> None:
        if cols_new * rows_new < n or cols_new < 1 or rows_new < 1:
            return
        uw_new = min(uw, (span_w - gap_x * (cols_new - 1)) / cols_new)
        sx = min(uw_new / uw, 1.0)
        need_h = rows_new * py - (py - uh)
        sy = min(avail_h / need_h, 1.0) if need_h > 0 else 1.0
        if sx < 0.55 or sy < 0.6:
            return
        penalty = (1 - sx) * 1.0 + (1 - sy) * 1.3 + bias + 0.02 * (cols_new * rows_new - n)
        # «Сирота»: последний ряд из одного блока при широких рядах (3+3+1) хуже, чем 4+3
        last_row = n - cols_new * (rows_new - 1)
        if rows_new > 1 and cols_new >= 3 and last_row == 1:
            penalty += 0.3
        candidates.append((penalty, cols_new, rows_new, sx, sy))

    single_row_pref = R == 1 and (kind == "timeline" or C >= 3)
    _consider(C, math.ceil(n / C), bias=0.0 if not single_row_pref else 0.12)          # ряды
    if R == 1 or C > 1:
        _consider(math.ceil(n / R), R, bias=0.08 if not single_row_pref else 0.0)      # колонки
    for rows_new in range(R + 1, R + 3):                                                # ряд(ы) + колонки
        _consider(math.ceil(n / rows_new), rows_new, bias=0.1)
    if not candidates:
        logger.info("Блоки: %d не помещаются, ограничиваем до %d", n, N)
        return {}
    _, cols_new, rows_new, scale_x, scale_y = min(candidates)

    slide_ids = _all_ids(slide)
    next_id = max(slide_ids | {1}) + 1
    tree = slide.shapes._spTree
    new_units: list[list] = []
    virtual_slots = [(name, slot) for name, slot in (variant.get("slots") or {}).items() if slot.get("virtual")]
    for k in range(extra):
        id_map: dict[int, int] = {}
        clones = []
        for e in template_els:
            clone = copy.deepcopy(e)
            next_id = _renumber(clone, next_id, id_map)
            tree.append(clone)
            clones.append(clone)
        new_units.append(clones)
        unit_index = N + k
        for name, slot in virtual_slots:
            if slot.get("unit") == unit_index:
                src_id = (variant["slots"].get(slot.get("virtual_from")) or {}).get("shape_id")
                if src_id in id_map:
                    virtual_map[name] = id_map[src_id]
    all_units = unit_elements + new_units

    px_new = int(px * scale_x) if scale_x != 1.0 else px
    py_new = int(py * scale_y) if scale_y != 1.0 else py
    if scale_x != 1.0 or scale_y != 1.0:
        for els in all_units:
            box = _union_bbox(els)
            if box is None:
                continue
            for e in els:
                if scale_x != 1.0:
                    _scale_x(e, box[0], scale_x)
                if scale_y != 1.0:
                    _scale_y(e, box[1], scale_y)
                factor = min(scale_x, scale_y)
                if factor < 0.92:
                    _scale_fonts(e, max(factor, 0.72))
    positions = [(x0 + (i % cols_new) * px_new, y0 + (i // cols_new) * py_new) for i in range(n)]
    _place(positions, all_units)
    return virtual_map


def _union_bbox(elements) -> tuple[int, int, int, int] | None:
    boxes = [b for b in (_bbox(e) for e in elements) if b is not None]
    if not boxes:
        return None
    x0 = min(b[0] for b in boxes)
    y0 = min(b[1] for b in boxes)
    x1 = max(b[0] + b[2] for b in boxes)
    y1 = max(b[1] + b[3] for b in boxes)
    return x0, y0, x1 - x0, y1 - y0


# ---------------------------------------------------------------------------
# Таблицы и графики
# ---------------------------------------------------------------------------


def fill_table(slide, shape_id: int, headers: list[str], rows: list[list[str]]) -> None:
    frame = find_by_id(slide, shape_id)
    if frame is None:
        return
    tbl = frame.find(f".//{_A}tbl")
    if tbl is None:
        return
    grid = tbl.find(f"{_A}tblGrid")
    cols = grid.findall(f"{_A}gridCol")
    trs = tbl.findall(f"{_A}tr")
    if not trs or not cols:
        return
    target_cols = max(len(headers), 1)
    total_w = sum(int(c.get("w")) for c in cols)
    # колонки
    while len(cols) > target_cols:
        grid.remove(cols[-1])
        for tr in trs:
            tcs = tr.findall(f"{_A}tc")
            tr.remove(tcs[-1])
        cols = grid.findall(f"{_A}gridCol")
    while len(cols) < target_cols:
        grid.append(copy.deepcopy(cols[-1]))
        for tr in trs:
            tcs = tr.findall(f"{_A}tc")
            tr.append(copy.deepcopy(tcs[-1]))
        cols = grid.findall(f"{_A}gridCol")
    for c in cols:
        c.set("w", str(total_w // len(cols)))
    # строки: первая — шапка, остальные — данные (стиль берём у 2-й строки, если есть)
    data_template = trs[1] if len(trs) > 1 else trs[0]
    target_rows = 1 + len(rows)
    while len(trs) > target_rows:
        tbl.remove(trs[-1])
        trs = tbl.findall(f"{_A}tr")
    while len(trs) < target_rows:
        tbl.append(copy.deepcopy(data_template))
        trs = tbl.findall(f"{_A}tr")
    # Высота строк: вписываем в исходную высоту рамки
    frame_box = _bbox(frame)
    if frame_box:
        row_h = max(int(frame_box[3] / target_rows), int(0.7 * CM))
        for tr in trs:
            tr.set("h", str(row_h))
    values = [headers] + rows
    for tr, row in zip(trs, values):
        tcs = tr.findall(f"{_A}tc")
        for i, tc in enumerate(tcs):
            text = str(row[i]) if i < len(row) else ""
            body = tc.find(f"{_A}txBody")
            if body is None:
                continue
            paragraphs = body.findall(f"{_A}p")
            for extra in paragraphs[1:]:
                body.remove(extra)
            p = paragraphs[0]
            template_run = _first_run_anywhere(body)
            for child in list(p):
                if child.tag in {f"{_A}r", f"{_A}br", f"{_A}fld"}:
                    p.remove(child)
            run = _make_run(template_run, text)
            end = p.find(f"{_A}endParaRPr")
            if end is not None:
                end.addprevious(run)
            else:
                p.append(run)


def replace_with_chart(slide, shape_id: int, chart_block, manifest, chart_type_map) -> bool:
    element = find_by_id(slide, shape_id)
    if element is None:
        return False
    box = _bbox(element)
    if box is None:
        return False
    from app.pipeline.styling import style_chart

    left, top, width, height = box
    chart_data = CategoryChartData()
    chart_data.categories = chart_block.categories
    for series in chart_block.series:
        chart_data.add_series(series.name, series.values)
    frame = slide.shapes.add_chart(chart_type_map[chart_block.chart_type.value], Emu(left), Emu(top), Emu(width), Emu(height), chart_data)
    if manifest is not None:
        try:
            style_chart(frame.chart, manifest, slide=slide)
        except Exception:  # noqa: BLE001
            logger.exception("Не удалось стилизовать график")
    parent = element.getparent()
    if parent is not None:
        # Новый график ставим на место картинки по z-порядку
        index = list(parent).index(element)
        parent.remove(element)
        parent.remove(frame._element)
        parent.insert(index, frame._element)
    return True


def replace_with_picture(slide, shape_id: int, image_path: str) -> bool:
    element = find_by_id(slide, shape_id)
    if element is None:
        return False
    box = _bbox(element)
    if box is None:
        return False
    left, top, width, height = box
    pic = slide.shapes.add_picture(image_path, Emu(left), Emu(top))
    # Заливка рамки с сохранением пропорций (обрезка по центру)
    iw, ih = pic.width, pic.height
    if iw and ih:
        target = width / height
        actual = iw / ih
        if actual > target:
            crop = (1 - target / actual) / 2
            pic.crop_left = pic.crop_right = crop
        else:
            crop = (1 - actual / target) / 2
            pic.crop_top = pic.crop_bottom = crop
    pic.left, pic.top, pic.width, pic.height = Emu(left), Emu(top), Emu(width), Emu(height)
    # Скругление/геометрия исходной фигуры
    src_geom = element.find(f"{_P}spPr/{_A}prstGeom")
    if src_geom is not None:
        dst_sppr = pic._element.find(f"{_P}spPr")
        dst_geom = dst_sppr.find(f"{_A}prstGeom")
        if dst_geom is not None:
            dst_sppr.replace(dst_geom, copy.deepcopy(src_geom))
    parent = element.getparent()
    if parent is not None:
        index = list(parent).index(element)
        parent.remove(element)
        parent.remove(pic._element)
        parent.insert(index, pic._element)
    return True


# ---------------------------------------------------------------------------
# Главная точка: собрать слайд по образцу
# ---------------------------------------------------------------------------


def render_sample_slide(prs, template: dict[str, Any], variant: dict[str, Any], slide_plan: dict[str, Any], blocks, manifest, chart_type_map, fit_font) -> Any:
    """Клонирует слайд-образец и заполняет его по плану. `blocks` — визуальные
    блоки ContentPlan (графики/таблицы) для этого слайда."""
    source = prs.slides[variant["slide_index"]]
    slide = clone_slide(prs, source)
    for group_id in (variant.get("composition") or {}).get("ungroup") or []:
        _ungroup(slide, group_id)
    slide_w, slide_h = prs.slide_width, prs.slide_height
    content: dict[str, Any] = {k: v for k, v in (slide_plan.get("content") or {}).items() if isinstance(v, str)}
    comp = variant.get("composition") or {}
    slots = variant.get("slots") or {}

    # Автонумерация блоков, если планировщик не дал номера
    for names in comp.get("unit_slots") or []:
        for name in names:
            if slots.get(name, {}).get("semantic_role") == "unit_number" and not (content.get(name) or "").strip():
                siblings = [n for n in names if n != name]
                if any((content.get(s) or "").strip() for s in siblings):
                    number = (slots[name].get("unit") or 0) + 1
                    sample = (slots[name].get("template_text") or "").strip()
                    # формат как в образце: «1» → без ведущего нуля, «01» → с нулём, «1.» → с точкой
                    padded = bool(re.match(r"^0\d", sample)) or not sample
                    value = f"{number:02d}" if padded else str(number)
                    if sample.endswith("."):
                        value += "."
                    content[name] = value

    virtual_map = adapt_units(slide, variant, content, slide_w, slide_h)

    filled_ids: set[int] = set()
    filled_segments: dict[int, set[int]] = {}  # shape_id → заполненные абзацы (сегментные слоты)
    for name, text in content.items():
        slot = slots.get(name)
        if slot is None:
            continue
        shape_id = virtual_map.get(name) if slot.get("virtual") else slot.get("shape_id")
        if shape_id is None:
            continue
        element = find_by_id(slide, shape_id)
        if element is None:
            continue
        clean = text.strip()
        if not clean:
            continue
        if slot.get("paragraph") is not None:
            set_segment_text(element, int(slot["paragraph"]), int(slot.get("line") or 0), clean)
            filled_segments.setdefault(shape_id, set()).add(int(slot["paragraph"]))
        else:
            set_shape_text(element, clean)
        filled_ids.add(shape_id)
        _set_name(element, name if slot.get("paragraph") is None else f"{name}_shape")
        # shrink-to-fit для целых фигур
        if slot.get("paragraph") is None and fit_font is not None:
            box = _bbox(element)
            size_pt = _font_pt(element) or _inherited_font_pt(slide, element) or (slot.get("resolved_font") or {}).get("size_pt")
            if box and size_pt:
                height = box[3]
                if slot.get("semantic_role") == "unit_title":
                    # заголовку карточки доступно место до следующей фигуры под ним
                    height = max(box[3], _room_below(slide, element, box), int(0.8 * CM))
                if slot.get("semantic_role") in _TITLE_ROLES:
                    box = _avoid_decor_right(slide, element, box, clean, float(size_pt))
                    height = max(_room_below(slide, element, box), int(1.2 * CM))
                    box = _widen_title(slide, element, box, clean, float(size_pt))
                    box = _avoid_decor_right(slide, element, box, clean, float(size_pt))  # расширение не должно залезть на декор
                start_pt = float(size_pt)
                if slot.get("semantic_role") == "unit_title" and _METRIC_VALUE_RE.match(clean):
                    # число-метрика в заголовке карточки: укрупняем, пока помещается
                    start_pt = float(size_pt) * 1.8
                elif slot.get("semantic_role") == "unit_title" and slot.get("unit") is not None and not any(
                    (content.get(n) or "").strip() for n, sl in slots.items()
                    if sl.get("unit") == slot.get("unit") and sl.get("semantic_role") in {"unit_text", "unit_label"} and n != name
                ):
                    # карточка только с заголовком (лаконичный вариант): заголовок
                    # заметнее — до 1.35× кегля образца, если хватает места
                    start_pt = float(size_pt) * 1.35
                if slot.get("semantic_role") in {"metric_value", "unit_number"} and "\n" not in clean:
                    # Значение метрики — строго одной строкой («150 млн», не «150 / млн»):
                    # кегль по ширине бокса с запасом на широкий шрифт, но не ниже 55 %
                    inner_w = max(1, box[2] - 2 * 91440)
                    one_line = inner_w / max(1, len(clean)) / (0.62 * 12700)
                    fitted = max(float(size_pt) * 0.55, min(start_pt, math.floor(one_line)))
                else:
                    fitted = fit_font(clean, box[2], height, start_pt, check_height=True)
                if fitted is not None and abs(fitted - float(size_pt)) > 0.5:
                    _set_font_pt(element, fitted)

    # Незаполненные слоты: убираем плашки-теги/заметки, чистим текст-заглушки
    removable = set(comp.get("removable_when_empty") or [])
    for name, slot in slots.items():
        if slot.get("virtual") or name in content and content[name].strip():
            continue
        element = find_by_id(slide, slot.get("shape_id"))
        if element is None:
            continue
        if slot.get("shape_id") in filled_ids:
            # Другой сегмент этой фигуры заполнен — чистим только этот сегмент
            if slot.get("paragraph") is not None:
                set_segment_text(element, int(slot["paragraph"]), int(slot.get("line") or 0), "")
            continue
        if name in removable:
            _remove_with_plate(slide, element, filled_ids)
            continue
        if slot.get("paragraph") is not None:
            set_segment_text(element, int(slot["paragraph"]), int(slot.get("line") or 0), "")
        else:
            clear_shape_text(element)

    # Таблица
    table_id = comp.get("table_shape_id")
    table_blocks = [b for b in blocks if getattr(b, "table", None) is not None]
    if table_id and table_blocks:
        tb = table_blocks[0].table
        fill_table(slide, table_id, tb.headers[:6], [r[:6] for r in tb.rows[:8]])

    # Графики: картинки-заглушки → нативные графики; лишние картинки убираем
    chart_blocks = [b for b in blocks if getattr(b, "chart", None) is not None]
    for i, pic_id in enumerate(comp.get("chart_pictures") or []):
        if i < len(chart_blocks):
            replace_with_chart(slide, pic_id, chart_blocks[i].chart, manifest, chart_type_map)
        else:
            element = find_by_id(slide, pic_id)
            if element is not None and element.getparent() is not None:
                element.getparent().remove(element)

    # Изображения
    images = slide_plan.get("image_content") or {}
    for name, spec in (variant.get("image_slots") or {}).items():
        path = images.get(name)
        if path and Path(str(path)).is_file():
            replace_with_picture(slide, spec["shape_id"], str(path))
        elif spec.get("source_kind") == "placeholder_shape":
            element = find_by_id(slide, spec["shape_id"])
            if element is not None and element.getparent() is not None:
                _remove_with_plate(slide, element, filled_ids)
        # Подпись-заглушку («Вставить фото») убираем всегда — и с картинкой, и без
        if spec.get("caption_shape_id"):
            caption = find_by_id(slide, spec["caption_shape_id"])
            if caption is not None and caption.getparent() is not None:
                caption.getparent().remove(caption)

    # Частично заполненные фигуры: абзацы образца, не ставшие слотами, убираем
    for shape_id, paragraphs in filled_segments.items():
        element = find_by_id(slide, shape_id)
        body = _txBody(element) if element is not None else None
        if body is None:
            continue
        all_ps = body.findall(f"{_A}p")
        for p_index, paragraph in enumerate(all_ps):
            if p_index in paragraphs:
                continue
            if "".join(t.text or "" for t in paragraph.iter(f"{_A}t")).strip() and len(body.findall(f"{_A}p")) > 1:
                body.remove(paragraph)  # абзац образца целиком (индексы слотов уже не нужны)

    # Остатки текста образца в незаполненных фигурах: любой текст, который мы
    # не заполняли, — это содержимое шаблона-примера, ему не место в результате
    for node in list(slide.shapes._spTree.iter(f"{_P}sp")):
        text = _shape_text(node).strip()
        if not text:
            continue
        cnv = node.find(f"{_P}nvSpPr/{_P}cNvPr")
        sid = int(cnv.get("id")) if cnv is not None and cnv.get("id", "").isdigit() else None
        if sid in filled_ids:
            continue
        clear_shape_text(node)
    return slide


def _avoid_decor_right(slide, element, box, text: str = "", size_pt: float | None = None):
    """Зона исключения: декор (картинки, плашки без текста), перекрывающий
    правую часть заголовка, — сужаем рамку заголовка до его левого края.

    Не сужаем так, чтобы самое длинное слово перестало помещаться в строку
    даже при минимальном кегле: разрыв слова хуже, чем наезд на свечение
    картинки (у декора часто прозрачные поля).
    """
    x0, y0, w, h = box
    prs = slide.part.package.presentation_part.presentation
    slide_area = int(prs.slide_width) * int(prs.slide_height)
    new_right = x0 + w
    candidates = list(slide.shapes._spTree)
    try:  # декор макета (не заполнители) тоже занимает место
        candidates += [
            node for node in slide.slide_layout.shapes._spTree
            if node.tag in {f"{_P}sp", f"{_P}pic", f"{_P}grpSp"} and node.find(f".//{_P}nvPr/{_P}ph") is None
        ]
    except Exception:  # noqa: BLE001
        pass
    for other in candidates:
        if other is element or other.tag not in {f"{_P}sp", f"{_P}pic", f"{_P}grpSp"}:
            continue
        ob = _bbox(other)
        if ob is None or ob[2] * ob[3] > 0.6 * slide_area or ob[2] < int(1.0 * CM) or ob[3] < int(1.0 * CM):
            continue
        if _shape_text(other).strip():
            continue
        if ob[1] + ob[3] <= y0 + int(0.2 * CM) or ob[1] >= y0 + h - int(0.2 * CM):
            continue  # не в полосе заголовка
        if ob[0] <= x0 + 0.3 * w or ob[0] >= x0 + w:
            continue  # либо слева/под текстом целиком, либо уже справа
        new_right = min(new_right, ob[0] - int(0.5 * CM))
    if text and size_pt:
        longest = max((len(word) for word in text.split()), default=0)
        min_size = float(size_pt) * 0.7
        needed = int(longest * min_size * 0.62 * 12700) + 2 * 91440
        if new_right - x0 < needed:
            new_right = min(x0 + w, x0 + needed)
    if new_right < x0 + w and new_right - x0 >= 0.4 * w:
        xfrm = _xfrm(element)
        ext = xfrm.find(f"{_A}ext") if xfrm is not None else None
        if ext is not None:
            ext.set("cx", str(new_right - x0))
            return (x0, y0, new_right - x0, h)
    return box


def _widen_title(slide, element, box, text: str, size_pt: float):
    """Узкий заголовок, справа от которого пусто, растягиваем до правого поля."""
    from app.pipeline.template_engine import _title_height

    prs_w = int(slide.part.package.presentation_part.presentation.slide_width)
    x0, y0, w, h = box
    if w >= prs_w * 0.6 or _title_height(text, max(1, w - 2 * 91440), size_pt) <= h:
        return box
    new_right = prs_w - max(x0, int(prs_w * 0.04))
    for other in slide.shapes._spTree:
        if other is element or other.tag not in {f"{_P}sp", f"{_P}pic", f"{_P}grpSp", f"{_P}graphicFrame", f"{_P}cxnSp"}:
            continue
        ob = _bbox(other)
        if ob is None or ob[2] < int(0.5 * CM) or ob[3] < int(0.3 * CM):
            continue
        if ob[2] * ob[3] > 0.3 * prs_w * prs_w * 9 / 16:
            continue
        if ob[0] < x0 + w - int(0.1 * CM):
            continue
        if ob[1] + ob[3] <= y0 or ob[1] >= y0 + h:
            continue
        new_right = min(new_right, ob[0] - int(0.5 * CM))
    for holder in (slide.slide_layout, slide.slide_layout.slide_master):
        for shape in holder.shapes:
            if shape.is_placeholder or shape.left is None:
                continue
            if int(shape.width) * int(shape.height) > 0.3 * prs_w * prs_w * 9 / 16:
                continue
            if int(shape.left) < x0 + w or int(shape.top) + int(shape.height) <= y0 or int(shape.top) >= y0 + h:
                continue
            new_right = min(new_right, int(shape.left) - int(0.5 * CM))
    # Фоновая картинка макета с нарисованной плашкой справа (белый «пузырь» финала
    # VK Tech): пикселей-фигур нет, но растягивать заголовок на плашку нельзя
    new_right = min(new_right, _background_flat_until(slide, x0 + w, y0, y0 + h, new_right))
    if new_right - x0 <= w * 1.15:
        return box
    xfrm = _xfrm(element)
    ext = xfrm.find(f"{_A}ext") if xfrm is not None else None
    if ext is None:
        return box
    ext.set("cx", str(new_right - x0))
    return (x0, y0, new_right - x0, h)


def _background_flat_until(slide, x_from: int, y0: int, y1: int, x_to: int) -> int:
    """Правая граница ровного фона в полосе y0..y1 начиная с x_from: анализирует
    фоновые картинки макета/мастера (≥ 75 % слайда). Если в полосе справа фон
    неровный (плашка, объект), возвращает x, где он перестаёт быть ровным."""
    try:
        from io import BytesIO

        from PIL import Image
    except ImportError:  # pragma: no cover
        return x_to
    prs = slide.part.package.presentation_part.presentation
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    limit = x_to
    for holder in (slide.slide_layout, slide.slide_layout.slide_master):
        for shape in holder.shapes:
            if shape.is_placeholder or "PICTURE" not in str(shape.shape_type) or shape.left is None:
                continue
            pw, ph = int(shape.width), int(shape.height)
            if pw * ph < 0.75 * sw * sh:
                continue
            try:
                image = Image.open(BytesIO(shape.image.blob)).convert("L")
            except Exception:  # noqa: BLE001
                continue
            gx = image.width / pw
            gy = image.height / ph
            py0 = max(0, int((y0 - int(shape.top)) * gy))
            py1 = min(image.height, int((y1 - int(shape.top)) * gy))
            if py1 - py0 < 2:
                continue
            step = max(1, int(0.3 * CM * gx))  # шаг ~3 мм
            px_from = max(0, int((x_from - int(shape.left)) * gx))
            px_to = min(image.width, int((x_to - int(shape.left)) * gx))
            base = list(image.crop((max(0, px_from - 12 * step), py0, max(1, px_from), py1)).getdata())
            if not base:
                continue
            base_sorted = sorted(base)
            median = base_sorted[len(base_sorted) // 2]
            if base_sorted[-1] - base_sorted[0] > 40:
                continue  # под самим заголовком фон неровный (градиент, объект) — судить нечего
            for px in range(px_from, px_to, step):
                column = list(image.crop((px, py0, min(image.width, px + step), py1)).getdata())
                if not column:
                    break
                col_median = sorted(column)[len(column) // 2]
                spread = max(column) - min(column)
                if abs(col_median - median) > 10 or spread > 40:
                    limit = min(limit, int(shape.left) + int(px / gx) - int(0.4 * CM))
                    break
    return limit


def _set_font_pt(element, size_pt: float) -> None:
    """Явно задаёт кегль всем ранам фигуры (в т.ч. когда размер наследуется от макета)."""
    from lxml import etree
    body = _txBody(element)
    if body is None:
        return
    sz = str(int(round(size_pt * 100)))
    for run in body.iter(f"{_A}r"):
        rpr = run.find(f"{_A}rPr")
        if rpr is None:
            rpr = etree.Element(f"{_A}rPr")
            run.insert(0, rpr)
        rpr.set("sz", sz)
    for rpr in body.iter(f"{_A}endParaRPr"):
        rpr.set("sz", sz)


def _set_name(element, name: str) -> None:
    cnv = element.find(f".//{_P}cNvPr")
    if cnv is not None:
        cnv.set("name", name)


def _remove_with_plate(slide, element, filled_ids: set[int]) -> None:
    """Удаляет фигуру слота и «плашку» под ней: пустые фигуры без текста,
    чей прямоугольник целиком содержит фигуру слота (фон заметки, иконка в ней)."""
    box = _bbox(element)
    tree = slide.shapes._spTree
    top = _top_level(element, slide)
    victims = [top if top is not None else element]
    if box is not None:
        x0, y0, w, h = box
        slide_area = int(slide.part.package.presentation_part.presentation.slide_width) * int(slide.part.package.presentation_part.presentation.slide_height)
        for other in list(tree):
            if other is top or other.tag not in {f"{_P}sp", f"{_P}pic", f"{_P}grpSp"}:
                continue
            ob = _bbox(other)
            if ob is None or ob[2] * ob[3] > 0.6 * slide_area:
                continue
            cnv = other.find(f".//{_P}cNvPr")
            oid = int(cnv.get("id")) if cnv is not None and cnv.get("id", "").isdigit() else None
            if oid in filled_ids or _shape_text(other).strip():
                continue
            tol = int(0.35 * CM)
            contains = ob[0] <= x0 + tol and ob[1] <= y0 + tol and ob[0] + ob[2] >= x0 + w - tol and ob[1] + ob[3] >= y0 + h - tol
            inside = ob[0] >= x0 - tol and ob[1] >= y0 - tol and ob[0] + ob[2] <= x0 + w + tol and ob[1] + ob[3] <= y0 + h + tol
            if contains or inside:
                victims.append(other)
        # плашка нашлась — заодно убираем мелкие фигуры внутри плашки (иконки)
        plates = [v for v in victims[1:] if _bbox(v) and _bbox(v)[2] * _bbox(v)[3] >= w * h]
        for plate in plates:
            pb = _bbox(plate)
            for other in list(tree):
                if other in victims or other.tag not in {f"{_P}sp", f"{_P}pic", f"{_P}grpSp"}:
                    continue
                ob = _bbox(other)
                if ob is None or _shape_text(other).strip():
                    continue
                if ob[0] >= pb[0] - tol and ob[1] >= pb[1] - tol and ob[0] + ob[2] <= pb[0] + pb[2] + tol and ob[1] + ob[3] <= pb[1] + pb[3] + tol:
                    victims.append(other)
    for victim in victims:
        parent = victim.getparent()
        if parent is not None:
            parent.remove(victim)


def _room_below(slide, element, box) -> int:
    """Высота от верха фигуры до ближайшей фигуры ниже с перекрытием по X (минус зазор)."""
    x0, y0, w, h = box
    below: list[int] = []
    for other in slide.shapes._spTree:
        if other is element or other.tag not in {f"{_P}sp", f"{_P}pic", f"{_P}grpSp", f"{_P}graphicFrame"}:
            continue
        ob = _bbox(other)
        if ob is None or ob[2] < int(0.8 * CM) or ob[3] < int(0.5 * CM):
            continue
        overlap = min(x0 + w, ob[0] + ob[2]) - max(x0, ob[0])
        if overlap < max(int(1.0 * CM), min(0.4 * w, 0.6 * ob[2])):
            continue  # фигура лишь слегка заходит под заголовок
        if ob[1] >= y0 + int(0.6 * CM):
            below.append(ob[1] - int(0.25 * CM))
    return (min(below) if below else y0 + h) - y0


def _font_pt(element) -> float | None:
    for rpr in element.iter(f"{_A}rPr"):
        if rpr.get("sz"):
            return int(rpr.get("sz")) / 100.0
    for rpr in element.iter(f"{_A}defRPr", f"{_A}endParaRPr"):
        if rpr.get("sz"):
            return int(rpr.get("sz")) / 100.0
    return None


def _inherited_font_pt(slide, element) -> float | None:
    """Кегль плейсхолдера, унаследованный от макета/мастера (когда в самой
    фигуре размер не задан): по idx/type заполнителя макета, затем стили
    заголовка/тела мастера."""
    ph = element.find(f".//{_P}nvPr/{_P}ph")
    if ph is None:
        return None
    idx, ph_type = ph.get("idx"), ph.get("type") or "body"
    try:
        layout = slide.slide_layout
        for shape in layout.placeholders:
            lp = shape._element.find(f".//{_P}nvPr/{_P}ph")
            if lp is None:
                continue
            same_idx = idx is not None and lp.get("idx") == idx
            same_type = (lp.get("type") or "body") == ph_type and ph_type in {"title", "ctrTitle", "subTitle"}
            if same_idx or same_type:
                size = _font_pt(shape._element)
                if size:
                    return size
        master = layout.slide_master
        style_tag = "titleStyle" if ph_type in {"title", "ctrTitle"} else "bodyStyle"
        style = master._element.find(f".//{_P}txStyles/{_P}{style_tag}/{_A}lvl1pPr/{_A}defRPr")
        if style is not None and style.get("sz"):
            return int(style.get("sz")) / 100.0
    except Exception:  # noqa: BLE001
        return None
    return None


_TITLE_ROLES = {"slide_title", "presentation_title", "closing_title", "section_title"}
_METRIC_VALUE_RE = re.compile(r"^[<>~≈±+\-−×x]?\s*\d[\d\s.,]*\s*(%|‰|×|x|раз|тыс\.?|млн|млрд|₽|\$|€|руб\.?|ч|мин|дн\.?|k|K|M)?\s*[₽$€%]?$|^×\s*\d+([.,]\d+)?$", re.I)
_ = (re, Pt, math)
