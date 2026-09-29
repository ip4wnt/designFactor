"""HTTP API поверх Job-оркестратора."""
from __future__ import annotations

import asyncio
import logging
import tempfile
import uuid
from pathlib import Path

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel

from app.config import get_settings
from app.job_store import get as get_job
from app.job_store import save as save_job
from app.orchestrator import (
    SLIDE_EDIT_ACTIONS,
    edit_slide,
    fix_audit_issues,
    rerun_assembly,
    run_audit_for_selected_variant,
    run_pipeline,
    select_variant,
    skip_audit,
)
from app.pipeline import parser as template_parser
from app.pipeline import template_engine
from app.pipeline.design_system import DesignSystem, build_design_system
from app.pipeline.excel_import import ExcelImportError, import_excel_block, list_sheet_names
from app.pipeline.export import ExportError, export_presentation
from app.pipeline.audit_fix import FIXABLE_CODES
from app.pipeline.audit_overlay import draw_issue_overlay
from app.pipeline.preview import PreviewError, render_variant_previews
from app.schemas.content_plan import ContentBlock, ContentBlockType, ContentPlan
from app.schemas.design_manifest import DesignManifest
from app.schemas.job import Job, JobStatus
from app.schemas.template_analysis import TemplateAnalysis, TemplateAnalysisStatus
from app.template_store import get as get_template_analysis
from app.template_store import save as save_template_analysis

logger = logging.getLogger(__name__)

router = APIRouter()

_EXPORT_MEDIA_TYPES = {
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
    "html": "text/html",
}


class FixRequest(BaseModel):
    issue_ids: list[str]


class SelectVariantRequest(BaseModel):
    variant: str


class SlideEditRequest(BaseModel):
    action: str  # recompose | take_from | rewrite
    source: str | None = None  # вариант-источник для take_from


class AddContentBlockRequest(BaseModel):
    """Второй вариант ввода данных таблиц/графиков — готовый JSON вручную
    (без загрузки Excel), для тест-фронта и будущей модели-генератора контента."""

    block: ContentBlock


@router.post("/templates/analyze")
async def analyze_template(template: UploadFile = File(...)) -> TemplateAnalysis:
    """Отдельная стадия «анализ шаблона»: парсит .pptx в DesignManifest и
    сохраняет результат как самостоятельную сущность TemplateAnalysis, вне
    Job/пайплайна генерации. Позволяет посмотреть/отредактировать
    дизайн-систему до того, как запускать LLM-контент-агента и сборку —
    см. GET .../design-system и PUT .../manifest.
    """
    settings = get_settings()
    template_analysis_id = str(uuid.uuid4())

    template_path = settings.templates_dir / f"{template_analysis_id}_{template.filename}"
    template_path.write_bytes(await template.read())

    analysis = TemplateAnalysis(
        template_analysis_id=template_analysis_id,
        original_filename=template.filename or "template.pptx",
        template_path=str(template_path),
    )
    try:
        analysis.manifest = template_parser.parse_template(template_path)
        # Template JSON движка строится здесь же и кэшируется в памяти по
        # (путь, mtime) — GET .../design-system и создание job переиспользуют его.
        await asyncio.to_thread(template_engine.build_template_json, template_path)
        analysis.status = TemplateAnalysisStatus.READY
    except Exception as exc:  # noqa: BLE001 — некорректный/повреждённый .pptx от клиента, не наша ошибка
        analysis.status = TemplateAnalysisStatus.FAILED
        analysis.error = str(exc)

    save_template_analysis(analysis)
    return analysis


@router.get("/templates/{template_analysis_id}")
async def get_template_analysis_endpoint(template_analysis_id: str) -> TemplateAnalysis:
    analysis = get_template_analysis(template_analysis_id)
    if analysis is None:
        raise HTTPException(status_code=404, detail="Анализ шаблона не найден")
    return analysis


@router.get("/templates/{template_analysis_id}/design-system")
async def get_design_system(template_analysis_id: str) -> DesignSystem:
    """Обогащённая, самообъясняющаяся проекция DesignManifest для UI —
    роли цветов, раскрытые имена шрифтов, сгруппированные макеты. См.
    app/pipeline/design_system.py. Не источник истины для сборки — тот
    остаётся в analysis.manifest, отдаваемом тем же ответом целиком.
    """
    analysis = get_template_analysis(template_analysis_id)
    if analysis is None:
        raise HTTPException(status_code=404, detail="Анализ шаблона не найден")
    if analysis.manifest is None:
        raise HTTPException(status_code=409, detail=f"Анализ ещё не готов (status={analysis.status.value})")
    try:
        template_json = await asyncio.to_thread(template_engine.build_template_json, analysis.template_path)
    except Exception as exc:  # noqa: BLE001 — дизайн-система должна открываться даже без Template JSON
        logger.warning("Template JSON для %s недоступен: %s", template_analysis_id, exc)
        template_json = None
    return build_design_system(template_analysis_id, analysis.manifest, template_json)


@router.put("/templates/{template_analysis_id}/manifest")
async def update_template_manifest(template_analysis_id: str, manifest: DesignManifest) -> TemplateAnalysis:
    """Ручная правка дизайн-манифеста поверх авто-разбора — например,
    поправить неверно определённую роль плейсхолдера или подставить точный
    hex вместо унаследованного из темы. Валидируется тем же контрактом
    DesignManifest, что и вывод парсера, так что сборка (assembly.py) не
    видит разницы между авто- и ручным манифестом.
    """
    analysis = get_template_analysis(template_analysis_id)
    if analysis is None:
        raise HTTPException(status_code=404, detail="Анализ шаблона не найден")

    analysis.manifest = manifest
    analysis.manually_edited = True
    analysis.status = TemplateAnalysisStatus.READY
    analysis.error = None
    save_template_analysis(analysis)
    return analysis


@router.post("/jobs")
async def create_job(
    template: UploadFile | None = File(None),
    template_analysis_id: str | None = Form(
        None,
        description=(
            "Опционально: id уже проанализированного шаблона (см. POST "
            "/templates/analyze), включая ручные правки его манифеста. "
            "Когда передан, повторный парсинг .pptx пропускается — "
            "переиспользуется сохранённый DesignManifest как есть, и "
            "поле `template` можно не передавать."
        ),
    ),
    brief: str = Form(...),
    purpose: str = Form(...),
    content_plan: str | None = Form(
        None,
        description=(
            "Опционально: готовый ContentPlan в виде JSON-строки (включая "
            "content_blocks с ChartData/TableData). Когда передан, планирование "
            "через LLM-контент-агента пропускается — это путь для сценария "
            "«модель/пользователь уже подготовили данные»."
        ),
    ),
    case_prompt_addition: str = Form(
        "",
        description=(
            "Опционально: текстовая добавка к промпту контент-агента от выбранного на "
            "фронте кейса-пресета (хранится в localStorage браузера, на backend приходит "
            "уже готовым текстом)."
        ),
    ),
    slide_count: int | None = Form(
        None,
        description="Опционально: целевое число слайдов (или нижняя граница диапазона).",
    ),
    slide_count_max: int | None = Form(
        None,
        description="Опционально: верхняя граница диапазона слайдов «от–до»; модель выбирает число внутри.",
    ),
) -> dict[str, str]:
    settings = get_settings()
    job_id = str(uuid.uuid4())

    preparsed_manifest: DesignManifest | None = None
    if template_analysis_id is not None:
        analysis = get_template_analysis(template_analysis_id)
        if analysis is None:
            raise HTTPException(status_code=404, detail="template_analysis_id не найден")
        if analysis.manifest is None:
            raise HTTPException(status_code=409, detail="Анализ шаблона ещё не готов")
        template_path = Path(analysis.template_path)
        preparsed_manifest = analysis.manifest
    elif template is not None:
        template_path = settings.templates_dir / f"{job_id}_{template.filename}"
        template_path.write_bytes(await template.read())
    else:
        raise HTTPException(status_code=400, detail="Нужно передать либо template, либо template_analysis_id")

    parsed_plan: ContentPlan | None = None
    if content_plan is not None:
        try:
            parsed_plan = ContentPlan.model_validate_json(content_plan)
        except Exception as exc:  # noqa: BLE001 — невалидный JSON от клиента, не наша ошибка
            raise HTTPException(status_code=400, detail=f"content_plan невалиден: {exc}") from exc

    job = Job(job_id=job_id, brief=brief, purpose=purpose, template_path=str(template_path))
    if preparsed_manifest is not None:
        job.design_manifest = preparsed_manifest
    save_job(job)

    # Не дожидаемся генерации — задача летит в фоне того же процесса
    # (asyncio-таск, без Celery/Redis), эндпоинт сразу возвращает job_id.
    asyncio.create_task(
        run_pipeline(
            job,
            content_plan=parsed_plan,
            case_prompt_addition=case_prompt_addition,
            slide_count=slide_count,
            slide_count_max=slide_count_max,
        )
    )

    return {"job_id": job_id}


@router.get("/jobs/{job_id}")
async def get_job_status(job_id: str) -> Job:
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    return job


@router.post("/jobs/{job_id}/fix")
async def fix_job(job_id: str, payload: FixRequest) -> dict[str, str]:
    """Пересборка всех трёх вариантов после исправления ContentPlan/манифеста.

    Возвращает job в AWAITING_VARIANT_CHOICE — выбор варианта и решение об
    аудите нужно сделать заново.
    """
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.status not in (JobStatus.DONE, JobStatus.FAILED, JobStatus.AWAITING_VARIANT_CHOICE, JobStatus.AWAITING_AUDIT_CHOICE):
        raise HTTPException(status_code=409, detail="Job ещё выполняется")
    if job.content_plan is None or job.design_manifest is None:
        raise HTTPException(status_code=409, detail="Job не дошёл до стадии сборки")

    known_ids = {
        issue.issue_id for issues in job.audit_issues.values() for issue in issues
    }
    unknown = set(payload.issue_ids) - known_ids
    if unknown:
        raise HTTPException(status_code=400, detail=f"Неизвестные issue_id: {sorted(unknown)}")

    asyncio.create_task(rerun_assembly(job))
    return {"job_id": job_id, "status": job.status.value}


@router.post("/jobs/{job_id}/select-variant")
async def select_variant_endpoint(job_id: str, payload: SelectVariantRequest) -> dict[str, str]:
    """Пользователь выбрал один из трёх собранных вариантов вёрстки на экране
    сравнения. Дальше пользователю предлагается решить, запустить ли аудит для
    выбранного варианта (POST .../audit) или пропустить его (POST .../skip-audit).
    """
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.status != JobStatus.AWAITING_VARIANT_CHOICE:
        raise HTTPException(status_code=409, detail="Job ещё не готов к выбору варианта или уже прошёл этот шаг")
    if payload.variant not in job.variant_paths:
        raise HTTPException(status_code=400, detail=f"Неизвестный вариант: {payload.variant}")

    await select_variant(job, payload.variant)
    save_job(job)
    return {"job_id": job_id, "status": job.status.value, "selected_variant": payload.variant}


@router.post("/jobs/{job_id}/audit")
async def run_audit_endpoint(job_id: str) -> dict[str, str]:
    """Пользователь выбрал «запустить аудит» для выбранного варианта.

    Асинхронно запускает детерминированный аудит (audit.py) — клиент должен опрашивать
    GET /jobs/{id} до статуса DONE/FAILED, после чего смотреть audit_issues.
    """
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    # DONE тоже допустим: повторный аудит после ручной доработки слайдов (audit_stale)
    if job.status not in (JobStatus.AWAITING_AUDIT_CHOICE, JobStatus.DONE) or job.selected_variant is None:
        raise HTTPException(status_code=409, detail="Job не готов к аудиту (вариант ещё не выбран)")

    asyncio.create_task(run_audit_for_selected_variant(job))
    return {"job_id": job_id, "status": job.status.value}


@router.post("/jobs/{job_id}/variants/{variant}/slides/{index}/edit")
async def edit_slide_endpoint(job_id: str, variant: str, index: int, payload: SlideEditRequest) -> dict[str, object]:
    """Ручная доработка одного слайда собранного варианта: другая компоновка,
    слайд из другого варианта или переписанные тексты (см. orchestrator.edit_slide).
    Синхронно: вариант перерендерен к моменту ответа, превью обновится по mtime."""
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.status not in (JobStatus.AWAITING_VARIANT_CHOICE, JobStatus.AWAITING_AUDIT_CHOICE, JobStatus.DONE):
        raise HTTPException(status_code=409, detail="Job сейчас занят — дождитесь завершения текущей операции")
    if variant not in job.variant_paths:
        raise HTTPException(status_code=404, detail=f"У job нет собранного варианта '{variant}'")
    if payload.action not in SLIDE_EDIT_ACTIONS:
        raise HTTPException(status_code=400, detail=f"Неизвестное действие: {payload.action}")
    if payload.action == "take_from" and (payload.source not in job.variant_paths or payload.source == variant):
        raise HTTPException(status_code=400, detail="Укажите другой собранный вариант-источник")
    try:
        result = await edit_slide(job, variant, index, payload.action, payload.source)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001 — сбой рендера не должен ронять задачу целиком
        logger.exception("Доработка слайда %s/%s#%s упала", job_id, variant, index)
        raise HTTPException(status_code=500, detail=f"Не удалось пересобрать слайд: {exc}") from exc
    save_job(job)
    return {"job_id": job_id, "status": job.status.value, "audit_stale": job.audit_stale, **result}


@router.post("/jobs/{job_id}/skip-audit")
async def skip_audit_endpoint(job_id: str) -> dict[str, str]:
    """Пользователь выбрал «пропустить аудит» и готов сразу к экспорту."""
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.status != JobStatus.AWAITING_AUDIT_CHOICE:
        raise HTTPException(status_code=409, detail="Job не готов к этой операции")

    skip_audit(job)
    save_job(job)
    return {"job_id": job_id, "status": job.status.value}


@router.post("/jobs/{job_id}/slides/{slide_id}/import-excel")
async def import_excel_into_slide(
    job_id: str,
    slide_id: str,
    file: UploadFile = File(...),
    sheet_name: str | None = Form(None),
    kind: str | None = Form(None),
) -> dict[str, object]:
    """Читает лист Excel, конвертирует в ChartData/TableData и добавляет
    получившийся content_block к указанному слайду ContentPlan'а.

    Не пересобирает .pptx сама — после успешного импорта нужно вызвать
    POST /jobs/{job_id}/fix, чтобы Компоновщик учёл новый content_block.
    Это то же самое разделение, что уже применяется для audit fix: сначала
    правится план, потом одна явная пересборка.
    """
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.content_plan is None:
        raise HTTPException(status_code=409, detail="Job ещё не дошёл до стадии планирования")
    if kind not in (None, "table", "chart"):
        raise HTTPException(status_code=400, detail="kind должен быть 'table', 'chart' или не задан")

    slide_spec = next((s for s in job.content_plan.slides if s.slide_id == slide_id), None)
    if slide_spec is None:
        raise HTTPException(status_code=404, detail=f"Слайд {slide_id} не найден в ContentPlan")

    suffix = Path(file.filename or "upload.xlsx").suffix or ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(await file.read())
        tmp_path = Path(tmp.name)

    try:
        resolved_kind, data = import_excel_block(tmp_path, sheet_name=sheet_name, kind=kind)
    except ExcelImportError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    finally:
        tmp_path.unlink(missing_ok=True)

    if resolved_kind == "table":
        block = ContentBlock(type=ContentBlockType.TABLE, table=data)
    else:
        block = ContentBlock(type=ContentBlockType.CHART, chart=data)

    slide_spec.content_blocks.append(block)
    save_job(job)

    return {"job_id": job_id, "slide_id": slide_id, "kind": resolved_kind, "block": block.model_dump()}


@router.post("/jobs/{job_id}/slides/{slide_id}/content-block")
async def add_content_block(job_id: str, slide_id: str, payload: AddContentBlockRequest) -> dict[str, object]:
    """Добавляет готовый ContentBlock (ChartData/TableData в виде JSON) к слайду.

    Симметричен import-excel: там лист Excel конвертируется в тот же
    ContentBlock и добавляется так же — оба пути сходятся в один контракт.
    Аналогично требует последующего POST /jobs/{job_id}/fix для пересборки.
    """
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.content_plan is None:
        raise HTTPException(status_code=409, detail="Job ещё не дошёл до стадии планирования")

    slide_spec = next((s for s in job.content_plan.slides if s.slide_id == slide_id), None)
    if slide_spec is None:
        raise HTTPException(status_code=404, detail=f"Слайд {slide_id} не найден в ContentPlan")

    slide_spec.content_blocks.append(payload.block)
    save_job(job)

    return {"job_id": job_id, "slide_id": slide_id, "block": payload.block.model_dump()}


@router.post("/excel/sheets")
async def preview_excel_sheets(file: UploadFile = File(...)) -> dict[str, list[str]]:
    """Возвращает имена листов загруженного Excel-файла для выбора перед импортом."""
    suffix = Path(file.filename or "upload.xlsx").suffix or ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(await file.read())
        tmp_path = Path(tmp.name)
    try:
        return {"sheets": list_sheet_names(tmp_path)}
    except Exception as exc:  # noqa: BLE001 — файл может быть не .xlsx вовсе
        raise HTTPException(status_code=400, detail=f"Не удалось прочитать файл: {exc}") from exc
    finally:
        tmp_path.unlink(missing_ok=True)


# Статусы, на которых job уже гарантированно имеет хотя бы один собранный .pptx в variant_paths:
# от выбора варианта и дальше. Раньше экспорт/превью требовали строго DONE, из-за
# чего экран сравнения вариантов не мог ни показать превью, ни отдать файл до выбора аудита.
_VARIANT_READY_STATUSES = {
    JobStatus.AWAITING_VARIANT_CHOICE,
    JobStatus.AWAITING_AUDIT_CHOICE,
    JobStatus.AUDITING,
    JobStatus.DONE,
}


@router.get("/jobs/{job_id}/variants/{variant}/preview")
async def preview_variant(job_id: str, variant: str) -> dict[str, object]:
    """Список URL-ов PNG-превью каждого слайда указанного варианта.

    Доступен уже на этапе AWAITING_VARIANT_CHOICE — именно тогда пользователю нужно
    увидеть все три варианта вёрстки послайдово, чтобы выбрать один из них.
    """
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.status not in _VARIANT_READY_STATUSES:
        raise HTTPException(status_code=409, detail="Варианты ещё не собраны")
    pptx_path = job.variant_paths.get(variant)
    if pptx_path is None:
        raise HTTPException(status_code=404, detail=f"У job нет собранного варианта '{variant}'")

    settings = get_settings()
    cache_dir = settings.outputs_dir / job.job_id / f"{variant}_preview"
    try:
        image_paths = await render_variant_previews(pptx_path, cache_dir)
    except PreviewError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    slide_urls = [
        f"/jobs/{job_id}/variants/{variant}/preview/{idx + 1}" for idx in range(len(image_paths))
    ]
    return {"job_id": job_id, "variant": variant, "slide_count": len(image_paths), "slides": slide_urls}


@router.get("/jobs/{job_id}/variants/{variant}/preview/{slide_number}")
async def preview_variant_slide(job_id: str, variant: str, slide_number: int) -> FileResponse:
    """Отдаёт одну PNG-картинку слайда из кэша, построенного выше через /preview."""
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.status not in _VARIANT_READY_STATUSES:
        raise HTTPException(status_code=409, detail="Варианты ещё не собраны")
    pptx_path = job.variant_paths.get(variant)
    if pptx_path is None:
        raise HTTPException(status_code=404, detail=f"У job нет собранного варианта '{variant}'")

    settings = get_settings()
    cache_dir = settings.outputs_dir / job.job_id / f"{variant}_preview"
    image_path = cache_dir / f"slide_{slide_number:02d}.png"
    if not image_path.exists():
        # Кэш мог ещё не быть построен — строим его сейчас же, чтобы прямая ссылка на картинку тоже работала.
        try:
            await render_variant_previews(pptx_path, cache_dir)
        except PreviewError as exc:
            raise HTTPException(status_code=502, detail=str(exc)) from exc
    if not image_path.exists():
        raise HTTPException(status_code=404, detail=f"Слайд {slide_number} не найден в превью")
    return FileResponse(image_path, media_type="image/png")


@router.get("/jobs/{job_id}/variants/{variant}/audit-overlay/{slide_number}")
async def audit_overlay_slide(job_id: str, variant: str, slide_number: int) -> FileResponse:
    """PNG слайда с рамками вокруг мест, которые отметил детерминированный
    аудит (см. audit_issues в GET /jobs/{id}) — чтобы было видно, ГДЕ именно
    исправлять, а не только читать текстовое описание проблемы.

    Рисуется поверх того же PNG-кэша, что и .../preview/{slide_number}, но
    сохраняется отдельным файлом — не портит кэш превью для сравнения вариантов.
    Если для этого варианта аудит ещё не запускался, отдаёт слайд без рамок
    (issue_ids будет пуст, а не ошибка) — эндпоинт не требует статуса DONE.
    """
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.status not in _VARIANT_READY_STATUSES:
        raise HTTPException(status_code=409, detail="Варианты ещё не собраны")
    pptx_path = job.variant_paths.get(variant)
    if pptx_path is None:
        raise HTTPException(status_code=404, detail=f"У job нет собранного варианта '{variant}'")

    settings = get_settings()
    cache_dir = settings.outputs_dir / job.job_id / f"{variant}_preview"
    image_path = cache_dir / f"slide_{slide_number:02d}.png"
    if not image_path.exists():
        try:
            await render_variant_previews(pptx_path, cache_dir)
        except PreviewError as exc:
            raise HTTPException(status_code=502, detail=str(exc)) from exc
    if not image_path.exists():
        raise HTTPException(status_code=404, detail=f"Слайд {slide_number} не найден в превью")

    slide_id = f"slide-{slide_number}"
    slide_issues = [issue for issue in job.audit_issues.get(variant, []) if issue.slide_id == slide_id]

    overlay_path = cache_dir / f"slide_{slide_number:02d}_audit.png"
    draw_issue_overlay(image_path, slide_issues, overlay_path)
    return FileResponse(overlay_path, media_type="image/png")


@router.get("/jobs/{job_id}/variants/{variant}/audit-map")
async def audit_map(job_id: str, variant: str) -> dict:
    """Карта аудита для экрана результатов: по каждому слайду — все
    содержательные фигуры (текст, картинки, таблицы, графики) с нормализованными
    рамками и привязанные к ним замечания. Фронтенд рисует рамки поверх
    превью: красные там, где есть замечания (кликабельны, с кнопкой
    «Исправить»), зелёные — где проверки прошли; замечания уровня слайда
    (без фигуры) выводятся списком под слайдом."""
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.status not in _VARIANT_READY_STATUSES:
        raise HTTPException(status_code=409, detail="Варианты ещё не собраны")
    pptx_path = job.variant_paths.get(variant)
    if pptx_path is None:
        raise HTTPException(status_code=404, detail=f"У job нет собранного варианта '{variant}'")

    def build() -> list[dict]:
        from pptx import Presentation

        prs = Presentation(str(pptx_path))
        sw, sh = prs.slide_width, prs.slide_height
        result = []
        for index, slide in enumerate(prs.slides, start=1):
            shapes = []
            for shape in slide.shapes:
                if None in (shape.left, shape.top, shape.width, shape.height):
                    continue
                kind = None
                if getattr(shape, "has_chart", False):
                    kind = "chart"
                elif getattr(shape, "has_table", False):
                    kind = "table"
                elif shape.shape_type == 13:
                    kind = "picture"
                elif getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                    kind = "text"
                if kind is None:
                    continue
                shapes.append(
                    {
                        "shape_id": shape.shape_id,
                        "name": shape.name,
                        "kind": kind,
                        "bbox": {
                            "x": shape.left / sw,
                            "y": shape.top / sh,
                            "width": shape.width / sw,
                            "height": shape.height / sh,
                        },
                    }
                )
            result.append({"slide_id": f"slide-{index}", "number": index, "shapes": shapes})
        return result

    slides = await asyncio.to_thread(build)
    issues = job.audit_issues.get(variant, [])
    by_slide: dict[str, list] = {}
    for issue in issues:
        payload = issue.model_dump(mode="json")
        payload["fixable"] = bool(issue.code in FIXABLE_CODES and issue.shape_id is not None)
        by_slide.setdefault(issue.slide_id, []).append(payload)
    for slide in slides:
        slide["issues"] = by_slide.get(slide["slide_id"], [])
    return {
        "job_id": job_id,
        "variant": variant,
        "status": job.status.value,
        "slides": slides,
        "deck_issues": by_slide.get("deck", []),
        "fixed_issue_ids": job.fixed_issue_ids,
    }


@router.post("/jobs/{job_id}/fix-issues")
async def fix_issues_endpoint(job_id: str, payload: FixRequest) -> dict:
    """Точечное исправление выбранных замечаний в .pptx выбранного варианта
    и повторный аудит (см. app/pipeline/audit_fix.py). Не пересобирает колоду
    и не сбрасывает выбор варианта. Выполняется синхронно (секунды; LLM
    задействуется только для сокращения текста)."""
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.status != JobStatus.DONE:
        raise HTTPException(status_code=409, detail="Исправления доступны после завершения аудита")
    if job.selected_variant is None:
        raise HTTPException(status_code=409, detail="Вариант не выбран")
    known = {issue.issue_id: issue for issue in job.audit_issues.get(job.selected_variant, [])}
    unknown = set(payload.issue_ids) - set(known)
    if unknown:
        raise HTTPException(status_code=400, detail=f"Неизвестные issue_id: {sorted(unknown)}")
    not_fixable = [i for i in payload.issue_ids if known[i].code not in FIXABLE_CODES or known[i].shape_id is None]
    applied = await fix_audit_issues(job, [i for i in payload.issue_ids if i not in not_fixable])
    return {
        "job_id": job_id,
        "status": job.status.value,
        "applied": applied,
        "not_fixable": not_fixable,
        "remaining_issues": len(job.audit_issues.get(job.selected_variant, [])),
    }


@router.get("/jobs/{job_id}/export")
async def export_job(job_id: str, format: str = "pptx", variant: str | None = None) -> FileResponse:
    """Без явно указанного `variant` экспортируется вариант, выбранный
    пользователем на экране сравнения (job.selected_variant), а не всегда variant_a.

    Доступен уже с AWAITING_VARIANT_CHOICE: на экране сравнения пользователь может скачать
    любой из трёх готовых вариантов, не дожидаясь аудита/завершения pipeline."""
    job = get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job не найден")
    if job.status not in _VARIANT_READY_STATUSES:
        raise HTTPException(status_code=409, detail="Job ещё не готов — варианты ещё не собраны")

    resolved_variant = variant or job.selected_variant or "variant_a"
    pptx_path = job.variant_paths.get(resolved_variant)
    if pptx_path is None:
        raise HTTPException(status_code=404, detail=f"У job нет собранного варианта '{resolved_variant}'")

    settings = get_settings()
    output_dir = settings.outputs_dir / job.job_id / resolved_variant
    try:
        result_path = await export_presentation(pptx_path, output_dir, format)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except ExportError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    job.export_paths[f"{resolved_variant}_{format}"] = str(result_path)

    media_type = _EXPORT_MEDIA_TYPES.get(format, "application/octet-stream")
    return FileResponse(result_path, media_type=media_type, filename=result_path.name)
