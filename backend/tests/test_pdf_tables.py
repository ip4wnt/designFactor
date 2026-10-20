"""Real PDF geometry -> prose/cards -> analytical tools -> exact planner cells."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import AsyncMock, patch

from agents.tool_context import ToolContext
from fastapi.testclient import TestClient
from pdfminer.pdfdocument import PDFPasswordIncorrect
from pdfplumber.utils.exceptions import PdfminerException
from pptx import Presentation

from app.analysis import worker
from app.api import routes
from app.config import Settings
from app.job_store import get as get_job
from app.main import app
from app.pipeline import briefing, docx_tables, pdf_tables, table_analysis
from app.pipeline.materials import MaterialError
from app.schemas.job import Job
from test_docx_tables import slide


def pdf_bytes(pages=None):
    """Small vector PDF fixture, with no optional PDF authoring dependency."""
    pages = pages if pages is not None else [{"tables": [
        [["Code", "Amount, RUB"], ["007", "1 234,50"], ["008", "0"]]]}]
    objects = [b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 3 >>\nstream\n\x00\x00\x00\nendstream"]
    kids = []
    for index, spec in enumerate(pages):
        page_number = len(objects) + 1
        content_number = page_number + 1
        kids.append(f"{page_number} 0 R")
        commands = []
        def text(x, y, value):
            value = str(value).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"BT /F1 10 Tf {x:g} {y:g} Td ({value}) Tj ET")
        if spec.get("scan"):
            commands.append("q 520 0 0 700 40 40 cm /Im1 Do Q")
        else:
            text(40, 750, spec.get("before", "Before tables"))
            top = 690
            for ordinal, rows in enumerate(spec.get("tables", [])):
                cols = max(len(row) for row in rows)
                width, height = 520 / cols, 16
                bottom = top - height * len(rows)
                if spec.get("grid", True):
                    for col in range(cols + 1):
                        x = 40 + col * width
                        commands.append(f"{x:g} {bottom:g} m {x:g} {top:g} l S")
                    for row in range(len(rows) + 1):
                        y = top - row * height
                        commands.append(f"40 {y:g} m 560 {y:g} l S")
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        if value is not None:
                            text(44 + c * width, top - r * height - 12, value)
                text(40, bottom - 20, f"Between table {ordinal + 1} and next block")
                top = bottom - 45
            text(40, top - 20, spec.get("after", "After tables"))
        stream = "\n".join(commands).encode("latin-1")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 3 0 R >> /XObject << /Im1 4 0 R >> >> /Contents {content_number} 0 R >>".encode())
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = f"<< /Type /Pages /Count {len(pages)} /Kids [{' '.join(kids)}] >>".encode()
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for n, value in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{n} 0 obj\n".encode() + value + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(output)


def settings_at(directory):
    return Settings(_env_file=None, STORAGE_DIR=Path(directory), LLM_BASE_URL="http://model.invalid/v1",
                    LLM_MODEL_NAME="test", SKILLS_DIR=Path(__file__).resolve().parents[1] / "skills")


class ExtractionTests(unittest.TestCase):
    def test_ordered_prose_references_exact_cells_and_pdf_persistence(self):
        content = pdf_bytes([{"tables": [
            [["Code", "Amount, RUB"], ["007", "1 234,50"], ["008", "0"]],
            [["Name", "Count"], ["Missing", ""], ["Present", "2"]]]}])
        document = pdf_tables.extract_pdf("report.PDF", content)
        self.assertEqual(len(document.tables), 2)
        first, second = document.tables
        self.assertEqual(first.data.rows, [["007", "1 234,50"], ["008", "0"]])
        self.assertEqual(second.data.rows, [["Missing", ""], ["Present", "2"]])
        self.assertEqual(first.page_number, 1)
        self.assertEqual(len(first.bounding_box), 4)
        self.assertLess(document.text.index("Before tables"), document.text.index(first.table_id))
        self.assertLess(document.text.index(first.table_id), document.text.index("Between table 1"))
        self.assertLess(document.text.index("Between table 1"), document.text.index(second.table_id))
        self.assertLess(document.text.index(second.table_id), document.text.index("After tables"))
        self.assertNotIn("1 234,50", document.text)
        with TemporaryDirectory() as tmp:
            settings = settings_at(tmp)
            docx_tables.save_material(document, content, settings)
            loaded = docx_tables.load_materials([document.material_id], settings)[0]
            self.assertEqual(loaded, document)
            self.assertNotIn("[[", loaded.text)
            self.assertIn("Страница 1", loaded.text)
            self.assertEqual((Path(tmp)/"materials"/document.material_id/"source.pdf").read_bytes(), content)

    def test_large_table_keeps_every_row_for_analysis_and_page_provenance(self):
        rows = [["Code", "Amount"], *[[f"{n:03}", str(n)] for n in range(30)]]
        document = pdf_tables.extract_pdf("large.pdf", pdf_bytes([{}, {"tables": [rows]}]))
        table = document.tables[0]
        self.assertEqual(table.page_number, 2)
        self.assertFalse(table.available)
        self.assertEqual(table.data.rows, rows[1:])
        public = docx_tables.public_table(table)
        self.assertIsNone(public["data"])
        self.assertTrue(public["analysis_available"])
        self.assertEqual(public["row_count"], 30)

    def test_empty_geometric_gutter_is_removed_but_blank_rows_and_cells_remain(self):
        rows = [["Code", "", "Amount"], ["007", "", "0"], ["", "", ""], ["008", "", "12"]]
        table = pdf_tables.extract_pdf("gutter.pdf", pdf_bytes([{"tables": [rows]}])).tables[0]
        self.assertEqual(table.data.headers, ["Code", "Amount"])
        self.assertEqual(table.data.rows, [["007", "0"], ["", ""], ["008", "12"]])

    def test_bad_headers_are_unavailable_instead_of_fabricated(self):
        for headers in (["Code", "Code"], ["Code", ""]):
            with self.subTest(headers=headers):
                table = pdf_tables.extract_pdf("bad.pdf", pdf_bytes([{"tables": [[headers, ["007", "1"]]]}])).tables[0]
                self.assertIsNone(table.data)
                self.assertFalse(table.available)
                self.assertIn("заголовки", table.limitation)

    def test_plain_text_and_borderless_content_are_preserved_without_invented_tables(self):
        for pages in ([{"before": "A normal prose document with amounts 100 and 200."}],
                      [{"grid": False, "tables": [[["Region", "Amount"], ["North", "007"], ["South", "0"]]]}]):
            document = pdf_tables.extract_pdf("text.pdf", pdf_bytes(pages))
            self.assertEqual(document.tables, [])
            self.assertTrue(document.text)
            if pages[0].get("grid") is False:
                self.assertIn("007", document.text)
                self.assertIn("North", document.text)

    def test_corrupt_password_scans_and_limits_fail_explicitly(self):
        for content in (b"not a pdf", pdf_bytes([{"scan": True}])):
            with self.subTest(size=len(content)), self.assertRaises(MaterialError):
                pdf_tables.extract_pdf("bad.pdf", content)
        with patch.object(pdf_tables.pdfplumber, "open", side_effect=PDFPasswordIncorrect()), self.assertRaisesRegex(MaterialError, "паролем"):
            pdf_tables.extract_pdf("locked.pdf", b"pdf")
        with patch.object(pdf_tables.pdfplumber, "open", side_effect=PdfminerException(PDFPasswordIncorrect())), self.assertRaisesRegex(MaterialError, "паролем"):
            pdf_tables.extract_pdf("locked.pdf", b"pdf")
        with patch.object(pdf_tables, "MAX_PAGES", 1), self.assertRaisesRegex(MaterialError, "страниц"):
            pdf_tables.extract_pdf("long.pdf", pdf_bytes([{}, {}]))
        rows = [["Code", "Amount"], ["007", "1"]]
        with patch.object(pdf_tables, "MAX_TABLES", 1), self.assertRaisesRegex(MaterialError, "таблиц"):
            pdf_tables.extract_pdf("long.pdf", pdf_bytes([{"tables": [rows, rows]}]))


class ApiTests(unittest.TestCase):
    def test_upload_returns_text_summary_and_id_then_job_loads_exact_server_table(self):
        async def describe(document, settings):
            document.tables[0].summary = "Amounts by code"
        with TemporaryDirectory() as tmp:
            settings = settings_at(tmp)
            settings.templates_dir.mkdir()
            with patch.object(routes, "get_settings", return_value=settings), \
                 patch.object(docx_tables, "describe_tables", new=AsyncMock(side_effect=describe)) as summaries, \
                 patch.object(routes, "run_planning", new=AsyncMock()):
                with TestClient(app) as client:
                    extracted = client.post("/materials/extract", files={"file": ("report.pdf", pdf_bytes())}).json()
                    self.assertEqual(extracted["status"], "ok", extracted)
                    summaries.assert_awaited_once()
                    self.assertIn("Before tables", extracted["text"])
                    self.assertIn("Amounts by code", extracted["text"])
                    self.assertIn(extracted["tables"][0]["table_id"], extracted["text"])
                    template = BytesIO(); Presentation().save(template)
                    response = client.post("/jobs", data={"brief": "Report", "idea": "Compare amounts",
                        "material_ids_json": json.dumps([extracted["material_id"]])},
                        files={"template": ("template.pptx", template.getvalue())})
                    self.assertEqual(response.status_code, 200, response.text)
                    job = get_job(response.json()["job_id"])
                    self.assertEqual(job.source_tables[0].data.rows[0], ["007", "1 234,50"])
                    self.assertEqual(job.source_tables[0].page_number, 1)
                    failed = client.post("/materials/extract", files={"file": ("scan.pdf", pdf_bytes([{"scan": True}]))}).json()
                    self.assertEqual(failed["status"], "error")
                    self.assertIn("OCR", failed["note"])


class AnalysisTests(unittest.IsolatedAsyncioTestCase):
    async def test_mixed_scan_warning_survives_summary_failure(self):
        document = pdf_tables.extract_pdf("mixed.pdf", pdf_bytes([{"scan": True}, {"tables": [
            [["Code", "Amount"], ["007", "1"]]]}]))
        self.assertIn("1", document.note)
        self.assertIn("OCR", document.note)
        self.assertIn("OCR", document.text)
        client = AsyncMock()
        client.__aenter__.return_value = client
        client.chat_structured.side_effect = TimeoutError()
        with TemporaryDirectory() as tmp, patch.object(docx_tables, "LLMClient", return_value=client), \
             patch("app.skills.get_settings", return_value=settings_at(tmp)):
            await docx_tables.describe_tables(document, settings_at(tmp))
        self.assertIn("OCR", document.note)
        self.assertIn("LLM-описания недоступны", document.note)
        source = json.loads(client.chat_structured.call_args.kwargs["user_input"])["tables"][0]
        self.assertEqual(source["page_number"], 2)
        self.assertEqual(source["table_id"], document.tables[0].table_id)

    async def test_actual_sql_uses_all_pdf_rows_and_planner_gets_result_id(self):
        rows = [["Code", "Amount, RUB"], ["007", "1 234,50"], *[["007", "0"] for _ in range(29)]]
        document = pdf_tables.extract_pdf("sales.pdf", pdf_bytes([{"tables": [rows]}]))
        original = deepcopy(document.tables)
        job = Job(job_id="pdf-test", brief="Report", purpose="Compare", template_path="unused", source_tables=document.tables)
        with TemporaryDirectory() as tmp:
            settings = settings_at(tmp)
            settings.TABLE_ANALYSIS_MAX_RESULTS = 1
            with patch("app.skills.get_settings", return_value=settings):
                request = table_analysis.analysis_request(job, settings)
            async def fake_run(agent, user, *, context, **kwargs):
                metadata = json.loads(user)["tables"][0]
                self.assertEqual(metadata["page_number"], 1)
                self.assertEqual(metadata["rows"], 30)
                self.assertEqual(len(metadata["bounding_box"]), 4)
                tools = {tool.name: tool for tool in agent.tools}
                async def invoke(name, args):
                    ctx = ToolContext(context=context, tool_name=name, tool_call_id="test", tool_arguments=json.dumps(args))
                    result = await tools[name].on_invoke_tool(ctx, json.dumps(args))
                    self.assertNotIn("error", json.loads(result), result)
                await invoke("create_view", {"name": "totals", "sql":
                    'SELECT "Code" AS code, SUM(CAST(replace(replace("Amount, RUB", \' \', \'\'), \',\', \'.\') AS DOUBLE)) AS revenue '
                    'FROM table_1 GROUP BY "Code"'})
                await invoke("prepare_slide_data", {"source": "totals", "title": "Amounts by code", "kind": "table",
                    "reason": "Sum of all 30 rows", "x": None, "y": None,
                    "presentation_columns": [{"key": "code", "label": "Код", "unit": ""},
                                             {"key": "revenue", "label": "Выручка", "unit": "руб."}]})
            with patch.object(worker.Runner, "run", new=fake_run):
                report = await worker.run_agent(request, Path(tmp)/"report.json", settings)
            self.assertEqual(report.status, "completed", report.notes)
            result = report.artifacts[0]
            self.assertEqual(result.data.rows, [["007", "1234.5"]])
            self.assertEqual(document.tables, original)
            job.table_analysis = report
            job.source_tables += report.artifacts
            inputs = briefing._plan_inputs(job, 3, None)
            self.assertIn(result.table_id, inputs["materials"])
            self.assertIn('"007", "1234.5"', inputs["materials"])
            specs = [slide("cover"), slide("data", source_table_id=result.table_id), slide("end")]
            docx_tables.bind_source_tables(specs, job.source_tables)
            self.assertEqual(specs[1].content_blocks[-1].table, result.data)
            self.assertIn("sales.pdf", specs[1].data_note)
            self.assertIn("страница 1", specs[1].data_note)


if __name__ == "__main__":
    unittest.main()
