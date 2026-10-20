"""Regression checks for annotations and native horizontal charts."""
import json
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app.pipeline.charts import add_chart, use_horizontal_bars
from app.pipeline.assembly import _add_native_chart
from app.pipeline.slide_clone import replace_with_chart
from app.schemas.content_plan import ChartData, ChartSeries, ChartType, ContentBlock, ContentBlockType
from app.schemas.design_manifest import DesignManifest, Typography, TypographyStyle


def manifest():
    return DesignManifest(template_id="annotations", palette={"dk1":"18233B", "lt1":"FFFFFF", "accent1":"0077FF"},
        typography=Typography(title=TypographyStyle(font="Arial",size=28),body=TypographyStyle(font="Arial",size=18)),
        layouts=[],slide_width_emu=12192000,slide_height_emu=6858000)


def fixture(**kwargs):
    fields=dict(chart_type=ChartType.BAR, categories=["Москва","Казань"],
        series=[ChartSeries(name="Продажи",values=[120,80])], category_axis_title="Регион",
        value_axis_title="Продажи", units="млн руб.", source="Условные данные")
    fields.update(kwargs)
    return ChartData(**fields)


class ChartAnnotationsTest(unittest.TestCase):
    def setUp(self):
        self.prs=Presentation();self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.box=tuple(Inches(value) for value in (.5, .5, 9, 6))

    def test_annotations_source_and_units_survive_save_reload(self):
        frame=add_chart(self.slide,fixture(),self.box,manifest())
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"axes.pptx";self.prs.save(path)
            slide=Presentation(path).slides[0];chart=slide.shapes[0].chart
            self.assertEqual(chart.category_axis.axis_title.text_frame.text,"Регион")
            self.assertEqual(chart.value_axis.axis_title.text_frame.text,"Продажи, млн руб.")
            self.assertEqual(slide.shapes[1].text,"Источник: Условные данные")
            self.assertLessEqual(frame.top+frame.height,slide.shapes[1].top)
            self.assertLessEqual(slide.shapes[1].top+slide.shapes[1].height,self.box[1]+self.box[3])

    def test_horizontal_preserves_order_values_and_semantic_axes(self):
        data=fixture(categories=["Департамент корпоративных продаж", "Департамент клиентской поддержки"])
        frame=add_chart(self.slide,data,self.box,manifest());chart=frame.chart
        self.assertEqual(chart.chart_type,XL_CHART_TYPE.BAR_CLUSTERED)
        self.assertTrue(chart.category_axis.reverse_order)
        self.assertEqual(chart._chartSpace.xpath(".//c:catAx/c:crosses")[0].get("val"),"autoZero")
        self.assertEqual(chart._chartSpace.xpath(".//c:valAx/c:crosses")[0].get("val"),"max")
        self.assertEqual([c.label for c in chart.plots[0].categories],data.categories)
        self.assertEqual(list(chart.series[0].values),[120,80])
        self.assertEqual(chart.value_axis.axis_title.text_frame.text,"Продажи, млн руб.")
        self.assertFalse(use_horizontal_bars(data.model_copy(update={"orientation":"vertical"})))
        self.assertTrue(use_horizontal_bars(fixture(orientation="horizontal")))
        self.assertFalse(use_horizontal_bars(data.model_copy(update={"chart_type":ChartType.LINE})))

    def test_legacy_payload_and_missing_metadata(self):
        data=ChartData.model_validate_json(json.dumps({"chart_type":"bar","categories":["A"],"series":[{"name":"Count","values":[4]}]}))
        chart=add_chart(self.slide,data,self.box,manifest()).chart
        self.assertFalse(chart.category_axis.has_title)
        self.assertEqual(chart.value_axis.axis_title.text_frame.text,"Count")
        self.assertEqual(len(self.slide.shapes),1)
        self.assertEqual(chart.chart_type,XL_CHART_TYPE.COLUMN_CLUSTERED)
        chart=add_chart(self.slide,fixture(value_axis_title="Продажи, млн руб."),self.box,manifest()).chart
        self.assertEqual(chart.value_axis.axis_title.text_frame.text,"Продажи, млн руб.")

    def test_pie_has_source_without_attempting_axes(self):
        chart=add_chart(self.slide,fixture(chart_type=ChartType.PIE),self.box,manifest()).chart
        self.assertEqual(chart.chart_type,XL_CHART_TYPE.PIE)
        self.assertTrue(chart.plots[0].data_labels.show_percentage)
        self.assertEqual(self.slide.shapes[1].text,"Источник: Условные данные")

    def test_assembly_and_template_replacement_use_annotations(self):
        block=ContentBlock(type=ContentBlockType.CHART,chart=fixture())
        _add_native_chart(self.slide,self.box,block,manifest())
        self.assertTrue(self.slide.shapes[0].chart.value_axis.has_title)
        replacement_slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        placeholder=replacement_slide.shapes.add_textbox(*self.box)
        types={"bar":XL_CHART_TYPE.COLUMN_CLUSTERED}
        self.assertTrue(replace_with_chart(replacement_slide,placeholder.shape_id,fixture(),manifest(),types))
        self.assertEqual(len(replacement_slide.shapes),2)
        self.assertTrue(replacement_slide.shapes[0].has_chart)
        self.assertEqual(replacement_slide.shapes[1].text,"Источник: Условные данные")

    def test_dense_chart_keeps_labels_disabled_and_dark_source_readable(self):
        self.slide.background.fill.solid()
        from pptx.dml.color import RGBColor
        self.slide.background.fill.fore_color.rgb=RGBColor.from_string("18233B")
        data=fixture(categories=[f"Длинная категория номер {i}" for i in range(15)],series=[ChartSeries(name="Count",values=list(range(15)))])
        chart=add_chart(self.slide,data,self.box,manifest()).chart
        self.assertFalse(chart.plots[0].has_data_labels)
        self.assertEqual(str(self.slide.shapes[1].text_frame.paragraphs[0].font.color.rgb),"FFFFFF")

    def test_excel_api_keeps_filename_units_headers_and_source_override(self):
        from io import BytesIO
        from types import SimpleNamespace
        from unittest.mock import patch
        from openpyxl import Workbook
        from fastapi.testclient import TestClient
        from app.main import app
        wb=Workbook();ws=wb.active;ws.append(["Регион","Выручка"])
        ws.append(["Москва",120]);ws.append(["Казань",80])
        stream=BytesIO();wb.save(stream);wb.close()
        spec=SimpleNamespace(slide_id="slide-1",purpose="Сравнить регионы",content_blocks=[])
        job=SimpleNamespace(content_plan=SimpleNamespace(slides=[spec]))
        with patch("app.api.routes.get_job",return_value=job),patch("app.api.routes.save_job"):
            client=TestClient(app)
            response=client.post("/jobs/demo/slides/slide-1/import-excel",files={"file":("sales.xlsx",stream.getvalue())},data={"units":"млн руб."})
            self.assertEqual(response.status_code,200,response.text)
            data=response.json()["block"]["chart"]
            self.assertEqual(data["source"],"sales.xlsx")
            self.assertEqual(data["units"],"млн руб.")
            self.assertEqual(data["category_axis_title"],"Регион")
            self.assertEqual(data["value_axis_title"],"Выручка")
            response=client.post("/jobs/demo/slides/slide-1/import-excel",files={"file":("sales.xlsx",stream.getvalue())},data={"source":"Внутренний отчёт"})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(response.json()["block"]["chart"]["source"],"Внутренний отчёт")

    def test_template_style_failure_keeps_native_chart_and_annotations(self):
        from unittest.mock import patch
        with patch("app.pipeline.charts.slide_is_dark",side_effect=ValueError("unsupported background")), \
             patch("app.pipeline.charts.style_chart",side_effect=ValueError("unsupported style")), \
             patch("app.pipeline.charts.logger"):
            frame=add_chart(self.slide,fixture(),self.box,manifest())
        self.assertTrue(frame.has_chart)
        self.assertEqual(frame.chart.value_axis.axis_title.text_frame.text,"Продажи, млн руб.")
        self.assertEqual(self.slide.shapes[1].text,"Источник: Условные данные")
