"""Recommended views travel through tools, planning and editable PPTX export."""
from copy import deepcopy
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import AsyncMock, patch
from agents.tool_context import ToolContext
from agents import ModelBehaviorError
from pptx import Presentation
from app.analysis import worker
from app.config import Settings
from app.pipeline import briefing, docx_tables, dramaturgy, table_analysis, template_engine
from app.schemas.content_plan import ContentPlan, TableData
from app.schemas.job import Job
from app.schemas.source_tables import SourceTable, TableAnalysisReport
from test_docx_tables import slide

META = [dict(key="city", label="Город", unit=""), dict(key="value", label="Выручка", unit="руб.")]


def result(kind="bar", count=2, missing=False):
    item = dict(kind=kind, title="Сравнение", reason="Сумма продаж; диаграмма наглядно сравнивает города",
        x="city", y=["value"], columns=["city", "value"], data=[{"city": str(i), "value": i * 100} for i in range(count)],
        presentation_columns=META, parent_table_ids=["input"], operations=[], notes=[])
    if missing:
        item["data"][1]["value"] = None
    return worker.compile_artifacts([item], {"key": "test"})[0]


class RecommendationTests(unittest.TestCase):
    def test_numeric_example_is_not_used_as_text_cards(self):
        template = {"variants": [dict(id="numeric", source="sample", composition=dict(kind="cards"),
            slots={"value": dict(semantic_role="unit_title", template_text="180 +")})]}
        self.assertEqual(dramaturgy.catalog_compositions(template), [])
        template["variants"][0]["slots"]["value"]["template_text"] = "Инструменты"
        self.assertEqual(len(dramaturgy.catalog_compositions(template)), 1)

    def test_chart_named_sample_without_chart_slot_gets_usable_native_fallback(self):
        from app.pipeline import sample_slides
        augment = sample_slides.augment_template
        def add_unusable_sample(template, path):
            augment(template, path)
            sample = deepcopy(template["variants"][0])
            sample.update(id="Образец графика без объекта", source="sample", composition={"kind": "chart", "chart_pictures": []})
            template["variants"].append(sample)
        with TemporaryDirectory() as tmp:
            path = Path(tmp)/"template.pptx"; Presentation().save(path)
            with patch.object(sample_slides, "augment_template", side_effect=add_unusable_sample):
                template = template_engine.build_template_json(path)
            source = result()
            plan = ContentPlan(brief="Отчет", purpose="Сравнить", slides=[slide("cover"), slide("data", source_table_id=source.table_id), slide("end")])
            docx_tables.bind_source_tables(plan.slides, [source])
            structured = [dramaturgy.structure_fallback(i, s, 3) for i, s in enumerate(plan.slides)]
            for variant in template_engine.VALID_VARIANTS:
                layout = dramaturgy.build_variant_plan(template, plan, structured, variant, "", template_engine._shorten)
                self.assertEqual(layout["slides"][1]["composition_kind"], "chart")
                self.assertNotEqual(layout["slides"][1]["variant"], "Образец графика без объекта")

    def test_layout_with_static_example_table_is_not_a_text_composition(self):
        prs = Presentation(); example = prs.slides.add_slide(prs.slide_layouts[1])
        table = example.shapes.add_table(2, 2, 100000, 100000, 3000000, 2000000)
        example.slide_layout.shapes._spTree.append(deepcopy(table._element))
        with TemporaryDirectory() as tmp:
            path = Path(tmp)/"static-table.pptx"; prs.save(path)
            template = template_engine.build_template_json(path)
            static_layouts = {v["id"] for v in template["variants"] if v.get("source", "layout") == "layout" and any(d.get("table") for d in v.get("decor", []))}
            self.assertTrue(static_layouts)
            self.assertFalse(static_layouts & {c["id"] for c in dramaturgy.catalog_compositions(template)})

    def test_recommendation_keeps_values_and_missing_points(self):
        source = result("line", 12)
        self.assertEqual(source.recommended_presentation, "chart")
        self.assertEqual(len(source.chart.categories), 12)
        self.assertEqual(source.chart.series[0].values, [i * 100 for i in range(12)])
        self.assertEqual(source.chart.units, "руб.")
        self.assertIn("сравнивает", source.presentation_reason)
        self.assertEqual(result("line", missing=True).chart.series[0].values, [0.0, None])
        self.assertEqual(result("table").recommended_presentation, "table")

    def test_all_chart_types_roundtrip_all_designs_without_model_supplied_values(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp)/"template.pptx"; Presentation().save(path)
            template = template_engine.build_template_json(path)
            from app.pipeline.parser import parse_template
            manifest = parse_template(path)
            for kind in ("bar", "column", "line", "pie"):
                source = result(kind); before = deepcopy(source)
                plan = ContentPlan(brief="Отчет", purpose="Сравнить", slides=[slide("cover"), slide("data", source_table_id=source.table_id), slide("end")])
                docx_tables.bind_source_tables(plan.slides, [source])
                self.assertEqual(plan.slides[1].content_blocks[-1].chart, source.chart)
                structured = [dramaturgy.structure_fallback(i, s, 3) for i, s in enumerate(plan.slides)]
                for variant in template_engine.VALID_VARIANTS:
                    with self.subTest(kind=kind, variant=variant):
                        layout = dramaturgy.build_variant_plan(template, plan, structured, variant, plan.brief, template_engine._shorten)
                        self.assertEqual(layout["slides"][1]["composition_kind"], "chart")
                        output = Path(tmp)/f"{kind}-{variant}.pptx"
                        template_engine.render_variant(template, layout, path, output, manifest, plan, variant)
                        saved = Presentation(output)
                        docx_tables.verify_rendered_tables(saved, plan, layout)
                        charts = [sh.chart for s in saved.slides for sh in s.shapes if sh.has_chart]
                        self.assertEqual(len(charts), 1)
                        self.assertEqual(charts[0].series[0].values, (0.0, 100.0))
                self.assertEqual(source, before)


class IntegrationTests(unittest.IsolatedAsyncioTestCase):
    async def test_tool_validates_views_and_planner_uses_recommendation(self):
        with TemporaryDirectory() as tmp:
            settings = Settings(_env_file=None, STORAGE_DIR=Path(tmp), SKILLS_DIR=Path(__file__).resolve().parents[1]/"skills")
            settings.TABLE_ANALYSIS_MAX_RESULTS = 4
            source = SourceTable(table_id="input", source_name="sales.xlsx", ordinal=1, summary="Продажи", available=True,
                data=TableData(headers=["Город", "Выручка"], rows=[["Москва", "100"], ["Казань", "200"]]))
            job = Job(job_id="test", brief="Отчет", purpose="Сравнить", template_path="unused", source_tables=[source])
            with patch("app.skills.get_settings", return_value=settings):
                request = table_analysis.analysis_request(job, settings)
            async def fake_run(agent, user, *, context, **kwargs):
                await agent.model.get_response()
                tools = {t.name: t for t in agent.tools}
                self.assertEqual(tools["prepare_slide_data"].params_json_schema["properties"]["kind"]["enum"], ["table", "bar", "column", "line", "pie"])
                async def call(name, args):
                    ctx = ToolContext(context=context, tool_name=name, tool_call_id="test", tool_arguments=json.dumps(args))
                    return json.loads(await tools[name].on_invoke_tool(ctx, json.dumps(args)))
                self.assertNotIn("error", await call("create_view", {"name": "totals", "sql":
                    'SELECT "Город" AS city, row_number() OVER () AS period, CAST("Выручка" AS DOUBLE) AS value, 10 AS orders FROM table_1'}))
                args = dict(kind="bar", title="Выручка", reason="Полосы лучше сравнивают города", source="totals", x="city", y=["value", "orders"],
                    presentation_columns=[*META, dict(key="orders", label="Заказы", unit="шт.")])
                self.assertIn("error", await call("prepare_slide_data", args)); self.assertEqual(context.slides, [])
                args.update(y=["value"], presentation_columns=META)
                for kind in ("bar", "column", "line", "pie"):
                    args["kind"] = kind
                    args["x"] = "period" if kind == "line" else "city"
                    args["presentation_columns"] = [dict(key="period", label="Период", unit=""), META[1]] if kind == "line" else META
                    self.assertNotIn("error", await call("prepare_slide_data", args))
            responses = AsyncMock(side_effect=[ModelBehaviorError("empty truncated model response"), object()])
            with patch.object(worker.Runner, "run", new=fake_run), patch.object(worker.OpenAIChatCompletionsModel, "get_response", responses):
                report = await worker.run_agent(request, Path(tmp)/"report.json", settings)
            self.assertEqual(responses.await_count, 2)
            self.assertEqual(report.model_calls, 2)
            self.assertEqual(report.status, "completed", report.notes)
            self.assertEqual([t.chart.chart_type.value for t in report.artifacts], ["bar", "bar", "line", "pie"])
            job.source_tables += report.artifacts; job.table_analysis = report
            client = AsyncMock(); client.__aenter__.return_value = client
            client.chat_structured.side_effect = lambda **kw: briefing.PlanOut(brief="Отчет", purpose="Сравнить", story="Вывод", slides=[
                slide("cover"), slide("data", source_table_id=report.artifacts[0].table_id), slide("end")])
            with patch.object(briefing, "_client", return_value=client), patch("app.skills.get_settings", return_value=settings):
                variants = await briefing.generate_plan_variants(job, 3, None)
            self.assertEqual(set(variants), {"A", "B", "C"})
            for variant in variants.values():
                self.assertEqual(variant.plan.slides[1].content_blocks[-1].chart, report.artifacts[0].chart)
            for call in client.chat_structured.call_args_list:
                self.assertIn('"recommended_presentation": "chart"', call.kwargs["user_input"])
                self.assertIn(report.artifacts[0].presentation_reason, call.kwargs["user_input"])
