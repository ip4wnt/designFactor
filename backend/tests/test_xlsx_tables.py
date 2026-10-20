"""XLSX upload -> persistent IDs -> actual analytical tools -> planner and PPTX cells."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import AsyncMock, patch
from zipfile import ZipFile

from agents.tool_context import ToolContext
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation

from app.analysis import worker
from app.api import routes
from app.config import Settings
from app.job_store import get as get_job
from app.main import app
from app.pipeline import briefing, docx_tables, table_analysis, xlsx_tables
from app.pipeline.materials import MaterialError
from app.schemas.job import Job
from test_docx_tables import slide


def book_bytes(sheets=None, configure=None):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in (sheets or {"Продажи": [["Город", "Выручка, руб."], ["Орёл", 125], ["Казань", 0]]}).items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    if configure:
        configure(workbook)
    out = BytesIO()
    workbook.save(out)
    workbook.close()
    return out.getvalue()


def settings_at(directory):
    return Settings(_env_file=None, STORAGE_DIR=Path(directory), LLM_BASE_URL="http://model.invalid/v1",
                    LLM_MODEL_NAME="test", SKILLS_DIR=Path(__file__).resolve().parents[1] / "skills")


class ExtractionTests(unittest.TestCase):
    def test_full_sheets_persist_with_exact_missing_values_dates_and_formats(self):
        sheets = {"Продажи": [[None, None, None], [None, "Код", "Дата", "Доля", "Заказы"],
                              [None, "007", datetime(2025, 1, 2), .23, 0],
                              [None, None, None, None, None], [None, "008", None, .5, 1]],
                  "Пустой": [], "Справочник": [["Ключ", "Имя"], ["007", "Орёл"]]}
        def configure(wb):
            wb["Продажи"]["D3"].number_format = "0%"
            wb["Продажи"]["D5"].number_format = "0%"
        content = book_bytes(sheets, configure)
        document = xlsx_tables.extract_xlsx("report.XLSX", content)
        self.assertEqual(len(document.tables), 2)
        table = document.tables[0]
        self.assertEqual(table.sheet_name, "Продажи")
        self.assertEqual(table.cell_range, "B2:E5")
        self.assertEqual(table.data.rows, [["007", "2025-01-02T00:00:00", "0.23", "0"],
                                           ["", "", "", ""], ["008", "", "0.5", "1"]])
        self.assertEqual(table.number_formats["Доля"], ["0%"])
        with TemporaryDirectory() as tmp:
            settings = settings_at(tmp)
            docx_tables.save_material(document, content, settings)
            loaded = docx_tables.load_materials([document.material_id], settings)[0]
            self.assertEqual(document, loaded)
            self.assertIn(table.table_id, loaded.text)
            self.assertIn("Продажи", loaded.text)
            self.assertNotIn("[[", loaded.text)
            self.assertEqual((Path(tmp)/"materials"/document.material_id/"source.xlsx").read_bytes(), content)

    def test_more_than_80_rows_and_six_sheets_are_not_truncated(self):
        sheets = {f"Данные {i}": [["Код", "Значение"], *[[str(n), n] for n in range(205)]] for i in range(7)}
        document = xlsx_tables.extract_xlsx("large.xlsx", book_bytes(sheets))
        self.assertEqual(len(document.tables), 7)
        self.assertEqual(len({t.table_id for t in document.tables}), 7)
        for table in document.tables:
            self.assertEqual(len(table.data.rows), 205)
            self.assertEqual(table.data.rows[-1], ["204", "204"])
            self.assertFalse(table.available)
            public = docx_tables.public_table(table)
            self.assertTrue(public["analysis_available"])
            self.assertIsNone(public["data"])
            self.assertEqual(public["row_count"], 205)

    def test_invalid_headers_merged_formula_gaps_and_excel_errors_are_explicit(self):
        cases = [({"Лист": [["A", "A"], [1, 2]]}, None, "заголовки"),
                 ({"Лист": [["A", "B"], [1, 2]]}, lambda w: w.active.merge_cells("A2:B2"), "Объединенные"),
                 ({"Лист": [["A", "B"], [1, "=A2*2"]]}, None, "B2"),
                 ({"Лист": [["A", "B"], [1, "#DIV/0!"]]}, None, "ошибки Excel")]
        for sheets, configure, message in cases:
            with self.subTest(message=message):
                table = xlsx_tables.extract_xlsx("bad.xlsx", book_bytes(sheets, configure)).tables[0]
                self.assertFalse(table.available)
                self.assertIsNone(table.data)
                self.assertIn(message, table.limitation)
                self.assertFalse(docx_tables.public_table(table)["analysis_available"])

    def test_saved_formula_result_is_used_without_executing_formula(self):
        source = book_bytes({"Лист": [["Количество", "Удвоено"], [3, "=A2*2"]]})
        out = BytesIO()
        with ZipFile(BytesIO(source)) as original, ZipFile(out, "w") as target:
            for entry in original.infolist():
                data = original.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<f>A2*2</f><v></v>", b"<f>A2*2</f><v>6</v>")
                target.writestr(entry, data)
        table = xlsx_tables.extract_xlsx("cached.xlsx", out.getvalue()).tables[0]
        self.assertTrue(table.available)
        self.assertEqual(table.data.rows, [["3", "6"]])

    def test_corrupt_empty_and_over_limit_books_fail_without_truncation(self):
        for content in (b"not xlsx", book_bytes({"Пустой": []}),
                        book_bytes({f"Лист {n}": [["A"], [1]] for n in range(11)})):
            with self.subTest(size=len(content)), self.assertRaises(MaterialError):
                xlsx_tables.extract_xlsx("bad.xlsx", content)
        with patch.object(xlsx_tables, "MAX_ROWS", 2):
            table = xlsx_tables.extract_xlsx("long.xlsx", book_bytes({"Лист": [["A"], [1], [2], [3]]})).tables[0]
            self.assertIsNone(table.data)
            self.assertIn("Превышен лимит", table.limitation)


class ApiTests(unittest.TestCase):
    def test_upload_summary_id_and_job_use_full_server_cells(self):
        rows = [["Город", "Выручка, руб."], *[["Орёл", n] for n in range(205)]]
        content = book_bytes({"Продажи": rows})
        async def describe(document, settings):
            document.tables[0].summary = "Выручка по городам за период"
        with TemporaryDirectory() as tmp:
            settings = settings_at(tmp)
            settings.templates_dir.mkdir()
            with patch.object(routes, "get_settings", return_value=settings), \
                 patch.object(docx_tables, "describe_tables", new=AsyncMock(side_effect=describe)) as summaries, \
                 patch.object(routes, "run_planning", new=AsyncMock()):
                with TestClient(app) as client:
                    extracted = client.post("/materials/extract", files={"file": ("sales.xlsx", content)}).json()
                    self.assertEqual(extracted["status"], "ok", extracted)
                    summaries.assert_awaited_once()
                    self.assertIn("Выручка по городам", extracted["text"])
                    self.assertIn(extracted["tables"][0]["table_id"], extracted["text"])
                    self.assertTrue(extracted["tables"][0]["analysis_available"])
                    template = BytesIO(); Presentation().save(template)
                    response = client.post("/jobs", data={"brief": "Отчет", "idea": "Сравнить города",
                        "material_ids_json": json.dumps([extracted["material_id"]])},
                        files={"template": ("template.pptx", template.getvalue())})
                    self.assertEqual(response.status_code, 200, response.text)
                    job = get_job(response.json()["job_id"])
                    self.assertEqual(len(job.source_tables[0].data.rows), 205)
                    self.assertEqual(job.source_tables[0].data.rows[-1], ["Орёл", "204"])


class AnalysisTests(unittest.IsolatedAsyncioTestCase):
    async def test_xlsx_summaries_keep_sheet_context_and_exact_ids(self):
        document = xlsx_tables.extract_xlsx("sales.xlsx", book_bytes())
        before = deepcopy(document.tables[0].data)
        client = AsyncMock()
        client.__aenter__.return_value = client
        client.chat_structured.return_value = docx_tables.TableSummaries(tables=[
            docx_tables.TableSummary(table_id=document.tables[0].table_id, summary="Продажи по городам")])
        with TemporaryDirectory() as tmp, patch.object(docx_tables, "LLMClient", return_value=client), \
             patch("app.skills.get_settings", return_value=settings_at(tmp)):
            await docx_tables.describe_tables(document, settings_at(tmp))
        data = json.loads(client.chat_structured.call_args.kwargs["user_input"])["tables"][0]
        self.assertEqual(data["table_id"], document.tables[0].table_id)
        self.assertEqual(data["sheet_name"], "Продажи")
        self.assertEqual(document.tables[0].data, before)
        self.assertIn("Продажи по городам", docx_tables.table_card(document.tables[0]))

    async def test_agent_reads_all_rows_joins_sheets_and_sends_result_ids_to_planner(self):
        document = xlsx_tables.extract_xlsx("sales.xlsx", book_bytes({
            "Продажи": [["Код", "Выручка, руб."], *[["007", 1] for _ in range(205)]],
            "Города": [["Код", "Город"], ["007", "Орёл"]]}))
        job = Job(job_id="xlsx-test", brief="Выручка городов", purpose="Сравнить",
                  template_path="unused", source_tables=document.tables)
        before = deepcopy(document.tables)
        with TemporaryDirectory() as tmp:
            settings = settings_at(tmp)
            settings.TABLE_ANALYSIS_MAX_RESULTS = 1
            with patch("app.skills.get_settings", return_value=settings):
                request = table_analysis.analysis_request(job, settings)
            async def fake_run(agent, user, *, context, **kwargs):
                metadata = json.loads(user)["tables"]
                self.assertEqual([t["sheet_name"] for t in metadata], ["Продажи", "Города"])
                self.assertEqual(metadata[0]["rows"], 205)
                tools = {tool.name: tool for tool in agent.tools}
                async def invoke(name, args):
                    ctx = ToolContext(context=context, tool_name=name, tool_call_id="test", tool_arguments=json.dumps(args))
                    result = await tools[name].on_invoke_tool(ctx, json.dumps(args))
                    self.assertNotIn("error", json.loads(result), result)
                await invoke("create_view", {"name": "totals", "sql":
                    'SELECT b."Город" AS city, SUM(CAST(a."Выручка, руб." AS DOUBLE)) AS revenue '
                    'FROM table_1 a JOIN table_2 b ON a."Код"=b."Код" GROUP BY b."Город"'})
                await invoke("prepare_slide_data", {"source": "totals", "title": "Выручка города", "kind": "table",
                    "reason": "Сумма всех 205 строк продаж после соединения со справочником", "x": None, "y": None,
                    "presentation_columns": [{"key": "city", "label": "Город", "unit": ""},
                                             {"key": "revenue", "label": "Выручка", "unit": "руб."}]})
            with patch.object(worker.Runner, "run", new=fake_run):
                report = await worker.run_agent(request, Path(tmp)/"report.json", settings)
            self.assertEqual(report.status, "completed", report.notes)
            result = report.artifacts[0]
            self.assertEqual(result.data.rows, [["Орёл", "205"]])
            self.assertEqual(result.data.headers, ["Город", "Выручка, руб."])
            self.assertEqual(set(result.parent_table_ids), {t.table_id for t in document.tables})
            self.assertEqual(document.tables, before)
            job.table_analysis = report
            job.source_tables += report.artifacts
            inputs = briefing._plan_inputs(job, 3, None)
            self.assertIn(result.table_id, inputs["materials"])
            self.assertIn('"Орёл", "205"', inputs["materials"])
            for table in document.tables:
                self.assertIn(table.table_id, inputs["materials"])
                self.assertIn(table.sheet_name, inputs["materials"])
            specs = [slide("cover"), slide("data", source_table_id=result.table_id), slide("end")]
            docx_tables.bind_source_tables(specs, job.source_tables)
            self.assertEqual(specs[1].content_blocks[-1].table, result.data)
            self.assertIn("sales.xlsx", specs[1].data_note)
            self.assertIn("Продажи", specs[1].data_note)


if __name__ == "__main__":
    unittest.main()
