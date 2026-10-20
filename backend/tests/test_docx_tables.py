"""DOCX -> server registry -> model ID -> exact editable PPTX, without an API key."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch
from xml.sax.saxutils import escape
from zipfile import ZipFile

from fastapi.testclient import TestClient
from pptx import Presentation

from app.api import routes
from app.main import app
from app.pipeline import briefing, docx_tables, dramaturgy, template_engine
from app.pipeline.materials import MaterialError
from app.schemas.content_plan import ContentBlock, ContentBlockType, ContentPlan, SlideSpec
from app.schemas.job import Job
from app.orchestrator import apply_manual_text


def docx_bytes(tables=None, merged=False):
    tables = tables if tables is not None else [
        [["Регион", "Выручка, ₽"], ["Орёл", "1 234,50"], ["Казань", "0"]],
        [["Показатель", "Значение"], ["Счёт", "007"]],
    ]
    def p(value):
        return '<w:p><w:r><w:t xml:space="preserve">' + escape(value) + '</w:t></w:r></w:p>'
    parts = [p("До таблиц")]
    for rows in tables:
        table = "<w:tbl><w:tblGrid>" + "<w:gridCol/>" * len(rows[0]) + "</w:tblGrid>"
        for row in rows:
            table += "<w:tr>"
            for index, value in enumerate(row):
                prop = '<w:tcPr><w:gridSpan w:val="2"/></w:tcPr>' if merged and index == 0 else ""
                table += "<w:tc>" + prop + p(value) + "</w:tc>"
            table += "</w:tr>"
        parts += [table + "</w:tbl>", p("После таблицы")]
    xml = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + "".join(parts) + "</w:body></w:document>"
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("word/document.xml", xml)
    return output.getvalue()


def slide(slide_id, **kwargs):
    return SlideSpec(slide_id=slide_id, title=slide_id, purpose="Показать результат", **kwargs)


def plan_for(table):
    slides = [slide("cover"), slide("data", source_table_id=table.table_id), slide("end")]
    docx_tables.bind_source_tables(slides, [table])
    return ContentPlan(brief="Отчёт", purpose="Сравнить регионы", slides=slides)


class ExtractionTests(unittest.TestCase):
    def test_order_exact_cells_and_persistent_registry(self):
        original = docx_bytes()
        document = docx_tables.extract_docx("report.docx", original)
        self.assertEqual(document.tables[0].data.rows[0], ["Орёл", "1 234,50"])
        self.assertEqual(document.tables[1].data.rows[0], ["Счёт", "007"])
        self.assertNotIn("1 234,50", document.text)
        self.assertLess(document.text.index("До таблиц"), document.text.index(document.tables[0].table_id))
        with TemporaryDirectory() as tmp:
            settings = SimpleNamespace(STORAGE_DIR=Path(tmp))
            docx_tables.save_material(document, original, settings)
            loaded = docx_tables.load_materials([document.material_id], settings)[0]
            self.assertEqual(loaded, document)
            self.assertNotIn("[[", loaded.text)
            self.assertIn("Колонки:", loaded.text)
            self.assertEqual((Path(tmp) / "materials" / document.material_id / "source.docx").read_bytes(), original)

    def test_large_merged_and_ragged_tables_are_not_silently_shortened(self):
        large = [["Категория", "Число"]] + [[str(i), str(i)] for i in range(7)]
        document = docx_tables.extract_docx("report.docx", docx_bytes([large]))
        self.assertFalse(document.tables[0].available)
        self.assertEqual(document.tables[0].data.rows, large[1:])
        self.assertIsNone(docx_tables.public_table(document.tables[0])["data"])
        self.assertTrue(docx_tables.public_table(document.tables[0])["analysis_available"])
        cases = [docx_bytes(merged=True), docx_bytes([[["A", "B"], ["a"]]])]
        for content in cases:
            with self.subTest(size=len(content)):
                document = docx_tables.extract_docx("report.docx", content)
                self.assertFalse(document.tables[0].available)
                self.assertIsNone(document.tables[0].data)
                self.assertTrue(document.tables[0].limitation)

    def test_invalid_documents_and_ids(self):
        with self.assertRaises(MaterialError):
            docx_tables.extract_docx("bad.docx", b"not a zip")
        for ids in [["../../secret"], [1], ["material_" + "a" * 32] * 2]:
            with self.subTest(ids=ids), self.assertRaises((MaterialError, TypeError)):
                docx_tables.load_materials(ids, SimpleNamespace(STORAGE_DIR=Path("unused")))

    def test_reference_resolution_replaces_model_cells_and_rejects_invalid_selection(self):
        table = docx_tables.extract_docx("report.docx", docx_bytes()).tables[0]
        plan = plan_for(table)
        plan.slides[1].content_blocks[0].table.rows[0][1] = "999999"
        docx_tables.bind_source_tables(plan.slides, [table])
        self.assertEqual(plan.slides[1].content_blocks[0].table, table.data)
        self.assertIsNot(plan.slides[1].content_blocks[0].table, table.data)
        for position, reference in [(0, table.table_id), (2, table.table_id), (1, "unknown")]:
            copy = plan.model_copy(deep=True)
            copy.slides[position].source_table_id = reference
            with self.subTest(position=position), self.assertRaises(ValueError):
                docx_tables.bind_source_tables(copy.slides, [table])

    def test_manual_text_edit_preserves_source_and_cells(self):
        table = docx_tables.extract_docx("report.docx", docx_bytes()).tables[0]
        spec = plan_for(table).slides[1]
        edited = apply_manual_text(spec, "Новый заголовок", "Комментарий", None)
        self.assertEqual(edited.source_table_id, table.table_id)
        self.assertEqual(edited.content_blocks[-1].table, table.data)

    def test_source_table_is_never_dropped_as_a_section(self):
        table = docx_tables.extract_docx("report.docx", docx_bytes()).tables[0]
        plan = plan_for(table)
        plan.slides[1].purpose = "Раздел с результатами"
        for variant in template_engine.VALID_VARIANTS:
            self.assertNotIn(1, dramaturgy.dropped_slides(plan, variant))


class ModelTests(unittest.IsolatedAsyncioTestCase):
    def settings(self):
        return SimpleNamespace(LLM_BASE_URL="http://model.invalid/v1", LLM_MODEL_NAME="test",
                               LLM_API_KEY=None, TABLE_SUMMARY_TIMEOUT_S=1, LLM_PLANNING_TIMEOUT_S=30,
                               PLAN_REPAIR_ROUNDS=0, PLAN_REPAIR_TIMEOUT_S=60)

    async def test_summaries_change_only_descriptions_and_report_fallback(self):
        document = docx_tables.extract_docx("report.docx", docx_bytes())
        before = deepcopy(document.tables[0].data)
        client = AsyncMock()
        client.__aenter__.return_value = client
        client.chat_structured.return_value = docx_tables.TableSummaries(tables=[
            docx_tables.TableSummary(table_id=t.table_id, summary="Продажи по регионам") for t in document.tables])
        with patch.object(docx_tables, "LLMClient", return_value=client), \
             patch.object(docx_tables, "load_skill", return_value={"system_prompt": "Describe"}):
            await docx_tables.describe_tables(document, self.settings())
            self.assertEqual(document.tables[0].summary, "Продажи по регионам")
            self.assertEqual(document.tables[0].data, before)
            client.chat_structured.side_effect = TimeoutError()
            await docx_tables.describe_tables(document, self.settings())
            self.assertIn("LLM-описания недоступны", document.note)
            self.assertEqual(document.tables[0].data, before)

    async def test_three_variants_resolve_ids_after_narrative_normalization(self):
        document = docx_tables.extract_docx("report.docx", docx_bytes())
        table = document.tables[0]
        job = Job(job_id="test", brief="Отчёт", purpose="Сравнить", template_path="unused",
                  source_tables=document.tables, materials_text="x" * 20_000)
        client = AsyncMock()
        client.__aenter__.return_value = client
        client.chat_structured.side_effect = lambda **kw: briefing.PlanOut(
            brief="Отчёт", purpose="Сравнить", story="Сравнение", slides=[slide("cover"),
                slide("data", source_table_id=table.table_id), slide("end")])
        skill_dir = Path(__file__).resolve().parents[1] / "skills"
        from app.config import Settings
        settings = Settings(_env_file=None, SKILLS_DIR=skill_dir)
        with patch.object(briefing, "LLMClient", return_value=client), \
             patch.object(briefing, "get_settings", return_value=self.settings()), \
             patch("app.skills.get_settings", return_value=settings):
            variants = await briefing.generate_plan_variants(job, 3, None)
        self.assertEqual(set(variants), {"A", "B", "C"})
        for variant in variants.values():
            self.assertEqual(variant.plan.slides[1].content_blocks[-1].table, table.data)
        user = client.chat_structured.call_args.kwargs["user_input"]
        self.assertIn(document.tables[-1].table_id, user)
        self.assertNotIn("1 234,50", user)

    async def test_invalid_model_reference_gets_corrective_attempt(self):
        table = docx_tables.extract_docx("report.docx", docx_bytes()).tables[0]
        client = AsyncMock()
        client.__aenter__.return_value = client
        client.chat_structured.side_effect = [slide("data", source_table_id="unknown"), slide("data", source_table_id=table.table_id)]
        with patch.object(briefing, "_client", return_value=client):
            result = await briefing._structured("system", "user", SlideSpec, max_tokens=500,
                validate=lambda s: docx_tables.bind_source_tables([s], [table], single_slide=True))
        self.assertEqual(client.chat_structured.await_count, 2)
        self.assertEqual(result.content_blocks[-1].table, table.data)
        self.assertIn("Ошибка проверки", client.chat_structured.call_args.kwargs["user_input"])


class ApiTests(unittest.TestCase):
    def test_upload_and_create_job_use_server_registry(self):
        from app.config import Settings
        from app.job_store import get
        with TemporaryDirectory() as tmp:
            settings = Settings(_env_file=None, STORAGE_DIR=Path(tmp), SKILLS_DIR=Path(__file__).resolve().parents[1] / "skills")
            settings.templates_dir.mkdir()
            with patch.object(routes, "get_settings", return_value=settings), \
                 patch.object(docx_tables, "describe_tables", new=AsyncMock()), \
                 patch.object(routes, "run_planning", new=AsyncMock()):
                with TestClient(app) as client:
                    response = client.post("/materials/extract", files={"file": ("report.docx", docx_bytes())})
                    self.assertEqual(response.status_code, 200)
                    extracted = response.json()
                    self.assertEqual(extracted["status"], "ok")
                    self.assertEqual(len(extracted["tables"]), 2)
                    template = BytesIO(); Presentation().save(template)
                    response = client.post("/jobs", data={"brief": "Отчёт", "idea": "Сравнить регионы",
                        "material_ids_json": json.dumps([extracted["material_id"]])},
                        files={"template": ("template.pptx", template.getvalue())})
                    self.assertEqual(response.status_code, 200, response.text)
                    job = get(response.json()["job_id"])
                    self.assertEqual(job.source_tables[0].data.rows[0], ["Орёл", "1 234,50"])
                    response = client.post("/jobs", data={"brief": "Bad", "material_ids_json": '["../../other"]'})
                    self.assertEqual(response.status_code, 400)


class RenderTests(unittest.TestCase):
    def test_all_three_designs_keep_exact_editable_source_tables(self):
        from app.pipeline.parser import parse_template
        table = docx_tables.extract_docx("report.docx", docx_bytes()).tables[0]
        plan = plan_for(table)
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "template.pptx"
            prs = Presentation()
            prs.save(path)
            template = template_engine.build_template_json(path)
            manifest = parse_template(path)
            structured = [dramaturgy.structure_fallback(i, s, len(plan.slides)) for i, s in enumerate(plan.slides)]
            for variant in template_engine.VALID_VARIANTS:
                layout = dramaturgy.build_variant_plan(template, plan, structured, variant, plan.brief, template_engine._shorten)
                output = Path(tmp) / f"{variant}.pptx"
                template_engine.render_variant(template, layout, path, output, manifest, plan, variant)
                loaded = Presentation(output)
                docx_tables.verify_rendered_tables(loaded, plan, layout)
                cells = [[c.text for c in row.cells] for shape in loaded.slides[1].shapes if shape.has_table for row in shape.table.rows]
                self.assertEqual(cells, [table.data.headers, *table.data.rows])

    def test_missing_table_is_an_export_error(self):
        table = docx_tables.extract_docx("report.docx", docx_bytes()).tables[0]
        prs = Presentation()
        for _ in range(3):
            prs.slides.add_slide(prs.slide_layouts[6])
        with self.assertRaisesRegex(ValueError, "не перенесена целиком"):
            docx_tables.verify_rendered_tables(prs, plan_for(table), {"slides": [{}, {}, {}]})


if __name__ == "__main__":
    unittest.main()
