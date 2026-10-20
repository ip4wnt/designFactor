"""Publication names and units are independent of SQL keys and cell values."""
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import json,unittest
from unittest.mock import AsyncMock,patch

from agents.tool_context import ToolContext
from pptx import Presentation

from app.analysis import worker
from app.config import Settings
from app.pipeline import briefing,docx_tables,dramaturgy,table_analysis,template_engine
from app.schemas.content_plan import ContentPlan,TableData
from app.schemas.source_tables import SourceTable
from test_docx_tables import slide


KEYS=['city','rev_year','orders','aov','share']
META=[dict(key=k,label=label,unit=unit) for k,label,unit in [
    ('city','Город',''),('rev_year','Годовая выручка','тыс. руб.'),('orders','Заказы','шт.'),
    ('aov','Средний чек','тыс. руб.'),('share','Доля выручки','%')]]
HEADERS=['Город','Годовая выручка, тыс. руб.','Заказы, шт.','Средний чек, тыс. руб.','Доля выручки, %']


def prepared():
    return dict(kind='table',title='Сравнение городов',reason='Годовые показатели',columns=KEYS,
        data=[dict(zip(KEYS,['Москва',29319.9,19667,1.49,16.5]))],notes=[],
        parent_table_ids=['input'],operations=[dict(tool='create_view',arguments={'sql':'SELECT ...'})],
        presentation_columns=deepcopy(META))


class ColumnLabelsTests(unittest.TestCase):
    def test_headers_get_labels_and_units_without_changing_source_keys_or_values(self):
        item=prepared(); before=deepcopy(item)
        artifact=worker.compile_artifacts([item],{'key':'test'})[0]
        self.assertEqual(artifact.data.headers,HEADERS)
        self.assertEqual(artifact.data.rows,[['Москва','29319.9','19667','1.49','16.5']])
        self.assertEqual([c.key for c in artifact.column_metadata],KEYS)
        self.assertEqual(item,before)

    def test_metadata_order_does_not_reorder_data_and_units_are_not_duplicated(self):
        item=prepared(); item['presentation_columns'].reverse()
        item['presentation_columns'][0]['label']='Доля выручки, %'
        artifact=worker.compile_artifacts([item],{'key':'test'})[0]
        self.assertEqual(artifact.data.headers,HEADERS)
        self.assertEqual(artifact.data.rows[0][-1],'16.5')

    def test_incomplete_duplicate_or_technical_headers_cannot_be_published(self):
        invalid=[None,META[:-1],META+[META[0]],
            [{**c,'label':c['key']} for c in META],
            [{**c,'label':'Одинаково','unit':''} for c in META]]
        for metadata in invalid:
            with self.subTest(metadata=metadata),self.assertRaises(ValueError):
                worker.publication_columns(KEYS,metadata)

    def test_unknown_units_are_not_invented_and_string_values_are_preserved(self):
        item=prepared(); item['columns']=['code','value']; item['data']=[{'code':'000007','value':'1 234,50'}]
        item['presentation_columns']=[dict(key='code',label='Код',unit=''),dict(key='value',label='Измерение',unit='')]
        artifact=worker.compile_artifacts([item],{'key':'test'})[0]
        self.assertEqual(artifact.data.headers,['Код','Измерение'])
        self.assertEqual(artifact.data.rows,[['000007','1 234,50']])

    def test_all_three_pptx_designs_receive_presentation_headers_and_exact_rows(self):
        table=worker.compile_artifacts([prepared()],{'key':'test'})[0]
        plan=ContentPlan(brief='Результаты',purpose='Сравнить города',slides=[slide('cover'),
            slide('data',source_table_id=table.table_id),slide('end')])
        docx_tables.bind_source_tables(plan.slides,[table])
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'template.pptx'; Presentation().save(path)
            template=template_engine.build_template_json(path)
            from app.pipeline.parser import parse_template
            manifest=parse_template(path)
            structured=[dramaturgy.structure_fallback(i,s,len(plan.slides)) for i,s in enumerate(plan.slides)]
            for variant in template_engine.VALID_VARIANTS:
                layout=dramaturgy.build_variant_plan(template,plan,structured,variant,plan.brief,template_engine._shorten)
                output=Path(tmp)/(variant+'.pptx')
                template_engine.render_variant(template,layout,path,output,manifest,plan,variant)
                saved=Presentation(output)
                tables=[[[c.text for c in row.cells] for row in shape.table.rows]
                    for s in saved.slides for shape in s.shapes if shape.has_table]
                self.assertIn([HEADERS,*table.data.rows],tables)


class ToolPublicationTests(unittest.IsolatedAsyncioTestCase):
    async def test_tool_requires_metadata_and_failed_publication_keeps_no_partial_slide(self):
        with TemporaryDirectory() as tmp:
            settings=Settings(_env_file=None,STORAGE_DIR=Path(tmp),SKILLS_DIR=Path(__file__).resolve().parents[1]/'skills')
            table=SourceTable(table_id='input',source_name='source.docx',ordinal=1,summary='Продажи',available=False,
                data=TableData(headers=['Город','Выручка, тыс. руб.'],rows=[['Москва','29319,9']]))
            from app.schemas.job import Job
            job=Job(job_id='test',brief='Отчет',purpose='Сравнить',template_path='unused',source_tables=[table])
            with patch('app.skills.get_settings',return_value=settings):
                request=table_analysis.analysis_request(job,settings)
            async def fake_run(agent,user,*,context,**kwargs):
                tools={t.name:t for t in agent.tools}
                self.assertIn('presentation_columns',tools['prepare_slide_data'].params_json_schema['required'])
                async def call(name,args):
                    ctx=ToolContext(context=context,tool_name=name,tool_call_id='test',tool_arguments=json.dumps(args))
                    return json.loads(await tools[name].on_invoke_tool(ctx,json.dumps(args)))
                result=await call('create_view',{'name':'sales','sql':
                    'SELECT col_01_gorod AS city, TRY_CAST(replace(col_02_vyruchka_tys_rub, \',\', \'.\') AS DOUBLE) AS rev_year FROM table_1_sql'})
                self.assertNotIn('error',result)
                args=dict(title='Годовая выручка',kind='table',reason='Сравнить города',x=None,y=None,source='sales')
                self.assertIn('error',await call('prepare_slide_data',args))
                self.assertEqual(context.slides,[])
                args['presentation_columns']=META[:2]
                result=await call('prepare_slide_data',args)
                self.assertEqual(result['display_headers'],HEADERS[:2])
                self.assertEqual(list(context.views['sales'].columns),['city','rev_year'])
            with patch.object(worker.Runner,'run',new=fake_run):
                report=await worker.run_agent(request,Path(tmp)/'progress.json',settings)
            self.assertEqual(report.status,'completed',report.notes)
            self.assertEqual(report.artifacts[0].data.headers,HEADERS[:2])
            self.assertEqual(report.artifacts[0].data.rows,[['Москва','29319.9']])
