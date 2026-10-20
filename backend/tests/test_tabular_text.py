"""CSV, табличный JSON и markdown-таблицы в тексте приводятся к SourceTable, как листы XLSX."""
import json
import unittest

from pptx.util import Pt  # noqa: F401  (проверка, что окружение с python-pptx)
from pptx.dml.color import RGBColor

from app.pipeline import tabular_text
from app.pipeline.materials import MaterialError
from app.pipeline.styling import contrast_palette


class TabularTextTests(unittest.TestCase):
    def test_csv_semicolon_and_comma_decimal(self):
        raw = "Месяц;Выручка, млн руб.\nЯнварь;12,5\nФевраль;13,1\n".encode("utf-8")
        document = tabular_text.extract_csv("revenue.csv", raw)
        self.assertEqual(len(document.tables), 1)
        table = document.tables[0]
        self.assertTrue(table.available)
        self.assertEqual(table.data.headers, ["Месяц", "Выручка, млн руб."])
        self.assertEqual(table.data.rows[1], ["Февраль", "13,1"])
        self.assertIn(f"[[{table.table_id}]]", document.text)

    def test_csv_large_table_kept_for_analysis_but_not_for_slide(self):
        rows = ["id,channel,amount"] + [f"{i},web,{i * 10}" for i in range(1, 1201)]
        document = tabular_text.extract_csv("orders.csv", "\n".join(rows).encode("utf-8"))
        table = document.tables[0]
        self.assertFalse(table.available)
        self.assertEqual(len(table.data.rows), 1200)
        self.assertIn("допустимо до 6 строк", table.limitation)

    def test_json_records_and_columns_rows(self):
        payload = {"shares": [{"Продукт": "A", "Доля": 0.5}, {"Продукт": "B", "Доля": 0.5}],
                   "matrix": {"columns": ["Регион", "План"], "rows": [["Москва", 100], ["Казань", 90]]}}
        document = tabular_text.extract_json("data.json", json.dumps(payload, ensure_ascii=False).encode("utf-8"))
        self.assertEqual(len(document.tables), 2)
        self.assertEqual(document.tables[0].data.headers, ["Продукт", "Доля"])
        self.assertEqual(document.tables[1].data.rows[0], ["Москва", "100"])
        with self.assertRaises(MaterialError):
            tabular_text.extract_json("x.json", b'{"a": 1}')

    def test_markdown_table_inside_text(self):
        text = ("Отчет за квартал\n\nОбращения в поддержку\n\n| Канал | Обращений | SLA, % |\n|---|---|---|\n"
                "| Чат | 420 | 97 |\n| Почта | 180 | 91 |\n\nВыводы ниже.\n")
        document = tabular_text.extract_text_tables("report.txt", text.encode("utf-8"))
        self.assertIsNotNone(document)
        table = document.tables[0]
        self.assertTrue(table.available)
        self.assertEqual(table.data.headers, ["Канал", "Обращений", "SLA, %"])
        self.assertIn("Обращения в поддержку", table.summary)
        self.assertIn(f"[[{table.table_id}]]", document.text)
        self.assertIn("Выводы ниже.", document.text)
        self.assertIsNone(tabular_text.extract_text_tables("plain.txt", "просто текст без таблиц".encode("utf-8")))


class ContrastPaletteTests(unittest.TestCase):
    def test_white_series_dropped_on_light_background(self):
        palette = [RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x00, 0x77, 0xFF), RGBColor(0xF5, 0xF5, 0xF5)]
        self.assertEqual(contrast_palette(palette, dark=False), [RGBColor(0x00, 0x77, 0xFF)])

    def test_black_series_dropped_on_dark_background_and_fallback(self):
        palette = [RGBColor(0x00, 0x00, 0x00), RGBColor(0x1A, 0x1A, 0x1A)]
        self.assertEqual(len(contrast_palette(palette, dark=True)), 6)
        self.assertEqual(contrast_palette([RGBColor(0xFF, 0xC0, 0x00)], dark=True), [RGBColor(0xFF, 0xC0, 0x00)])


if __name__ == "__main__":
    unittest.main()
