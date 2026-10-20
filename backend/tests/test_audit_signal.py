"""Аудит ловит обрывы фраз и не шумит на свойствах самого шаблона."""
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from app.pipeline import audit
from app.schemas.design_manifest import DesignManifest, Typography, TypographyStyle


def _manifest():
    return DesignManifest(template_id="audit", palette={"dk1": "18233B", "lt1": "FFFFFF", "accent1": "0077FF"},
                          typography=Typography(title=TypographyStyle(font="Arial", size=24), body=TypographyStyle(font="Arial", size=18)),
                          layouts=[], slide_width_emu=12192000, slide_height_emu=6858000)


def _deck(texts, title_text="Заголовок на своем месте", title_size=24):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only: плейсхолдер заголовка от макета
    slide.shapes.title.text = title_text
    for run in slide.shapes.title.text_frame.paragraphs[0].runs:
        run.font.size = Pt(title_size)
    top = Inches(2)
    for text in texts:
        box = slide.shapes.add_textbox(Inches(1), top, Inches(8), Inches(0.8))
        box.text_frame.text = text
        for run in box.text_frame.paragraphs[0].runs:
            run.font.size = Pt(21)
        top += Inches(1)
    return prs


class TruncationTests(unittest.TestCase):
    def _issues(self, texts, expected=None, **kw):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            _deck(texts, **kw).save(path)
            return audit.audit_presentation(path, allowed_sizes=[18, 24, 32], expected_texts=expected or [])

    def test_hanging_tail_detected_with_bbox(self):
        issues = [i for i in self._issues(["Анализ роста, региональных разрывов и"]) if i.code == "text_truncated"]
        self.assertEqual(len(issues), 1)
        self.assertIn("обрывается на слове «и»", issues[0].description)
        self.assertIsNotNone(issues[0].bbox)

    def test_prefix_of_plan_phrase_detected_and_heads_allowed(self):
        plan = ["Рост выручки подтвержден, но без усиления слабых каналов потенциал не будет реализован в ближайшие годы",
                "Квартальная динамика: выручка растет каждый квартал"]
        cut = "Рост выручки подтвержден, но без усиления слабых каналов потенциал"
        issues = [i for i in self._issues([cut, "Квартальная динамика"], plan) if i.code == "text_truncated"]
        self.assertEqual(len(issues), 1)
        self.assertIn("начало более длинной фразы", issues[0].description)

    def test_complete_sentences_are_clean(self):
        issues = [i for i in self._issues(["Выручка выросла на 44% за год.", "Согласуйте бюджет до конца квартала"])
                  if i.code == "text_truncated"]
        self.assertEqual(issues, [])


class TemplateNoiseTests(unittest.TestCase):
    def _codes(self, **kw):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            _deck(["Обычный абзац текста на слайде."], **kw).save(path)
            return [(i.code, i.description) for i in audit.audit_presentation(path, manifest=_manifest(), allowed_sizes=[18, 24, 32])]

    def test_in_range_font_size_is_not_reported(self):
        codes = self._codes()
        self.assertFalse(any(c == "font_size_off_scale" for c, _ in codes), codes)

    def test_out_of_range_font_size_is_reported(self):
        codes = self._codes(title_size=60)
        self.assertTrue(any(c == "font_size_off_scale" and "60pt" in d for c, d in codes), codes)

    def test_neutral_gray_is_not_palette_violation(self):
        self.assertTrue(audit._is_neutral_gray("8F8F8F"))
        self.assertTrue(audit._is_neutral_gray("1B1B1F"))
        self.assertFalse(audit._is_neutral_gray("FF8833"))

    def test_layout_placeholder_near_edge_is_not_reported(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            prs = _deck([])
            slide = prs.slides[0]
            slide.shapes.title.left = Inches(0.05)  # как «у края» поставил бы сам макет
            layout_title = [s for s in slide.slide_layout.placeholders if s.placeholder_format.idx == slide.shapes.title.placeholder_format.idx][0]
            layout_title.left = Inches(0.05)
            prs.save(path)
            issues = audit.audit_presentation(path, allowed_sizes=[18, 24, 32])
            self.assertFalse(any(i.code == "edge_margin" for i in issues), [i.description for i in issues])


if __name__ == "__main__":
    unittest.main()
