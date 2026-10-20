"""Run: python -m unittest discover -s tests -v (from backend)."""
import datetime
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from app.pipeline.excel_import import import_excel_block, sheet_to_chart_data, select_chart_type
from app.pipeline.styling import style_chart
from app.schemas.content_plan import ChartType
from app.schemas.design_manifest import DesignManifest, Typography, TypographyStyle


class ChartImprovementsTest(unittest.TestCase):
    def test_dates_choose_line_and_comparison_goal_overrides(self):
        rows = [["Date", "Revenue"], ["2026-02", 120], ["2026-01", 100]]
        self.assertEqual(select_chart_type(rows), ChartType.LINE)
        self.assertEqual(select_chart_type(rows, "сравнить месяцы"), ChartType.BAR)

    def test_pie_requires_small_nonnegative_complete_series(self):
        rows = [["Product", "Revenue"], ["A", 70], ["B", 30]]
        self.assertEqual(select_chart_type(rows, "показать доли"), ChartType.PIE)
        rows[1][1] = -1
        self.assertEqual(select_chart_type(rows, "показать доли"), ChartType.BAR)

    def test_missing_values_are_gaps_and_invalid_text_is_rejected(self):
        rows = [["Region", "Revenue"], ["A", "1 234,5"], ["B", None]]
        chart = sheet_to_chart_data("Data", rows)
        self.assertEqual(chart.series[0].values, [1234.5, None])
        rows[2][1] = "not a number"
        self.assertIsNone(sheet_to_chart_data("Data", rows))
        rows[2][1] = float("nan")
        self.assertIsNone(sheet_to_chart_data("Data", rows))

    def test_import_ignores_sheet_name_sorts_dates_and_preserves_reference(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/"data.xlsx"
            wb = Workbook(); ws = wb.active; ws.title = "Обычный лист"
            for row in [["Date", "Revenue"], [datetime.date(2026,2,1),120], [datetime.date(2026,1,1),100]]:
                ws.append(row)
            wb.save(path); wb.close()
            kind, chart = import_excel_block(path)
            self.assertEqual(kind, "chart")
            self.assertEqual(chart.chart_type, ChartType.LINE)
            self.assertEqual(chart.series[0].values, [100,120])
            wb = Workbook(); ws = wb.active
            ws.append(["Item", "Value"])
            for i in range(200): ws.append([f"Item {i}",i])
            wb.save(path); wb.close()
            kind, table = import_excel_block(path)
            self.assertEqual(kind, "table")
            self.assertEqual(len(table.rows), 200)

    def test_native_charts_save_reload_with_theme_and_gaps(self):
        manifest = DesignManifest(template_id="test", palette={"accent1":"1254AA", "accent2":"FF8833"},
            typography=Typography(title=TypographyStyle(font="Arial",size=28),body=TypographyStyle(font="Arial",size=18)),
            layouts=[],slide_width_emu=9144000,slide_height_emu=6858000)
        prs = Presentation()
        for kind in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.LINE, XL_CHART_TYPE.COLUMN_CLUSTERED):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            data = CategoryChartData(); data.categories = ["A","B","C"]
            data.add_series("Revenue", [10, None if kind == XL_CHART_TYPE.LINE else 20,30])
            chart = slide.shapes.add_chart(kind, Inches(1), Inches(1), Inches(8), Inches(5), data).chart
            style_chart(chart,manifest)
            self.assertEqual(chart.plots[0].data_labels.number_format, "0%" if kind == XL_CHART_TYPE.PIE else "#,##0")
            if kind == XL_CHART_TYPE.PIE:
                self.assertTrue(chart.has_legend)
                self.assertTrue(chart.plots[0].data_labels.show_percentage)
                self.assertNotEqual(chart.series[0].points[0].format.fill.fore_color.rgb,chart.series[0].points[1].format.fill.fore_color.rgb)
            if kind == XL_CHART_TYPE.LINE:
                self.assertEqual(str(chart.series[0].format.line.color.rgb), "1254AA")
                self.assertEqual(chart._chartSpace.xpath(".//c:dispBlanksAs")[0].get("val"),"gap")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/"charts.pptx"; prs.save(path)
            loaded = Presentation(path)
            self.assertEqual(len(loaded.slides),3)
            self.assertIsNone(loaded.slides[1].shapes[0].chart.series[0].values[1])

if __name__ == "__main__":
    unittest.main()
