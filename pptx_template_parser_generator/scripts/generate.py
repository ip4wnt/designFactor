
import os
import argparse
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from PIL import ImageFont

from pptx_template_parser import build_template
from pptx_template_parser.template.export import save_template_json


# ============================================================
# НАСТРОЙКИ
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "output"



# ============================================================
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ============================================================

def hex_to_rgb(hex_color):
    """
    '#FAFCFF' -> RGBColor(...)
    """

    if not hex_color:
        return RGBColor(0, 0, 0)

    hex_color = str(
        hex_color
    ).lstrip("#")

    if len(hex_color) != 6:
        raise ValueError(
            f"Некорректный HEX-цвет: {hex_color}"
        )

    return RGBColor(
        int(hex_color[0:2], 16),
        int(hex_color[2:4], 16),
        int(hex_color[4:6], 16),
    )


def cm_to_px(cm):
    """
    Перевод сантиметров в пиксели при 96 DPI.
    """

    return cm * 96 / 2.54


# ============================================================
# ШРИФТ
# ============================================================

def find_font(font_family):
    """
    Ищет шрифт в C:\\Windows\\Fonts.

    Если шрифт не найден — возвращает None.
    """

    if not font_family:
        return None

    fonts_folder = r"C:\Windows\Fonts"

    possible_names = [
        f"{font_family}.ttf",
        f"{font_family}-Regular.ttf",
        f"{font_family.lower()}.ttf",
        f"{font_family.lower()}-regular.ttf",
        f"{font_family}.otf",
        f"{font_family}-Regular.otf",
        f"{font_family.lower()}.otf",
        f"{font_family.lower()}-regular.otf",
    ]

    for name in possible_names:
        path = os.path.join(
            fonts_folder,
            name,
        )

        if os.path.exists(path):
            return path

    return None


def text_width(text, font):
    """
    Получает ширину строки в пикселях.
    """

    bbox = font.getbbox(text)

    return bbox[2] - bbox[0]


def text_fits(
    text,
    slot,
    font_config,
    margins,
):
    """
    Проверяет, помещается ли текст по ширине.

    Если физического файла шрифта нет,
    проверка пропускается.
    """

    if not text:
        return True

    if not font_config:
        return True

    family = font_config.get("family")
    size = font_config.get("size")

    if not family or not size:
        return True

    font_path = find_font(family)

    if font_path is None:
        print(
            f"  ⚠ Шрифт '{family}' не найден "
            f"в C:\\Windows\\Fonts"
        )

        print(
            "    → проверка ширины текста пропущена"
        )

        return True

    font = ImageFont.truetype(
        font_path,
        int(size),
    )

    max_width = cm_to_px(
        slot["w"]
        - margins.get("left", 0)
        - margins.get("right", 0)
    )

    for line in str(text).split("\n"):
        width = text_width(
            line,
            font,
        )

        if width > max_width:
            return False

    return True


# ============================================================
# ВАЛИДАЦИЯ
# ============================================================

def validate_slot(
    slot,
    slide_width,
    slide_height,
):
    """
    Проверяет, находится ли слот внутри слайда.
    """

    return (
        slot["x"] >= 0
        and slot["y"] >= 0
        and slot["w"] >= 0
        and slot["h"] >= 0
        and slot["x"] + slot["w"] <= slide_width
        and slot["y"] + slot["h"] <= slide_height
    )


def get_variant(
    template,
    variant_id,
):
    """
    Находит вариант по его реальному id,
    извлеченному из PPTX.
    """

    for variant in template["variants"]:
        if variant["id"] == variant_id:
            return variant

    available = [
        variant["id"]
        for variant in template["variants"]
    ]

    raise ValueError(
        f"В шаблоне нет варианта: {variant_id}\n\n"
        f"Доступные варианты:\n"
        + "\n".join(
            f"- {value}"
            for value in available
        )
    )


# ============================================================
# ВЫРАВНИВАНИЕ
# ============================================================

def get_horizontal_alignment(value):
    """
    Преобразует значение из шаблона
    в значение python-pptx.
    """

    if not value:
        return PP_ALIGN.LEFT

    value = str(value).lower()

    if "center" in value:
        return PP_ALIGN.CENTER

    if "right" in value:
        return PP_ALIGN.RIGHT

    if "justify" in value:
        return PP_ALIGN.JUSTIFY

    if "distributed" in value:
        return PP_ALIGN.DISTRIBUTE

    return PP_ALIGN.LEFT


def get_vertical_alignment(value):
    """
    Преобразует vertical_align из шаблона.
    """

    if not value:
        return MSO_ANCHOR.TOP

    value = str(value).lower()

    if value == "top":
        return MSO_ANCHOR.TOP

    if value == "bottom":
        return MSO_ANCHOR.BOTTOM

    if value in (
        "middle",
        "center",
    ):
        return MSO_ANCHOR.MIDDLE

    return MSO_ANCHOR.TOP


# ============================================================
# РЕСУРСЫ ИЗ PPTX
# ============================================================

def get_asset_path(
    target,
    assets_dir,
):
    """
    Преобразует внутренний путь PPTX:

        ppt/media/image4.png

    в реальный путь внутри распакованного PPTX.
    """

    if not target:
        return None

    target = str(target).lstrip(
        "/\\"
    )

    path = os.path.join(
        assets_dir,
        *target.split("/"),
    )

    if os.path.isfile(path):
        return path

    return None


# ============================================================
# ЗАГЛУШКА ДЛЯ ОТСУТСТВУЮЩЕГО ИЗОБРАЖЕНИЯ
# ============================================================

def add_image_placeholder(
    slide,
    image,
    text="IMAGE PLACEHOLDER",
):
    """
    Создает прямоугольник на месте отсутствующего
    изображения.

    Координаты и размер берутся из шаблона.
    """

    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Cm(image["x"]),
        Cm(image["y"]),
        Cm(image["w"]),
        Cm(image["h"]),
    )

    # --------------------------------------------------------
    # Фон
    # --------------------------------------------------------

    shape.fill.solid()

    shape.fill.fore_color.rgb = RGBColor(
        80,
        80,
        80,
    )

    # --------------------------------------------------------
    # Рамка
    # --------------------------------------------------------

    shape.line.color.rgb = RGBColor(
        150,
        150,
        150,
    )

    # --------------------------------------------------------
    # Текст
    # --------------------------------------------------------

    text_frame = shape.text_frame

    text_frame.clear()

    text_frame.vertical_anchor = (
        MSO_ANCHOR.MIDDLE
    )

    paragraph = text_frame.paragraphs[0]

    paragraph.alignment = (
        PP_ALIGN.CENTER
    )

    run = paragraph.add_run()

    run.text = text

    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(
        230,
        230,
        230,
    )

    return shape


# ============================================================
# BACKGROUND
# ============================================================

def add_background(
    slide,
    background,
    assets_dir,
):
    """
    Добавляет настоящий background из исходного PPTX.

    Если ресурс действительно отсутствует,
    вставляется заглушка.
    """

    if not background:
        return

    if background.get("color"):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = hex_to_rgb(background["color"])

    if not background.get("target"):
        return

    image_path = get_asset_path(
        background["target"],
        assets_dir,
    )

    if image_path:
        slide.shapes.add_picture(
            image_path,
            Cm(background["x"]),
            Cm(background["y"]),
            width=Cm(background["w"]),
            height=Cm(background["h"]),
        )

        return

    print(
        "  ⚠ Background не найден:"
        f" {background['target']}"
    )

    print(
        "    → вставлена заглушка"
    )

    add_image_placeholder(
        slide,
        background,
        "BACKGROUND PLACEHOLDER",
    )


# ============================================================
# ОБЫЧНОЕ ИЗОБРАЖЕНИЕ
# ============================================================

def add_image(
    slide,
    image,
    assets_dir,
):
    """
    Добавляет изображение из исходного PPTX.

    Если ресурс отсутствует — вставляет заглушку.
    """

    image_path = get_asset_path(
        image["target"],
        assets_dir,
    )

    if image_path:
        slide.shapes.add_picture(
            image_path,
            Cm(image["x"]),
            Cm(image["y"]),
            width=Cm(image["w"]),
            height=Cm(image["h"]),
        )

        return

    print(
        "  ⚠ Изображение не найдено:"
        f" {image['target']}"
    )

    print(
        "    → вставлена заглушка"
    )

    add_image_placeholder(
        slide,
        image,
    )


# ============================================================
# ЦВЕТ ТЕКСТА
# ============================================================

def get_text_color(
    slot,
    template,
):
    """
    Получает цвет текста.

    Приоритет:

    1. цвет самого текстового слота;
    2. text_primary;
    3. lt1;
    4. первый доступный цвет;
    5. черный.
    """

    font = slot.get("font") or {}

    color = font.get("color")

    if color:
        return color

    colors = template.get(
        "colors",
        {},
    )

    for key in (
        "text_primary",
        "lt1",
    ):
        value = colors.get(key)

        if value:
            return value

    for value in colors.values():
        if value:
            return value

    return "#000000"


# ============================================================
# ШРИФТ СЛОТА
# ============================================================

def get_slot_font(
    slot,
    template,
):
    """
    Берет шрифт непосредственно
    из конкретного текстового слота.

    Не используем title-шрифт как fallback,
    потому что у разных слотов могут быть
    разные параметры.
    """

    font = slot.get("font")

    if not font:
        raise ValueError(
            "У текстового слота отсутствует "
            "информация о шрифте:\n"
            f"{slot}"
        )

    return font


# ============================================================
# ДОБАВЛЕНИЕ ТЕКСТА
# ============================================================

def add_textbox(
    slide,
    text,
    slot_name,
    slot,
    template,
):
    """
    Создает текстовое поле строго
    по параметрам, извлеченным из шаблона.
    """

    slide_width = template[
        "slide_size"
    ]["width"]

    slide_height = template[
        "slide_size"
    ]["height"]

    if not validate_slot(
        slot,
        slide_width,
        slide_height,
    ):
        raise ValueError(
            "Текстовый слот выходит "
            "за границы слайда:\n\n"
            f"Slot: {slot_name}\n"
            f"{slot}"
        )

    margins = template.get(
        "text_box_defaults",
        {},
    ).get(
        "margins",
        {},
    )

    font_config = get_slot_font(
        slot,
        template,
    )

    # --------------------------------------------------------
    # Проверка ширины
    # --------------------------------------------------------

    if not text_fits(
        text,
        slot,
        font_config,
        margins,
    ):
        raise ValueError(
            "\n"
            "Текст не помещается в слот.\n\n"
            f"Slot: {slot_name}\n"
            f"Text: {text}\n"
            f"Font: "
            f"{font_config.get('family')} "
            f"{font_config.get('size')} pt\n"
        )

    # --------------------------------------------------------
    # TextBox
    # --------------------------------------------------------

    textbox = slide.shapes.add_textbox(
        Cm(slot["x"]),
        Cm(slot["y"]),
        Cm(slot["w"]),
        Cm(slot["h"]),
    )

    text_frame = textbox.text_frame

    text_frame.clear()

    # --------------------------------------------------------
    # Отступы
    # --------------------------------------------------------

    text_frame.margin_left = Cm(
        margins.get("left", 0)
    )

    text_frame.margin_right = Cm(
        margins.get("right", 0)
    )

    text_frame.margin_top = Cm(
        margins.get("top", 0)
    )

    text_frame.margin_bottom = Cm(
        margins.get("bottom", 0)
    )

    # --------------------------------------------------------
    # Вертикальное выравнивание
    # --------------------------------------------------------

    text_frame.vertical_anchor = (
        get_vertical_alignment(
            slot.get("vertical_align")
        )
    )

    # --------------------------------------------------------
    # Paragraph
    # --------------------------------------------------------

    paragraph = text_frame.paragraphs[0]

    paragraph.alignment = (
        get_horizontal_alignment(
            slot.get("horizontal_align")
        )
    )

    # --------------------------------------------------------
    # Run
    # --------------------------------------------------------

    run = paragraph.add_run()

    run.text = str(text)

    # --------------------------------------------------------
    # Шрифт
    # --------------------------------------------------------

    family = font_config.get(
        "family"
    )

    if family:
        run.font.name = family

    size = font_config.get(
        "size"
    )

    if size:
        run.font.size = Pt(size)

    weight = font_config.get(
        "weight",
        "regular",
    )

    run.font.bold = (
        weight == "bold"
    )

    run.font.italic = (
        weight == "italic"
    )

    # --------------------------------------------------------
    # Цвет
    # --------------------------------------------------------

    color = get_text_color(
        slot,
        template,
    )

    if color:
        run.font.color.rgb = (
            hex_to_rgb(color)
        )

    return textbox


# ============================================================
# СОЗДАНИЕ ОДНОГО СЛАЙДА
# ============================================================

def create_slide(
    prs,
    template,
    variant_id,
    content,
    assets_dir,
):
    """
    Создает один слайд по варианту шаблона.
    """

    variant = get_variant(
        template,
        variant_id,
    )

    # --------------------------------------------------------
    # Пустой layout
    # --------------------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    # --------------------------------------------------------
    # Background
    # --------------------------------------------------------

    add_background(
        slide,
        variant.get("background"),
        assets_dir,
    )

    # --------------------------------------------------------
    # Images
    # --------------------------------------------------------

    for image in variant.get(
        "images",
        [],
    ):
        add_image(
            slide,
            image,
            assets_dir,
        )

    # --------------------------------------------------------
    # Text
    # --------------------------------------------------------

    slots = variant.get(
        "slots",
        {},
    )

    for slot_name, text in content.items():

        if text is None:
            continue

        if slot_name not in slots:
            raise ValueError(
                f"В варианте '{variant_id}' "
                f"нет текстового слота "
                f"'{slot_name}'.\n\n"
                "Доступные слоты:\n"
                + "\n".join(
                    f"- {name}"
                    for name in slots
                )
            )

        slot = slots[slot_name]

        add_textbox(
            slide=slide,
            text=text,
            slot_name=slot_name,
            slot=slot,
            template=template,
        )

    return slide


def create_slide_from_layout(
    prs, template, variant_id, content, image_content=None,
    variant=None, layouts=None,
):
    """Use the original PowerPoint layout, including groups and logos."""
    variant = variant if variant is not None else get_variant(template, variant_id)
    if layouts is None:
        layouts = {
            str(layout.part.partname).lstrip("/"): layout
            for master in prs.slide_masters
            for layout in master.slide_layouts
        }
    layout = layouts[variant["layout_file"]]
    slide = prs.slides.add_slide(layout)

    # python-pptx omits footer placeholders when adding a slide. Some designs
    # place an editable footer at the *top*, so copy only explicitly declared
    # footer slots from their layout, preserving size and text styling.
    for slot in variant["slots"].values():
        idx = slot.get("placeholder_idx")
        if idx is None or any(shape.placeholder_format.idx == idx
                              for shape in slide.placeholders):
            continue
        layout_shape = next(
            (shape for shape in layout.shapes
             if shape.is_placeholder and shape.placeholder_format.idx == idx
             and shape.placeholder_format.type == PP_PLACEHOLDER.FOOTER),
            None,
        )
        if layout_shape is None:
            continue
        copied = deepcopy(layout_shape._element)
        copied.nvSpPr.cNvPr.set("id", str(slide.shapes._next_shape_id))
        slide.shapes._spTree.append(copied)

    for slot_name, value in content.items():
        if value is None:
            continue
        slot = variant["slots"].get(slot_name)
        if slot is None:
            raise ValueError(f"В макете '{variant_id}' нет слота '{slot_name}'")
        placeholder_idx = slot.get("placeholder_idx")
        if placeholder_idx is not None:
            slide.placeholders[placeholder_idx].text = str(value)
        else:
            add_textbox(slide, str(value), slot_name, slot, template)

    for slot_name, path in (image_content or {}).items():
        if path is None:
            continue
        slot = variant.get("image_slots", {}).get(slot_name)
        if slot is None:
            raise ValueError(f"В макете '{variant_id}' нет слота изображения '{slot_name}'")
        image_path = Path(path)
        if not image_path.is_file():
            raise FileNotFoundError(f"Изображение не найдено: {image_path}")
        slide.placeholders[slot["placeholder_idx"]].insert_picture(str(image_path))
    return slide


# ============================================================
# СОЗДАНИЕ ПРЕЗЕНТАЦИИ
# ============================================================

def _replace_picture_with_placeholder(slide, shape_index, placeholder_idx):
    """Convert an authored picture to a native, clickable picture placeholder."""
    picture = slide.shapes[shape_index]
    if picture.shape_type != MSO_SHAPE_TYPE.PICTURE:
        return
    image_rid = picture._element.blipFill.blip.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    )
    old = picture._element
    replacement = OxmlElement("p:sp")
    nv = OxmlElement("p:nvSpPr")
    cnv = OxmlElement("p:cNvPr")
    cnv.set("id", str(picture.shape_id))
    cnv.set("name", f"Изображение {placeholder_idx}")
    nv.append(cnv)
    nv.append(OxmlElement("p:cNvSpPr"))
    nvpr = OxmlElement("p:nvPr")
    ph = OxmlElement("p:ph")
    ph.set("type", "pic")
    ph.set("idx", str(placeholder_idx))
    nvpr.append(ph)
    nv.append(nvpr)
    replacement.append(nv)
    sppr = OxmlElement("p:spPr")
    if old.spPr.xfrm is not None:
        sppr.append(deepcopy(old.spPr.xfrm))
    geom = OxmlElement("a:prstGeom")
    geom.set("prst", "rect")
    geom.append(OxmlElement("a:avLst"))
    sppr.append(geom)
    replacement.append(sppr)
    body = OxmlElement("p:txBody")
    body.append(OxmlElement("a:bodyPr"))
    body.append(OxmlElement("a:lstStyle"))
    body.append(OxmlElement("a:p"))
    replacement.append(body)
    old.getparent().replace(old, replacement)
    # Remove the unused slide relationship so the source photo does not remain
    # silently embedded in the generated PPTX (shared references stay intact).
    if image_rid and not slide._element.xpath(
        f'.//*[@r:embed="{image_rid}"]'
    ):
        slide.part.drop_rel(image_rid)


def _replace_slide_text(shape, label):
    """Keep text-box formatting while removing all source words."""
    first_run = None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if first_run is None:
                first_run = run
            run.text = ""
    if first_run is not None:
        first_run.text = label
    else:
        shape.text_frame.paragraphs[0].add_run().text = label


def _prepare_slide_variants(prs, variants):
    """Replace slide content with editable slots without modifying template JSON."""
    for variant in variants:
        slide = prs.slides[variant["slide_index"]]
        used_indices = {shape.placeholder_format.idx for shape in slide.shapes
                        if shape.is_placeholder}
        next_idx = max(used_indices | {99}) + 1
        for slot in variant.get("image_slots", {}).values():
            if slot["placeholder_idx"] is None:
                _replace_picture_with_placeholder(slide, slot["shape_index"], next_idx)
                next_idx += 1
        for name, slot in variant["slots"].items():
            _replace_slide_text(slide.shapes[slot["shape_index"]], name)


def _fill_slide_variant(slide, variant, content, image_content):
    variant_id = variant["id"]
    for slot_name, value in content.items():
        if value is None:
            continue
        slot = variant["slots"].get(slot_name)
        if slot is None:
            raise ValueError(f"На слайде '{variant_id}' нет слота '{slot_name}'")
        _replace_slide_text(slide.shapes[slot["shape_index"]], str(value))
    for slot_name, path in image_content.items():
        slot = variant.get("image_slots", {}).get(slot_name)
        if slot is None:
            raise ValueError(f"На слайде '{variant_id}' нет слота изображения '{slot_name}'")
        if path is not None:
            image_path = Path(path)
            if not image_path.is_file():
                raise FileNotFoundError(f"Изображение не найдено: {image_path}")
            slide.shapes[slot["shape_index"]].insert_picture(str(image_path))


def create_presentation(
    template,
    slides,
    output_file,
    assets_dir,
    source_file=None,
):
    """
    Создает всю презентацию.
    """

    if Path(output_file).exists():
        raise FileExistsError(f"Презентация уже существует: {output_file}")

    prs = Presentation(source_file) if source_file else Presentation()

    if source_file and template.get("source_mode") != "slides":
        # Only remove slide instances in memory; retain all masters/layouts.
        for slide_id in list(prs.slides._sldIdLst):
            prs.part.drop_rel(slide_id.rId)
            prs.slides._sldIdLst.remove(slide_id)

    # --------------------------------------------------------
    # Размер слайда
    # --------------------------------------------------------

    slide_width = template[
        "slide_size"
    ]["width"]

    slide_height = template[
        "slide_size"
    ]["height"]

    if not source_file:
        prs.slide_width = Cm(slide_width)
        prs.slide_height = Cm(slide_height)

    variant_by_id = {variant["id"]: variant for variant in template["variants"]}
    layouts = (
        {str(layout.part.partname).lstrip("/"): layout
         for master in prs.slide_masters for layout in master.slide_layouts}
        if source_file else None
    )

    if source_file and template.get("source_mode") == "slides":
        _prepare_slide_variants(prs, template["variants"])

    # --------------------------------------------------------
    # Слайды
    # --------------------------------------------------------

    for index, slide_data in enumerate(
        slides,
        start=1,
    ):
        variant_id = slide_data[
            "variant"
        ]

        content = slide_data.get(
            "content",
            {},
        )
        image_content = slide_data.get("image_content", {})

        print(
            f"[{index}/{len(slides)}] "
            f"Создание: {variant_id}"
        )

        if source_file and template.get("source_mode") == "slides":
            try:
                variant = variant_by_id[variant_id]
            except KeyError as exc:
                raise ValueError(f"Неизвестный слайд: {variant_id}") from exc
            _fill_slide_variant(prs.slides[variant["slide_index"]], variant,
                                content, image_content)
        elif source_file:
            try:
                variant = variant_by_id[variant_id]
            except KeyError as exc:
                raise ValueError(f"Неизвестный макет: {variant_id}") from exc
            create_slide_from_layout(
                prs, template, variant_id, content, image_content,
                variant=variant, layouts=layouts,
            )
        else:
            if image_content:
                raise ValueError("Для вставки изображений нужен source_file")
            create_slide(
                prs=prs,
                template=template,
                variant_id=variant_id,
                content=content,
                assets_dir=assets_dir,
            )

    # --------------------------------------------------------
    # Сохранение
    # --------------------------------------------------------

    prs.save(
        output_file
    )

    print()

    print(
        f"Презентация создана: "
        f"{output_file}"
    )


def next_output_pair(requested_pptx):
    """Choose a free PPTX/JSON pair without replacing an earlier run."""
    requested_pptx = Path(requested_pptx)
    if requested_pptx.suffix.lower() != ".pptx":
        raise ValueError("Выходной файл должен иметь расширение .pptx")

    number = 1
    while True:
        pptx_path = (
            requested_pptx if number == 1 else
            requested_pptx.with_name(
                f"{requested_pptx.stem}_{number}{requested_pptx.suffix}"
            )
        )
        json_path = pptx_path.with_suffix(".json")
        if not pptx_path.exists() and not json_path.exists():
            return pptx_path, json_path
        number += 1


# ============================================================
# ПЕЧАТЬ ИНФОРМАЦИИ О ШАБЛОНЕ
# ============================================================

def print_template_info(template):
    """
    Показывает, что реально
    извлек parser из PPTX.
    """

    print()
    print("=" * 70)
    print("ШАБЛОН")
    print("=" * 70)

    print(
        f"Размер: "
        f"{template['slide_size']['width']} x "
        f"{template['slide_size']['height']} "
        f"{template['slide_size'].get('unit', '')}"
    )

    print()

    print("Варианты:")

    for variant in template.get(
        "variants",
        [],
    ):
        print()
        print(
            f"  {variant['id']}"
        )

        print("  Слоты:")

        for slot_name, slot in variant.get(
            "slots",
            {},
        ).items():

            print(
                f"    - {slot_name}: "
                f"x={slot['x']}, "
                f"y={slot['y']}, "
                f"w={slot['w']}, "
                f"h={slot['h']}, "
                f"h_align={slot.get('horizontal_align')}, "
                f"v_align={slot.get('vertical_align')}"
            )

        background = variant.get(
            "background"
        )

        if background:
            if background.get("target"):
                print(f"  Background: {background['target']}")
            if background.get("color"):
                print(f"  Background color: {background['color']}")

        images = variant.get(
            "images",
            [],
        )

        if variant.get("image_slots"):
            print("  Слоты изображений:")
            for name, slot in variant["image_slots"].items():
                print(f"    - {name}: idx={slot['placeholder_idx']}")

        if images:
            print(
                "  Images:"
            )

            for image in images:
                print(
                    f"    - {image['target']}"
                )

    print()
    print("=" * 70)
    print()


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Создать слайд для каждого макета PPTX")
    parser.add_argument("source", help="Путь к исходному PPTX")
    parser.add_argument("--output", help="Путь к выходному PPTX")
    parser.add_argument("--verbose", action="store_true", help="Показать все макеты и слоты")
    args = parser.parse_args()
    source_file = Path(args.source)
    if not source_file.is_file():
        parser.error(f"Шаблон не найден: {source_file}")
    requested_output = (Path(args.output) if args.output else
                        DEFAULT_OUTPUT_DIR / f"{source_file.stem}_generated.pptx")
    requested_output.parent.mkdir(parents=True, exist_ok=True)
    if source_file.resolve() == requested_output.resolve():
        parser.error("Исходный и выходной файлы должны различаться")
    try:
        output_file, json_file = next_output_pair(requested_output)
    except ValueError as exc:
        parser.error(str(exc))

    template = build_template(source_file)
    if args.verbose:
        print_template_info(template)
    else:
        print(f"Найдено макетов: {len(template['variants'])}; "
              f"слотов изображений: "
              f"{sum(len(v['image_slots']) for v in template['variants'])}")

    # Пример: {"Слайд со спикером": {"image": "speaker.jpg"}}.
    # Если файл не указан, в PPTX остаётся заполнитель для ручной вставки.
    IMAGE_BY_VARIANT = {}

    # При работе со слайдами исходное содержимое заменяется внутри генератора.
    slides = [
        {"variant": variant["id"],
         "content": ({slot: slot for slot, spec in variant["slots"].items()
                     if spec["placeholder_idx"] is not None}
                     if template.get("source_mode") != "slides" else {}),
         "image_content": IMAGE_BY_VARIANT.get(variant["id"], {})}
        for variant in template["variants"]
    ]
    print("Источник вариантов: " + ("слайды" if template.get("source_mode") == "slides"
                                    else "макеты образца"))
    create_presentation(template, slides, output_file, None,
                        source_file=source_file)
    save_template_json(template, json_file)
    print(f"JSON шаблона сохранён: {json_file}")

