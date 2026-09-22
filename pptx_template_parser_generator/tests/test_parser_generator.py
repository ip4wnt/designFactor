import unittest
import tempfile
import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from pptx_template_parser.template.normalizer import normalize_variant, normalize_images
from scripts.generate import add_background, create_presentation, next_output_pair
from pptx_template_parser import build_template
from pptx_template_parser.layout.variants import parse_variants
from scripts.export_json import main as save_template_json


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SAMPLE_TEMPLATES = PROJECT_ROOT / "sample_templates"


class BackgroundTests(unittest.TestCase):
    def test_default_master_uses_authored_slides_as_empty_templates(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.pptx"
            output = Path(directory) / "result.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Настоящий заголовок"
            slide.placeholders[1].text = "Настоящий текст"
            presentation.save(source)
            template = build_template(source)
            self.assertEqual(template["source_mode"], "slides")
            self.assertEqual(len(template["variants"]), 1)
            self.assertNotIn("Настоящий текст", json.dumps(template, ensure_ascii=False))
            create_presentation(template, [{"variant": "slide_001", "content": {}}],
                                output, None, source_file=source)
            generated = Presentation(output)
            self.assertEqual(generated.slides[0].shapes.title.text, "title")
            self.assertNotIn("Настоящий текст", " ".join(
                shape.text for shape in generated.slides[0].shapes if shape.has_text_frame))

    def test_custom_master_background_keeps_layout_mode(self):
        from pptx.dml.color import RGBColor
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "custom.pptx"
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[0])
            presentation.slide_master.background.fill.solid()
            presentation.slide_master.background.fill.fore_color.rgb = RGBColor(20, 30, 40)
            presentation.save(source)
            self.assertEqual(build_template(source)["source_mode"], "layouts")

    def test_supplied_plain_slide_deck_uses_slides(self):
        matches = list(SAMPLE_TEMPLATES.glob("Шаблон коты без образцов слайдов.pptx"))
        if not matches:
            self.skipTest("Пример с обычными слайдами отсутствует")
        template = build_template(matches[0])
        self.assertEqual(template["source_mode"], "slides")
        self.assertEqual(len(template["variants"]), len(Presentation(matches[0]).slides))
        self.assertTrue(any(v["image_slots"] for v in template["variants"]))
        self.assertTrue(any(v["constraints"]["image_regions_cm"]
                            for v in template["variants"]))
        self.assertTrue(all(v["background"] and v["background"].get("target")
                            for v in template["variants"]))
        self.assertEqual(template["variants"][0]["background"]["shape_index"], 0)
        template_before_generation = json.dumps(template, sort_keys=True)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output.pptx"
            create_presentation(template, [{"variant": v["id"], "content": {}}
                                           for v in template["variants"]],
                                output, None, source_file=matches[0])
            self.assertEqual(json.dumps(template, sort_keys=True), template_before_generation)
            generated = Presentation(output)
            self.assertEqual(len(generated.slides), len(Presentation(matches[0]).slides))
            for slide in generated.slides:
                self.assertEqual(sum(shape.shape_type == 13 for shape in slide.shapes), 1)
                background = slide.shapes[0]
                self.assertEqual((background.left, background.top, background.width,
                                  background.height),
                                 (0, 0, generated.slide_width, generated.slide_height))
            image_slot = next(v["image_slots"] for v in template["variants"]
                              if v["image_slots"])
            self.assertTrue(image_slot)
            placeholder = next(shape for shape in generated.slides[0].shapes
                               if shape.is_placeholder and shape.placeholder_format.type ==
                               PP_PLACEHOLDER.PICTURE)
            self.assertGreater(placeholder.width, 0)
            from io import BytesIO
            original_picture = next(shape for index, shape in enumerate(
                Presentation(matches[0]).slides[0].shapes)
                if index > 0 and shape.shape_type == 13)
            placeholder.insert_picture(BytesIO(original_picture.image.blob))

    def test_same_image_at_distinct_positions_is_not_lost(self):
        images = [
            {"target": "ppt/media/logo.png", "x": x, "y": 1,
             "w": 2, "h": 2, "name": "logo"}
            for x in (1, 8, 1)
        ]
        result = normalize_images(images, 25, 14)
        self.assertEqual([image["x"] for image in result["images"]], [1, 8])

    def test_generator_creates_new_pptx_and_json_on_each_run(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "template.pptx"
            output = Path(directory) / "result.pptx"
            Presentation().save(source)
            command = [
                sys.executable,
                "-m", "scripts.generate",
                str(source), "--output", str(output),
            ]
            for _ in range(2):
                subprocess.run(command, check=True, capture_output=True, text=True,
                               cwd=PROJECT_ROOT)

            for name in ("result", "result_2"):
                generated = output.with_name(f"{name}.pptx")
                metadata = output.with_name(f"{name}.json")
                self.assertEqual(len(Presentation(generated).slides), 11)
                self.assertEqual(len(json.loads(metadata.read_text(encoding="utf-8"))
                                     ["variants"]), 11)
            self.assertEqual(next_output_pair(output)[0].name, "result_3.pptx")

    def test_generator_does_not_overwrite_existing_output(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "template.pptx"
            output = Path(directory) / "existing.pptx"
            Presentation().save(source)
            output.write_bytes(b"existing user data")
            template = build_template(source)
            with self.assertRaises(FileExistsError):
                create_presentation(template, [], output, None, source_file=source)
            self.assertEqual(output.read_bytes(), b"existing user data")

    def test_json_contains_actual_layout_types_and_constraints(self):
        presentation = Presentation()
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "input.pptx"
            output = Path(directory) / "rules.json"
            presentation.save(source)
            save_template_json([str(source), "--output", str(output)])
            data = json.loads(output.read_text(encoding="utf-8"))

        self.assertEqual(data["type"], "presentation_template")
        self.assertEqual(len(data["variants"]), len(presentation.slide_layouts))
        title_layout = data["variants"][0]
        self.assertEqual(title_layout["type"], "title_text")
        self.assertIn("title", title_layout["constraints"]["text_regions_cm"])
        self.assertEqual(title_layout["constraints"]["slide_bounds_cm"]["w"],
                         data["slide_size"]["width"])

    def test_all_masters_are_parsed_without_slides(self):
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[0])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.pptx"
            presentation.save(path)
            variants = parse_variants(path)
        self.assertEqual(len(variants), len(presentation.slide_master.slide_layouts))
        self.assertNotEqual(len(variants), len(presentation.slides))

    def test_supplied_template_includes_second_master(self):
        path = SAMPLE_TEMPLATES / "VK Tech шаблон 1.pptx"
        if not path.exists():
            self.skipTest("Шаблон не найден")
        presentation = Presentation(path)
        variants = parse_variants(path)
        self.assertEqual(
            len(variants),
            sum(len(master.slide_layouts) for master in presentation.slide_masters),
        )
        self.assertEqual(len({variant["name"] for variant in variants}), len(variants))

    def test_generated_layout_preserves_grouped_logo_and_text(self):
        path = SAMPLE_TEMPLATES / "VK Tech шаблон 1.pptx"
        if not path.exists():
            self.skipTest("Шаблон не найден")
        template = build_template(path)
        variant = template["variants"][3]
        source = Presentation(path)
        layout = source.slide_masters[1].slide_layouts[0]
        self.assertTrue(any(shape.shape_type == 6 for shape in layout.shapes))
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            create_presentation(
                template,
                [{"variant": variant["id"], "content": {"title": "Проверка"}}],
                output,
                None,
                source_file=path,
            )
            generated = Presentation(output)
            self.assertEqual(len(generated.slides), 1)
            self.assertEqual(generated.slides[0].slide_layout.name, layout.name)
            self.assertEqual(generated.slides[0].shapes.title.text, "Проверка")
            self.assertTrue(any(
                shape.shape_type == 6
                for shape in generated.slides[0].slide_layout.shapes
            ))

    def test_workspace_template_uses_its_own_layouts(self):
        path = SAMPLE_TEMPLATES / "VK Tech шаблон 2.pptx"
        if not path.exists():
            path = SAMPLE_TEMPLATES / (
                "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx"
            )
        if not path.exists():
            self.skipTest("Второй шаблон не найден")
        template = build_template(path)
        self.assertEqual(
            len(template["variants"]),
            sum(len(master.slide_layouts) for master in Presentation(path).slide_masters),
        )
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            slides = [
                {"variant": variant["id"],
                 "content": {name: name for name, slot in variant["slots"].items()
                             if slot["placeholder_idx"] is not None}}
                for variant in template["variants"]
            ]
            create_presentation(template, slides, output, None, source_file=path)
            generated = Presentation(output)
            self.assertEqual(len(generated.slides), len(slides))
            self.assertEqual(
                [slide.slide_layout.name for slide in generated.slides],
                [variant["id"] for variant in template["variants"]],
            )

    def test_top_footer_is_recreated_as_an_editable_slot(self):
        candidates = [
            path for path in SAMPLE_TEMPLATES.glob("*3.pptx")
            if not path.name.startswith("~$")
        ]
        if not candidates:
            self.skipTest("Третий шаблон не найден")
        path = candidates[0]
        template = build_template(path)
        presentation = Presentation(path)
        self.assertEqual(len(template["variants"]), sum(
            len(master.slide_layouts) for master in presentation.slide_masters
        ))
        footer_variants = [variant for variant in template["variants"]
                           if "footer" in variant["slots"]]
        self.assertTrue(footer_variants)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "footer.pptx"
            create_presentation(
                template,
                [{"variant": variant["id"], "content": {"footer": "footer"}}
                 for variant in footer_variants],
                output, None, source_file=path,
            )
            generated = Presentation(output)
            for variant, slide in zip(footer_variants, generated.slides):
                slot = variant["slots"]["footer"]
                footer = slide.placeholders[slot["placeholder_idx"]]
                self.assertEqual(footer.placeholder_format.type, PP_PLACEHOLDER.FOOTER)
                self.assertEqual(footer.text, "footer")
                self.assertAlmostEqual(footer.top / 360000, slot["y"], places=2)

    def test_picture_placeholder_is_not_replaced_with_text(self):
        paths = [p for p in SAMPLE_TEMPLATES.glob("*3.pptx")
                 if not p.name.startswith("~$")]
        if not paths:
            self.skipTest("Третий шаблон не найден")
        source = paths[0]
        template = build_template(source)
        variant = next(v for v in template["variants"] if v["image_slots"])
        self.assertNotIn("image", variant["slots"])
        image_idx = variant["image_slots"]["image"]["placeholder_idx"]
        self.assertNotIn(image_idx, [s["placeholder_idx"] for s in
                                     variant["slots"].values()])

        text = {name: name for name, slot in variant["slots"].items()
                if slot["placeholder_idx"] is not None}
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "picture-placeholder.pptx"
            create_presentation(template, [{"variant": variant["id"],
                                            "content": text}], output,
                                None, source_file=source)
            picture = Presentation(output).slides[0].placeholders[image_idx]
            self.assertEqual(picture.placeholder_format.type, PP_PLACEHOLDER.PICTURE)
            self.assertEqual(picture.text, "")

            photo = PROJECT_ROOT / "assets" / "vktech" / "logo.png"
            if photo.is_file():
                picture_output = Path(directory) / "with-picture.pptx"
                create_presentation(template, [{"variant": variant["id"],
                                                "content": text,
                                                "image_content": {"image": photo}}],
                                    picture_output, None, source_file=source)
                self.assertGreater(
                    len(Presentation(picture_output).slides[0]
                        .placeholders[image_idx].image.blob),
                    0,
                )

    def test_master_color_survives_without_full_slide_image(self):
        variant = normalize_variant(
            {"name": "test", "slots_raw": [], "images_raw": [],
             "background_raw": {"color": "#123456"}},
            25.4, 14.29,
        )
        self.assertEqual(variant["background"], {"color": "#123456"})

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        add_background(slide, variant["background"], "")
        self.assertEqual(str(slide.background.fill.fore_color.rgb), "123456")

    def test_master_image_becomes_full_slide_background(self):
        variant = normalize_variant(
            {"name": "test", "slots_raw": [], "images_raw": [],
             "background_raw": {"target": "ppt/media/bg.png"}},
            25.4, 14.29,
        )
        self.assertEqual(variant["background"]["w"], 25.4)
        self.assertEqual(variant["background"]["target"], "ppt/media/bg.png")


if __name__ == "__main__":
    unittest.main()
