from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from pptx_template_parser.layout.shapes import (
    parse_layout_shapes,
)


def parse_variants(pptx_path):
    presentation = (
        pptx_path if hasattr(pptx_path, "slide_masters")
        else Presentation(pptx_path)
    )

    variants = []

    used_names = set()
    # presentation.slide_layouts exposes only layouts of the first master.
    for master_index, master in enumerate(presentation.slide_masters, start=1):
        for index, layout in enumerate(master.slide_layouts):
            parsed = parse_layout_shapes(layout)
            name = layout.name
            if name in used_names:
                name = f"{name} [образец {master_index}]"
            if name in used_names:
                name = f"{name} [образец {master_index}, макет {index + 1}]"
            suffix = 2
            while name in used_names:
                name = f"{layout.name} [образец {master_index}, макет {index + 1}, {suffix}]"
                suffix += 1
            used_names.add(name)

            variants.append({
                "index": index,
                "name": name,
                "layout_name": layout.name,
                "master_index": master_index,
                "master_file": str(master.part.partname).lstrip("/"),
                "layout_file": str(layout.part.partname).lstrip("/"),
                "slots_raw": parsed["text_shapes"],
                "images_raw": parsed["images"],
                "image_placeholders_raw": parsed["image_placeholders"],
            })

    return variants


def has_default_layouts(presentation):
    """Recognize the untouched built-in layout structure conservatively.

    Names differ by Office language; placeholder types do not. An ambiguous
    template stays in layout mode so custom designs are never discarded.
    """
    if len(presentation.slide_masters) != 1:
        return False
    master = presentation.slide_masters[0]
    layouts = list(master.slide_layouts)
    if len(layouts) not in (11, 12):
        return False
    if any(not shape.is_placeholder for shape in master.shapes):
        return False
    if not _default_background(master.element):
        return False
    essential = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
    for layout in layouts:
        if any(not shape.is_placeholder for shape in layout.shapes):
            return False
        if not _default_background(layout.element):
            return False
        kinds = {shape.placeholder_format.type for shape in layout.shapes}
        if not essential.issubset(kinds):
            return False
    return True


def _default_background(element):
    bg = element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
    if bg is None:
        return True
    children = list(bg)
    if len(children) != 1 or children[0].tag != '{http://schemas.openxmlformats.org/presentationml/2006/main}bgRef':
        return False
    ref = children[0]
    return (ref.get('idx') == '1001' and len(ref) == 1 and
            ref[0].tag == '{http://schemas.openxmlformats.org/drawingml/2006/main}schemeClr'
            and ref[0].get('val') == 'bg1')


def parse_slide_variants(presentation):
    variants = []
    for index, slide in enumerate(presentation.slides):
        layout = slide.slide_layout
        parsed = parse_layout_shapes(slide)
        variants.append({
            'index': index,
            'slide_index': index,
            'name': f'slide_{index + 1:03d}',
            'layout_name': layout.name,
            'master_index': 1,
            'master_file': str(layout.slide_master.part.partname).lstrip('/'),
            'layout_file': str(layout.part.partname).lstrip('/'),
            'slide_file': str(slide.part.partname).lstrip('/'),
            'slots_raw': parsed['text_shapes'],
            'images_raw': parsed['images'],
            'image_placeholders_raw': parsed['image_placeholders'],
        })
    return variants
